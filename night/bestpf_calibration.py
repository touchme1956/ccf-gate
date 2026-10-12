#!/usr/bin/env python3
"""night/bestpf_calibration.py — R8 world_calibration（4種類の『予想の作り方』を過去に当てはめ、どれがどれだけ当たったかを測る）

読むだけ。門・採点・配分（portfolio.json）には触れない。事前登録 out/bestpf_prereg.json（7c4b672）の modules.R8_world_calibration どおり。

問い: 4つの世界（A1 直近26年の実績・A2 直近15年の実績・B 100年の平均・C/D 今の値段からの予想）は、
      過去に予想として使っていたら、どれがどれだけ当たったか → 最終の選択で使う世界の重み π を出す。

パネル（どれも『その時点までに公開されていたデータだけ』で予想を作る）
  P1 French 49業種（時価加重の月次・1926-07〜・ドル建て名目）。割安さ＝log(Sum of BE / Sum of ME)
     ＝年次の表の y 行（y年6月に組む: BE は y−1年の会計年度・ME は y−1年12月末）→ y年6月に分かったとみなす。
  P2 French の国（F-F_International_Countries の Value-Weight Dollar Returns・All 4 Data Items Not Reqd の Mkt・1975〜・
     ドル建て名目）。割安さ＝log(配当利回り Yld)（同じファイルの『Average of Annual Value-Weight Averages』表）。
     r 行は r−1年12月末に組む表＝r−1年12月末に分かったとみなす。米国は French の国のファイルに無いので入らない。
  P3 JST マクロ歴史データ R6（国ごとの株の総リターン eq_tr・配当利回り eq_dp・年次・1870〜2020）。
     主は『現地通貨の実質』（cpi で割る＝国ごとのインフレの差を消す）。ドル建て（xrusd）は追加の点検だけ。
  米国市場（Shiller の ie_data.xls・CAPE・実質の配当込み）は診断だけ（π には使わない）。

時点: 5年ごと（年が0か5で終わる年の12月末）。予想の先: 15年（主）・10年と20年（感度）。先が全部データの中にある時点だけ。

予想の作り方（その時点 t までのデータだけ）
  F_A1 … t までの26年の年率（26年がそろわない資産は欠測）
  F_A2 … t までの15年の年率
  F_B  … t までに分かる全期間の年率（10年以上）を、その時点のパネルの平均へ50%縮める
  F_V  … パネルの平均（F_B と同じ錨）＋ γ×(その時点の割安さ − パネルの平均の割安さ)。
         γ は『t までに結果が出そろった過去の時点 s（s＋先 ≦ t）』だけで、各時点の平均を引いた形で当てはめる（横断の傾きを束ねた最小二乗）。
         γ を当てはめる過去が無い初めの時点は欠測。

測るもの: 二乗誤差の平均（パネルごと・まとめて）・各時点の順位相関の平均と t・実現を予想に回帰した傾き。
          ★傾きは『横断の中の傾き』（各時点の平均を引く）で読む。3パネルを束ねた傾きは名目（P1・P2）と実質（P3）の水準の差を含む
          見かけの値（2026-10-12 検証役の指摘）。水準の時期の当たりは level_timing（パネルごと・各時点の横断の平均どうし）で別に測る。
          4種類を同じ（資産×時点）の集合で比べる（主＝4つともそろう行。2つずつの比較はその組でそろう行）。
重み: 主＝π_w ∝ 1/MSE（A1→A1・A2→A2・B→B・V→C と D で半分ずつ）。
      感度（事前登録）＝順位相関の正の部分に比例・非負で合計1の最小二乗の重ね（学習は1990年までの時点・検証は1991年以後）・先10年・先20年。
      計算できなければ等しく（各1/5）。
先読みの機械検査: 各（パネル×時点×先）で、データを t で切ってから予想を計算し直し、1ビットも変わらないことを確かめる。

出力: out/bestpf_calibration.json
"""
import datetime, hashlib, io, json, math, os, random, statistics as S, sys, zipfile

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(BASE, 'night'))
import mw_common as M   # noqa: E402

CACHE = os.path.join(BASE, 'out', '_bestpf_cache', 'R8')
os.makedirs(CACHE, exist_ok=True)
M.CACHE = CACHE          # French の zip の置き場所をこのモジュールの場所へ（mw_common.get はモジュールの CACHE を呼ぶたびに読む）
OUT = os.path.join(BASE, 'out', 'bestpf_calibration.json')
PREREG = 'out/bestpf_prereg.json（7c4b672）modules.R8_world_calibration'

FR_49 = '49_Industry_Portfolios'
FR_CTRY_URL = 'https://mba.tuck.dartmouth.edu/pages/faculty/ken.french/ftp/F-F_International_Countries.zip'
JST_URL = 'https://www.macrohistory.net/app/download/9834512569/JSTdatasetR6.xlsx'
SHILLER_URL = 'http://www.econ.yale.edu/~shiller/data/ie_data.xls'

FCS = ['A1', 'A2', 'B', 'V']
WORLDS = ['A1', 'A2', 'B', 'C', 'D']
HMAIN, HSENS = 15, (10, 20)
SHRINK = 0.5
B_MIN_YEARS = 10           # F_B の『全期間』は10年以上（事前登録は下限を決めていない＝ここで決めた・結果を見る前）
CS_MIN = 5                 # 1つの時点の横断で順位相関・γ の当てはめに使う資産の最低数
TRAIN_END = 1990           # 重ねの学習は1990年までの時点・検証は1991年以後
DEU_HYPER = (1922, 1923)   # JST: 年平均の物価・為替では10億倍の物価の年を割れない → 欠測（理由は data_notes）


# ───────────────────────── 小道具 ─────────────────────────
def sha16(b):
    return hashlib.sha1(b).hexdigest()[:16]


def fetch(url, name, max_age_days=3650):
    """M.get と同じ（キャッシュはこのモジュールの場所）"""
    return M.get(url, name=name, max_age_days=max_age_days)


def rank(x):
    o = sorted(range(len(x)), key=lambda i: x[i])
    r = [0.0] * len(x)
    i = 0
    while i < len(o):
        j = i
        while j + 1 < len(o) and x[o[j + 1]] == x[o[i]]:
            j += 1
        for k in range(i, j + 1):
            r[o[k]] = (i + j) / 2 + 1
        i = j + 1
    return r


def pearson(a, b):
    n = len(a)
    if n < 3:
        return None
    ma, mb = sum(a) / n, sum(b) / n
    va = sum((x - ma) ** 2 for x in a); vb = sum((y - mb) ** 2 for y in b)
    if va <= 0 or vb <= 0:
        return None
    return sum((x - ma) * (y - mb) for x, y in zip(a, b)) / math.sqrt(va * vb)


def spearman(a, b):
    return pearson(rank(a), rank(b))


def nw_t(x, lag):
    """Newey-West（Bartlett）の t。短い列でも使えるように（n≥4）。重なる窓の自己相関で t が膨らまないように"""
    n = len(x)
    if n < 4:
        return None
    m = sum(x) / n
    e = [v - m for v in x]
    s = sum(v * v for v in e) / n
    for L in range(1, min(lag, n - 1) + 1):
        s += 2 * (1 - L / (lag + 1)) * sum(e[i] * e[i - L] for i in range(L, n)) / n
    return m / math.sqrt(s / n) if s > 0 else None


def tstat(x):
    n = len(x)
    if n < 3:
        return None
    sd = S.stdev(x)
    return (sum(x) / n) / (sd / math.sqrt(n)) if sd > 0 else None


def ols(x, y):
    n = len(x)
    if n < 3:
        return None
    mx, my = sum(x) / n, sum(y) / n
    sxx = sum((a - mx) ** 2 for a in x)
    if sxx <= 0:
        return None
    b = sum((a - mx) * (c - my) for a, c in zip(x, y)) / sxx
    return {'slope': b, 'intercept': my - b * mx}


def r4(v, k=4):
    return None if v is None else round(v, k)


# ───────────────────────── パネル ─────────────────────────
class Panel:
    """ret: {資産: {日付: リターン(小数)}}（月次は yyyymm・年次は yyyy）
       val: {資産: {分かった日: 割安さ(log)}}（同じ日付の数え方）
       grid: 予想を作る時点（年）。end(t)＝t年12月末"""

    def __init__(self, name, freq, ret, val, grid, unit, notes):
        self.name, self.freq, self.ret, self.val, self.grid, self.unit, self.notes = name, freq, ret, val, grid, unit, notes
        self.per_year = 12 if freq == 'M' else 1

    def end(self, y):
        return y * 100 + 12 if self.freq == 'M' else y

    def start(self, y):
        return y * 100 + 1 if self.freq == 'M' else y

    def last_data(self):
        return max(max(s) for s in self.ret.values() if s)

    def cut(self, t):
        """t年12月末で切ったパネル（先読みの機械検査用）"""
        e = self.end(t)
        ret = {a: {d: v for d, v in s.items() if d <= e} for a, s in self.ret.items()}
        val = {a: {d: v for d, v in s.items() if d <= e} for a, s in self.val.items()}
        return Panel(self.name, self.freq, ret, val, self.grid, self.unit, self.notes)

    def periods(self, y0, y1):
        if self.freq == 'M':
            return [y * 100 + m for y in range(y0, y1 + 1) for m in range(1, 13)]
        return list(range(y0, y1 + 1))

    def ann(self, a, y0, y1):
        """y0年1月〜y1年12月の年率（幾何）。1つでも欠ければ None（欠測を0と読まない）"""
        s = self.ret.get(a, {})
        ps = self.periods(y0, y1)
        g = 0.0
        for p in ps:
            v = s.get(p)
            if v is None or v <= -1:
                return None
            g += math.log1p(v)
        return math.exp(g * self.per_year / len(ps)) - 1

    def ann_all(self, a, t):
        """t年12月末までに分かる全期間の年率（欠けた期間は飛ばす＝0と読まない）。B_MIN_YEARS 未満なら None"""
        e = self.end(t)
        xs = [v for d, v in self.ret.get(a, {}).items() if d <= e]
        if len(xs) < B_MIN_YEARS * self.per_year or any(v <= -1 for v in xs):
            return None
        return math.exp(math.fsum(math.log1p(v) for v in xs) * self.per_year / len(xs)) - 1

    def val_at(self, a, t):
        """t年12月末の割安さ＝分かった日が (t年12月末−12か月, t年12月末] の値（古すぎる値は使わない）"""
        e = self.end(t)
        lo = (t - 1) * 100 + 12 if self.freq == 'M' else t - 1
        c = [(d, v) for d, v in self.val.get(a, {}).items() if lo < d <= e and v is not None]
        return max(c)[1] if c else None


def forecasts(P, t, H):
    """時点 t の4種類の予想（P は t で切ってあってもなくても同じ答えを返さなければならない）
    → {資産: {'A1','A2','B','V'}}, 補足 {m_hat, gamma, n_fit, fit_points}"""
    assets = sorted(P.ret)
    hist = {a: P.ann_all(a, t) for a in assets}
    hv = [v for v in hist.values() if v is not None]
    m_hat = sum(hv) / len(hv) if len(hv) >= CS_MIN else None
    # γ: s＋H ≦ t の過去の時点だけ（結果が t までに出そろっている）
    num = den = 0.0
    n_fit, pts = 0, []
    for s in P.grid:
        if s + H > t:
            continue
        pair = []
        for a in assets:
            v, y = P.val_at(a, s), P.ann(a, s + 1, s + H)
            if v is not None and y is not None:
                pair.append((v, y))
        if len(pair) < CS_MIN:
            continue
        mv = sum(p[0] for p in pair) / len(pair); my = sum(p[1] for p in pair) / len(pair)
        num += sum((p[0] - mv) * (p[1] - my) for p in pair)
        den += sum((p[0] - mv) ** 2 for p in pair)
        n_fit += len(pair); pts.append(s)
    gamma = num / den if den > 0 else None
    vt = {a: P.val_at(a, t) for a in assets}
    vv = [v for v in vt.values() if v is not None]
    vbar = sum(vv) / len(vv) if len(vv) >= CS_MIN else None
    out = {}
    for a in assets:
        f = {'A1': P.ann(a, t - 25, t), 'A2': P.ann(a, t - 14, t)}
        f['B'] = (1 - SHRINK) * hist[a] + SHRINK * m_hat if hist[a] is not None and m_hat is not None else None
        f['V'] = (m_hat + gamma * (vt[a] - vbar)) if (m_hat is not None and gamma is not None and vbar is not None and vt[a] is not None) else None
        out[a] = f
    return out, {'m_hat': m_hat, 'gamma': gamma, 'n_fit': n_fit, 'fit_points': pts, 'vbar': vbar}


# ───────────────────────── データの読み込み ─────────────────────────
def load_p1(price_updated=False):
    T = M.french_tables(FR_49)
    rt = T['Average Value Weighted Returns -- Monthly']
    bm = T['Sum of BE / Sum of ME']
    cols = rt['cols']
    ret = {c: {} for c in cols}
    for d, row in rt['data'].items():
        for c, x in zip(cols, row):
            if x is not None:
                ret[c][d] = x / 100
    val, nonpos = {c: {} for c in cols}, []
    for y, row in bm['data'].items():
        for c, x in zip(cols, row):
            if x is None:
                continue
            if x <= 0:                       # 簿価が0以下（自社株買いで負の簿価など）→ log が無い＝欠測
                nonpos.append((c, y, x)); continue
            if not price_updated:
                val[c][y * 100 + 6] = math.log(x)     # y年6月に分かる
            else:
                # 追加の点検: y年12月の値段へ更新（y年の配当込みリターンで割る＝配当の分だけ近似）。分かるのは y年12月
                g = [ret[c].get(y * 100 + m) for m in range(1, 13)]
                if all(v is not None for v in g):
                    val[c][y * 100 + 12] = math.log(x) - math.fsum(math.log1p(v) for v in g)
    b = open(os.path.join(CACHE, f'fr_{FR_49}.zip'), 'rb').read()
    notes = {'source': M.FR.format(FR_49), 'sha1_16': sha16(b), 'returns': 'Average Value Weighted Returns -- Monthly（%→小数）',
             'valuation': 'Sum of BE / Sum of ME（年次の y 行を y年6月に分かったとみなす）' + ('・y年の配当込みリターンで12月の値段へ更新' if price_updated else ''),
             'nonpositive_BE_ME（欠測にした）': [f'{c} {y}: {x}' for c, y, x in nonpos],
             'first_month': min(min(s) for s in ret.values() if s), 'last_month': max(max(s) for s in ret.values() if s)}
    grid = list(range(1930, 2026, 5))
    return Panel('P1_French49', 'M', ret, val, grid, 'ドル建て名目（配当込み）', notes)


def ctry_tables():
    b = fetch(FR_CTRY_URL, 'F-F_International_Countries.zip')
    z = zipfile.ZipFile(io.BytesIO(b))
    out = {}
    for fn in z.namelist():
        if not fn.lower().endswith('.dat'):
            continue
        L = z.read(fn).decode('latin-1').split('\n')
        i = next(k for k, l in enumerate(L) if 'Value-Weight Dollar Returns' in l and 'Not Reqd' in l)  # 最初＝月次
        assert L[i + 2].split()[0] == 'Mkt', (fn, L[i + 2])
        ret = {}
        started = False
        for l in L[i + 3:]:
            p = l.split()
            if not p or not (p[0].isdigit() and len(p[0]) == 6):
                if started:
                    break
                continue
            started = True
            v = float(p[1])
            if v > -99:
                ret[int(p[0])] = v / 100
        j = next(k for k, l in enumerate(L) if l.startswith('Average of Annual Value-Weight Averages'))
        assert 'Not Reqd' in L[j + 2], (fn, L[j + 2])
        hk = next(k for k in range(j, j + 8) if L[k].split()[:5] == ['Firms', 'B/M', 'E/P', 'CE/P', 'Yld'])
        rat = {}
        for l in L[hk + 1:]:
            p = l.split()
            if not p or not p[0].isdigit():
                break
            rat[int(p[0])] = {'BM': float(p[2]), 'EP': float(p[3]), 'CEP': float(p[4]), 'Yld': float(p[5])}
        out[fn[:-4]] = (ret, rat)
    return out, b


def load_p2(ratio='Yld', lag_years=0):
    tabs, b = ctry_tables()
    ret, val, bad = {}, {}, []
    for c, (r, rat) in tabs.items():
        ret[c] = r
        val[c] = {}
        for y, d in rat.items():
            x = d[ratio]
            if x is None or x <= 0:
                bad.append((c, y, x)); continue
            val[c][(y - 1 + lag_years) * 100 + 12] = math.log(x / 100)    # r 行＝ r−1年12月末の値
    notes = {'source': FR_CTRY_URL, 'sha1_16': sha16(b),
             'returns': 'Value-Weight Dollar Returns・All 4 Data Items Not Reqd の Mkt（月次・%→小数）',
             'valuation': f'Average of Annual Value-Weight Averages（Not Reqd）の {ratio}（r 行を r−1年12月末に分かったとみなす' + (f'・さらに{lag_years}年遅らせた' if lag_years else '') + '）',
             'nonpositive（欠測にした）': [f'{c} {y}: {x}' for c, y, x in bad],
             'countries': {c: [min(r), max(r)] for c, (r, _) in tabs.items()},
             'caveat': 'French の国のファイルに米国は無い。日本は2008年の行から社数が645→3,026社へ変わる（出所の切替）'}
    grid = list(range(1975, 2026, 5))
    return Panel(f'P2_FrenchCountries_{ratio}' + (f'_lag{lag_years}' if lag_years else ''), 'M', ret, val, grid, 'ドル建て名目（配当込み）', notes)


def jst_raw():
    import openpyxl
    b = fetch(JST_URL, 'JSTdatasetR6.xlsx')
    wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
    rows = wb['Sheet1'].iter_rows(values_only=True)
    head = [str(c) for c in next(rows)]
    ix = {c: i for i, c in enumerate(head)}
    data = {}
    for r in rows:
        if r[ix['year']] is None:
            continue
        rec = {}
        for c in ('eq_tr', 'eq_dp', 'cpi', 'xrusd', 'eq_tr_interp', 'eq_dp_interp'):
            v = r[ix[c]]
            rec[c] = float(v) if isinstance(v, (int, float)) else None
        data.setdefault(str(r[ix['iso']]), {})[int(r[ix['year']])] = rec
    return data, b


def load_p3(unit='real'):
    data, b = jst_raw()
    ret, val, dropped, extreme = {}, {}, [], []
    for iso, d in sorted(data.items()):
        r_, v_ = {}, {}
        for y in sorted(d):
            rec = d[y]
            tr = rec['eq_tr']
            if tr is None:
                continue
            if iso == 'DEU' and y in DEU_HYPER:
                dropped.append(f'{iso} {y}'); continue
            p = d.get(y - 1)
            if unit == 'real':
                if p is None or rec['cpi'] is None or p['cpi'] is None or p['cpi'] <= 0:
                    continue
                x = (1 + tr) / (rec['cpi'] / p['cpi']) - 1
            else:
                if p is None or rec['xrusd'] is None or p['xrusd'] is None or rec['xrusd'] <= 0:
                    continue
                x = (1 + tr) * (p['xrusd'] / rec['xrusd']) - 1
            r_[y] = x
            if x > 1.0 or x < -0.6:
                extreme.append(f'{iso} {y}: {x:+.2f}')
        for y, rec in d.items():
            if rec['eq_dp'] is not None and rec['eq_dp'] > 0:
                v_[y] = math.log(rec['eq_dp'])
        if r_:
            ret[iso], val[iso] = r_, v_
    notes = {'source': JST_URL, 'sha1_16': sha16(b),
             'returns': '株の総リターン eq_tr を ' + ('cpi（年次）で割った現地通貨の実質' if unit == 'real' else 'xrusd（現地通貨/米ドル）でドルへ換えた名目'),
             'valuation': 'log(eq_dp)（y年の値を y年末に分かったとみなす・0以下は欠測）',
             'dropped': dropped + ['（ドイツの1922〜23年: 物価が年に10億倍の年で、年次の物価・為替では割れない）'],
             'extreme_years（残した・|実質|>100% か < −60%）': extreme,
             'countries': {iso: [min(s), max(s), len(s)] for iso, s in ret.items()},
             'no_equity': sorted(set(data) - set(ret))}
    grid = list(range(1870, 2021, 5))
    return Panel('P3_JST_' + unit, 'A', ret, val, grid, '現地通貨の実質（配当込み）' if unit == 'real' else 'ドル建て名目（配当込み）', notes)


# ───────────────────────── 行を作る ─────────────────────────
def build_rows(P, H, check_prefix=True):
    rows, meta, pref = [], [], {'checked': 0, 'identical': 0, 'failures': []}
    last = P.last_data()
    for t in P.grid:
        if P.end(t + H) > last:
            continue
        F, info = forecasts(P, t, H)
        if check_prefix:
            Fc, infoc = forecasts(P.cut(t), t, H)
            pref['checked'] += 1
            same = json.dumps(F, sort_keys=True) == json.dumps(Fc, sort_keys=True) and json.dumps(info, sort_keys=True) == json.dumps(infoc, sort_keys=True)
            if same:
                pref['identical'] += 1
            else:
                pref['failures'].append(t)
        meta.append({'t': t, 'gamma': r4(info['gamma'], 5), 'm_hat': r4(info['m_hat']), 'n_fit': info['n_fit'], 'fit_points': info['fit_points']})
        for a, f in F.items():
            y = P.ann(a, t + 1, t + H)
            if y is None:
                continue
            rows.append({'panel': P.name, 't': t, 'asset': a, 'Y': y, **f})
    return rows, meta, pref


# ───────────────────────── 測る ─────────────────────────
def subset(rows, fcs):
    return [r for r in rows if all(r[f] is not None for f in fcs)]


def evaluate(rows, fcs, H):
    """同じ行の集合で fcs を比べる。→ {予想: 指標}・横断の表"""
    if not rows:
        return None
    groups = {}
    for r in rows:
        groups.setdefault((r['panel'], r['t']), []).append(r)
    cs_keys = sorted(k for k, g in groups.items() if len(g) >= CS_MIN)
    lag = max(1, math.ceil(H / 5) - 1)
    res = {'n_rows': len(rows), 'n_cross_sections': len(cs_keys), 'panels': sorted({r['panel'] for r in rows}),
           'times': sorted({r['t'] for r in rows})}
    cs_table = []
    for k in cs_keys:
        g = groups[k]
        e = {'panel': k[0], 't': k[1], 'n': len(g)}
        for f in fcs:
            e[f] = r4(spearman([r[f] for r in g], [r['Y'] for r in g]), 3)
        cs_table.append(e)
    for f in fcs:
        err = [r[f] - r['Y'] for r in rows]
        mse = sum(e * e for e in err) / len(err)
        # 横断の平均を引いた後の誤差（相対の当たり）
        rel = []
        for k, g in groups.items():
            mf = sum(r[f] for r in g) / len(g); my = sum(r['Y'] for r in g) / len(g)
            rel += [((r[f] - mf) - (r['Y'] - my)) ** 2 for r in g]
        rho = [e[f] for e in cs_table if e[f] is not None]
        # パネルごとに時間の順で NW（重なる窓）
        nwp = {}
        for p in res['panels']:
            xs = [e[f] for e in cs_table if e['panel'] == p and e[f] is not None]
            nwp[p] = r4(nw_t(xs, lag), 2) if xs else None
        o = ols([r[f] for r in rows], [r['Y'] for r in rows])
        # 横断の中の傾き（各時点の平均を引いて束ねる）
        dx, dy = [], []
        for k, g in groups.items():
            mf = sum(r[f] for r in g) / len(g); my = sum(r['Y'] for r in g) / len(g)
            dx += [r[f] - mf for r in g]; dy += [r['Y'] - my for r in g]
        ow = ols(dx, dy)
        # パネルの平均だけを引いた傾き（名目と実質のパネルの水準の差を消す・時期の差は残す）
        px, py = [], []
        for p in res['panels']:
            g = [r for r in rows if r['panel'] == p]
            mf = sum(r[f] for r in g) / len(g); my = sum(r['Y'] for r in g) / len(g)
            px += [r[f] - mf for r in g]; py += [r['Y'] - my for r in g]
        op = ols(px, py)
        # 水準の時期の当たり: パネルごとに、各時点の横断の平均（予想）と横断の平均（実現）を時間の順に並べた相関と傾き
        lt = {}
        for p in res['panels']:
            tm = sorted({k[1] for k in groups if k[0] == p})
            mfs, mys = [], []
            for t_ in tm:
                g = groups[(p, t_)]
                mfs.append(sum(r[f] for r in g) / len(g)); mys.append(sum(r['Y'] for r in g) / len(g))
            ok = len(tm) >= 5
            ol = ols(mfs, mys) if ok else None
            lt[p] = {'n_times': len(tm), 'corr': r4(pearson(mfs, mys), 3) if ok else None, 'slope': r4(ol['slope'], 3) if ol else None,
                     'note': None if ok else '時点が5未満なので出さない'}
        n_eff = len(rho) * 5 / H if rho else 0
        res[f] = {'MSE_pt2': r4(mse * 1e4, 2), 'RMSE_pt': r4(math.sqrt(mse) * 100, 2), 'bias_pt': r4(sum(err) / len(err) * 100, 2),
                  'MSE_relative_pt2': r4(sum(rel) / len(rel) * 1e4, 2), 'MSE_level_pt2': r4((mse - sum(rel) / len(rel)) * 1e4, 2),
                  'rank_corr_mean': r4(sum(rho) / len(rho), 3) if rho else None, 'rank_corr_n': len(rho),
                  't_naive': r4(tstat(rho), 2), 't_overlap_adj': r4((sum(rho) / len(rho)) / (S.stdev(rho) / math.sqrt(n_eff)), 2) if len(rho) >= 3 and S.stdev(rho) > 0 and n_eff > 0 else None,
                  'nw_t_by_panel': nwp, 'share_cs_positive': r4(sum(1 for x in rho if x > 0) / len(rho), 3) if rho else None,
                  'slope_pooled': r4(o['slope'], 3) if o else None, 'slope_within_cs': r4(ow['slope'], 3) if ow else None,
                  'slope_within_panel': r4(op['slope'], 3) if op else None, 'level_timing_by_panel': lt}
    return res, cs_table


def to_worlds(w4):
    """{A1,A2,B,V} → {A1,A2,B,C,D}（V を C と D で半分ずつ）"""
    s = sum(w4.values())
    if not s or s <= 0 or any(v is None for v in w4.values()):
        return {w: 0.2 for w in WORLDS}, '計算できないので等しく（各1/5）'
    w = {k: v / s for k, v in w4.items()}
    return {'A1': w['A1'], 'A2': w['A2'], 'B': w['B'], 'C': w['V'] / 2, 'D': w['V'] / 2}, None


def pi_inv_mse(ev, fcs=FCS):
    if ev is None or any(ev[f]['MSE_pt2'] in (None, 0) for f in fcs):
        return to_worlds({f: None for f in fcs})
    return to_worlds({f: 1 / ev[f]['MSE_pt2'] for f in fcs})


def pi_rank(ev, fcs=FCS):
    if ev is None:
        return to_worlds({f: None for f in fcs})
    w = {f: max(0.0, ev[f]['rank_corr_mean'] or 0.0) for f in fcs}
    if sum(w.values()) <= 0:
        return {x: 0.2 for x in WORLDS}, '順位相関がどれも0以下なので等しく（各1/5）'
    return to_worlds(w)


def nnls_simplex(X, y):
    """min ||y − X w||²  s.t. w≥0, Σw=1（4列なので部分集合を全部試す＝厳密）"""
    import itertools
    import numpy as np
    X, y = np.asarray(X, float), np.asarray(y, float)
    k = X.shape[1]
    best = None
    for m in range(1, k + 1):
        for S_ in itertools.combinations(range(k), m):
            A = X[:, S_]
            G = A.T @ A
            b = A.T @ y
            try:
                Gi = np.linalg.inv(G)
            except np.linalg.LinAlgError:
                continue
            one = np.ones(m)
            lam = (one @ Gi @ b - 1) / (one @ Gi @ one)
            w = Gi @ (b - lam * one)
            if (w < -1e-12).any():
                continue
            w = np.clip(w, 0, None); w = w / w.sum()
            full = np.zeros(k); full[list(S_)] = w
            sse = float(((y - X @ full) ** 2).sum())
            if best is None or sse < best[0] - 1e-15:
                best = (sse, full)
    return best[1].tolist() if best else None


def stack_eval(rows):
    tr = [r for r in rows if r['t'] <= TRAIN_END]
    va = [r for r in rows if r['t'] > TRAIN_END]
    out = {'train_rows': len(tr), 'valid_rows': len(va), 'train_times': sorted({r['t'] for r in tr}), 'valid_times': sorted({r['t'] for r in va})}
    if len(tr) < 20:
        return out, None
    w = nnls_simplex([[r[f] for f in FCS] for r in tr], [r['Y'] for r in tr])
    wfull = nnls_simplex([[r[f] for f in FCS] for r in rows], [r['Y'] for r in rows])
    out['w_train'] = {f: r4(x, 4) for f, x in zip(FCS, w)}
    out['w_full_sample（参考）'] = {f: r4(x, 4) for f, x in zip(FCS, wfull)} if wfull else None
    # 学習期間の 1/MSE の重み（検証での比較用）
    mse_tr = {f: sum((r[f] - r['Y']) ** 2 for r in tr) / len(tr) for f in FCS}
    winv = {f: (1 / mse_tr[f]) / sum(1 / v for v in mse_tr.values()) for f in FCS}

    def mse(rs, wd):
        return sum((sum(wd[f] * r[f] for f in FCS) - r['Y']) ** 2 for r in rs) / len(rs) * 1e4 if rs else None
    wd = dict(zip(FCS, w))
    eq = {f: 0.25 for f in FCS}
    out['valid_MSE_pt2'] = {'stack(学習の重み)': r4(mse(va, wd), 2), '1/MSE(学習の重み)': r4(mse(va, winv), 2), '等しく': r4(mse(va, eq), 2),
                            **{f'{f}だけ': r4(mse(va, {g: (1.0 if g == f else 0.0) for g in FCS}), 2) for f in FCS}}
    out['train_MSE_pt2'] = {'stack': r4(mse(tr, wd), 2), **{f'{f}だけ': r4(mse_tr[f] * 1e4, 2) for f in FCS}}
    out['w_inv_mse_train'] = {f: r4(v, 4) for f, v in winv.items()}
    return out, wd


def pair_tests(rows, H, fcs=FCS):
    """2つの予想の二乗誤差の差（Diebold-Mariano 型）: 各横断（パネル×時点）の平均の差を時点の数で t にする（重なりの分だけ数を減らす）"""
    groups = {}
    for r in rows:
        groups.setdefault((r['panel'], r['t']), []).append(r)
    out = {}
    for i, a in enumerate(fcs):
        for b in fcs[i + 1:]:
            d = [sum((r[a] - r['Y']) ** 2 - (r[b] - r['Y']) ** 2 for r in g) / len(g) * 1e4 for g in groups.values() if len(g) >= CS_MIN]
            pooled = sum((r[a] - r['Y']) ** 2 - (r[b] - r['Y']) ** 2 for r in rows) / len(rows) * 1e4
            n_eff = len(d) * 5 / H
            t = (sum(d) / len(d)) / (S.stdev(d) / math.sqrt(n_eff)) if len(d) >= 3 and S.stdev(d) > 0 else None
            out[f'{a}_minus_{b}'] = {'pooled_diff_pt2': r4(pooled, 2), 'cs_mean_diff_pt2': r4(sum(d) / len(d), 2) if d else None,
                                     'n_cs': len(d), 't_overlap_adj': r4(t, 2), 'share_cs_a_better': r4(sum(1 for x in d if x < 0) / len(d), 3) if d else None}
    return out


def bootstrap_pi(rows, n=2000, seed=20261012):
    """横断（パネル×時点）を戻して引き直し、π の主（1/MSE）の幅と『どの予想の MSE が最小か』の割合を出す（重なる窓の分だけ幅は狭く出る）"""
    rnd = random.Random(seed)
    groups = {}
    for r in rows:
        groups.setdefault((r['panel'], r['t']), []).append(r)
    keys = sorted(groups)
    sse = {k: {f: sum((r[f] - r['Y']) ** 2 for r in groups[k]) for f in FCS} for k in keys}
    cnt = {k: len(groups[k]) for k in keys}
    draws = {w: [] for w in WORLDS}
    best = {f: 0 for f in FCS}
    for _ in range(n):
        pick = [rnd.choice(keys) for _ in keys]
        N = sum(cnt[k] for k in pick)
        mse = {f: sum(sse[k][f] for k in pick) / N for f in FCS}
        best[min(FCS, key=lambda f: mse[f])] += 1
        w, _ = to_worlds({f: 1 / mse[f] for f in FCS})
        for x in WORLDS:
            draws[x].append(w[x])
    q = lambda v, p: sorted(v)[int(p * (len(v) - 1))]
    return {'n_draws': n, 'unit': '横断（パネル×時点）', 'pi_p10_p50_p90': {w: [r4(q(v, .1)), r4(q(v, .5)), r4(q(v, .9))] for w, v in draws.items()},
            'share_lowest_MSE': {f: r4(c / n, 3) for f, c in best.items()},
            'note': '15年の窓が5年ずつ重なるので、本当の幅はこれより広い'}


def timing_checks(P1):
    """割安さの『分かった日』の割り当てを数字で確かめる（先読みの機械検査では見えない部分）"""
    out = {}
    # 国: r 行→r+1 行の log(Yld) の変化が、年 r のリターンと強く負に相関すれば r+1 行＝ r年12月末の値段
    tabs, _ = ctry_tables()
    xs, y0, y1 = [], [], []
    for c, (ret, rat) in tabs.items():
        def yr(R):
            g = [ret.get(R * 100 + m) for m in range(1, 13)]
            return None if None in g else math.fsum(math.log1p(v) for v in g)
        for r in sorted(rat):
            if r + 1 in rat and rat[r]['Yld'] > 0 and rat[r + 1]['Yld'] > 0 and yr(r) is not None and yr(r + 1) is not None:
                xs.append(math.log(rat[r + 1]['Yld'] / rat[r]['Yld'])); y0.append(yr(r)); y1.append(yr(r + 1))
    out['P2_country_Yld'] = {'n': len(xs), 'corr_with_year_r_return': r4(pearson(xs, y0), 3), 'corr_with_year_r+1_return': r4(pearson(xs, y1), 3),
                             'reading': 'r+1 行の配当利回りは r年12月末の値段で割っている（年 r と強い負の相関・年 r+1 とは0）→ r年12月末に分かったとみなすのが正しい'}
    # 業種: y 行→y+1 行の log(BE/ME) の変化と、年 y・年 y−1 のリターン
    T = M.french_tables(FR_49)
    bm = T['Sum of BE / Sum of ME']
    cols = bm['cols']
    xs, a0, a1 = [], [], []
    for y in sorted(bm['data']):
        if y + 1 not in bm['data']:
            continue
        for j, c in enumerate(cols):
            u, v = bm['data'][y][j], bm['data'][y + 1][j]
            if not (u and v and u > 0 and v > 0):
                continue
            g0 = [P1.ret[c].get(y * 100 + m) for m in range(1, 13)]
            g1 = [P1.ret[c].get((y - 1) * 100 + m) for m in range(1, 13)]
            if None in g0 or None in g1:
                continue
            xs.append(math.log(v / u)); a0.append(math.fsum(math.log1p(x) for x in g0)); a1.append(math.fsum(math.log1p(x) for x in g1))
    out['P1_industry_BEME'] = {'n': len(xs), 'corr_with_year_y_return': r4(pearson(xs, a0), 3), 'corr_with_year_y-1_return': r4(pearson(xs, a1), 3),
                               'reading': 'y+1 行の ME は y年12月末（年 y と強い負の相関）＝ y 行は y−1年12月末の ME と y−1年度の BE → y年6月に分かる。t年12月の予想には y=t 行を使う（t+1 行は使わない）'}
    # JST: eq_div_rtn_t = eq_dp_t × (1+eq_capgain_t) が成り立てば eq_dp_t = D_t / P_t（t年末の値段）
    import openpyxl
    b = fetch(JST_URL, 'JSTdatasetR6.xlsx')
    wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
    rows = wb['Sheet1'].iter_rows(values_only=True)
    head = [str(c) for c in next(rows)]
    ix = {c: i for i, c in enumerate(head)}
    diffs = []
    for r in rows:
        dr, dp, cg = r[ix['eq_div_rtn']], r[ix['eq_dp']], r[ix['eq_capgain']]
        if all(isinstance(x, (int, float)) for x in (dr, dp, cg)):
            diffs.append(abs(dr - dp * (1 + cg)))
    out['P3_JST_eq_dp'] = {'n': len(diffs), 'share_identity_holds(<1e-4)': r4(sum(1 for x in diffs if x < 1e-4) / len(diffs), 4),
                           'reading': 'eq_div_rtn_t ＝ eq_dp_t ×(1＋eq_capgain_t) が全行で成り立つ → eq_dp_t は t年の配当 ÷ t年末の値段＝ t年末に分かる'}
    return out


# ───────────────────────── 米国市場（Shiller）の診断 ─────────────────────────
def shiller_diag():
    import xlrd
    b = fetch(SHILLER_URL, 'ie_data.xls')
    sh = xlrd.open_workbook(file_contents=b).sheet_by_name('Data')
    # 見出しの確認（列9＝Real Total Return Price・列12＝CAPE）
    hdr = ' | '.join(str(sh.cell_value(r, c)) for r in range(4, 8) for c in (0, 9, 12))
    dec = {}
    for i in range(8, sh.nrows):
        r = sh.row_values(i)
        if not isinstance(r[0], float) or not isinstance(r[9], float):
            continue
        y, m = int(r[0]), int(round((r[0] - int(r[0])) * 100))
        if m == 12:
            dec[y] = (r[9], r[12] if isinstance(r[12], float) else None)
    ys = sorted(dec)
    tr = {y: dec[y][0] for y in ys}
    cape = {y: dec[y][1] for y in ys}

    def ann(a, b_):   # a年12月末→b年12月末
        return (tr[b_] / tr[a]) ** (1 / (b_ - a)) - 1 if a in tr and b_ in tr else None
    first = ys[0]
    out = {'source': SHILLER_URL, 'sha1_16': sha16(b), 'header_check': hdr[:400], 'first_dec': first, 'last_dec': ys[-1], 'unit': '実質（配当再投資）'}
    for H in (10, 15, 20):
        rows, pre_ok, pre_n = [], 0, 0
        for t in range(first + 26, ys[-1] - H + 1):
            fA1, fA2, fB = ann(t - 26, t), ann(t - 15, t), ann(first, t)
            # V: s＋H ≦ t の過去の12月だけで 実現 ~ a + b log(CAPE) を当てはめる（10点以上）
            pts = [(math.log(cape[s]), ann(s, s + H)) for s in ys if s + H <= t and cape.get(s)]
            fV = None
            if len(pts) >= 10 and cape.get(t):
                o = ols([p[0] for p in pts], [p[1] for p in pts])
                fV = o['intercept'] + o['slope'] * math.log(cape[t]) if o else None
            y = ann(t, t + H)
            rows.append({'t': t, 'Y': y, 'A1': fA1, 'A2': fA2, 'B': fB, 'V': fV})
        cm = [r for r in rows if all(r[k] is not None for k in ('Y', 'A1', 'A2', 'B', 'V'))]
        e = {'n_years': len(cm), 'years': [cm[0]['t'], cm[-1]['t']] if cm else None}
        for f in FCS:
            if not cm:
                continue
            err = [r[f] - r['Y'] for r in cm]
            e[f] = {'RMSE_pt': r4(math.sqrt(sum(x * x for x in err) / len(err)) * 100, 2), 'bias_pt': r4(sum(err) / len(err) * 100, 2),
                    'corr_time_series': r4(pearson([r[f] for r in cm], [r['Y'] for r in cm]), 3)}
        # 5年ごとの点だけ（重なりを減らす）
        cm5 = [r for r in cm if r['t'] % 5 == 0]
        e['every5y'] = {f: {'RMSE_pt': r4(math.sqrt(sum((r[f] - r['Y']) ** 2 for r in cm5) / len(cm5)) * 100, 2),
                            'corr': r4(pearson([r[f] for r in cm5], [r['Y'] for r in cm5]), 3)} for f in FCS} if len(cm5) >= 4 else None
        e['n_every5y'] = len(cm5)
        out[f'H{H}'] = e
    out['note'] = '1つの市場の時間の並び（横断ではない）。12月ごとの窓は重なる。π には使わない（診断だけ）。'
    return out


# ───────────────────────── 主 ─────────────────────────
def run_panels(panels, H, check_prefix=True):
    rows, metas, prefs = [], {}, {}
    for P in panels:
        r, m, p = build_rows(P, H, check_prefix)
        rows += r; metas[P.name] = m; prefs[P.name] = p
    return rows, metas, prefs


def describe_all(rows, H):
    out = {}
    common = subset(rows, FCS)
    ev, cs = evaluate(common, FCS, H)
    out['common4_pooled'] = ev
    out['common4_cross_sections'] = cs
    out['common4_by_panel'] = {}
    for p in sorted({r['panel'] for r in common}):
        e, _ = evaluate([r for r in common if r['panel'] == p], FCS, H)
        out['common4_by_panel'][p] = e
    # 2つずつ（その組でそろう行）
    pw = {}
    for i, a in enumerate(FCS):
        for b in FCS[i + 1:]:
            rs = subset(rows, [a, b])
            e, _ = evaluate(rs, [a, b], H) if rs else (None, None)
            if e:
                pw[f'{a}_vs_{b}'] = {'n_rows': e['n_rows'], 'times': [e['times'][0], e['times'][-1]], 'panels': e['panels'],
                                     a: {k: e[a][k] for k in ('MSE_pt2', 'rank_corr_mean', 't_overlap_adj', 'slope_within_cs', 'slope_pooled')},
                                     b: {k: e[b][k] for k in ('MSE_pt2', 'rank_corr_mean', 't_overlap_adj', 'slope_within_cs', 'slope_pooled')},
                                     'better_MSE': a if e[a]['MSE_pt2'] < e[b]['MSE_pt2'] else b}
    out['pairwise'] = pw
    out['pair_tests_common4（二乗誤差の差・負なら左が良い）'] = pair_tests(common, H) if common else None
    # 1つずつ（その予想が出せる行すべて）
    out['single_all_rows'] = {}
    for f in FCS:
        rs = subset(rows, [f])
        e, _ = evaluate(rs, [f], H) if rs else (None, None)
        if e:
            out['single_all_rows'][f] = {'n_rows': e['n_rows'], 'times': [e['times'][0], e['times'][-1]], **{k: e[f][k] for k in ('MSE_pt2', 'rank_corr_mean', 't_overlap_adj', 'slope_within_cs', 'slope_pooled', 'bias_pt')}}
    return out, common, ev


def main():
    t0 = datetime.datetime.now()
    P1 = load_p1(); P2 = load_p2('Yld'); P3 = load_p3('real')
    MAIN = [P1, P2, P3]
    res, pis, prefix_all = {}, {}, {}
    rows_by_H, common_by_H = {}, {}
    for H in (HMAIN,) + HSENS:
        rows, metas, prefs = run_panels(MAIN, H)
        d, common, ev = describe_all(rows, H)
        d['gamma_by_time'] = metas
        res[f'H{H}'] = d
        prefix_all[f'H{H}'] = prefs
        rows_by_H[H], common_by_H[H] = rows, common
        print(f'H{H}: 行 {len(rows)}・4つそろう行 {len(common)}', flush=True)
        if ev:
            for f in FCS:
                print(f"   {f}: MSE {ev[f]['MSE_pt2']}  RMSE {ev[f]['RMSE_pt']}pt  順位相関 {ev[f]['rank_corr_mean']} (t調整 {ev[f]['t_overlap_adj']})  横断の中の傾き {ev[f]['slope_within_cs']}（束ねた傾き {ev[f]['slope_pooled']}＝パネルの水準の差を含む）  偏り {ev[f]['bias_pt']}", flush=True)

    ev15 = res[f'H{HMAIN}']['common4_pooled']
    pi_main, why_main = pi_inv_mse(ev15)
    pi_sens = {}
    for H in HSENS:
        pi_sens[f'H{H}'], _ = pi_inv_mse(res[f'H{H}']['common4_pooled'])
    pi_sens['rank_pos'], why_rank = pi_rank(ev15)
    boot = bootstrap_pi(common_by_H[HMAIN])
    print('π の幅（引き直し）', boot['pi_p10_p50_p90'], '最小の割合', boot['share_lowest_MSE'], flush=True)
    st, wtr = stack_eval(common_by_H[HMAIN])
    if wtr:
        pi_sens['stack'], _ = to_worlds(wtr)
    else:
        pi_sens['stack'] = {w: 0.2 for w in WORLDS}
        st['note'] = '学習の行が足りないので等しく（各1/5）'

    # ── 追加の点検（事前登録の外・π_sens には入れない＝最終の頑健さの数に混ぜない） ──
    extra = {}
    ev_panel = res[f'H{HMAIN}']['common4_by_panel']
    if all(ev_panel.get(p.name) for p in MAIN):
        mse_eq = {f: sum(ev_panel[p.name][f]['MSE_pt2'] for p in MAIN) / len(MAIN) for f in FCS}
        extra['panel_equal（パネルを同じ重みで）'] = to_worlds({f: 1 / mse_eq[f] for f in FCS})[0]
    for p in MAIN:
        if ev_panel.get(p.name):
            extra[f'only_{p.name}'] = pi_inv_mse(ev_panel[p.name])[0]
    extra['relative_MSE（各時点の平均を引いた後）'] = to_worlds({f: 1 / ev15[f]['MSE_relative_pt2'] for f in FCS})[0]
    alt = {}
    for nm, panels in (('P3_usd（JST をドル建て）', [P1, P2, load_p3('usd')]),
                       ('P2_EP（国の割安さを E/P に）', [P1, load_p2('EP'), P3]),
                       ('P2_lag1y（国の割安さを1年遅らせる）', [P1, load_p2('Yld', 1), P3]),
                       ('P1_BM_price_updated（業種の BE/ME を12月の値段へ更新）', [load_p1(True), P2, P3])):
        rows, _, prefs = run_panels(panels, HMAIN)
        common = subset(rows, FCS)
        ev, _ = evaluate(common, FCS, HMAIN)
        extra[nm] = pi_inv_mse(ev)[0]
        alt[nm] = {'n_rows': ev['n_rows'], **{f: {k: ev[f][k] for k in ('MSE_pt2', 'rank_corr_mean', 't_overlap_adj', 'slope_within_cs', 'slope_pooled')} for f in FCS},
                   'prefix_ok': all(p['checked'] == p['identical'] for p in prefs.values())}
        print('点検', nm, {f: (ev[f]['MSE_pt2'], ev[f]['rank_corr_mean']) for f in FCS}, flush=True)

    # ── 先読みの機械検査の まとめ（主の3パネル×3つの先・追加の点検の分も） ──
    n_chk = sum(p['checked'] for d in prefix_all.values() for p in d.values())
    n_same = sum(p['identical'] for d in prefix_all.values() for p in d.values())
    fails = {f'{h}:{k}': p['failures'] for h, d in prefix_all.items() for k, p in d.items() if p['failures']}
    # 逆向きの検査: t より後のデータを壊しても予想が変わらないか（乱数で上書き）
    rnd = random.Random(20261012)
    n_scr = n_scr_same = 0
    for P in MAIN:
        for t in P.grid[::2]:
            if P.end(t + HMAIN) > P.last_data():
                continue
            F, info = forecasts(P, t, HMAIN)
            e = P.end(t)
            ret = {a: {d: (v if d <= e else rnd.uniform(-0.5, 0.5)) for d, v in s.items()} for a, s in P.ret.items()}
            val = {a: {d: (v if d <= e else rnd.uniform(-5, 1)) for d, v in s.items()} for a, s in P.val.items()}
            Q = Panel(P.name, P.freq, ret, val, P.grid, P.unit, P.notes)
            F2, info2 = forecasts(Q, t, HMAIN)
            n_scr += 1
            n_scr_same += json.dumps([F, info], sort_keys=True) == json.dumps([F2, info2], sort_keys=True)
    prefix = {'method': '各（パネル×時点×先）で、データを t年12月末で切ったパネルから予想を計算し直し、json の文字列で1ビットも違わないか（γ・錨も含む）',
              'checked': n_chk, 'identical': n_same, 'failures': fails,
              'scramble_after_t': {'method': 't より後のリターンと割安さを乱数で上書きしても予想が同じか（主の3パネル・先15年・1つおきの時点）', 'checked': n_scr, 'identical': n_scr_same},
              'not_checked': '割安さの「分かった日」の割り当て（French の BE/ME＝y年6月・国の比率＝r−1年12月・JST＝y年末）が正しいかは機械では確かめられない（data_notes に根拠）'}
    print('先読みの検査', n_same, '/', n_chk, '・乱数で上書き', n_scr_same, '/', n_scr, flush=True)

    # ── 事前登録の予想の答え合わせ ──
    mse = {f: ev15[f]['MSE_pt2'] for f in FCS}
    order = sorted(FCS, key=lambda f: mse[f])
    pc = {'prediction': '（事前登録）A2 型は誤差が最も大きく、順位相関は0以下。B 型と V 型が最も当たる',
          'mse_order_best_to_worst': order,
          'A2_largest_MSE': order[-1] == 'A2',
          'A2_rank_corr_le_0': (ev15['A2']['rank_corr_mean'] or 0) <= 0,
          'B_and_V_best_two_by_MSE': set(order[:2]) == {'B', 'V'},
          'rank_corr_order_best_to_worst': sorted(FCS, key=lambda f: -(ev15[f]['rank_corr_mean'] or -9)),
          'by_panel': {p: {'mse_order': sorted(FCS, key=lambda f: e[f]['MSE_pt2']), 'rank_corr': {f: e[f]['rank_corr_mean'] for f in FCS}}
                       for p, e in ev_panel.items() if e},
          'by_horizon': {f'H{H}': {'mse_order': sorted(FCS, key=lambda f: res[f'H{H}']['common4_pooled'][f]['MSE_pt2']),
                                   'rank_corr': {f: res[f'H{H}']['common4_pooled'][f]['rank_corr_mean'] for f in FCS}} for H in (HMAIN,) + HSENS}}
    print('予想の答え合わせ', json.dumps({k: pc[k] for k in ('mse_order_best_to_worst', 'A2_largest_MSE', 'A2_rank_corr_le_0', 'B_and_V_best_two_by_MSE', 'rank_corr_order_best_to_worst')}, ensure_ascii=False), flush=True)

    tchk = timing_checks(P1)
    print('割安さの日付の確認', json.dumps({k: {kk: vv for kk, vv in v.items() if kk != 'reading'} for k, v in tchk.items()}, ensure_ascii=False), flush=True)
    sh = shiller_diag()
    print('Shiller 診断 H15', json.dumps(sh.get('H15'), ensure_ascii=False)[:600], flush=True)

    # 全表: 主（先15年）の行（資産×時点の予想と実現・%）
    rows15 = [{'panel': r['panel'], 't': r['t'], 'asset': r['asset'], 'Y%': r4(r['Y'] * 100, 2), **{f: (r4(r[f] * 100, 2) if r[f] is not None else None) for f in FCS}}
              for r in rows_by_H[HMAIN]]
    out = {
        'module': 'R8', 'name': 'world_calibration', 'prereg': PREREG, 'generated': t0.isoformat(timespec='seconds'),
        'question': '4つの世界の『予想の作り方』を過去に当てはめたら、どれがどれだけ当たったか → 世界の重み π',
        'pi_main': {w: round(v, 4) for w, v in pi_main.items()},
        'pi_main_rule': '主: 先15年・3パネルを行でまとめ・4つの予想がそろう行で π_w ∝ 1/MSE（A1→A1・A2→A2・B→B・V→C と D で半分ずつ）' + (f'・{why_main}' if why_main else ''),
        'pi_main_bootstrap': boot,
        'pi_sens': {k: {w: round(x, 4) for w, x in v.items()} for k, v in pi_sens.items()},
        'pi_sens_rule': {'H10': '先10年・1/MSE', 'H20': '先20年・1/MSE', 'rank_pos': '先15年・順位相関の平均の正の部分に比例' + (f'（{why_rank}）' if why_rank else ''),
                         'stack': f'先15年・非負で合計1の最小二乗の重ね（学習は{TRAIN_END}年までの時点）'},
        'pi_extra_checks（事前登録の外・bestpf_eval の頑健さには入れない）': {k: {w: round(x, 4) for w, x in v.items()} for k, v in extra.items()},
        'extra_checks_detail': alt,
        'forecaster_to_world': {'A1': 'F_A1（26年の実績を延ばす）→ 世界 A1', 'A2': 'F_A2（15年の実績を延ばす）→ 世界 A2',
                                'B': 'F_B（全期間の平均を50%縮める）→ 世界 B', 'V': 'F_V（割安さからの予想）→ 世界 C と D に半分ずつ'},
        'prediction_check': pc,
        'stacking': st,
        'prefix_check': prefix,
        'valuation_timing_checks': tchk,
        'results': res,
        'shiller_us_diagnostic': sh,
        'design': {'panels': {P.name: {'unit': P.unit, 'grid': P.grid, 'n_assets': len(P.ret)} for P in MAIN},
                   'horizons': {'main': HMAIN, 'sens': list(HSENS)}, 'shrink_B': SHRINK, 'B_min_years': B_MIN_YEARS, 'cs_min_assets': CS_MIN,
                   'valuation_window': '割安さは「分かった日」が (t−12か月, t年12月末] の値だけ（古すぎる値は使わない）',
                   'gamma_fit': '過去の時点 s（s＋先 ≦ t・5年ごとの同じ格子）の横断を、各時点の平均を引いて束ねた最小二乗（切片なし）。錨＝F_B と同じパネルの平均',
                   'pooled': '3パネルの行をそのまままとめる（業種が49あり行の多くは P1）。パネルを同じ重みにした版は追加の点検',
                   't_overlap_adj': '順位相関の平均 ÷（標準偏差÷√(横断の数×5/先)）＝重なる窓の分だけ数を減らした t',
                   'slope_pooled': '束ねた最小二乗の傾き。3パネルを束ねると名目と実質の水準の差を含む見かけの値になるので読まない（slope_within_cs を読む・slope_diagnostics）',
                   'units': '年率は小数で計算し、MSE は pt²（1pt＝1%）・RMSE と偏りは pt'},
        'data_notes': {P.name: P.notes for P in MAIN},
        'rows_H15（%）': rows15,
    }
    e, pm = ev15, out['pi_main']
    pp = res[f'H{HMAIN}']['common4_by_panel']
    PN = {'P1_French49': '業種49（French・名目ドル）', P2.name: '国（French・名目ドル）', 'P3_JST_real': '国（JST・現地通貨の実質）'}
    f2 = lambda v: '—' if v is None else f'{v:+.2f}'

    # ── 傾きの読み方（2026-10-12 検証役の指摘で足した）──
    # 束ねた傾き slope_pooled は、名目のパネル（P1・P2）と実質のパネル（P3）の水準の差を含む＝どの予想も
    # 「どのパネルの行か」を知っているだけで正の傾きが出る（シンプソンのパラドックス）。予想の腕前には
    # 横断の中の傾き（各時点の平均を引く）と、パネルごとの水準の時期の当たりを使う。
    lt = {p: {f: pp[p][f]['level_timing_by_panel'].get(p) for f in FCS} for p in pp if pp[p]}
    lvl_rmse = {f: r4(math.sqrt(max(e[f]['MSE_level_pt2'], 0)), 2) for f in FCS}
    neg_level = {p: [f for f in FCS if (lt[p][f] or {}).get('corr') is not None and lt[p][f]['corr'] < 0] for p in lt}
    out['slope_diagnostics'] = {
        'reading': ('束ねた傾き（slope_pooled）は名目のパネル（業種・French の国）と実質のパネル（JST）を混ぜたことで出る見かけの値。'
                    'パネルの中では4つとも負か0に近い。予想の腕前は「横断の中の傾き」（同じ時点の中で並べる力）と「水準の時期の当たり」で読む。'),
        'slope_pooled_raw（読まない）': {f: e[f]['slope_pooled'] for f in FCS},
        'slope_within_cs_pooled（同じ時点の中で並べる力）': {f: e[f]['slope_within_cs'] for f in FCS},
        'slope_within_panel_pooled（パネルの平均だけ引く）': {f: e[f]['slope_within_panel'] for f in FCS},
        'slope_raw_by_panel': {p: {f: pp[p][f]['slope_pooled'] for f in FCS} for p in pp if pp[p]},
        'slope_within_cs_by_panel': {p: {f: pp[p][f]['slope_within_cs'] for f in FCS} for p in pp if pp[p]},
        'level_timing_by_panel（各時点の横断の平均どうし・時間の順）': lt,
        'level_part_of_error_RMSE_pt（時点ごとの水準の外れだけ）': lvl_rmse,
        'relative_part_of_error_RMSE_pt（時点の平均を引いた後）': {f: r4(math.sqrt(e[f]['MSE_relative_pt2']), 2) for f in FCS},
        'how': ('slope_within_cs＝各（パネル×時点）の平均を引いて束ねた最小二乗の傾き。slope_within_panel＝パネルの平均だけを引いた傾き。'
                'level_timing＝パネルごとに、各時点の予想の横断の平均と実現の横断の平均の相関・傾き（時点が5未満のパネルは出さない）。'
                '水準の外れ＝MSE − 時点の平均を引いた後の MSE（MSE_level_pt2）の平方根')}
    # 横断の中の傾きの頑健さ（パネルごと・予想の差の端の2.5%ずつを除いた傾きと、予想の差のばらつき）
    rob = {}
    for p in pp:
        g = {}
        for r in common_by_H[HMAIN]:
            if r['panel'] == p:
                g.setdefault(r['t'], []).append(r)
        rob[p] = {}
        for f in FCS:
            dx, dy = [], []
            for gg in g.values():
                mf = sum(r[f] for r in gg) / len(gg); my = sum(r['Y'] for r in gg) / len(gg)
                dx += [r[f] - mf for r in gg]; dy += [r['Y'] - my for r in gg]
            n = len(dx)
            o = sorted(range(n), key=lambda i: dx[i]); k = int(n * 0.025)
            keep = o[k:n - k]
            sxx = sum(dx[i] ** 2 for i in keep)
            rob[p][f] = {'slope_within_cs': pp[p][f]['slope_within_cs'],
                         'slope_within_cs_trim5%': r4(sum(dx[i] * dy[i] for i in keep) / sxx, 3) if sxx > 0 else None,
                         'sd_forecast_within_cs_pt': r4(S.pstdev(dx) * 100, 2) if n > 1 else None, 'n': n}
    out['slope_diagnostics']['within_cs_robustness_by_panel（予想の差の端を両側2.5%ずつ除く）'] = rob
    p1 = pp['P1_French49']; p3 = pp['P3_JST_real']
    nw1 = {f: (p1[f]['nw_t_by_panel'] or {}).get('P1_French49') for f in FCS}
    out['rank_corr_by_panel_reading'] = {
        'P1_French49': {f: {'rank_corr_mean': p1[f]['rank_corr_mean'], 'nw_t': nw1[f]} for f in FCS},
        'P3_JST_real': {f: {'rank_corr_mean': p3[f]['rank_corr_mean'], 'nw_t': (p3[f]['nw_t_by_panel'] or {}).get('P3_JST_real')} for f in FCS},
        'reading': ('まとめた順位相関（V・B が上）は、ほとんど国の長期データ（JST）から来ている。業種49では4つのどれも順位をはっきり当てていない'
                    f"（V {p1['V']['rank_corr_mean']}・B {p1['B']['rank_corr_mean']}・A1 {p1['A1']['rank_corr_mean']}・A2 {p1['A2']['rank_corr_mean']}・"
                    f"NW の t の絶対値はどれも {max(abs(x) for x in nw1.values() if x is not None):.2f} 以下）。")}
    shd = sh.get(f'H{HMAIN}') or {}
    out['shiller_reading'] = (
        f"米国（Shiller・先{HMAIN}年・{(shd.get('years') or ['?', '?'])[0]}〜{(shd.get('years') or ['?', '?'])[1]}年の12月）では、CAPE で水準を時期ごとに当てる予想が最も当たった"
        f"（RMSE {shd.get('V', {}).get('RMSE_pt')}pt・時系列の相関 {f2(shd.get('V', {}).get('corr_time_series'))}）。"
        f"これはパネルの V（水準は過去の平均に固定し、割安さは同じ時点の中の並べ替えにだけ使う）とは別物で、パネルの V の水準は時期を当てていない。"
        f"長い平均を延ばす B の時系列の相関は {f2(shd.get('B', {}).get('corr_time_series'))}・A1 {f2(shd.get('A1', {}).get('corr_time_series'))}・A2 {f2(shd.get('A2', {}).get('corr_time_series'))}＝"
        "過去の実績や平均が高い時期ほど、その後の15年は低かった。米国1つの市場の時間の並びで、窓は大きく重なる（π には使わない）。")

    wcs = {f: e[f]['slope_within_cs'] for f in FCS}
    all_neg = all((lt[p][f] or {}).get('corr') is None or lt[p][f]['corr'] < 0 for p in lt for f in FCS)
    lt_txt = '・'.join(f"{PN.get(p, p)}: " + ' '.join(f"{f} {f2((lt[p][f] or {}).get('corr'))}" for f in FCS)
                       for p in lt if any((lt[p][f] or {}).get('corr') is not None for f in FCS))
    out['summary_ja'] = (
        f"4種類の予想の作り方を、業種（49・1955〜2010年の12時点）・国（French 14〜15か国・2000〜2010年）・国（JST 10〜16か国・1900〜2005年・物価の上がり分を除いた実質）の"
        f"『その時点から15年の年率』に当てはめた（4つがそろう {e['n_rows']} 行・{e['n_cross_sections']} の横断〔パネル×時点〕・{len(e['times'])} の年）。"
        f"外れの大きさ（2乗の平均の平方根）は 割安さからの予想 V {e['V']['RMSE_pt']}pt ・100年の平均型 B {e['B']['RMSE_pt']}pt ・26年の実績 A1 {e['A1']['RMSE_pt']}pt ・15年の実績 A2 {e['A2']['RMSE_pt']}pt。"
        f"順位の当たり（各時点の順位相関の平均）は V {e['V']['rank_corr_mean']}・B {e['B']['rank_corr_mean']}・A1 {e['A1']['rank_corr_mean']}・A2 {e['A2']['rank_corr_mean']}。"
        f"ただしこの差はほとんど国（JST: V {p3['V']['rank_corr_mean']}・B {p3['B']['rank_corr_mean']}）から来ていて、"
        f"業種49では4つのどれも順位をはっきり当てていない（V {p1['V']['rank_corr_mean']}・B {p1['B']['rank_corr_mean']}・A1 {p1['A1']['rank_corr_mean']}・A2 {p1['A2']['rank_corr_mean']}）。"
        f"同じ時点の中で並べる力（予想の差のうち実現に出た割合＝横断の中の傾き）は V {wcs['V']}・B {wcs['B']}・A1 {wcs['A1']}・A2 {wcs['A2']}。"
        f"一方、パネルの中ではどの予想も、先15年の年率が高い時期・低い時期を当てられなかった（各時点の平均どうしの相関 {lt_txt}）。"
        f"時期ごとの水準の外れだけで RMSE V {lvl_rmse['V']}・B {lvl_rmse['B']}・A1 {lvl_rmse['A1']}・A2 {lvl_rmse['A2']}pt ある。"
        f"（3パネルを束ねて実現を予想に回帰した傾き〔A2 {e['A2']['slope_pooled']}・V {e['V']['slope_pooled']}〕は、名目と実質のパネルを混ぜたことで出る見かけの値なので使わない。）"
        f"主の重み π（1/MSE）は A1 {pm['A1']}・A2 {pm['A2']}・B {pm['B']}・C {pm['C']}・D {pm['D']}。"
        f"重ね（学習 〜{TRAIN_END}年）は C と D に合わせて {round(out['pi_sens']['stack']['C'] + out['pi_sens']['stack']['D'], 3)}・A2 {out['pi_sens']['stack']['A2']}。")
    out['feeds_final_notes'] = [
        'π は 1/MSE（事前登録の主）のまま。下の注意は読み方の申し送りで、π の値は変えない。',
        (f"同じ時点の中で並べる力は V {wcs['V']}・B {wcs['B']}・A1 {wcs['A1']}・A2 {wcs['A2']}（横断の中の傾き）。割安さ型の差は7割ほど、長い平均の差は3割ほどしか実現に出ず、直近15年の実績の差は実現に出なかった。"
         f"パネルで分けると、国（JST）では V {p3['V']['slope_within_cs']}・B {p3['B']['slope_within_cs']} と両方よく出たが、"
         f"業種49では V {p1['V']['slope_within_cs']}・B {p1['B']['slope_within_cs']}。業種の V は予想の差がほとんど無く（同じ時点の中のばらつき {rob['P1_French49']['V']['sd_forecast_within_cs_pt']}pt）、"
         f"端の5%を除くと傾きは {rob['P1_French49']['V']['slope_within_cs_trim5%']}＝業種の並べ替えの力は確かめられない。"),
        f"業種49では4つのどれも順位をはっきり当てていない（V {p1['V']['rank_corr_mean']}・B {p1['B']['rank_corr_mean']}）。割安さ型 V の優位は国の長期データ（JST）から来ている。業種の ETF（SMH・XLK・ITA・XLE など）の間の差を、C・D（割安さ型）の世界で強く読まない。",
        f"パネルの中では、どの予想も先15年の水準の高低を時期ごとに当てられなかった（各時点の平均どうしの相関{'はどれも負' if all_neg else ''}: {lt_txt}）。時期ごとの水準の外れだけで RMSE {min(lvl_rmse.values())}〜{max(lvl_rmse.values())}pt。最終の主の物差し P15π（15%以上の確率）は絶対の水準で決まるので、どの世界の水準も「時期の当たりは無い」と読み、水準には時期ごとの外れ（RMSE 4〜6pt）があると見込む。長期平均を錨にし、直近の高い実績で水準を上げない（長期平均を延ばす B でさえ、時期の相関は負）。",
        f"A2 世界の数字（直近15年の実績を延ばす）は、外れが最大で、同じ時点の中で並べる力も {wcs['A2']}。半導体などの直近の実績を延ばす世界には割り引きが要る。",
        '米国の CAPE 型（Shiller）が水準の時期を当てたのは、パネルの V とは別の作り方（1つの市場の時間の並び）。C・D の世界の水準の当たりの証拠としては使わない。']
    out['fixes_applied'] = [
        {'issue': '検証役（medium）: 束ねた傾き V 0.60・B 0.53・A2 0.07 は名目と実質のパネルを混ぜたことによる見かけの値（シンプソンのパラドックス）',
         'what_changed': ('evaluate() に slope_within_panel（パネルの平均だけ引く）と level_timing_by_panel（各時点の横断の平均どうしの相関・傾き）を足し、'
                          'slope_diagnostics を新設。summary_ja と申し送りの傾きを横断の中の傾き（slope_within_cs）に替え、'
                          '「パネルの中ではどの予想も水準の時期を当てていない」と時期ごとの水準の外れ（RMSE）を書いた。'
                          '表の中の slope_pooled は残したが、隣に slope_within_cs を置き、design に「読まない」と書いた。'
                          'Shiller の「同じ向き」をやめ、shiller_reading で CAPE 型はパネルの V と別物・B の時系列の相関が負と書いた。'
                          'あわせて横断の中の傾きの頑健さ（パネルごと・予想の差の端を両側2.5%ずつ除く）を足した: 業種の V の 0.63 は端を除くと 0.05 に消える。'),
         'pi_changed': False},
        {'issue': '検証役（low）: 順位の当たり V・B の優位はほとんど JST から来ており、業種49ではどれもほぼ0',
         'what_changed': 'rank_corr_by_panel_reading を新設し、summary_ja と feeds_final_notes に業種49の数字（V・B もほぼ0・NW の t は小さい）を書いた。',
         'pi_changed': False},
        {'issue': '検証役（low）: 「37時点」は「37の横断（23の年）」の書き違い',
         'what_changed': 'summary_ja を「37 の横断〔パネル×時点〕・23 の年」に直した（数はコードから出す）。',
         'pi_changed': False}]
    out['deviations_and_choices'] = [
        '国の割安さは D/P（配当利回り・log）を主にした（事前登録は「E/P か D/P」・計算の前に決めた）。E/P は追加の点検（ほぼ同じ）',
        'French の国のファイルに米国が無いので P2 は米国抜き21か国（15年の比較で4つそろうのは 2000・2005・2010 年の14〜15か国）',
        'JST の単位は事前登録に無かったので、国ごとのインフレの差を消すため「現地通貨の実質」を主にした（ドル建ては追加の点検・ほぼ同じ）',
        'JST のドイツ1922〜23年は欠測にした（物価が年に10億倍の年で、年次の物価・為替では割れない）',
        'F_B の「全期間」は10年以上・割安さは t の12か月以内の値・横断は5資産以上・V の錨は F_B と同じパネルの平均・γ は5年ごとの過去の時点だけ（事前登録に細部が無かった点を計算の前に決めた）',
        'π_sens には事前登録の感度（rank_pos・stack・先10年・先20年）だけを入れた。パネルを同じ重み・パネルごと・単位・日付の遅れは pi_extra_checks に分けた（bestpf_eval の頑健さの数に混ぜない）',
        'まとめの MSE は3パネルの行をそのまま束ねた（行の61%が業種）。パネルを同じ重みにしても π はほぼ同じ（B 0.29・A2 0.19）']
    json.dump(out, open(OUT, 'w'), ensure_ascii=False, indent=1)
    print('π 主', out['pi_main'])
    print('π 感度', out['pi_sens'])
    print('→', OUT, f'（{(datetime.datetime.now() - t0).seconds}秒）')


if __name__ == '__main__':
    main()
