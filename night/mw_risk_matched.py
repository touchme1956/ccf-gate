#!/usr/bin/env python3
"""night/mw_risk_matched.py — 『市場に勝てる歴史検証』の角度 risk_matched（読むだけ・門の判定には不使用）

問い: 市場よりシャープレシオの高い『買いだけ』のポートフォリオ（低ベータ・低ボラ・質・収益性）を、
      市場と同じぶれになるまで借入でふくらませる（ぶれが大きいものは現金で薄める）と、
      S&P500 型の純粋な市場（French Mkt）に 2007年以降も勝てたか。
事前登録: out/mw_risk_matched_prereg.json（線は out/mw_prereg.json・判定は mw_common.grade）
出力: out/mw_risk_matched.json

族
- P（主・Holm 14本）: 低リスク群の上位3源＋質群の上位3源（訓練シャープで選ぶ）＋その6本の等分混合 × {static, dynamic}
- U（参照）: P の7候補の倍率1
- X1（探索）: 残り49源の代表 × {static, dynamic}
- XU（探索・参照）: X1 の倍率1
- S（報告のみ）: P を 借入 RF+2% と 費用3倍 で
事前登録2（out/mw_risk_matched_prereg2.json・探索）
- X2: 低リスク21源をベータ1へ合わせる（beta_static / beta_dynamic）
- X3/X3U: 訓練だけで組む混合（上位10の等分・ぶれの逆数・買いだけの最大シャープ・最小分散）
- X4/X4U: 大型株 ME5（25_Portfolios_ME_*_5x5 の BIG 行）
- S2（報告のみ）: P を倍率の上限 1.5倍で
"""
import json, math, os, re, statistics as S, subprocess, sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import mw_common as M  # noqa: E402

_FT, _orig_ft = {}, M.french_tables


def _ft_cached(name):  # French の zip を一度だけ読む（mw_common は変えない）
    if name not in _FT:
        _FT[name] = _orig_ft(name)
    return _FT[name]


M.french_tables = _ft_cached

BASE = M.BASE
ANGLE = 'risk_matched'
PREREG = 'mw_risk_matched_prereg.json'
PREREG2 = 'mw_risk_matched_prereg2.json'
LMAX, SPREAD, COST, WIN = 2.0, 0.01, 0.001, 36
RANK_A, RANK_Z = 196307, 200612
REGIONS = ['developed', 'world_ex_us', 'jpn', 'emerging']
NONUS = ['world_ex_us', 'jpn', 'emerging']

LOW_RISK = ["beta_60m", "beta_dimson_21d", "betabab_1260d", "betadown_252d", "earnings_variability", "ivol_capm_21d",
            "ivol_capm_252d", "ivol_ff3_21d", "ivol_hxz4_21d", "ocfq_saleq_std", "rmax1_21d", "rmax5_21d", "rvol_21d",
            "seas_6_10na", "turnover_126d", "zero_trades_126d", "zero_trades_21d", "zero_trades_252d"]
QUALITY = ["at_turnover", "cop_at", "cop_atl1", "dgp_dsale", "gp_at", "gp_atl1", "mispricing_perf", "ni_inc8q", "niq_at",
           "op_at", "op_atl1", "opex_at", "qmj", "qmj_growth", "qmj_prof", "qmj_safety", "sale_bev", "dolvol_var_126d",
           "ebit_bev", "ebit_sale", "f_score", "ni_be", "niq_be", "o_score", "ocf_at", "ope_be", "ope_bel1", "turnover_var_126d"]

JKP_TO = {}
for k in ["beta_60m", "betabab_1260d", "betadown_252d", "ivol_capm_252d"]:
    JKP_TO[k] = 0.4
for k in ["rvol_21d", "ivol_capm_21d", "ivol_ff3_21d", "ivol_hxz4_21d", "rmax1_21d", "rmax5_21d", "beta_dimson_21d", "zero_trades_21d"]:
    JKP_TO[k] = 1.5
for k in ["turnover_126d", "zero_trades_126d", "zero_trades_252d", "dolvol_var_126d", "turnover_var_126d"]:
    JKP_TO[k] = 0.6
for k in ["earnings_variability", "ocfq_saleq_std"]:
    JKP_TO[k] = 0.3
JKP_TO["seas_6_10na"] = 1.0
for k in ["gp_at", "gp_atl1", "cop_at", "cop_atl1", "op_at", "op_atl1", "ope_be", "ope_bel1", "ebit_bev", "ebit_sale",
          "ocf_at", "sale_bev", "at_turnover", "f_score", "o_score", "opex_at", "dgp_dsale", "ni_be"]:
    JKP_TO[k] = 0.4
for k in ["niq_at", "niq_be", "ni_inc8q"]:
    JKP_TO[k] = 0.8
for k in ["qmj", "qmj_prof", "qmj_safety", "qmj_growth", "mispricing_perf"]:
    JKP_TO[k] = 0.6
# C5 の地域で使う（French の源の対応先）
for k in ["at_gr1"]:
    JKP_TO[k] = 0.6
JKP_TO["be_me"] = 0.3
JKP_TO["ret_12_1"] = 1.5

FRENCH_SRC = {
    # 源: (ファイル, 列の選び方, 回転率, 群, 公表年, C5 の対応特徴, 意味)
    'BETA': ('Portfolios_Formed_on_BETA', 'all', 0.4, '低リスク', 2014, 'beta_60m', 'ベータ（60か月）'),
    'VAR': ('Portfolios_Formed_on_VAR', 'all', 1.5, '低リスク', 2006, 'rvol_21d', '分散（60日）'),
    'RESVAR': ('Portfolios_Formed_on_RESVAR', 'all', 1.5, '低リスク', 2006, 'ivol_ff3_21d', '残差分散（60日・FF3）'),
    'OP': ('Portfolios_Formed_on_OP', 'all', 0.4, '質', 2015, 'ope_be', '営業利益÷自己資本'),
    'INV': ('Portfolios_Formed_on_INV', 'all', 0.6, '質', 2008, 'at_gr1', '総資産の伸び'),
    'BIG_OP': ('6_Portfolios_ME_OP_2x3', 'big', 0.4, '質', 2015, 'ope_be', '大型株の営業利益÷自己資本'),
    'BIG_INV': ('6_Portfolios_ME_INV_2x3', 'big', 0.6, '質', 2008, 'at_gr1', '大型株の総資産の伸び'),
    'BIG_BM': ('6_Portfolios_2x3', 'big', 0.3, 'その他', 1992, 'be_me', '大型株の簿価÷時価'),
    'BIG_PRIOR': ('6_Portfolios_ME_Prior_12_2', 'big', 1.5, 'その他', 1993, 'ret_12_1', '大型株の勢い（12-2）'),
}


DEVIATIONS = [
    'C5 の地域は指示どおり developed・world_ex_us・jpn・emerging の4地域（3以上で合格）。developed は米国を含み独立ではないので、米国外3地域だけの数（nonus_positive/nonus_regions）も併記した。4地域で3以上なら米国外3地域でも2以上になるので、判定は厳しい側',
    'AQR BAB（参考のみ・判定しない）の読み取りで日付の形（mm/dd/yyyy）を読めず初回は取れなかった。事前登録2の追加と同時に読み取りを直した（判定に関係しない）',
    'JKP の ivol_hxz4_21d の3本と ni_inc8q の三分位2 は月が途切れている。順位づけの窓（196307〜200612）が欠ける候補は順位から外し、評価は連続した最長の区間だけで行う（欠測を0で埋めない）',
    'seas_6_10na（6〜10年前の同じでない月のリターン）は JKP のクラスタ表で Low Risk に入っているので、登録どおり低リスク群として扱った（中身は季節性の信号）',
    'シャープレシオ（C8）は mw_common.sharpe ではなく、戦略と市場を同じ月だけで比べる sharpe_same で計算した（mw_common.sharpe を市場に使うと市場だけ1926年からの月で測ってしまうため）',
    '第1族の結果を見た後で事前登録2（X2・X3・X4・S2）を登録した。事前登録2の族はすべて探索',
]


def sha_of(path):
    try:
        out = subprocess.run(['git', 'log', '--format=%H', '--', path], cwd=BASE, capture_output=True, text=True).stdout.strip().split('\n')
        return out[-1] or None  # 最初に入ったコミット（事前登録）
    except Exception:  # noqa
        return None


def next_m(ym):
    y, m = divmod(ym, 100)
    return ym + 1 if m < 12 else (y + 1) * 100 + 1


def contiguous(d):
    ks = sorted(d)
    return all(next_m(a) == b for a, b in zip(ks, ks[1:]))


# ───────────────────────── データ ─────────────────────────
def col_center(col):
    """French の列名 → 分位の中心（%）"""
    c = col.strip()
    m = re.match(r'^(Lo|Hi) (\d+)$', c)
    if m:
        w = int(m.group(2))
        return w / 2 if m.group(1) == 'Lo' else 100 - w / 2
    if c == 'Med 40':
        return 50.0
    m = re.match(r'^Qnt (\d)$', c)
    if m:
        return (int(m.group(1)) - 0.5) * 20
    m = re.match(r'^Dec (\d+)$', c) or re.match(r'^(\d+)-Dec$', c)
    if m:
        return (int(m.group(1)) - 0.5) * 10
    if c.startswith('BIG Lo'):
        return 15.0
    if c.startswith('ME2'):
        return 50.0
    if c.startswith('BIG Hi'):
        return 85.0
    raise ValueError(col)


def tercile_of(center):
    return '1.0' if center < 100 / 3 else ('2.0' if center < 200 / 3 else '3.0')


def load_candidates(RF):
    cands = {}
    for src, (fname, how, to, grp, pub, cmap, mean) in FRENCH_SRC.items():
        s = M.french_series(fname, 'Value Weight')
        cols = [c for c in s if (how == 'all' or c.startswith('BIG') or c.startswith('ME2'))]
        for c in cols:
            name = f'FR_{src}[{c}]'
            cands[name] = {'name': name, 'kind': 'french', 'source': 'FR_' + src, 'col': c, 'group': grp, 'pub': pub,
                           'to': to, 'r': dict(s[c]), 'map': (cmap, tercile_of(col_center(c))),
                           'desc': f'French {mean}の {c}（時価加重）'}
    details = jkp_details()
    for grp, keys in (('低リスク', LOW_RISK), ('質', QUALITY)):
        for k in keys:
            try:
                p = M.jkp_portfolios('usa', k, 'vw')
            except Exception as e:  # noqa
                print('  JKP 取得失敗', k, e)
                continue
            for pf in ('1.0', '2.0', '3.0'):
                if pf not in p:
                    continue
                tot = {ym: v + RF[ym] for ym, v in p[pf].items() if ym in RF}
                name = f'JKP_{k}[{pf[0]}]'
                cands[name] = {'name': name, 'kind': 'jkp', 'source': 'JKP_' + k, 'col': pf, 'group': grp,
                               'pub': details.get(k, {}).get('pub'), 'to': JKP_TO[k], 'r': tot, 'map': (k, pf),
                               'desc': f"JKP 米国 {details.get(k, {}).get('name', k)} の三分位{pf[0]}（1=低い側・時価加重）"}
    return cands


def jkp_details():
    out = {}
    try:
        import openpyxl
        wb = openpyxl.load_workbook(os.path.join(M.CACHE, 'jkp_factor_details.xlsx'), read_only=True)
        rows = list(wb['details'].iter_rows(values_only=True))
        h = rows[0]
        ia, ic, iname = h.index('abr_jkp'), h.index('cite'), h.index('name_new')
        for r in rows[1:]:
            if not r[ia]:
                continue
            m = re.search(r'(19|20)\d{2}', str(r[ic] or ''))
            out[r[ia]] = {'pub': int(m.group(0)) if m else None, 'name': r[iname], 'cite': r[ic]}
    except Exception as e:  # noqa
        print('  Factor Details 読めず', e)
    return out


# ───────────────────────── 倍率 ─────────────────────────
def lever(r_ex, m_ex, rf, rule, L_static=None, to=0.4, cost=COST, spread=SPREAD, lmax=LMAX, win=WIN):
    """r_ex, m_ex: 月次の超過（小数）。rf: 月次 RF（重みの漂いと総リターンに使う）。
    戻り: (gross_ex, net_ex, info) 超過リターン（借入の上乗せ込み）。総リターン = 超過 + rf"""
    ks = sorted(k for k in set(r_ex) & set(m_ex) & set(rf))
    if not contiguous({k: 1 for k in ks}):  # 抜けがあれば連続した最長の区間だけ（欠測を0で埋めない）
        runs, cur = [], [ks[0]]
        for a, b in zip(ks, ks[1:]):
            if next_m(a) == b:
                cur.append(b)
            else:
                runs.append(cur); cur = [b]
        runs.append(cur)
        ks = max(runs, key=len)
    gross, net, Ls, trades = {}, {}, [], []
    hp, hm = [], []
    prevL = prev_r = prev_R = None
    for k in ks:
        if rule in ('static', 'beta_static'):
            L = L_static
        elif rule == 'unlevered':
            L = 1.0
        elif rule == 'beta_dynamic':
            L = None
            if len(hp) >= win:
                b = beta_of(hp[-win:], hm[-win:])
                L = lmax if b <= 0 else min(lmax, 1 / b)
        else:
            L = None
            if len(hp) >= win:
                sp, sm = S.stdev(hp[-win:]), S.stdev(hm[-win:])
                L = min(lmax, sm / sp) if sp > 0 else None
        hp.append(r_ex[k]); hm.append(m_ex[k])  # t の倍率を決めた後に t を履歴へ
        if L is None:
            continue
        g = L * r_ex[k] - max(L - 1, 0) * spread / 12
        c_under = L * to / 12 * cost
        if prevL is not None:
            w_drift = prevL * (1 + prev_r) / (1 + prev_R) if (1 + prev_R) > 0 else L
            tr = abs(L - w_drift)
        else:
            tr = 0.0
        trades.append(tr)
        gross[k] = g
        net[k] = g - c_under - tr * cost
        Ls.append(L)
        prevL, prev_r, prev_R = L, r_ex[k] + rf[k], g + rf[k]
    info = {'L_mean': round(S.mean(Ls), 3) if Ls else None, 'L_min': round(min(Ls), 3) if Ls else None,
            'L_max': round(max(Ls), 3) if Ls else None,
            'lev_turnover_ann': round(S.mean(trades) * 12, 3) if trades else None,
            'months': len(gross)}
    return gross, net, info


def static_L(r_ex, m_ex, z=M.TRAIN_END, lmax=LMAX):
    ks = sorted(k for k in set(r_ex) & set(m_ex) if k <= z)
    if len(ks) < 60:
        return None
    return min(lmax, S.stdev([m_ex[k] for k in ks]) / S.stdev([r_ex[k] for k in ks]))


def beta_of(x, m):
    mx, mm = S.mean(x), S.mean(m)
    vm = sum((v - mm) ** 2 for v in m)
    return sum((a - mx) * (b - mm) for a, b in zip(x, m)) / vm if vm else 0.0


def static_beta_L(r_ex, m_ex, z=M.TRAIN_END, lmax=LMAX):
    ks = sorted(k for k in set(r_ex) & set(m_ex) if k <= z)
    if len(ks) < 60:
        return None
    b = beta_of([r_ex[k] for k in ks], [m_ex[k] for k in ks])
    return lmax if b <= 0 else min(lmax, 1 / b)


def to_total(ex, rf):
    return {k: v + rf[k] for k, v in ex.items() if k in rf}


def sharpe_same(s_tot, b_tot, rf, a=None, z=None):
    ks = sorted(k for k in set(s_tot) & set(b_tot) & set(rf) if (a is None or k >= a) and (z is None or k <= z))
    if len(ks) < 24:
        return None
    def sh(d):
        x = [d[k] - rf[k] for k in ks]
        sd = S.stdev(x)
        return round(S.mean(x) / sd * math.sqrt(12), 3) if sd else None
    return (sh(s_tot), sh(b_tot))


def train_sharpe(r_tot, rf, a=RANK_A, z=RANK_Z):
    ks = [k for k in sorted(r_tot) if a <= k <= z and k in rf]
    need = (z // 100 - a // 100) * 12 + (z % 100) - (a % 100) + 1
    if len(ks) < need:
        return None
    x = [r_tot[k] - rf[k] for k in ks]
    return S.mean(x) / S.stdev(x) * math.sqrt(12)


# ───────────────────────── 地域（C5） ─────────────────────────
_REG = {}


def reg_series(region, key, pf):
    if (region, key) not in _REG:
        try:
            _REG[(region, key)] = M.jkp_portfolios(region, key, 'vw')
        except Exception as e:  # noqa
            _REG[(region, key)] = None
            print('  地域 取得失敗', region, key, str(e)[:80])
    p = _REG[(region, key)]
    return None if p is None else p.get(pf)


_RMKT = {}


def reg_mkt(region):
    if region not in _RMKT:
        _RMKT[region] = M.jkp_mkt(region, 'vw')
    return _RMKT[region]


def replicate(maps, rule, RF, to, weights=None, renorm=False, lmax=LMAX):
    """maps: [(key, pf), ...]（混合なら複数）。weights: 同じ長さの重み（無ければ等分）。
    renorm=False: 1つでも欠けたらその地域は N/A（第1族）。renorm=True: 欠けた重みの合計 ≤ 50% なら残りを正規化（第2族 X3）。
    各地域で同じ規則。戻り: {'regions','positive','per_region',...}"""
    per = {}
    w_all = weights or [1 / len(maps)] * len(maps)
    for reg in REGIONS:
        parts = [(reg_series(reg, k, pf), w) for (k, pf), w in zip(maps, w_all)]
        miss = sum(w for p, w in parts if p is None)
        if (not renorm and miss > 0) or (renorm and miss > 0.5 * sum(w_all)):
            per[reg] = None
            continue
        parts = [(p, w) for p, w in parts if p is not None]
        ws = sum(w for _, w in parts)
        ks = sorted(set.intersection(*[set(p) for p, _ in parts]))
        if len(ks) < 60:
            per[reg] = None
            continue
        r_ex = {k: sum(p[k] * w for p, w in parts) / ws for k in ks}
        m_ex = reg_mkt(reg)
        cost = 0.003 if reg == 'emerging' else COST
        L = static_L(r_ex, m_ex, lmax=lmax) if rule == 'static' else (static_beta_L(r_ex, m_ex, lmax=lmax) if rule == 'beta_static' else None)
        if rule in ('static', 'beta_static') and L is None:
            per[reg] = None
            continue
        # 地域の月の抜け（JKP の地域は連続のはず）を確かめる。抜けがあれば連続した最長の区間だけ
        kk = sorted(set(r_ex) & set(m_ex) & set(RF))
        if not contiguous({k: 1 for k in kk}):
            runs, cur = [], [kk[0]]
            for a, b in zip(kk, kk[1:]):
                if next_m(a) == b:
                    cur.append(b)
                else:
                    runs.append(cur); cur = [b]
            runs.append(cur)
            best = max(runs, key=len)
            r_ex = {k: r_ex[k] for k in best}
        g, n, info = lever(r_ex, m_ex, RF, rule, L_static=L, to=to, cost=cost, lmax=lmax)
        full = M.excess_stats(n, m_ex)
        hold = M.excess_stats(n, m_ex, a=M.HOLD_START)
        if full is None:
            per[reg] = None
            continue
        per[reg] = {'from': full['from'], 'to': full['to'], 'years': full['years'], 'L': round(L, 3) if L else info['L_mean'],
                    'net_ex_ann': full['ex_ann'], 't': full['t'], 'cagr_diff': full['cagr_diff'],
                    'positive': full['ex_ann'] > 0,
                    'hold_net_ex_ann': hold['ex_ann'] if hold else None, 'hold_t': hold['t'] if hold else None}
    got = [r for r in REGIONS if per.get(r)]
    pos = [r for r in got if per[r]['positive']]
    nonus = [r for r in NONUS if per.get(r)]
    return {'regions': len(got), 'positive': len(pos), 'nonus_regions': len(nonus),
            'nonus_positive': sum(1 for r in nonus if per[r]['positive']), 'per_region': per}


# ───────────────────────── 評価 ─────────────────────────
def evaluate(c, rule, MKT, MKTRF, RF, spread=SPREAD, cost=COST, lmax=LMAX):
    r_ex = {k: v - RF[k] for k, v in c['r'].items() if k in RF}
    L = static_L(r_ex, MKTRF, lmax=lmax) if rule == 'static' else (static_beta_L(r_ex, MKTRF, lmax=lmax) if rule == 'beta_static' else None)
    g_ex, n_ex, info = lever(r_ex, MKTRF, RF, rule, L_static=L, to=c['to'], cost=cost, spread=spread, lmax=lmax)
    g, n = to_total(g_ex, RF), to_total(n_ex, RF)
    first = min(g)
    b = M.window(MKT, first)
    e = {'full': M.excess_stats(g, MKT), 'train': M.excess_stats(g, MKT, z=M.TRAIN_END),
         'hold': M.excess_stats(g, MKT, a=M.HOLD_START), 'recent': M.excess_stats(g, MKT, a=M.RECENT_START),
         'net_cost_full': M.excess_stats(n, MKT), 'net_cost_hold': M.excess_stats(n, MKT, a=M.HOLD_START)}
    e['roll20_net'] = M.rolling(n, MKT, 20)
    e['dca20_net'] = M.dca(n, MKT, 20)
    e['sharpe_pair'] = {'train': sharpe_same(n, MKT, RF, z=M.TRAIN_END), 'hold': sharpe_same(n, MKT, RF, a=M.HOLD_START)}
    e['L_static'] = round(L, 3) if L is not None else None
    e['lever_info'] = info
    ks = sorted(set(n) & set(b))
    e['maxdd'] = {'strategy': round(M.maxdd({k: n[k] for k in ks}) * 100, 1), 'mkt': round(M.maxdd({k: b[k] for k in ks}) * 100, 1)}
    if c.get('pub'):
        e['post_publication'] = M.excess_stats(g, MKT, a=(c['pub'] + 1) * 100 + 1)
    return e, g, n


RULE_JA = {'static': '訓練のぶれ比で固定倍率', 'dynamic': '36か月のぶれ比で毎月倍率', 'unlevered': '倍率1（借入なし）',
           'beta_static': '訓練のベータの逆数で固定倍率', 'beta_dynamic': '36か月のベータの逆数で毎月倍率'}

X4_FILES = {'BETA': ('25_Portfolios_ME_BETA_5x5', 'beta_60m', 0.4), 'VAR': ('25_Portfolios_ME_VAR_5x5', 'rvol_21d', 1.5),
            'RESVAR': ('25_Portfolios_ME_RESVAR_5x5', 'ivol_ff3_21d', 1.5), 'OP': ('25_Portfolios_ME_OP_5x5', 'ope_be', 0.4),
            'INV': ('25_Portfolios_ME_INV_5x5', 'at_gr1', 0.6), 'PRIOR': ('25_Portfolios_ME_Prior_12_2', 'ret_12_1', 1.5),
            'AC': ('25_Portfolios_ME_AC_5x5', 'oaccruals_at', 0.8), 'NI': ('25_Portfolios_ME_NI_5x5', 'chcsho_12m', 0.6),
            'BM': ('25_Portfolios_5x5', 'be_me', 0.3)}
X4_PUB = {'BETA': 2014, 'VAR': 2006, 'RESVAR': 2006, 'OP': 2015, 'INV': 2008, 'PRIOR': 1993, 'AC': 1996, 'NI': 2008, 'BM': 1992}


def x4_tercile(src, col):
    c = col.strip()
    if src == 'NI':
        if 'NegNI' in c or 'ZeroNI' in c:
            return '1.0'
        if 'LoNI' in c or c.endswith('NI2'):
            return '2.0'
        return '3.0'
    if c.startswith('BIG Lo'):
        q = 1
    elif c.startswith('BIG Hi'):
        q = 5
    else:
        q = int(c[-1])
    return tercile_of((q - 0.5) * 20)


def load_x4(RF):
    out = {}
    for src, (fname, key, to) in X4_FILES.items():
        s = M.french_series(fname, 'Value Weight')
        for c in s:
            if not (c.startswith('BIG') or c.startswith('ME5')):
                continue
            name = f'FR5_{src}[{c}]'
            out[name] = {'name': name, 'kind': 'french5', 'source': 'FR5_' + src, 'col': c, 'group': '大型株', 'pub': X4_PUB[src],
                         'to': to, 'r': dict(s[c]), 'map': (key, x4_tercile(src, c)),
                         'desc': f'French 大型株（NYSE 80%点超）の {c}（{fname}・時価加重）'}
    return out


def qp_simplex(Sig, a, iters=200):
    """min x'Σx s.t. a'x = 1, x ≥ 0 を有効制約法で解く → x（numpy）。a=μ で最大シャープ（接点）、a=1 で最小分散"""
    import numpy as np
    n = len(a)
    F = [i for i in range(n) if a[i] > 0]
    for _ in range(iters):
        idx = np.array(F)
        Sff = Sig[np.ix_(idx, idx)]
        z = np.linalg.solve(Sff, a[idx])
        x = np.zeros(n)
        x[idx] = z / (a[idx] @ z)
        if (x[idx] < -1e-12).any():
            j = idx[int(np.argmin(x[idx]))]
            F.remove(j)
            continue
        lam = 2 * x @ Sig @ x
        nu = 2 * Sig @ x - lam * a
        out = [i for i in range(n) if i not in F and a[i] > 0 and nu[i] < -1e-10]
        if not out:
            x[x < 0] = 0
            return x
        F.append(min(out, key=lambda i: nu[i]))
    raise RuntimeError('qp_simplex 収束せず')


def main():
    ff = M.ff_factors()
    MKT, MKTRF, RF = ff['mkt'], ff['mktrf'], ff['rf']
    sanity = {'french_mkt_cagr_full': round(M.cagr(MKT) * 100, 2), 'french_mkt_cagr_2007': round(M.cagr(M.window(MKT, M.HOLD_START)) * 100, 2),
              'french_mkt_range': [min(MKT), max(MKT)]}
    print('sanity', sanity)
    cands = load_candidates(RF)
    print('候補', len(cands))
    # 月の連続
    sanity['noncontiguous_candidates'] = [n for n, c in cands.items() if not contiguous(c['r'])]
    # JKP 三分位の順序（ぶれ）
    order = {}
    for k in ['beta_60m', 'betabab_1260d', 'rvol_21d', 'ivol_capm_252d']:
        vs = {pf: round(S.stdev([v for ym, v in cands[f'JKP_{k}[{pf[0]}]']['r'].items() if RANK_A <= ym <= RANK_Z]) * math.sqrt(12) * 100, 1)
              for pf in ('1.0', '2.0', '3.0')}
        order[k] = vs
    sanity['jkp_tercile_vol_order_train'] = order
    sanity['jkp_pf1_is_low'] = all(v['1.0'] < v['2.0'] < v['3.0'] for v in order.values())
    jm = M.jkp_mkt('usa', 'vw')
    ks = sorted(set(jm) & set(MKTRF))
    sanity['jkp_usa_vw_mkt_vs_french_mktrf_corr'] = round(M.corr([jm[k] for k in ks], [MKTRF[k] for k in ks]), 4)

    # 訓練のシャープ（順位づけ）
    ranking = {}
    for n, c in cands.items():
        c['train_sharpe'] = train_sharpe(c['r'], RF)
        ranking[n] = round(c['train_sharpe'], 4) if c['train_sharpe'] is not None else None
    mkt_train_sharpe = train_sharpe(MKT, RF)
    best = {}
    for n, c in cands.items():
        if c['train_sharpe'] is None:
            continue
        s = c['source']
        if s not in best or c['train_sharpe'] > cands[best[s]]['train_sharpe']:
            best[s] = n
    grp_of = {s: cands[n]['group'] for s, n in best.items()}
    lr = sorted([s for s in best if grp_of[s] == '低リスク'], key=lambda s: -cands[best[s]]['train_sharpe'])
    ql = sorted([s for s in best if grp_of[s] == '質'], key=lambda s: -cands[best[s]]['train_sharpe'])
    P_src = lr[:3] + ql[:3]
    X_src = [s for s in sorted(best, key=lambda s: -cands[best[s]]['train_sharpe']) if s not in P_src]
    print('P', [best[s] for s in P_src])
    # 混合
    parts = [cands[best[s]] for s in P_src]
    ks = sorted(set.intersection(*[set(p['r']) for p in parts]))
    blend = {'name': 'BLEND6[P]', 'kind': 'blend', 'source': 'BLEND6', 'col': '+'.join(best[s] for s in P_src), 'group': '混合',
             'pub': None, 'to': S.mean([p['to'] for p in parts]) + 0.1,
             'r': {k: S.mean([p['r'][k] for p in parts]) for k in ks}, 'maps': [p['map'] for p in parts],
             'desc': 'P の6本（低リスク3＋質3）の等分混合（毎月戻す）'}
    blend['train_sharpe'] = train_sharpe(blend['r'], RF)
    P_c = [cands[best[s]] for s in P_src] + [blend]
    X_c = [cands[best[s]] for s in X_src]

    # ── 第2族の候補（事前登録2） ──
    import numpy as np
    all_src = sorted(best, key=lambda s: -cands[best[s]]['train_sharpe'])
    LR_c = [cands[best[s]] for s in lr]
    top10 = [cands[best[s]] for s in all_src[:10]]

    def mk_blend(name, comps, ws, desc):
        ws = [float(w) / float(sum(ws)) for w in ws]  # numpy の型を JSON に出さない
        ks = sorted(set.intersection(*[set(p['r']) for p in comps]))
        b = {'name': name, 'kind': 'blend', 'source': name, 'col': '+'.join(p['name'] for p in comps), 'group': '混合',
             'pub': None, 'to': sum(w * p['to'] for w, p in zip(ws, comps)) + 0.1,
             'r': {k: sum(w * p['r'][k] for w, p in zip(ws, comps)) for k in ks}, 'maps': [p['map'] for p in comps],
             'weights': ws, 'desc': desc,
             'composition': [(p['name'], round(w, 4)) for p, w in zip(comps, ws) if w > 1e-6]}
        b['train_sharpe'] = train_sharpe(b['r'], RF)
        return b

    x3 = [mk_blend('X3a_EW10', top10, [1] * 10, '訓練シャープ上位10源の等分混合')]
    sd10 = [S.stdev([p['r'][k] - RF[k] for k in p['r'] if RANK_A <= k <= RANK_Z]) for p in top10]
    x3.append(mk_blend('X3b_IV10', top10, [1 / v for v in sd10], '訓練シャープ上位10源をぶれの逆数で重みづけ'))
    U55 = [cands[best[s]] for s in all_src]
    kk = [k for k in sorted(RF) if RANK_A <= k <= RANK_Z]
    X = np.array([[p['r'][k] - RF[k] for k in kk] for p in U55])
    mu, Sig = X.mean(axis=1), np.cov(X)
    x_t = qp_simplex(Sig, mu)
    x_m = qp_simplex(Sig, np.ones(len(U55)))
    wt, wm = x_t / x_t.sum(), x_m / x_m.sum()
    keep_t = [i for i in range(len(U55)) if wt[i] > 1e-6]
    keep_m = [i for i in range(len(U55)) if wm[i] > 1e-6]
    x3.append(mk_blend('X3c_TAN', [U55[i] for i in keep_t], [wt[i] for i in keep_t], '55源の買いだけの最大シャープ（訓練の平均・共分散）'))
    x3.append(mk_blend('X3d_MINV', [U55[i] for i in keep_m], [wm[i] for i in keep_m], '55源の買いだけの最小分散（訓練の共分散）'))
    c5x = load_x4(RF)
    for n_, c in c5x.items():
        c['train_sharpe'] = train_sharpe(c['r'], RF)
    best5 = {}
    for n_, c in c5x.items():
        if c['train_sharpe'] is None:
            continue
        if c['source'] not in best5 or c['train_sharpe'] > c5x[best5[c['source']]]['train_sharpe']:
            best5[c['source']] = n_
    X4_c = [c5x[best5[s]] for s in sorted(best5)]
    print('X3', [(b['name'], len(b['composition'])) for b in x3], 'X4', [c['name'] for c in X4_c])

    def run(c, rule, fam, primary, spread=SPREAD, cost=COST, tag='', lmax=LMAX):
        e, g, n = evaluate(c, rule, MKT, MKTRF, RF, spread=spread, cost=cost, lmax=lmax)
        name = f"{fam}_{c['name']}_{rule}{tag}"
        rec = {'name': name, 'family': fam, 'primary': primary, 'rule': rule, 'candidate': c['name'], 'source': c['source'],
               'group': c['group'], 'description': f"{c['desc']} を {RULE_JA[rule]}{tag}",
               'train_sharpe_unlevered_196307_200612': round(c['train_sharpe'], 3) if c.get('train_sharpe') else None,
               'turnover_ann': round(c['to'], 3), 'pub_year': c.get('pub'), 'lmax': lmax}
        if c.get('composition'):
            rec['composition'] = c['composition']
        rec.update(e)
        rec['_c'] = c
        return rec

    fams = {f: [] for f in ('P', 'U', 'X1', 'XU', 'S', 'X2', 'X3', 'X3U', 'X4', 'X4U', 'S2')}
    for c in P_c:
        for rule in ('static', 'dynamic'):
            fams['P'].append(run(c, rule, 'P', True))
        fams['U'].append(run(c, 'unlevered', 'U', False))
        for rule in ('static', 'dynamic'):
            fams['S'].append(run(c, rule, 'S', False, spread=0.02, tag='・借入RF+2%'))
            fams['S'].append(run(c, rule, 'S', False, cost=0.003, tag='・費用3倍'))
            fams['S2'].append(run(c, rule, 'S2', False, lmax=1.5, tag='・上限1.5倍'))
    for c in X_c:
        for rule in ('static', 'dynamic'):
            fams['X1'].append(run(c, rule, 'X1', False))
        fams['XU'].append(run(c, 'unlevered', 'XU', False))
    for c in LR_c:
        for rule in ('beta_static', 'beta_dynamic'):
            fams['X2'].append(run(c, rule, 'X2', False))
    for c in x3:
        for rule in ('static', 'dynamic'):
            fams['X3'].append(run(c, rule, 'X3', False))
        fams['X3U'].append(run(c, 'unlevered', 'X3U', False))
    for c in X4_c:
        for rule in ('static', 'dynamic'):
            fams['X4'].append(run(c, rule, 'X4', False))
        fams['X4U'].append(run(c, 'unlevered', 'X4U', False))

    # C5
    print('C5 ...')
    rep_cache = {}
    GRADED = ('P', 'U', 'X1', 'XU', 'X2', 'X3', 'X3U', 'X4', 'X4U')
    for fam in GRADED:
        renorm = fam in ('X3', 'X3U')
        for rec in fams[fam]:
            c = rec['_c']
            maps = c['maps'] if c['kind'] == 'blend' else [c['map']]
            ws = c.get('weights')
            key = (tuple(maps), tuple(ws) if ws else None, rec['rule'], renorm)
            if key not in rep_cache:
                rep_cache[key] = replicate(maps, rec['rule'], RF, c['to'], weights=ws, renorm=renorm)
            rec['repl'] = rep_cache[key]
            rec['repl_map'] = [list(m) for m in maps]
    for f in fams:
        for rec in fams[f]:
            rec.pop('_c', None)

    # Holm と判定
    for fam in GRADED:
        ps = {r['name']: (r['hold']['p'] if r['hold'] else None) for r in fams[fam]}
        hp = M.holm(ps)
        for r in fams[fam]:
            r['family_holm_p'] = hp.get(r['name'])
            lev = r['rule'] != 'unlevered'
            sp = r['sharpe_pair']
            g, crit = M.grade(r['full'], r['train'], r['hold'], r['roll20_net'], cost_hold=r['net_cost_hold'],
                              repl={'regions': r['repl']['regions'], 'positive': r['repl']['positive']} if r.get('repl') else None,
                              family_holm_p=r['family_holm_p'], sharpe_pair=sp if lev else None, leveraged_or_timing=lev)
            r['grade'], r['criteria'] = g, crit
    for f in ('S', 'S2'):
        for r in fams[f]:
            r['grade'], r['criteria'] = None, None
            r['note'] = '報告のみ（判定しない）'

    # AQR BAB（参考）
    bab = None
    try:
        import openpyxl, io
        b = M.get('https://www.aqr.com/-/media/AQR/Documents/Insights/Data-Sets/Betting-Against-Beta-Equity-Factors-Monthly.xlsx',
                  name='aqr_bab_monthly.xlsx', max_age_days=60)
        wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
        ws = wb['BAB Factors']
        rows = list(ws.iter_rows(values_only=True))
        hi = next(i for i, r in enumerate(rows) if r and r[0] == 'DATE')
        h = rows[hi]
        iu = h.index('USA')
        d = {}
        for r in rows[hi + 1:]:
            if r[0] is None or r[iu] is None:
                continue
            dt = r[0]
            if hasattr(dt, 'year'):
                ym = dt.year * 100 + dt.month
            else:  # 'mm/dd/yyyy'
                mm, dd, yy = str(dt).split('/')
                ym = int(yy) * 100 + int(mm)
            d[ym] = float(r[iu])
        z = {k: 0.0 for k in d}
        bab = {'full': M.excess_stats(d, z), 'train': M.excess_stats(d, z, z=M.TRAIN_END), 'hold': M.excess_stats(d, z, a=M.HOLD_START),
               'note': 'AQR の BAB（買い−売り・ゼロコスト・米国）。参考のみ・判定しない。ex_ann = BAB の平均（年率）'}
    except Exception as e:  # noqa
        bab = {'error': str(e)[:200]}

    # 静的倍率の検算: 訓練期間のぶれ
    chk = []
    for r in fams['P']:
        if r['rule'] == 'static' and r['train']:
            chk.append({'name': r['name'], 'L': r['L_static'], 'train_vol_s': r['train']['vol_s'], 'train_vol_b': r['train']['vol_b']})
    sanity['static_train_vol_check'] = chk
    sanity['mkt_train_sharpe_196307_200612'] = round(mkt_train_sharpe, 3)

    all_rec = [r for f in fams for r in fams[f]]
    out = {
        'angle': ANGLE, 'prereg': PREREG, 'prereg_commit': sha_of(f'out/{PREREG}'),
        'prereg2': PREREG2, 'prereg2_commit': sha_of(f'out/{PREREG2}'),
        'family_labels': {'P': '主（事前登録1）', 'U': '参照・倍率1（事前登録1）', 'X1': '探索（事前登録1）', 'XU': '探索・倍率1（事前登録1）',
                          'S': '報告のみ（借入RF+2%・費用3倍）', 'X2': '探索（事前登録2）ベータ合わせ', 'X3': '探索（事前登録2）訓練で組む混合',
                          'X3U': '探索・倍率1（事前登録2）', 'X4': '探索（事前登録2）大型株ME5', 'X4U': '探索・倍率1（事前登録2）',
                          'S2': '報告のみ（上限1.5倍）'},
        'global_prereg': 'mw_prereg.json', 'global_prereg_commit': sha_of('out/mw_prereg.json'),
        'benchmark': 'French Mkt（Mkt-RF+RF・総リターン）。地域は JKP vw 市場（超過）',
        'sanity': sanity,
        'selection': {'ranking_window': [RANK_A, RANK_Z], 'n_candidates_ranked': sum(1 for v in ranking.values() if v is not None),
                      'n_candidates_total': len(ranking), 'train_sharpe_unlevered': ranking,
                      'source_best': {s: best[s] for s in best}, 'P_sources': P_src, 'X1_sources': X_src,
                      'low_risk_order': [(s, best[s], round(cands[best[s]]['train_sharpe'], 3)) for s in lr],
                      'quality_order': [(s, best[s], round(cands[best[s]]['train_sharpe'], 3)) for s in ql],
                      'blend_train_sharpe': round(blend['train_sharpe'], 3) if blend['train_sharpe'] else None},
        'aqr_bab_reference': bab,
        'deviations': DEVIATIONS,
        'x3_weights': {b['name']: b['composition'] for b in x3},
        'n_tested': len(all_rec), 'n_graded': sum(1 for r in all_rec if r['grade']),
        'grade_counts': {f: {g: sum(1 for r in fams[f] if r['grade'] == g) for g in 'SABC'} for f in GRADED},
        'tested': all_rec,
    }
    json.dumps(out, ensure_ascii=False)  # 書く前に全部が JSON にできるか確かめる（途中で壊れたファイルを残さない）
    p = M.save('mw_risk_matched.json', out)
    print('saved', p, 'n_tested', len(all_rec))
    for f in ('P', 'U', 'X2', 'X3', 'X3U', 'X4', 'X4U'):
        for r in fams[f]:
            h, hn, fu = r['hold'], r['net_cost_hold'], r['full']
            print(f"{r['name'][:60]:60s} {r['grade']} L={r['L_static'] or r['lever_info']['L_mean']} full {fu['ex_ann']:+.2f} t{fu['t']} | hold {h['ex_ann']:+.2f} t{h['t']} cagrΔ{h['cagr_diff']:+.2f} net {hn['ex_ann']:+.2f} | roll {r['roll20_net']['win_rate'] if r['roll20_net'] else None} | C5 {r['repl']['positive']}/{r['repl']['regions']} | SR {r['sharpe_pair']}")


# ───────────────────────── 頑健性の点検（事前登録3・格付けは変えない） ─────────────────────────
PREREG3 = 'mw_risk_matched_prereg3.json'
K2_DEV = ['gbr', 'deu', 'fra', 'che', 'aus', 'can', 'swe', 'nld', 'dnk', 'hkg', 'sgp', 'ita', 'esp', 'bel', 'nor', 'fin', 'isr', 'nzl', 'aut', 'irl', 'prt']
K2_EM = ['kor', 'twn', 'ind', 'chn', 'bra', 'zaf', 'mex', 'mys', 'tha', 'idn', 'tur', 'pol', 'chl', 'phl']
K3_QUAL = ['ocf_at', 'cop_at', 'cop_atl1', 'gp_at', 'gp_atl1', 'op_at', 'op_atl1', 'ope_be', 'ope_bel1', 'ebit_bev', 'ebit_sale', 'ni_be',
           'niq_at', 'niq_be', 'qmj', 'qmj_prof', 'qmj_growth', 'qmj_safety', 'sale_bev', 'at_turnover', 'f_score', 'mispricing_perf', 'o_score']
K3_VAR = ['rvol_21d', 'rmax1_21d', 'rmax5_21d', 'earnings_variability', 'ocfq_saleq_std', 'turnover_126d', 'zero_trades_21d', 'zero_trades_126d', 'zero_trades_252d']
K3_RESVAR = ['ivol_capm_21d', 'ivol_capm_252d', 'ivol_ff3_21d', 'ivol_hxz4_21d']
K3_BETA = ['beta_60m', 'betabab_1260d', 'betadown_252d', 'beta_dimson_21d']
_CTRY = {}


def ctry_series(ctry, key, pf, nmin=10):
    if (ctry, key) not in _CTRY:
        try:
            d = {}
            for x in M.jkp_rows(ctry, key, 'portfolios', 'vw'):
                if x['ret'] in ('', 'NA', 'na') or x['n'] in ('', 'NA') or float(x['n']) < nmin:
                    continue
                d.setdefault(x['pf'], {})[M._ym(x['date'])] = float(x['ret'])
            _CTRY[(ctry, key)] = d
        except Exception as e:  # noqa
            _CTRY[(ctry, key)] = None
    p = _CTRY[(ctry, key)]
    return None if p is None else p.get(pf)


def country_check(maps, weights, rule, RF, to):
    per = {}
    w_all = weights or [1 / len(maps)] * len(maps)
    for ctry in K2_DEV + K2_EM:
        parts = [(ctry_series(ctry, k, pf), w) for (k, pf), w in zip(maps, w_all)]
        miss = sum(w for p, w in parts if not p)
        if miss > 0.5 * sum(w_all):
            continue
        parts = [(p, w) for p, w in parts if p]
        ws = sum(w for _, w in parts)
        try:
            m_ex = M.jkp_mkt(ctry, 'vw')
        except Exception:  # noqa
            continue
        ks = sorted(set.intersection(*[set(p) for p, _ in parts]) & set(m_ex) & set(RF))
        if len(ks) < 120:
            continue
        r_ex = {k: sum(p[k] * w for p, w in parts) / ws for k in ks}
        cost = 0.003 if ctry in K2_EM else COST
        L = None
        if rule in ('static', 'beta_static'):
            L = static_L(r_ex, m_ex) if rule == 'static' else static_beta_L(r_ex, m_ex)
            if L is None:
                continue
        g, n, info = lever(r_ex, m_ex, RF, rule, L_static=L, to=to, cost=cost)
        if len(n) < 120:
            continue
        full = M.excess_stats(n, m_ex)
        hold = M.excess_stats(n, m_ex, a=M.HOLD_START)
        if not full:
            continue
        per[ctry] = {'from': full['from'], 'years': full['years'], 'net_ex_ann': full['ex_ann'], 't': full['t'],
                     'hold_net_ex_ann': hold['ex_ann'] if hold else None}
    k = len(per)
    pos = sum(1 for v in per.values() if v['net_ex_ann'] > 0)
    hk = [v for v in per.values() if v['hold_net_ex_ann'] is not None]
    # 片側の符号検定（正が多い向き）
    p = sum(math.comb(k, i) for i in range(pos, k + 1)) / 2 ** k if k else None
    return {'countries': k, 'positive': pos, 'share': round(pos / k, 3) if k else None, 'sign_p_one_sided': round(p, 4) if p is not None else None,
            'hold_countries': len(hk), 'hold_positive': sum(1 for v in hk if v['hold_net_ex_ann'] > 0),
            'robust': bool(k and pos / k >= 2 / 3 and p < 0.05), 'per_country': per}


def french_analog(key, pf, RF):
    """K3: JKP の特徴・三分位 → French の作り方の違うポートフォリオ（総リターン）のリスト [(名前, 系列)]"""
    t = int(float(pf))
    out = []
    if key in K3_QUAL:
        if key == 'o_score':
            t = 4 - t
        op = M.french_series('Portfolios_Formed_on_OP', 'Value Weight')
        out.append((f"FR OP {['Lo 30', 'Med 40', 'Hi 30'][t - 1]}", op[['Lo 30', 'Med 40', 'Hi 30'][t - 1]]))
        b = M.french_series('6_Portfolios_ME_OP_2x3', 'Value Weight')
        out.append((f"FR {['BIG LoOP', 'ME2 OP2', 'BIG HiOP'][t - 1]}", b[['BIG LoOP', 'ME2 OP2', 'BIG HiOP'][t - 1]]))
        return out
    fname = 'Portfolios_Formed_on_VAR' if key in K3_VAR else ('Portfolios_Formed_on_RESVAR' if key in K3_RESVAR else ('Portfolios_Formed_on_BETA' if key in K3_BETA else None))
    if not fname:
        return out
    s = M.french_series(fname, 'Value Weight')
    cols = [['Lo 20', 'Qnt 2'], ['Qnt 3'], ['Qnt 4', 'Hi 20']][t - 1]
    ks = sorted(set.intersection(*[set(s[c]) for c in cols]))
    out.append((f"FR {fname.split('_')[-1]} {'+'.join(cols)}", {k: S.mean([s[c][k] for c in cols]) for k in ks}))
    return out


def checks_main():
    ff = M.ff_factors()
    MKT, MKTRF, RF = ff['mkt'], ff['mktrf'], ff['rf']
    d = json.load(open(os.path.join(BASE, 'out', 'mw_risk_matched.json')))
    cands = load_candidates(RF)
    c5x = load_x4(RF)
    cands.update(c5x)
    sel = d['selection']
    P6 = [cands[sel['source_best'][s]] for s in sel['P_sources']]
    jm = M.jkp_mkt('usa', 'vw')

    def blend_of(name, comp):
        comps = [cands[n] for n, _ in comp]
        ws = [w for _, w in comp]
        ks = sorted(set.intersection(*[set(p['r']) for p in comps]))
        return {'name': name, 'kind': 'blend', 'source': name, 'group': '混合', 'pub': None,
                'to': sum(w * p['to'] for w, p in zip(ws, comps)) + 0.1, 'r': {k: sum(w * p['r'][k] for w, p in zip(ws, comps)) for k in ks},
                'maps': [p['map'] for p in comps], 'weights': ws, 'desc': name}

    targets = [r for r in d['tested'] if r['grade'] in ('S', 'A')]
    print('点検の対象', len(targets))
    res = {}
    for r in targets:
        nm = r['candidate']
        if nm == 'BLEND6[P]':
            c = blend_of(nm, [(p['name'], 1 / 6) for p in P6])
        elif r.get('composition'):
            c = blend_of(nm, r['composition'])
        else:
            c = cands[nm]
        rule = r['rule']
        e, g, n = evaluate(c, rule, MKT, MKTRF, RF)
        assert abs(e['net_cost_hold']['ex_ann'] - r['net_cost_hold']['ex_ann']) < 0.02, ('再現できない', r['name'])
        out = {}
        # K1
        subs = [(200701, 201212), (201301, 201912), (202001, None)]
        k1 = []
        for a, z in subs:
            st = M.excess_stats(n, MKT, a=a, z=z)
            k1.append({'from': a, 'to': z or (st['to'] if st else None), 'net_ex_ann': st['ex_ann'] if st else None, 't': st['t'] if st else None})
        out['K1_subperiods'] = {'periods': k1, 'positive': sum(1 for x in k1 if x['net_ex_ann'] and x['net_ex_ann'] > 0),
                                'robust': sum(1 for x in k1 if x['net_ex_ann'] and x['net_ex_ann'] > 0) >= 2}
        # K4
        ks = sorted(k for k in set(n) & set(MKT) if k >= M.HOLD_START)
        diff = sorted((n[k] - MKT[k] for k in ks), reverse=True)
        rest = diff[12:]
        v = S.mean(rest) * 12 * 100
        out['K4_drop_best12'] = {'hold_net_ex_ann_all': round(S.mean(diff) * 1200, 2), 'without_best12': round(v, 2), 'robust': v > 0}
        # K2
        maps = c['maps'] if c['kind'] == 'blend' else [c['map']]
        out['K2_unseen_countries'] = country_check(maps, c.get('weights'), rule, RF, c['to'])
        # K3・K5・K6（JKP 米国の単独の候補だけ）
        if c['kind'] == 'jkp':
            key, pf = c['map']
            k3 = []
            for lab, ser in french_analog(key, pf, RF):
                cc = {'name': lab, 'kind': 'french', 'r': ser, 'to': c['to'], 'pub': None}
                e3, g3, n3 = evaluate(cc, rule, MKT, MKTRF, RF)
                h3 = e3['net_cost_hold']
                k3.append({'analog': lab, 'hold_net_ex_ann': h3['ex_ann'], 't': h3['t'], 'full_net_ex_ann': e3['net_cost_full']['ex_ann'], 'full_t': e3['net_cost_full']['t']})
            out['K3_independent_data'] = {'analogs': k3, 'robust': bool(k3) and all(x['hold_net_ex_ann'] > 0 for x in k3)} if k3 else {'analogs': [], 'robust': None, 'note': '対応なし（N/A）'}
            try:
                pc = M.jkp_portfolios('usa', key, 'vw_cap')[pf]
                cc = {'name': nm + '_cap', 'kind': 'jkp', 'r': {k: v + RF[k] for k, v in pc.items() if k in RF}, 'to': c['to'], 'pub': None}
                e5, _, _ = evaluate(cc, rule, MKT, MKTRF, RF)
                out['K5_capped'] = {'hold_net_ex_ann': e5['net_cost_hold']['ex_ann'], 't': e5['net_cost_hold']['t'],
                                    'full_net_ex_ann': e5['net_cost_full']['ex_ann'], 'full_t': e5['net_cost_full']['t']}
            except Exception as ex:  # noqa
                out['K5_capped'] = {'error': str(ex)[:120]}
            n_ex = {k: v - RF[k] for k, v in n.items() if k in RF}
            st = M.excess_stats(n_ex, jm, a=M.HOLD_START)
            out['K6_same_universe'] = {'hold_net_ex_ann': st['ex_ann'], 't': st['t']} if st else None
        else:
            out['K3_independent_data'] = {'robust': None, 'note': 'French の源・混合は対象外（すでに French または混合）'}
        res[r['name']] = out
        k2 = out['K2_unseen_countries']
        print(f"{r['name'][:55]:55s} {r['grade']} K1 {out['K1_subperiods']['positive']}/3 K4 {out['K4_drop_best12']['without_best12']:+.2f} "
              f"K2 {k2['positive']}/{k2['countries']} p{k2['sign_p_one_sided']} hold {k2['hold_positive']}/{k2['hold_countries']} "
              f"K3 {[(x['analog'], x['hold_net_ex_ann'], x['t']) for x in out['K3_independent_data'].get('analogs', [])]} "
              f"K5 {out.get('K5_capped', {}).get('hold_net_ex_ann')} K6 {(out.get('K6_same_universe') or {}).get('hold_net_ex_ann')}")
    d['prereg3'] = PREREG3
    d['prereg3_commit'] = sha_of(f'out/{PREREG3}')
    d['robustness_checks'] = res
    json.dumps(d, ensure_ascii=False)
    M.save('mw_risk_matched.json', d)
    print('saved checks', len(res))


# ───────────────────────── 現実の答え合わせ（事前登録4） ─────────────────────────
PREREG4 = 'mw_risk_matched_prereg4.json'
R_VEH = [('USMV', 'SPY', 'usa', 'rvol_21d', '1.0', '米国の最小分散（MSCI USA Min Vol）'),
         ('SPLV', 'SPY', 'usa', 'rvol_21d', '1.0', '米国の低ボラ100銘柄（S&P 500 Low Volatility）'),
         ('LGLV', 'SPY', 'usa', 'rvol_21d', '1.0', '米国大型の低ボラ（SPDR）'),
         ('FDLO', 'SPY', 'usa', 'rvol_21d', '1.0', '米国の低ボラ（Fidelity）'),
         ('XMLV', 'SPY', 'usa', 'rvol_21d', '1.0', '米国中型の低ボラ（S&P MidCap 400 Low Volatility）'),
         ('QUAL', 'SPY', 'usa', 'ocf_at', '3.0', '米国の質（MSCI USA Sector Neutral Quality）'),
         ('SPHQ', 'SPY', 'usa', 'ocf_at', '3.0', '米国の質（S&P 500 Quality。2016年まで別の指数）'),
         ('JQUA', 'SPY', 'usa', 'ocf_at', '3.0', '米国の質（JPMorgan US Quality Factor）'),
         ('DGRW', 'SPY', 'usa', 'ocf_at', '3.0', '米国の質を加味した増配（WisdomTree）'),
         ('EFAV', 'EFA', 'world_ex_us', 'rvol_21d', '1.0', '先進国（米国外）の最小分散'),
         ('EEMV', 'EEM', 'emerging', 'rvol_21d', '1.0', '新興国の最小分散'),
         ('ACWV', 'ACWI', 'world', 'rvol_21d', '1.0', '全世界の最小分散')]


def reality_main():
    ff = M.ff_factors()
    MKT, MKTRF, RF = ff['mkt'], ff['mktrf'], ff['rf']
    d = json.load(open(os.path.join(BASE, 'out', 'mw_risk_matched.json')))
    END = max(RF)
    recs, twins = [], {}
    for etf, bench, reg, key, pf, ja in R_VEH:
        e = {k: v for k, v in M.yahoo(etf, '1mo').items() if k <= END and k in RF}
        b = {k: v for k, v in M.yahoo(bench, '1mo').items() if k <= END and k in RF}
        r_ex = {k: v - RF[k] for k, v in e.items() if k in b}
        m_ex = {k: v - RF[k] for k, v in b.items()}
        # 紙の対応（全期間で規則を回してから、同じ月だけを取り出す）
        pp = M.jkp_portfolios(reg, key, 'vw')[pf]
        pm = MKTRF if reg == 'usa' else M.jkp_mkt(reg, 'vw')
        for rule in ('dynamic', 'beta_dynamic', 'unlevered'):
            g_ex, n_ex, info = lever(r_ex, m_ex, RF, rule, to=0.0, cost=COST)
            g, n = to_total(g_ex, RF), to_total(n_ex, RF)
            bb = {k: b[k] for k in g}
            name = f"R_{etf}_{rule}"
            rec = {'name': name, 'family': 'R', 'primary': False, 'rule': rule, 'candidate': etf, 'source': 'Yahoo ' + etf,
                   'group': '現実のETF', 'description': f"{ja}（{etf}）を {RULE_JA[rule]}・相手 {bench}（どちらも Yahoo の分配込み）",
                   'turnover_ann': 0.0, 'pub_year': None, 'benchmark_vehicle': bench,
                   'full': M.excess_stats(g, b), 'train': M.excess_stats(g, b, z=M.TRAIN_END),
                   'hold': M.excess_stats(g, b, a=M.HOLD_START), 'recent': M.excess_stats(g, b, a=M.RECENT_START),
                   'net_cost_full': M.excess_stats(n, b), 'net_cost_hold': M.excess_stats(n, b, a=M.HOLD_START),
                   'roll20_net': M.rolling(n, b, 20), 'dca20_net': M.dca(n, b, 20),
                   'sharpe_pair': {'train': sharpe_same(n, b, RF, z=M.TRAIN_END), 'hold': sharpe_same(n, b, RF, a=M.HOLD_START)},
                   'L_static': None, 'lever_info': info, 'repl': None}
            ks = sorted(n)
            rec['maxdd'] = {'strategy': round(M.maxdd(n) * 100, 1), 'mkt': round(M.maxdd(bb) * 100, 1)} if n else None
            # 紙の双子
            pg, pn, pinfo = lever(pp, pm, RF, rule, to=JKP_TO.get(key, 0.4), cost=0.003 if reg == 'emerging' else COST)
            common = sorted(set(pn) & set(n))
            if len(common) >= 24:
                if reg == 'usa':
                    ps = M.excess_stats({k: pn[k] + RF[k] for k in common}, MKT)
                else:
                    ps = M.excess_stats({k: pn[k] for k in common}, pm)
                es = M.excess_stats({k: n[k] for k in common}, b)
                rec['paper_twin'] = {'paper': f'JKP {reg} {key} 三分位{pf[0]}（vw）・相手 ' + ('French Mkt' if reg == 'usa' else f'JKP {reg} vw 市場'),
                                     'from': common[0], 'to': common[-1], 'paper_net_ex_ann': ps['ex_ann'] if ps else None, 'paper_t': ps['t'] if ps else None,
                                     'etf_net_ex_ann_same_window': es['ex_ann'] if es else None, 'etf_t_same_window': es['t'] if es else None,
                                     'gap_paper_minus_etf': round(ps['ex_ann'] - es['ex_ann'], 2) if ps and es else None,
                                     'paper_L_mean': pinfo['L_mean']}
            recs.append(rec)
    hp = M.holm({r['name']: (r['hold']['p'] if r['hold'] else None) for r in recs})
    for r in recs:
        r['family_holm_p'] = hp.get(r['name'])
        lev = r['rule'] != 'unlevered'
        g, crit = M.grade(r['full'], r['train'], r['hold'], r['roll20_net'], cost_hold=r['net_cost_hold'], repl=None,
                          family_holm_p=r['family_holm_p'], sharpe_pair=r['sharpe_pair'] if lev else None, leveraged_or_timing=lev)
        r['grade'], r['criteria'] = g, crit
        h, pt = r['net_cost_hold'], r.get('paper_twin', {})
        print(f"{r['name']:24s} {g} {h['from'] if h else ''}〜 L={r['lever_info']['L_mean']} net {h['ex_ann'] if h else None:+} t{h['t'] if h else None} "
              f"cagrΔ {h['cagr_diff'] if h else None} SR {r['sharpe_pair']['hold']} | 紙 {pt.get('paper_net_ex_ann')} t{pt.get('paper_t')} vs ETF {pt.get('etf_net_ex_ann_same_window')} 差 {pt.get('gap_paper_minus_etf')}")
    d['tested'] = [r for r in d['tested'] if r['family'] != 'R'] + recs
    d['prereg4'] = PREREG4
    d['prereg4_commit'] = sha_of(f'out/{PREREG4}')
    d['n_tested'] = len(d['tested'])
    d['n_graded'] = sum(1 for r in d['tested'] if r['grade'])
    d['grade_counts']['R'] = {g: sum(1 for r in recs if r['grade'] == g) for g in 'SABC'}
    d.setdefault('family_labels', {})['R'] = '現実の答え合わせ（事前登録4・実在ETFを同じ規則で）'
    json.dumps(d, ensure_ascii=False)
    M.save('mw_risk_matched.json', d)
    print('saved reality', len(recs))


# ───────────────────────── 選び出しをしない混合（事前登録5） ─────────────────────────
PREREG5 = 'mw_risk_matched_prereg5.json'


def ew_available(series_list, min_share=0.5):
    """毎月、そろう系列だけを等分。そろう数が全体の min_share 未満の月は使わない（欠測を0で埋めない）"""
    months = sorted(set().union(*[set(x) for x in series_list]))
    need = math.ceil(len(series_list) * min_share)
    out = {}
    for k in months:
        v = [x[k] for x in series_list if k in x]
        if len(v) >= need and len(v) > 0:
            out[k] = S.mean(v)
    return out


def contiguous_tail(d):
    """最後の連続した区間だけ（途中で抜けた月をまたがない）"""
    ks = sorted(d)
    if not ks:
        return {}
    runs, cur = [], [ks[0]]
    for a, b in zip(ks, ks[1:]):
        if next_m(a) == b:
            cur.append(b)
        else:
            runs.append(cur); cur = [b]
    runs.append(cur)
    best = max(runs, key=len)
    return {k: d[k] for k in best}


def x5_main():
    ff = M.ff_factors()
    MKT, MKTRF, RF = ff['mkt'], ff['mktrf'], ff['rf']
    d = json.load(open(os.path.join(BASE, 'out', 'mw_risk_matched.json')))
    sides = {}
    for k in QUALITY + LOW_RISK:
        side, _ = M.jkp_good_side('usa', k, 'vw', upto=M.TRAIN_END)
        sides[k] = side
    print('良い側', sides)

    def build(region, keys, weighting='vw', nmin=None):
        ser = []
        for k in keys:
            if sides.get(k) is None:
                continue
            try:
                if region == 'usa':
                    p = M.jkp_portfolios('usa', k, weighting).get(sides[k])
                else:
                    p = reg_series(region, k, sides[k]) if weighting == 'vw' else None
            except Exception:  # noqa
                p = None
            if p:
                ser.append(p)
        return contiguous_tail(ew_available(ser)), len(ser)

    groups = {'X5_Q_ALL': QUALITY, 'X5_LR_ALL': LOW_RISK}
    to_g = {g: S.mean([JKP_TO[k] for k in keys]) + 0.1 for g, keys in groups.items()}
    us = {}
    for g, keys in groups.items():
        us[g], n = build('usa', keys)
        print(g, '構成', n, min(us[g]), max(us[g]))
    ks = sorted(set(us['X5_Q_ALL']) & set(us['X5_LR_ALL']))
    us['X5_QLR_ALL'] = {k: 0.5 * us['X5_Q_ALL'][k] + 0.5 * us['X5_LR_ALL'][k] for k in ks}
    to_g['X5_QLR_ALL'] = 0.5 * to_g['X5_Q_ALL'] + 0.5 * to_g['X5_LR_ALL'] + 0.1
    # 地域
    reg = {}
    for r_ in REGIONS:
        q, _ = build(r_, QUALITY)
        l, _ = build(r_, LOW_RISK)
        kk = sorted(set(q) & set(l))
        reg[r_] = {'X5_Q_ALL': q, 'X5_LR_ALL': l, 'X5_QLR_ALL': contiguous_tail({k: 0.5 * q[k] + 0.5 * l[k] for k in kk})}
    desc = {'X5_Q_ALL': '質群28特徴の良い側の三分位（JKP 米国 vw）の等分', 'X5_LR_ALL': '低リスク群18特徴の良い側の三分位の等分',
            'X5_QLR_ALL': '上の2つの 50/50'}
    recs = []
    for g in ('X5_Q_ALL', 'X5_LR_ALL', 'X5_QLR_ALL'):
        c = {'name': g, 'kind': 'blend', 'source': g, 'group': '混合（選び出しなし）', 'pub': None, 'to': to_g[g],
             'r': {k: v + RF[k] for k, v in us[g].items() if k in RF}, 'desc': desc[g]}
        c['train_sharpe'] = train_sharpe(c['r'], RF)
        for rule in ('static', 'dynamic', 'beta_static', 'beta_dynamic', 'unlevered'):
            e, gser, nser = evaluate(c, rule, MKT, MKTRF, RF)
            rec = {'name': f'X5_{g}_{rule}', 'family': 'X5', 'primary': False, 'rule': rule, 'candidate': g, 'source': g,
                   'group': c['group'], 'description': f"{desc[g]} を {RULE_JA[rule]}",
                   'train_sharpe_unlevered_196307_200612': round(c['train_sharpe'], 3) if c['train_sharpe'] else None,
                   'turnover_ann': round(c['to'], 3), 'pub_year': None}
            rec.update(e)
            # C5
            per = {}
            for r_ in REGIONS:
                r_ex = reg[r_][g]
                m_ex = reg_mkt(r_)
                if len(set(r_ex) & set(m_ex)) < 60:
                    per[r_] = None
                    continue
                cost = 0.003 if r_ == 'emerging' else COST
                L = static_L(r_ex, m_ex) if rule == 'static' else (static_beta_L(r_ex, m_ex) if rule == 'beta_static' else None)
                if rule in ('static', 'beta_static') and L is None:
                    per[r_] = None
                    continue
                _, nn, info = lever(r_ex, m_ex, RF, rule, L_static=L, to=c['to'], cost=cost)
                full = M.excess_stats(nn, m_ex)
                hold = M.excess_stats(nn, m_ex, a=M.HOLD_START)
                per[r_] = None if not full else {'from': full['from'], 'to': full['to'], 'years': full['years'], 'L': round(L, 3) if L else info['L_mean'],
                                                 'net_ex_ann': full['ex_ann'], 't': full['t'], 'cagr_diff': full['cagr_diff'], 'positive': full['ex_ann'] > 0,
                                                 'hold_net_ex_ann': hold['ex_ann'] if hold else None, 'hold_t': hold['t'] if hold else None}
            got = [x for x in REGIONS if per.get(x)]
            nonus = [x for x in NONUS if per.get(x)]
            rec['repl'] = {'regions': len(got), 'positive': sum(1 for x in got if per[x]['positive']), 'nonus_regions': len(nonus),
                           'nonus_positive': sum(1 for x in nonus if per[x]['positive']), 'per_region': per}
            rec['_n'] = nser
            rec['_c'] = c
            recs.append(rec)
    hp = M.holm({r['name']: (r['hold']['p'] if r['hold'] else None) for r in recs})
    checks = {}
    for r in recs:
        r['family_holm_p'] = hp.get(r['name'])
        lev = r['rule'] != 'unlevered'
        g, crit = M.grade(r['full'], r['train'], r['hold'], r['roll20_net'], cost_hold=r['net_cost_hold'],
                          repl={'regions': r['repl']['regions'], 'positive': r['repl']['positive']},
                          family_holm_p=r['family_holm_p'], sharpe_pair=r['sharpe_pair'] if lev else None, leveraged_or_timing=lev)
        r['grade'], r['criteria'] = g, crit
        n, c = r.pop('_n'), r.pop('_c')
        if g in ('S', 'A'):
            k1 = []
            for a, z in [(200701, 201212), (201301, 201912), (202001, None)]:
                st = M.excess_stats(n, MKT, a=a, z=z)
                k1.append({'from': a, 'net_ex_ann': st['ex_ann'] if st else None, 't': st['t'] if st else None})
            ks = sorted(k for k in set(n) & set(MKT) if k >= M.HOLD_START)
            diff = sorted((n[k] - MKT[k] for k in ks), reverse=True)
            capped, _ = build('usa', QUALITY if 'Q_ALL' in c['name'] else LOW_RISK, weighting='vw_cap') if c['name'] != 'X5_QLR_ALL' else (None, 0)
            if c['name'] == 'X5_QLR_ALL':
                qc, _ = build('usa', QUALITY, 'vw_cap'); lc, _ = build('usa', LOW_RISK, 'vw_cap')
                kk = sorted(set(qc) & set(lc)); capped = {k: 0.5 * qc[k] + 0.5 * lc[k] for k in kk}
            cc = dict(c); cc['r'] = {k: v + RF[k] for k, v in capped.items() if k in RF}
            e5, _, _ = evaluate(cc, r['rule'], MKT, MKTRF, RF)
            checks[r['name']] = {'K1_subperiods': {'periods': k1, 'positive': sum(1 for x in k1 if x['net_ex_ann'] and x['net_ex_ann'] > 0)},
                                 'K4_drop_best12': {'without_best12': round(S.mean(diff[12:]) * 1200, 2)},
                                 'K5_capped': {'hold_net_ex_ann': e5['net_cost_hold']['ex_ann'], 't': e5['net_cost_hold']['t']}}
        h, f = r['net_cost_hold'], r['full']
        print(f"{r['name']:32s} {g} L={r['L_static'] or r['lever_info']['L_mean']} full {f['ex_ann']:+.2f} t{f['t']} ({f['from']}) tr t{r['train']['t']} | hold net {h['ex_ann']:+.2f} t{h['t']} | roll {r['roll20_net']['win_rate'] if r['roll20_net'] else None} | C5 {r['repl']['positive']}/{r['repl']['regions']} | SR {r['sharpe_pair']} | {checks.get(r['name'])}")
    d['tested'] = [r for r in d['tested'] if r['family'] != 'X5'] + recs
    d['prereg5'] = PREREG5
    d['prereg5_commit'] = sha_of(f'out/{PREREG5}')
    d['x5_good_sides'] = sides
    d['n_tested'] = len(d['tested'])
    d['n_graded'] = sum(1 for r in d['tested'] if r['grade'])
    d['grade_counts']['X5'] = {g: sum(1 for r in recs if r['grade'] == g) for g in 'SABC'}
    d.setdefault('family_labels', {})['X5'] = '探索（事前登録5）選び出しをしない混合'
    d.setdefault('robustness_checks', {}).update(checks)
    json.dumps(d, ensure_ascii=False)
    M.save('mw_risk_matched.json', d)
    print('saved x5', len(recs))


if __name__ == '__main__':
    if '--checks' in sys.argv:
        checks_main()
    elif '--reality' in sys.argv:
        reality_main()
    elif '--x5' in sys.argv:
        x5_main()
    else:
        main()
