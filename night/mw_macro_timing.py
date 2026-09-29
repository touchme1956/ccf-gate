#!/usr/bin/env python3
"""night/mw_macro_timing.py — 『市場に勝てる歴史検証』角度 macro_timing（読むだけ・門の判定には不使用）

問い: Goyal-Welch の予言変数（dfy dfr tms tbl lty ntis infl svar ik csp dp ep）を、その時点までのデータだけで
毎月回帰し直した1か月先の予測（再帰的 out-of-sample）を等しく平均し（Rapach-Strauss-Zhou 2010）、
Campbell-Thompson 2008 の符号の制約を掛け、w = 予測 ÷ (3×直近60か月の分散) ∈ [0, 1.5] で株の割合を決める規則と、
『成長×トレンド』（10か月線の下 かつ 失業率が上向きの時だけ現金）は、市場（French Mkt）を買って持つだけに勝つか。

事前登録: out/mw_macro_timing_prereg.json（規則・線は測る前に固定。ここで動かさない）
出力    : out/mw_macro_timing.json

約束
- 月次リターンは小数。株と現金は総リターンどうし・超過どうしで比べる（混ぜない）。
- 月 m のリターンの判断は d(m) の月末: 1926-07 以降 d = m−1、1926-06 以前（株価が月中平均）d = m−2。
- 欠測を 0 で埋めない（絶対のルール7）。その変数はその月の予測に参加しない。
"""
import io, json, math, os, random, re, statistics as S, subprocess, sys, zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import mw_common as M  # noqa: E402

PREREG = 'mw_macro_timing_prereg.json'
OUT = 'mw_macro_timing.json'
END_M = 202608
FRENCH_START = 192607
GAMMA, WMAX, SPREAD, COST = 3.0, 1.5, 0.015, 0.001
BURN, VARWIN = 240, 60
TAX = 0.20315
SEED = 20260928
PREDS = ['dfy', 'dfr', 'tms', 'tbl', 'lty', 'ntis', 'infl', 'svar', 'ik', 'csp', 'dp', 'ep']
SIGN = {'dp': 1, 'ep': 1, 'dfy': 1, 'dfr': 1, 'tms': 1, 'tbl': -1, 'lty': -1, 'ntis': -1, 'infl': -1, 'svar': 1, 'ik': -1, 'csp': 1}
JST_PREDS = ['dp', 'tbl', 'lty', 'tms', 'infl']
JST_C5 = ['AUS', 'BEL', 'CHE', 'DEU', 'DNK', 'ESP', 'FIN', 'FRA', 'GBR', 'ITA', 'JPN', 'NLD', 'NOR', 'PRT', 'SWE']
POST = {'P': {'post_GW_CT_2008': 200901, 'post_RSZ_2010': 201101}, 'I': {'post_GW_CT_2008': 200901},
        'G': {'post_GrowthTrend_2016': 201701, 'post_Faber_2007': 200801}, 'R': {'post_GW_CT_2008': 200901, 'post_Faber_2007': 200801},
        'X1': {'post_GW_CT_2008': 200901, 'post_RSZ_2010': 201101}, 'X2': {'post_GrowthTrend_2016': 201701, 'post_Faber_2007': 200801}}
PREREG2 = 'mw_macro_timing_prereg2.json'
PREREG3 = 'mw_macro_timing_prereg3.json'
EXPLORATORY = {'X1': PREREG2, 'X2': PREREG2, 'X3': PREREG3}
POST['X3'] = {'post_JMJ_1996': 199701, 'post_GrowthTrend_2016': 201701}
POLICY_CC = {'GB': ['GB'], 'AU': ['AU'], 'CA': ['CA'], 'JP': ['JP'], 'DE': ['DE', 'EZ'], 'IT': ['IT', 'EZ'], 'ES': ['ES', 'EZ'],
             'FR': ['EZ'], 'NL': ['EZ'], 'BE': ['EZ'], 'AT': ['EZ'], 'FI': ['EZ'], 'IE': ['EZ']}
EURO = {'DE', 'IT', 'ES', 'FR', 'NL', 'BE', 'AT', 'FI', 'IE'}
PREREG4 = 'mw_macro_timing_prereg4.json'
NEW9 = ['avgcor', 'wtexas', 'dtoy', 'dtoat', 'skvw', 'tail', 'rdsp', 'lzrt', 'ndrbl']
SIGN.update({'avgcor': 1, 'wtexas': -1, 'dtoy': 1, 'dtoat': -1, 'skvw': -1, 'tail': 1, 'rdsp': -1, 'lzrt': 1, 'ndrbl': -1})
EXPLORATORY.update({'X4': PREREG4, 'X5': PREREG4})
POST['X4'] = {'post_GW_CT_2008': 200901, 'post_RSZ_2010': 201101}
POST['X5'] = {'post_GW_CT_2008': 200901}
POST['PH'] = POST['X3']
PREREG5 = 'mw_macro_timing_prereg5.json'
EXPLORATORY['X6'] = PREREG5
POST['X6'] = {'post_GW_CT_2008': 200901, 'post_RSZ_2010': 201101, 'post_GrowthTrend_2016': 201701}
PREREG6 = 'mw_macro_timing_prereg6.json'
ISO3 = {'GB': 'GBR', 'AT': 'AUT', 'AU': 'AUS', 'BE': 'BEL', 'CA': 'CAN', 'DK': 'DNK', 'FI': 'FIN', 'FR': 'FRA', 'DE': 'DEU', 'IE': 'IRL',
        'IT': 'ITA', 'JP': 'JPN', 'NL': 'NLD', 'NO': 'NOR', 'ES': 'ESP', 'SE': 'SWE'}
V1_PREDS = ['dp', 'ep', 'tbl', 'lty', 'tms', 'infl']
LOG = []


def log(*a):
    s = ' '.join(str(x) for x in a)
    LOG.append(s)
    print(s, flush=True)


def ym_add(k, n):
    y, m = divmod(k, 100)
    t = y * 12 + (m - 1) + n
    return (t // 12) * 100 + t % 12 + 1


def git_sha(path):
    try:
        return subprocess.run(['git', 'log', '-1', '--format=%H', '--', path], cwd=M.BASE, capture_output=True, text=True).stdout.strip()
    except Exception:  # noqa
        return None


# ───────────────────────── データ ─────────────────────────
def goyal():
    import openpyxl
    b = M.get('https://docs.google.com/spreadsheets/d/17mw_IpaiLFDrGnrPRQ2o1ugV5nJsZuD1/export?format=xlsx',
              'goyal_predictors_2025.xlsx', max_age_days=3650)
    wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
    rows = list(wb['Monthly'].iter_rows(values_only=True))
    h = list(rows[0])
    want = ['price', 'd12', 'e12', 'Rfree', 'tbl', 'lty', 'tms', 'dfy', 'dfr', 'infl', 'ntis', 'svar', 'csp', 'AAA', 'BAA', 'corpr', 'ltr',
            'avgcor', 'wtexas', 'dtoy', 'dtoat', 'skvw', 'tail', 'rdsp', 'lzrt', 'ndrbl']
    idx = {w: h.index(w) for w in want}
    out = {w: {} for w in want}
    for r in rows[1:]:
        if r[0] is None:
            continue
        k = int(r[0])
        for w, i in idx.items():
            if r[i] is not None and isinstance(r[i], (int, float)):
                out[w][k] = float(r[i])
    rq = list(wb['Quarterly'].iter_rows(values_only=True))
    hq = list(rq[0])
    iq = hq.index('i/k')
    ik = {}
    for r in rq[1:]:
        if r[0] is None or r[iq] is None:
            continue
        q = int(r[0])
        y, qq = divmod(q, 10)
        ik[y * 100 + 3 * qq] = float(r[iq])   # 四半期末の月をキーに
    out['ik_q'] = ik
    return out


def fred(sid, max_age=30):
    b = M.get('https://fred.stlouisfed.org/graph/fredgraph.csv?id=' + sid, f'fred_{sid}.csv', max_age_days=max_age).decode()
    out = {}
    for ln in b.strip().splitlines()[1:]:
        d, v = ln.split(',')[:2]
        if v in ('', '.'):
            continue
        out[int(d[:4]) * 100 + int(d[5:7])] = float(v)
    return out


def us_base(gy):
    """株の総リターン r・現金 c（1871-02〜2026-08）と French の Mkt・RF"""
    ff = M.ff_factors('monthly')
    mkt = {k: v for k, v in ff['mkt'].items() if k <= END_M}
    rf = {k: v for k, v in ff['rf'].items() if k <= END_M}
    r, c = {}, {}
    ks = sorted(gy['price'])
    for p, k in zip(ks, ks[1:]):
        if k >= FRENCH_START:
            break
        if k in gy['d12'] and gy['price'][p] > 0 and k in gy['Rfree']:
            r[k] = (gy['price'][k] + gy['d12'][k] / 12) / gy['price'][p] - 1
            c[k] = gy['Rfree'][k]
    for k in mkt:
        if k in rf:
            r[k] = mkt[k]
            c[k] = rf[k]
    return r, c, mkt, rf


def dtime(m):
    return ym_add(m, -1) if m >= FRENCH_START else ym_add(m, -2)


def pred_at(gy, name, d):
    """d の月末に分かる予言変数の値（None = 無い）"""
    try:
        if name == 'dp':
            return math.log(gy['d12'][d] / gy['price'][d])
        if name == 'ep':
            e = gy['e12'].get(ym_add(d, -3))
            return math.log(e / gy['price'][d]) if e and e > 0 else None
        if name in ('infl', 'ndrbl'):
            return gy[name].get(ym_add(d, -1))
        if name == 'ik':
            best = None
            for q in gy['_ikq']:
                if ym_add(q, 3) <= d:
                    best = q
                else:
                    break
            return gy['ik_q'][best] if best is not None else None
        return gy[name].get(d)
    except (KeyError, ValueError, ZeroDivisionError):
        return None


# ───────────────────────── 再帰的な予測 ─────────────────────────
def recursive_forecasts(months, z, er, dfun, min_pairs=BURN):
    """z: {m: 月 m の判断時点で分かる値}、er: {m: 超過}。→ {m: (raw, ct用の(a,b,pm), n)}
    推定に使うのは m' ≤ d(m) で z・er の両方がある月"""
    pair_ms = sorted(m for m in er if m in z)
    out = {}
    n = sx = sy = sxx = sxy = 0.0
    j = 0
    for m in months:
        d = dfun(m)
        while j < len(pair_ms) and pair_ms[j] <= d:
            x, y = z[pair_ms[j]], er[pair_ms[j]]
            n += 1; sx += x; sy += y; sxx += x * x; sxy += x * y
            j += 1
        if n < min_pairs or m not in z:
            continue
        vx = sxx - sx * sx / n
        if vx <= 0:
            continue
        b = (sxy - sx * sy / n) / vx
        a = (sy - b * sx) / n
        out[m] = {'raw': a + b * z[m], 'b': b, 'pm': sy / n, 'n': int(n)}
    return out


def ct_value(fc, sign):
    v = fc['raw'] if (fc['b'] * sign) > 0 else fc['pm']
    return max(0.0, v)


def prevailing_mean(months, er, dfun, min_obs=BURN):
    ks = sorted(er)
    out = {}
    j = n = s = 0
    for m in months:
        d = dfun(m)
        while j < len(ks) and ks[j] <= d:
            n += 1; s += er[ks[j]]; j += 1
        if n >= min_obs:
            out[m] = s / n
    return out


def trailing_var(months, er, dfun, win=VARWIN):
    ks = sorted(er)
    out = {}
    j = 0
    buf = []
    for m in months:
        d = dfun(m)
        while j < len(ks) and ks[j] <= d:
            buf.append(er[ks[j]]); j += 1
        if len(buf) >= win:
            out[m] = S.variance(buf[-win:])
    return out


def weights_from(fc, var, cap=WMAX):
    return {m: min(cap, max(0.0, f / (GAMMA * var[m]))) for m, f in fc.items() if m in var and var[m] > 0}


def run_w(W, r, c, spread=SPREAD, cost=COST):
    """W: {m: 株の割合} → gross, net, TO"""
    ks = sorted(m for m in W if m in r and m in c)
    gross, net, TO = {}, {}, {}
    for m in ks:
        w = W[m]
        g = c[m] + w * (r[m] - c[m]) - max(w - 1.0, 0.0) * spread / 12
        p = ym_add(m, -1)
        if p in gross:
            prev = W[p] * (1 + r[p]) / (1 + gross[p]) if (1 + gross[p]) > 0 else W[p]
        else:
            prev = w
        TO[m] = abs(w - prev)
        gross[m] = g
        net[m] = g - TO[m] * cost
    return gross, net, TO


# ───────────────────────── 統計（報告） ─────────────────────────
def aligned(s, b, rf, a=None, z=None):
    ks = sorted(k for k in s if k in b and k in rf and (a is None or k >= a) and (z is None or k <= z))
    return ks, [s[k] - rf[k] for k in ks], [b[k] - rf[k] for k in ks]


def sharpe_pair(s, b, rf, a=None, z=None, per_year=12):
    ks, xs, xb = aligned(s, b, rf, a, z)
    if len(ks) < 24:
        return None
    f = lambda x: S.mean(x) / S.stdev(x) * math.sqrt(per_year) if S.stdev(x) > 0 else None
    return (round(f(xs), 3), round(f(xb), 3))


def jk_memmel(s, b, rf, a=None, z=None):
    ks, xs, xb = aligned(s, b, rf, a, z)
    T = len(ks)
    if T < 24:
        return None
    sa, sb = S.mean(xs) / S.stdev(xs), S.mean(xb) / S.stdev(xb)
    rho = M.corr(xs, xb)
    V = (2 - 2 * rho + 0.5 * (sa * sa + sb * sb - 2 * sa * sb * rho * rho)) / T
    zz = (sa - sb) / math.sqrt(V) if V > 0 else None
    return {'z': round(zz, 2) if zz is not None else None, 'p_one_sided': round(0.5 * math.erfc(zz / math.sqrt(2)), 4) if zz is not None else None,
            'rho': round(rho, 3), 'months': T}


def block_boot(s, b, rf, a=None, z=None, B=2000, L=12):
    import numpy as np
    ks, xs, xb = aligned(s, b, rf, a, z)
    T = len(ks)
    if T < 48:
        return None
    xs, xb = np.array(xs), np.array(xb)
    rng = np.random.default_rng(SEED)
    nb = int(math.ceil(T / L))
    starts = rng.integers(0, T - L + 1, size=(B, nb))
    idx = (starts[:, :, None] + np.arange(L)[None, None, :]).reshape(B, -1)[:, :T]
    A, Bm = xs[idx], xb[idx]
    sra = A.mean(1) / A.std(1, ddof=1)
    srb = Bm.mean(1) / Bm.std(1, ddof=1)
    dif = (sra - srb) * math.sqrt(12)
    return {'frac_diff_le_0': round(float((dif <= 0).mean()), 4), 'diff_median': round(float(np.median(dif)), 3),
            'diff_p05': round(float(np.percentile(dif, 5)), 3), 'block': L, 'draws': B}


def oos_r2(er, f, base, a=None, z=None):
    ks = sorted(k for k in f if k in er and k in base and (a is None or k >= a) and (z is None or k <= z))
    if len(ks) < 24:
        return None
    e1 = [(er[k] - f[k]) ** 2 for k in ks]
    e0 = [(er[k] - base[k]) ** 2 for k in ks]
    cw = [(er[k] - base[k]) ** 2 - ((er[k] - f[k]) ** 2 - (base[k] - f[k]) ** 2) for k in ks]
    t = M.nw_t(cw, 12)
    return {'r2_os_pct': round((1 - sum(e1) / sum(e0)) * 100, 3), 'cw_t': round(t, 2) if t is not None else None, 'months': len(ks)}


def tax_sim(W, r, c, dy, keys, rate=TAX, cost=COST, spread=SPREAD):
    """日本の課税口座の近似（報告のみ）。1 を keys[0] に入れて W に従う。→ 課税後の最終額"""
    E = B = 0.0
    C = 1.0
    gains = {}
    carry = []  # [(年, 損失)]
    for i, m in enumerate(keys):
        V = E + C
        x = W[m] * V - E
        if x < -1e-12 and E > 0:
            sell = min(-x, E)
            g = sell - B * sell / E
            gains[m // 100] = gains.get(m // 100, 0.0) + g
            B -= B * sell / E; E -= sell
            C += sell - sell * cost
        elif x > 1e-12:
            E += x; B += x
            C -= x + x * cost
        dv = dy.get(m, 0.0)
        E0 = E
        E *= 1 + r[m] - dv
        C += E0 * dv * (1 - rate)
        if C >= 0:
            C *= 1 + c[m] * (1 - rate)
        else:
            C *= 1 + c[m] + spread / 12
        if m % 100 == 12 or i == len(keys) - 1:
            y = m // 100
            g = gains.pop(y, 0.0)
            if i == len(keys) - 1:
                g += E - B
            carry = [(yy, l) for yy, l in carry if y - yy <= 3]
            if g > 0:
                use = 0.0
                nc = []
                for yy, l in carry:
                    take = min(l, g - use)
                    use += take
                    if l - take > 1e-15:
                        nc.append((yy, l - take))
                carry = nc
                C -= (g - use) * rate
            elif g < 0:
                carry.append((y, -g))
    return E + C


def tax_report(W, r, c, dy, a, z):
    keys = [m for m in sorted(W) if m in r and m in c and m >= a and m <= z]
    if len(keys) < 24:
        return None
    ts = tax_sim(W, r, c, dy, keys)
    tb = tax_sim({m: 1.0 for m in keys}, r, c, dy, keys)
    yrs = len(keys) / 12
    return {'window': [keys[0], keys[-1]], 'after_tax_ratio': round(ts / tb, 4),
            'after_tax_cagr_diff': round(((ts ** (1 / yrs)) - (tb ** (1 / yrs))) * 100, 2)}


# ───────────────────────── JST（年次・C5） ─────────────────────────
def jst_data():
    import openpyxl
    b = M.get('https://www.macrohistory.net/app/download/9834512569/JSTdatasetR6.xlsx?t=1763503850', 'jst_R6.xlsx', max_age_days=3650)
    wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
    rows = list(wb['Sheet1'].iter_rows(values_only=True))
    h = rows[0]
    cols = ['eq_tr', 'eq_dp', 'bill_rate', 'cpi', 'ltrate']
    I = {c: h.index(c) for c in cols + ['year', 'iso']}
    out = {}
    for r in rows[1:]:
        d = {c: (float(r[I[c]]) if isinstance(r[I[c]], (int, float)) else None) for c in cols}
        out.setdefault(r[I['iso']], {})[int(r[I['year']])] = d
    return out


def jst_country(dat, preds, mode, cap, return_w=False):
    """mode: 'mean_ct' / 'mean_raw' / 'median_ct' / 'pm' / 'tilt_ratio' / 'tilt_dev'（探索2）。→ (戦略 net {年}, 相手 {年}, 除外した年)"""
    ys = sorted(dat)
    excl = set()
    for Y in ys:
        p = dat.get(Y - 1)
        if p and dat[Y]['cpi'] and p['cpi'] and dat[Y]['cpi'] / p['cpi'] - 1 >= 1.0:
            excl.add(Y)
    er = {}
    for Y in ys:
        d = dat[Y]
        if Y not in excl and d['eq_tr'] is not None and d['bill_rate'] is not None:
            er[Y] = d['eq_tr'] - d['bill_rate']

    def zval(name, Y):   # 年 Y のリターンに使う値（Y−1 年末に分かる）
        p = dat.get(Y - 1, {})
        try:
            if name == 'dp':
                return math.log(p['eq_dp']) if p.get('eq_dp') and p['eq_dp'] > 0 else None
            if name == 'tbl':
                return p.get('bill_rate')
            if name == 'lty':
                return p['ltrate'] / 100 if p.get('ltrate') is not None else None
            if name == 'tms':
                return p['ltrate'] / 100 - p['bill_rate'] if p.get('ltrate') is not None and p.get('bill_rate') is not None else None
            if name == 'infl':
                a, b = dat.get(Y - 2, {}).get('cpi'), dat.get(Y - 3, {}).get('cpi')
                return a / b - 1 if a and b else None
        except (TypeError, ValueError):
            return None
        return None

    dfun = lambda Y: Y - 1
    fcs = {}
    for pn in preds:
        z = {Y: v for Y in ys if (v := zval(pn, Y)) is not None}
        fcs[pn] = recursive_forecasts(ys, z, er, dfun, min_pairs=20)
    pm = prevailing_mean(ys, er, dfun, min_obs=20)
    var = trailing_var(ys, er, dfun, win=10)
    fc = {}
    for Y in ys:
        if mode == 'pm':
            if Y in pm:
                fc[Y] = pm[Y]
            continue
        vals = []
        for pn in preds:
            f = fcs[pn].get(Y)
            if f is None:
                continue
            vals.append(f['raw'] if mode == 'mean_raw' else ct_value(f, SIGN[pn]))
        if vals:
            fc[Y] = S.median(vals) if mode == 'median_ct' else S.mean(vals)
    if mode == 'tilt_ratio':
        W = {Y: min(cap, max(0.0, fc[Y] / pm[Y])) for Y in fc if Y in pm and pm[Y] > 0}
    elif mode == 'tilt_dev':
        W = {Y: min(cap, max(0.0, 1 + (fc[Y] - pm[Y]) / (GAMMA * var[Y]))) for Y in fc if Y in pm and Y in var and var[Y] > 0}
    else:
        W = weights_from(fc, var, cap)
    strat, bench = {}, {}
    wprev_drift = None
    prevY = None
    for Y in ys:
        d = dat[Y]
        if Y not in W or Y in excl or d['eq_tr'] is None or d['bill_rate'] is None:
            wprev_drift = None
            continue
        w = W[Y]
        R = d['bill_rate'] + w * (d['eq_tr'] - d['bill_rate']) - max(w - 1, 0) * SPREAD
        to = abs(w - wprev_drift) if wprev_drift is not None and prevY == Y - 1 else 0.0
        net = R - to * COST
        bench[Y] = d['eq_tr']
        if net <= -0.999:
            strat[Y] = -0.999
            break
        strat[Y] = net
        wprev_drift = w * (1 + d['eq_tr']) / (1 + R) if 1 + R > 0 else None
        prevY = Y
    if return_w:
        return strat, bench, sorted(excl), {Y: W[Y] for Y in strat}
    return strat, bench, sorted(excl)


def jst_block(jst, preds, mode, cap):
    det, pos, reg = {}, 0, 0
    for iso in JST_C5 + ['USA']:
        if iso not in jst:
            continue
        st, bm, excl, Wj = jst_country(jst[iso], preds, mode, cap, return_w=True)
        xs = M.excess_stats(st, bm, per_year=1, lag=2) if len(st) >= 24 else None
        row = {'years': len(st), 'first': min(st) if st else None, 'excluded_hyperinflation': excl, 'net': xs,
               'avg_w': round(S.mean(Wj.values()), 3) if Wj else None}
        bill = {Y: jst[iso][Y]['bill_rate'] for Y in st}
        if len(st) >= 24:
            row['sharpe_net_vs_bench_report'] = sharpe_pair(st, bm, bill, per_year=1)
        if iso != 'USA':
            if xs and len(st) >= 30:
                reg += 1
                ok = xs['ex_ann'] > 0 and xs['cagr_diff'] > 0
                pos += ok
                row['positive'] = ok
            else:
                row['positive'] = None
        det[iso] = row
    return {'regions': reg, 'positive': pos, 'detail': det}


# ───────────────────────── French の国別（G1・R2 の C5） ─────────────────────────
COUNTRIES = {  # French のファイル名: (ISO2, 表示名)
    'UK.Dat': ('GB', '英国'), 'Austria.Dat': ('AT', 'オーストリア'), 'Austrlia.Dat': ('AU', '豪州'), 'Belgium.Dat': ('BE', 'ベルギー'),
    'Canada.Dat': ('CA', 'カナダ'), 'Denmark.Dat': ('DK', 'デンマーク'), 'Finland.Dat': ('FI', 'フィンランド'), 'France.Dat': ('FR', 'フランス'),
    'Germany.Dat': ('DE', 'ドイツ'), 'Ireland.Dat': ('IE', 'アイルランド'), 'Italy.Dat': ('IT', 'イタリア'), 'Japan.Dat': ('JP', '日本'),
    'Nethrlnd.Dat': ('NL', 'オランダ'), 'NewZland.Dat': ('NZ', 'ニュージーランド'), 'Norway.Dat': ('NO', 'ノルウェー'), 'Spain.Dat': ('ES', 'スペイン'),
    'Sweden.Dat': ('SE', 'スウェーデン'), 'Swtzrlnd.Dat': ('CH', 'スイス'),
}
UNEMP_CC = ['GB', 'AT', 'AU', 'BE', 'CA', 'DK', 'FI', 'FR', 'DE', 'IE', 'IT', 'JP', 'NL', 'NO', 'ES', 'SE']


def _blocks(lines):
    out, cur = [], None
    for l in lines:
        s = l.strip()
        if re.match(r'^\d{4,6}\s', s):
            if cur is not None:
                cur['rows'].append(s.split())
            continue
        if s == '':
            if cur is not None and cur['rows']:
                out.append(cur)
                cur = None
            continue
        if cur is None or cur['rows']:
            if cur is not None and cur['rows']:
                out.append(cur)
            cur = {'hdr': [s], 'rows': []}
        else:
            cur['hdr'].append(s)
    if cur and cur['rows']:
        out.append(cur)
    return out


def french_countries():
    z = zipfile.ZipFile(io.BytesIO(M.get(M.FR.format('F-F_International_Countries'), name='fr_F-F_International_Countries.zip')))
    out = {}
    for n in z.namelist():
        if n not in COUNTRIES:
            continue
        bl = _blocks(z.read(n).decode('latin-1').splitlines())
        assert 'Local' in bl[1]['hdr'][0] and 'Not Reqd' in bl[1]['hdr'][0], n
        out[n] = {int(r[0]): float(r[1]) / 100 for r in bl[1]['rows'] if float(r[1]) > -99}
    return out


def country_gt(loc, cash, unemp, use_unemp, volmatch=False):
    """国の G1（use_unemp=True）/ R2 → (net, 相手)。volmatch=True は探索2の X2a（L をその国の訓練期間で決める）"""
    ks = sorted(k for k in loc if k in cash)
    tri, lv = {}, 1.0
    allk = sorted(loc)
    for k in allk:
        lv *= 1 + loc[k]
        tri[k] = lv
    W = {}
    for m in ks:
        d = ym_add(m, -1)
        win = [ym_add(d, -i) for i in range(10)]
        if not all(x in tri for x in win):
            continue
        down = tri[d] < S.mean(tri[x] for x in win)
        if use_unemp:
            uw = [ym_add(d, -2 - i) for i in range(12)]
            if not all(x in unemp for x in uw):
                continue
            rising = unemp[uw[0]] > S.mean(unemp[x] for x in uw)
            W[m] = 0.0 if (down and rising) else 1.0
        else:
            W[m] = 0.0 if down else 1.0
    g, n, TO = run_w(W, loc, cash)
    if volmatch:
        tk = [k for k in g if k <= M.TRAIN_END]
        if len(tk) < 60:
            return {}, {}
        L = min(WMAX, S.stdev([loc[k] for k in tk]) / S.stdev([g[k] for k in tk]))
        W = {m: (0.0 if w == 0 else L) for m, w in W.items()}
        g, n, TO = run_w(W, loc, cash)
    return n, {k: loc[k] for k in n}


def country_block(fc_all, use_unemp, volmatch=False):
    det, pos, reg = {}, 0, 0
    for fn, (cc, jp) in COUNTRIES.items():
        if fn not in fc_all:
            continue
        if use_unemp and cc not in UNEMP_CC:
            continue
        try:
            r3 = fred(f'IR3TIB01{cc}M156N')
            rc = fred(f'IRSTCI01{cc}M156N')
        except RuntimeError:
            det[cc] = 'N/A（現金の金利が取れない）'
            continue
        rate = dict(rc); rate.update(r3)
        loc = fc_all[fn]
        cash = {k: rate[ym_add(k, -1)] / 1200 for k in loc if ym_add(k, -1) in rate}
        un = {}
        if use_unemp:
            try:
                un = fred(f'LRHUTTTT{cc}M156S')
            except RuntimeError:
                det[cc] = 'N/A（失業率が取れない）'
                continue
        n, b = country_gt(loc, cash, un, use_unemp, volmatch)
        if len(n) < 120:
            det[cc] = f'N/A（評価できる月が {len(n)}）'
            continue
        xs = M.excess_stats(n, b)
        reg += 1
        ok = xs['ex_ann'] > 0 and xs['cagr_diff'] > 0
        pos += ok
        det[cc] = {'name': jp, 'window': [min(n), max(n)], 'net': xs, 'positive': ok,
                   'hold_net': M.excess_stats(n, b, a=M.HOLD_START)}
    return {'regions': reg, 'positive': pos, 'detail': det}


# ───────────────────────── 金融政策の向き（探索3） ─────────────────────────
def policy_regime(parts):
    """parts: [(系列 {yyyymm: 金利}, 始め, 終わり)] を順につなぐ。変更は同じ系列の中の前の値との差だけ（つなぎ目は数えない）。
    → {yyyymm: +1（最後の変更が引き下げ＝緩和）/ −1（引き上げ＝引き締め）}。最初の変更より前と、系列の無い月は作らない"""
    reg, cur = {}, None
    for s, a, z in parts:
        prev = None
        for k in sorted(k for k in s if (a is None or k >= a) and (z is None or k <= z)):
            if prev is not None:
                dv = s[k] - s[prev]
                if dv < -1e-9:
                    cur = 1
                elif dv > 1e-9:
                    cur = -1
            if cur is not None:
                reg[k] = cur
            prev = k
    return reg


def country_policy_block(fc_all, volmatch):
    det, pos, reg_n = {}, 0, 0
    ser = {}
    for cc in set(x for v in POLICY_CC.values() for x in v):
        try:
            ser[cc] = fred(f'INTDSR{cc}M193N')
        except RuntimeError:
            ser[cc] = None
    for fn, (cc, jp) in COUNTRIES.items():
        if cc not in POLICY_CC or fn not in fc_all:
            continue
        parts = []
        for src in POLICY_CC[cc]:
            if ser.get(src) is None:
                continue
            if src == 'EZ':
                parts.append((ser[src], 199901, None))
            else:
                parts.append((ser[src], None, 199812 if cc in EURO else None))
        rg = policy_regime(parts)
        try:
            r3 = fred(f'IR3TIB01{cc}M156N')
            rc = fred(f'IRSTCI01{cc}M156N')
        except RuntimeError:
            det[cc] = 'N/A（現金の金利が取れない）'
            continue
        rate = dict(rc); rate.update(r3)
        loc = fc_all[fn]
        cash = {k: rate[ym_add(k, -1)] / 1200 for k in loc if ym_add(k, -1) in rate}
        W = {m: (1.0 if rg[ym_add(m, -1)] == 1 else 0.0) for m in loc if m in cash and ym_add(m, -1) in rg}
        g, n, TO = run_w(W, loc, cash)
        if volmatch:
            tk = [k for k in g if k <= M.TRAIN_END]
            if len(tk) < 60:
                det[cc] = f'N/A（訓練期間が {len(tk)} か月）'
                continue
            L = min(WMAX, S.stdev([loc[k] for k in tk]) / S.stdev([g[k] for k in tk]))
            W = {m: (0.0 if w == 0 else L) for m, w in W.items()}
            g, n, TO = run_w(W, loc, cash)
        if len(n) < 120:
            det[cc] = f'N/A（評価できる月が {len(n)}）'
            continue
        b = {k: loc[k] for k in n}
        xs = M.excess_stats(n, b)
        reg_n += 1
        ok = xs['ex_ann'] > 0 and xs['cagr_diff'] > 0
        pos += ok
        det[cc] = {'name': jp, 'window': [min(n), max(n)], 'net': xs, 'positive': ok, 'pct_restrictive': round(sum(1 for w in W.values() if w == 0) / len(W), 3)}
    return {'regions': reg_n, 'positive': pos, 'detail': det}


# ───────────────────────── 検証（prereg6） ─────────────────────────
def french_countries_ann():
    """→ {ファイル名: {'loc': {yyyymm: 総リターン}, 'ann': {年: {'EP','Yld'}}}}（mw_valdca と同じ読み方）"""
    z = zipfile.ZipFile(io.BytesIO(M.get(M.FR.format('F-F_International_Countries'), name='fr_F-F_International_Countries.zip')))
    out = {}
    for n in z.namelist():
        if n not in COUNTRIES:
            continue
        bl = _blocks(z.read(n).decode('latin-1').splitlines())
        assert 'Local' in bl[1]['hdr'][0] and 'Not Reqd' in bl[1]['hdr'][0], n
        assert 'Average of Annual' in bl[8]['hdr'][0] and 'Not Reqd' in bl[8]['hdr'][2], n
        loc = {int(r[0]): float(r[1]) / 100 for r in bl[1]['rows'] if float(r[1]) > -99}
        ann = {}
        for r in bl[8]['rows']:
            ep, yld = float(r[3]), float(r[5])
            ann[int(r[0])] = {'EP': ep / 100 if ep > -99 else None, 'Yld': yld / 100 if yld > -99 else None}
        out[n] = {'loc': loc, 'ann': ann}
    return out


def v1_country(fc, cc):
    """X6a と部品（予測で傾けるだけ・トレンド×失業だけ）を1か国で → dict"""
    loc, ann = fc['loc'], fc['ann']
    r3 = fred(f'IR3TIB01{cc}M156N')
    try:
        rc = fred(f'IRSTCI01{cc}M156N')
    except RuntimeError:
        rc = {}
    rate = dict(rc); rate.update(r3)
    try:
        y10 = fred(f'IRLTLT01{cc}M156N')
    except RuntimeError:
        y10 = {}
    try:
        cpi = fred(f'{ISO3[cc]}CPIALLMINMEI')
    except RuntimeError:
        cpi = {}
    un = fred(f'LRHUTTTT{cc}M156S')
    ks = sorted(loc)
    cash = {k: rate[ym_add(k, -1)] / 1200 for k in ks if ym_add(k, -1) in rate}
    # 配当を除いた価格指数（年 Y の比率は Y 年7月から・mw_valdca と同じ近似）
    def known_year(k):
        y, m = divmod(k, 100)
        return y if m > 6 else y - 1
    PI, lv = {ym_add(ks[0], -1): 1.0}, 1.0
    for k in ks:
        a = ann.get(known_year(k)) or ann.get(k // 100)
        dyv = a['Yld'] if a and a['Yld'] is not None else 0.0   # 価格指数の近似のためだけ（年内の比にしか効かない）
        lv *= (1 + loc[k]) / (1 + dyv / 12)
        PI[k] = lv
    tri, lv = {}, 1.0
    for k in ks:
        lv *= 1 + loc[k]
        tri[k] = lv

    def zval(name, d):
        Yk = known_year(d)
        a = ann.get(Yk)
        dec = (Yk - 1) * 100 + 12
        if name in ('dp', 'ep'):
            if a is None or dec not in PI or d not in PI:
                return None
            v = a['Yld'] if name == 'dp' else a['EP']
            return math.log(v * PI[dec] / PI[d]) if v and v > 0 else None
        if name == 'tbl':
            return rate[d] / 100 if d in rate else None
        if name == 'lty':
            return y10[d] / 100 if d in y10 else None
        if name == 'tms':
            return y10[d] / 100 - rate[d] / 100 if d in y10 and d in rate else None
        if name == 'infl':
            a1, b1 = cpi.get(ym_add(d, -1)), cpi.get(ym_add(d, -2))
            return a1 / b1 - 1 if a1 and b1 else None
        return None

    dfun = lambda m: ym_add(m, -1)
    er = {k: loc[k] - cash[k] for k in ks if k in cash}
    FCc = {}
    for pn in V1_PREDS:
        z = {m: v for m in ks if (v := zval(pn, dfun(m))) is not None}
        FCc[pn] = recursive_forecasts(ks, z, er, dfun)
    pm = prevailing_mean(ks, er, dfun)
    f = {}
    for m in ks:
        v = [ct_value(FCc[pn][m], SIGN[pn]) for pn in V1_PREDS if m in FCc[pn]]
        if v:
            f[m] = S.mean(v)
    tilt = {m: min(WMAX, max(0.0, f[m] / pm[m])) for m in f if m in pm and pm[m] > 0}
    bad = {}
    for m in ks:
        d = dfun(m)
        win = [ym_add(d, -i) for i in range(10)]
        uw = [ym_add(d, -2 - i) for i in range(12)]
        if not all(x in tri for x in win) or not all(x in un for x in uw):
            continue
        bad[m] = (tri[d] < S.mean(tri[x] for x in win)) and (un[uw[0]] > S.mean(un[x] for x in uw))
    rows = {}
    for tag, W in (('X6a', {m: (0.0 if bad[m] else tilt[m]) for m in bad if m in tilt}),
                   ('tilt_only', dict(tilt)), ('gt_only', {m: (0.0 if bad[m] else 1.0) for m in bad})):
        g, n, TO = run_w(W, loc, cash)
        if len(n) < 120:
            rows[tag] = f'N/A（評価できる月が {len(n)}）'
            continue
        b = {k: loc[k] for k in n}
        xs = M.excess_stats(n, b)
        rows[tag] = {'window': [min(n), max(n)], 'net': xs, 'positive': bool(xs['ex_ann'] > 0 and xs['cagr_diff'] > 0),
                     'sharpe_net_vs_bench': sharpe_pair(n, b, cash), 'avg_w': round(S.mean(W[k] for k in n), 3),
                     'n_predictors_avg': round(S.mean(sum(1 for pn in V1_PREDS if m in FCc[pn]) for m in n), 1)}
    return rows


# ───────────────────────── 評価 ─────────────────────────
def evaluate(name, fam, W, r, c, mkt, rf, dy, er=None, fc=None, base=None):
    gross, net, TO = run_w(W, r, c)
    us_g = {k: v for k, v in gross.items() if k >= FRENCH_START}
    us_n = {k: v for k, v in net.items() if k >= FRENCH_START}
    if len(us_n) < 24:
        return {'name': name, 'family': fam, 'note': 'US の窓で評価できる月が無い'}
    ks = sorted(us_n)
    e = {'name': name, 'family': fam, 'window': [ks[0], ks[-1]],
         'full': M.excess_stats(us_g, mkt), 'train': M.excess_stats(us_g, mkt, z=M.TRAIN_END),
         'hold': M.excess_stats(us_g, mkt, a=M.HOLD_START), 'recent': M.excess_stats(us_g, mkt, a=M.RECENT_START),
         'net_full': M.excess_stats(us_n, mkt), 'net_train': M.excess_stats(us_n, mkt, z=M.TRAIN_END),
         'cost_hold': M.excess_stats(us_n, mkt, a=M.HOLD_START), 'net_recent': M.excess_stats(us_n, mkt, a=M.RECENT_START)}
    e['post_pub'] = {k: {'gross': M.excess_stats(us_g, mkt, a=a), 'net': M.excess_stats(us_n, mkt, a=a)} for k, a in POST[fam].items()}
    e['roll20'] = M.rolling(us_n, mkt, 20)
    e['dca20'] = M.dca(us_n, mkt, 20)
    e['sharpe'] = {w: sharpe_pair(us_n, mkt, rf, a, z) for w, (a, z) in
                   {'full': (None, None), 'train': (None, M.TRAIN_END), 'hold': (M.HOLD_START, None), 'recent': (M.RECENT_START, None)}.items()}
    e['sharpe_tests'] = {w: {'jk_memmel': jk_memmel(us_n, mkt, rf, a, z), 'block_bootstrap': block_boot(us_n, mkt, rf, a, z)}
                         for w, (a, z) in {'train': (None, M.TRAIN_END), 'hold': (M.HOLD_START, None)}.items()}
    wk = [W[k] for k in ks]
    wh = [W[k] for k in ks if k >= M.HOLD_START]
    e['weights'] = {'avg_full': round(S.mean(wk), 3), 'avg_hold': round(S.mean(wh), 3) if wh else None,
                    'pct_zero': round(sum(1 for x in wk if x <= 1e-9) / len(wk), 3),
                    'pct_levered': round(sum(1 for x in wk if x > 1 + 1e-9) / len(wk), 3),
                    'pct_at_cap': round(sum(1 for x in wk if x >= WMAX - 1e-9) / len(wk), 3)}
    e['turnover_per_year'] = round(S.mean(TO[k] for k in ks) * 12, 3)
    e['maxdd'] = {'s_net': round(M.maxdd(us_n) * 100, 1), 'b': round(M.maxdd({k: mkt[k] for k in us_n if k in mkt}) * 100, 1)}
    # 1891〜1925（報告）
    pre_g = {k: v for k, v in gross.items() if 189101 <= k < FRENCH_START}
    pre_n = {k: v for k, v in net.items() if 189101 <= k < FRENCH_START}
    if len(pre_n) >= 120:
        e['leak_free_1891_1926'] = {'gross': M.excess_stats(pre_g, r), 'net': M.excess_stats(pre_n, r),
                                    'sharpe_net_vs_bench': sharpe_pair(pre_n, r, c), 'avg_w': round(S.mean(W[k] for k in pre_n), 3)}
    # 課税（報告）
    e['japan_tax'] = {'hold': tax_report(W, r, c, dy, M.HOLD_START, ks[-1]), 'full': tax_report(W, r, c, dy, ks[0], ks[-1])}
    if fc is not None and er is not None and base is not None:
        e['oos_r2'] = {'train': oos_r2(er, fc, base, a=FRENCH_START, z=M.TRAIN_END), 'hold': oos_r2(er, fc, base, a=M.HOLD_START),
                       'pre1926': oos_r2(er, fc, base, a=189101, z=192606)}
    return e


def main():
    pre_sha = git_sha('out/' + PREREG)
    gy = goyal()
    gy['_ikq'] = sorted(gy['ik_q'])
    r, c, mkt, rf = us_base(gy)
    er = {k: r[k] - c[k] for k in r if k in c}
    months = sorted(k for k in r if k >= 187103)
    last_d = max(gy['price'])
    months = [m for m in months if dtime(m) <= last_d]
    dy = {}
    for m in months:
        p = ym_add(m, -1)
        if p in gy['d12'] and p in gy['price']:
            dy[m] = gy['d12'][p] / gy['price'][p] / 12
    # 検算
    san = {'french_mkt_cagr_1926_2026': round(M.cagr(mkt) * 100, 2), 'french_mkt_cagr_2007': round(M.cagr(M.window(mkt, M.HOLD_START)) * 100, 2),
           'shiller_tr_cagr_1871_1926': round(M.cagr({k: v for k, v in r.items() if k < FRENCH_START}) * 100, 2),
           'months': [months[0], months[-1]], 'last_predictor_month': last_d}
    log('検算', san)

    # 予言変数
    Z = {}
    for pn in PREDS:
        Z[pn] = {}
        for m in months:
            v = pred_at(gy, pn, dtime(m))
            if v is not None:
                Z[pn][m] = v
        log('変数', pn, min(Z[pn]) if Z[pn] else None, max(Z[pn]) if Z[pn] else None, len(Z[pn]))
    FC = {pn: recursive_forecasts(months, Z[pn], er, dtime) for pn in PREDS}
    PM = prevailing_mean(months, er, dtime)
    VAR = trailing_var(months, er, dtime)

    comb = {'mean_ct': {}, 'mean_raw': {}, 'median_ct': {}}
    npart = {}
    for m in months:
        ct = [ct_value(FC[pn][m], SIGN[pn]) for pn in PREDS if m in FC[pn]]
        raw = [FC[pn][m]['raw'] for pn in PREDS if m in FC[pn]]
        if ct:
            comb['mean_ct'][m] = S.mean(ct)
            comb['median_ct'][m] = S.median(ct)
            comb['mean_raw'][m] = S.mean(raw)
            npart[m] = len(ct)

    tested = []
    # P
    specs = [('P1_RSZ_mean_CT', 'mean_ct', WMAX), ('P2_RSZ_mean_raw', 'mean_raw', WMAX), ('P3_RSZ_median_CT', 'median_ct', WMAX),
             ('P4_RSZ_mean_CT_unlev', 'mean_ct', 1.0)]
    jst = jst_data()
    for nm, mode, cap in specs:
        W = weights_from(comb[mode], VAR, cap)
        e = evaluate(nm, 'P', W, r, c, mkt, rf, dy, er, comb[mode], PM)
        e['repl'] = jst_block(jst, JST_PREDS, mode, cap)
        tested.append(e)
    # I
    for pn in PREDS:
        fct = {m: ct_value(f, SIGN[pn]) for m, f in FC[pn].items()}
        W = weights_from(fct, VAR, WMAX)
        base_i = {m: f['pm'] for m, f in FC[pn].items()}
        e = evaluate('I_' + pn, 'I', W, r, c, mkt, rf, dy, er, fct, base_i)
        e['oos_r2_raw'] = {'train': oos_r2(er, {m: f['raw'] for m, f in FC[pn].items()}, base_i, a=FRENCH_START, z=M.TRAIN_END),
                           'hold': oos_r2(er, {m: f['raw'] for m, f in FC[pn].items()}, base_i, a=M.HOLD_START)}
        e['repl'] = jst_block(jst, [pn], 'mean_ct', WMAX) if pn in JST_PREDS else None
        tested.append(e)
    # G
    tri, lv = {}, 1.0
    for k in sorted(r):
        lv *= 1 + r[k]
        tri[k] = lv
    un = fred('UNRATE')
    ip = fred('INDPRO')

    def down_at(d):
        win = [ym_add(d, -i) for i in range(10)]
        if not all(x in tri for x in win):
            return None
        return tri[d] < S.mean(tri[x] for x in win)

    def g_w(cond):
        W = {}
        for m in months:
            d = dtime(m)
            dn = down_at(d)
            if dn is None:
                continue
            cd = cond(d)
            if cd is None:
                continue
            W[m] = 0.0 if (dn and cd) else 1.0
        return W

    def c_unrate(d):
        w = [ym_add(d, -1 - i) for i in range(12)]
        return un[w[0]] > S.mean(un[x] for x in w) if all(x in un for x in w) else None

    def c_ip(d):
        a, b = ym_add(d, -1), ym_add(d, -13)
        return ip[a] / ip[b] - 1 < 0 if a in ip and b in ip else None

    def c_earn(d):
        a, b = gy['e12'].get(ym_add(d, -3)), gy['e12'].get(ym_add(d, -15))
        return a / b - 1 < 0 if a and b else None

    fc_all = french_countries()
    for nm, cond in (('G1_GT_UNRATE', c_unrate), ('G2_GT_INDPRO', c_ip), ('G3_GT_EARN', c_earn)):
        W = g_w(cond)
        e = evaluate(nm, 'G', W, r, c, mkt, rf, dy)
        e['repl'] = country_block(fc_all, True) if nm == 'G1_GT_UNRATE' else None
        tested.append(e)
    # R
    W = weights_from(PM, VAR, WMAX)
    e = evaluate('R1_PM_vol', 'R', W, r, c, mkt, rf, dy)
    e['repl'] = jst_block(jst, [], 'pm', WMAX)
    tested.append(e)
    W = g_w(lambda d: True)
    e = evaluate('R2_SMA10', 'R', W, r, c, mkt, rf, dy)
    e['repl'] = country_block(fc_all, False)
    tested.append(e)

    # ── 探索2（prereg2）: X1 予測だけで傾ける／X2 トレンド×成長を市場のぶれまで持ち上げる
    f1 = comb['mean_ct']
    W = {m: min(WMAX, max(0.0, f1[m] / PM[m])) for m in f1 if m in PM and PM[m] > 0}
    e = evaluate('X1a_tilt_ratio', 'X1', W, r, c, mkt, rf, dy)
    e['repl'] = jst_block(jst, JST_PREDS, 'tilt_ratio', WMAX)
    tested.append(e)
    W = {m: min(WMAX, max(0.0, 1 + (f1[m] - PM[m]) / (GAMMA * VAR[m]))) for m in f1 if m in PM and m in VAR and VAR[m] > 0}
    e = evaluate('X1b_tilt_dev', 'X1', W, r, c, mkt, rf, dy)
    e['repl'] = jst_block(jst, JST_PREDS, 'tilt_dev', WMAX)
    tested.append(e)
    for nm, cond, src in (('X2a_G1_volmatch', c_unrate, 'G1'), ('X2b_G2_volmatch', c_ip, 'G2')):
        W0 = g_w(cond)
        g0, _, _ = run_w(W0, r, c)
        tk = [k for k in g0 if FRENCH_START <= k <= M.TRAIN_END and k in mkt]
        L = min(WMAX, S.stdev([mkt[k] for k in tk]) / S.stdev([g0[k] for k in tk]))
        W = {m: (0.0 if w == 0 else L) for m, w in W0.items()}
        e = evaluate(nm, 'X2', W, r, c, mkt, rf, dy)
        e['L_from_train'] = round(L, 4)
        e['L_train_window'] = [tk[0], tk[-1]]
        e['repl'] = country_block(fc_all, True, volmatch=True) if src == 'G1' else None
        tested.append(e)
    # ── 探索3（prereg3）: 信用スプレッドの拡大 × トレンド／金融政策の向き
    def c_credit(d):
        w = [ym_add(d, -i) for i in range(12)]
        return gy['dfy'][d] > S.mean(gy['dfy'][x] for x in w) if all(x in gy['dfy'] for x in w) else None

    def volmatch_w(W0):
        g0, _, _ = run_w(W0, r, c)
        tk = [k for k in g0 if FRENCH_START <= k <= M.TRAIN_END and k in mkt]
        L = min(WMAX, S.stdev([mkt[k] for k in tk]) / S.stdev([g0[k] for k in tk]))
        return {m: (0.0 if w == 0 else L) for m, w in W0.items()}, L, [tk[0], tk[-1]]

    W0 = g_w(c_credit)
    e = evaluate('X3a_GT_CREDIT', 'X3', W0, r, c, mkt, rf, dy)
    e['repl'] = None
    tested.append(e)
    W, L, tw = volmatch_w(W0)
    e = evaluate('X3b_GT_CREDIT_volmatch', 'X3', W, r, c, mkt, rf, dy)
    e['L_from_train'], e['L_train_window'], e['repl'] = round(L, 4), tw, None
    tested.append(e)
    pol = policy_regime([(fred('M13009USM156NNBR'), None, 196907), (fred('INTDSRUSM193N'), 196908, 200212), (fred('DPCREDIT'), 200301, None)])
    months_all = sorted(k for k in r if k >= 187103)
    W0 = {m: (1.0 if pol[dtime(m)] == 1 else 0.0) for m in months_all if dtime(m) in pol}
    e = evaluate('X3c_MONPOL', 'X3', W0, r, c, mkt, rf, dy)
    e['pct_restrictive'] = round(sum(1 for m, w in W0.items() if w == 0 and m >= FRENCH_START) / sum(1 for m in W0 if m >= FRENCH_START), 3)
    e['repl'] = country_policy_block(fc_all, False)
    tested.append(e)
    W, L, tw = volmatch_w(W0)
    e = evaluate('X3d_MONPOL_volmatch', 'X3', W, r, c, mkt, rf, dy)
    e['L_from_train'], e['L_train_window'] = round(L, 4), tw
    e['repl'] = country_policy_block(fc_all, True)
    tested.append(e)

    # ── 探索4（prereg4）: GWZ の新しい9変数
    for pn in NEW9:
        Z[pn] = {m: v for m in months if (v := pred_at(gy, pn, dtime(m))) is not None}
        log('変数', pn, min(Z[pn]) if Z[pn] else None, max(Z[pn]) if Z[pn] else None, len(Z[pn]))
        FC[pn] = recursive_forecasts(months, Z[pn], er, dtime)

    def comb_of(names):
        out = {}
        for m in months:
            v = [ct_value(FC[pn][m], SIGN[pn]) for pn in names if m in FC[pn]]
            if v:
                out[m] = S.mean(v)
        return out

    for tag, names in (('new9', NEW9), ('all21', PREDS + NEW9)):
        fx = comb_of(names)
        W = weights_from(fx, VAR, WMAX)
        e = evaluate(f"X4{'a' if tag == 'new9' else 'c'}_{tag}_mean_CT", 'X4', W, r, c, mkt, rf, dy, er, fx, PM)
        e['repl'] = None
        tested.append(e)
        W = {m: min(WMAX, max(0.0, fx[m] / PM[m])) for m in fx if m in PM and PM[m] > 0}
        e = evaluate(f"X4{'b' if tag == 'new9' else 'd'}_{tag}_tilt", 'X4', W, r, c, mkt, rf, dy)
        e['repl'] = None
        tested.append(e)
    for pn in NEW9:
        fct = {m: ct_value(f, SIGN[pn]) for m, f in FC[pn].items()}
        W = weights_from(fct, VAR, WMAX)
        base_i = {m: f['pm'] for m, f in FC[pn].items()}
        e = evaluate('X5_' + pn, 'X5', W, r, c, mkt, rf, dy, er, fct, base_i)
        e['repl'] = None
        tested.append(e)
    # ── 探索5（prereg5）: 予測で傾ける × 悪い状態の時だけ降りる
    tilt = {m: min(WMAX, max(0.0, f1[m] / PM[m])) for m in f1 if m in PM and PM[m] > 0}
    for nm, cond in (('X6a_tilt_x_G1', c_unrate), ('X6b_tilt_x_G2', c_ip), ('X6c_tilt_x_G3', c_earn)):
        Wg = g_w(cond)
        W = {m: (0.0 if Wg[m] == 0 else tilt[m]) for m in Wg if m in tilt}
        e = evaluate(nm, 'X6', W, r, c, mkt, rf, dy)
        e['repl'] = None
        tested.append(e)

    # ── 検証（prereg6）: X6a を疑う
    ver = {'prereg': PREREG6}
    fca = french_countries_ann()
    v1 = {}
    for fn, (cc, jp) in COUNTRIES.items():
        if cc not in ISO3 or fn not in fca:
            continue
        try:
            v1[cc] = v1_country(fca[fn], cc)
        except RuntimeError as ex:
            v1[cc] = f'N/A（取得失敗: {str(ex)[:60]}）'
    summ = {}
    for tag in ('X6a', 'tilt_only', 'gt_only'):
        ok = [v[tag] for v in v1.values() if isinstance(v, dict) and isinstance(v.get(tag), dict)]
        summ[tag] = {'regions': len(ok), 'positive': sum(1 for x in ok if x['positive']),
                     'sharpe_better': sum(1 for x in ok if x['sharpe_net_vs_bench'] and x['sharpe_net_vs_bench'][0] > x['sharpe_net_vs_bench'][1])}
    summ['X6a']['replicates_2of3'] = bool(summ['X6a']['regions'] and summ['X6a']['positive'] / summ['X6a']['regions'] >= 2 / 3)
    ver['V1_international'] = {'summary': summ, 'detail': v1}
    log('検証 V1', summ)

    def x6_us(sma_n=10, ma_n=12, lag_u=1, cap=WMAX, spread=SPREAD, cost=COST):
        W = {}
        for m in months:
            d = dtime(m)
            if m not in tilt and m not in f1:
                continue
            win = [ym_add(d, -i) for i in range(sma_n)]
            uw = [ym_add(d, -lag_u - i) for i in range(ma_n)]
            if not all(x in tri for x in win) or not all(x in un for x in uw) or m not in PM or PM[m] <= 0:
                continue
            bad = tri[d] < S.mean(tri[x] for x in win) and un[uw[0]] > S.mean(un[x] for x in uw)
            W[m] = 0.0 if bad else min(cap, max(0.0, f1[m] / PM[m]))
        g, n, TO = run_w(W, r, c, spread=spread, cost=cost)
        g = {k: v for k, v in g.items() if k >= FRENCH_START}
        n = {k: v for k, v in n.items() if k >= FRENCH_START}
        return {'window': [min(n), max(n)], 'full': M.excess_stats(g, mkt), 'train': M.excess_stats(g, mkt, z=M.TRAIN_END),
                'hold': M.excess_stats(g, mkt, a=M.HOLD_START), 'cost_hold': M.excess_stats(n, mkt, a=M.HOLD_START),
                'sharpe': {w: sharpe_pair(n, mkt, rf, a, z) for w, (a, z) in {'train': (None, M.TRAIN_END), 'hold': (M.HOLD_START, None)}.items()}}, g, n
    base_v, g6, n6 = x6_us()
    v2 = {'base': {k: base_v[k] for k in ('full', 'train', 'hold')}}
    for nm, kw in (('sma8', {'sma_n': 8}), ('sma12', {'sma_n': 12}), ('unrate_ma6', {'ma_n': 6}), ('unrate_ma18', {'ma_n': 18}),
                   ('unrate_lag2', {'lag_u': 2}), ('cap1.25', {'cap': 1.25}), ('cap2.0', {'cap': 2.0}),
                   ('spread3pct', {'spread': 0.03}), ('cost0.30pct', {'cost': 0.003})):
        v, _, _ = x6_us(**kw)
        v2[nm] = {'train': (v['train'] or {}).get('ex_ann'), 'train_t': (v['train'] or {}).get('t'),
                  'hold': (v['hold'] or {}).get('ex_ann'), 'hold_t': (v['hold'] or {}).get('t'),
                  'net_hold': (v['cost_hold'] or {}).get('ex_ann'), 'full_t': (v['full'] or {}).get('t'), 'sharpe': v['sharpe']}
    ver['V2_sensitivity_post_hoc'] = v2
    log('検証 V2', {k: (v.get('train'), v.get('train_t'), v.get('hold'), v.get('hold_t')) for k, v in v2.items() if k != 'base'})
    ex89 = lambda d: {k: v for k, v in d.items() if not (200801 <= k <= 200912)}
    yearly = {}
    for y in range(1950, 2026):
        kk = [k for k in g6 if k // 100 == y]
        if len(kk) >= 6:
            yearly[y] = round(((math.prod(1 + n6[k] for k in kk)) - math.prod(1 + mkt[k] for k in kk)) * 100, 1)
    ver['V3_concentration'] = {'hold_ex_2008_09_gross': M.excess_stats(ex89(g6), mkt, a=M.HOLD_START),
                               'full_ex_2008_09_gross': M.excess_stats(ex89(g6), mkt),
                               'yearly_net_minus_mkt_pct': yearly,
                               'by_decade_cagr_diff': {str(a): round((M.cagr({k: n6[k] for k in n6 if a * 100 <= k < (a + 10) * 100}) -
                                                                        M.cagr({k: mkt[k] for k in n6 if a * 100 <= k < (a + 10) * 100})) * 100, 2)
                                                       for a in range(1950, 2030, 10) if any(a * 100 <= k < (a + 10) * 100 for k in n6)},
                               'best_5_years': sorted(yearly.items(), key=lambda x: -x[1])[:5]}
    ver['V4_realtime'] = '測っていない（ALFRED の速報値）。既知の限界'

    # 事後（格付けしない）: X3c の 2003 年以降を FF 金利の誘導目標に替えた版
    pol2 = policy_regime([(fred('M13009USM156NNBR'), None, 196907), (fred('INTDSRUSM193N'), 196908, 200212),
                          (fred('DFEDTAR'), 200301, 200812), (fred('DFEDTARU'), 200901, None)])
    W0 = {m: (1.0 if pol2[dtime(m)] == 1 else 0.0) for m in months_all if dtime(m) in pol2}
    e = evaluate('X3c_posthoc_fftarget', 'PH', W0, r, c, mkt, rf, dy)
    e['post_hoc'] = True
    e['repl'] = None
    e['note'] = '事後（X3c の結果を見た後の変更）。格付けしない・Holm に入れない'
    tested.append(e)

    for e in tested:
        e['exploratory'] = e['family'] in EXPLORATORY
        if e['exploratory']:
            e['prereg_file'] = EXPLORATORY[e['family']]

    # Holm と格付け
    fams = {}
    for e in tested:
        if e.get('post_hoc'):
            continue
        fams.setdefault(e['family'], {})[e['name']] = (e.get('hold') or {}).get('p')
    hp = {f: M.holm(v) for f, v in fams.items()}
    for e in tested:
        if e.get('post_hoc'):
            e['grade'], e['criteria'], e['holm_p'] = None, None, None
            h = e.get('hold') or {}
            log(f"{e['name']:22s} 事後  保有 {h.get('ex_ann')} t{h.get('t')}  シャープ {(e.get('sharpe') or {}).get('hold')}")
            continue
        e['holm_p'] = hp[e['family']].get(e['name'])
        sp = e.get('sharpe') or {}
        pair = {'train': sp.get('train'), 'hold': sp.get('hold')}
        rp = e.get('repl')
        rpg = {'regions': rp['regions'], 'positive': rp['positive']} if rp else None
        g, cr = M.grade(e.get('full'), e.get('train'), e.get('hold'), e.get('roll20'), cost_hold=e.get('cost_hold'), repl=rpg,
                        family_holm_p=e['holm_p'], sharpe_pair=pair, leveraged_or_timing=True)
        e['grade'], e['criteria'] = g, cr
        h = e.get('hold') or {}
        log(f"{e['name']:22s} {g}  保有 {h.get('ex_ann')} t{h.get('t')}  訓練 {(e.get('train') or {}).get('ex_ann')} t{(e.get('train') or {}).get('t')}"
            f"  シャープ {pair}  再現 {rpg}")

    # 検算（報告）: RSZ 2010 の標本期間に近い 1965-01〜2005-12 の OOS R²（RSZ は四半期で mean の R²_OS 3.58%。月次なので小さく出るのが普通）
    san['oos_r2_1965_2005'] = {'P2_mean_raw': oos_r2(er, comb['mean_raw'], PM, a=196501, z=200512),
                               'P1_mean_CT': oos_r2(er, comb['mean_ct'], PM, a=196501, z=200512)}
    # 報告: JST の国ごとのシャープ（戦略 > 相手 の国の数）
    for e in tested:
        rp = e.get('repl')
        if rp and isinstance(rp.get('detail'), dict):
            sb = [v.get('sharpe_net_vs_bench_report') for k, v in rp['detail'].items() if isinstance(v, dict) and k != 'USA' and v.get('sharpe_net_vs_bench_report')]
            if sb:
                rp['sharpe_better_countries_report'] = f"{sum(1 for a, b in sb if a > b)}/{len(sb)}"
    graded = [e for e in tested if not e.get('post_hoc')]
    cnt = {}
    for e in graded:
        cnt[e['grade']] = cnt.get(e['grade'], 0) + 1
    near = []
    for e in graded:
        cr = e.get('criteria') or {}
        base_fail = [k for k in ('C1_train', 'C2_hold_sign', 'C6_net_cost', 'C8_sharpe') if cr.get(k) is False]
        if len(base_fail) <= 1:
            near.append({'name': e['name'], 'family': e['family'], 'exploratory': e['exploratory'], 'fails_for_B': base_fail,
                         'hold_ex': (e.get('hold') or {}).get('ex_ann'), 'hold_t': (e.get('hold') or {}).get('t'),
                         'train_ex': (e.get('train') or {}).get('ex_ann'), 'train_t': (e.get('train') or {}).get('t'),
                         'sharpe': {k: (e.get('sharpe') or {}).get(k) for k in ('train', 'hold')}})
    summary = {'n_graded': len(graded), 'n_post_hoc': len(tested) - len(graded), 'grades': cnt, 'near_misses_one_short_of_B': near}
    log('格付けの数', cnt, '一歩手前', [x['name'] for x in near])

    out = {'angle': 'macro_timing', 'prereg': PREREG, 'prereg_commit': pre_sha, 'summary': summary,
           'prereg2': PREREG2, 'prereg2_commit': git_sha('out/' + PREREG2),
           'prereg3': PREREG3, 'prereg3_commit': git_sha('out/' + PREREG3),
           'prereg4': PREREG4, 'prereg4_commit': git_sha('out/' + PREREG4),
           'prereg5': PREREG5, 'prereg5_commit': git_sha('out/' + PREREG5),
           'prereg6': PREREG6, 'prereg6_commit': git_sha('out/' + PREREG6), 'verification_X6a': ver, 'sanity': san,
           'participation': {str(y): round(S.mean(npart[m] for m in npart if m // 100 == y), 1) for y in range(1891, 2026, 5) if any(m // 100 == y for m in npart)},
           'n_tested': len(tested), 'tested': tested, 'log': LOG}
    # 数えた本数・正直に書く逸脱
    out['n_graded'] = summary['n_graded']
    out['n_evaluated_including_reports'] = len(tested) + (len(ver['V2_sensitivity_post_hoc']) - 1) + 3 * sum(
        1 for v in ver['V1_international']['detail'].values() if isinstance(v, dict))
    out['deviations'] = [
        'ブリーフは予言変数の終わりを 2024-12 と書いていたが、GW 2025年版の月次は 2025-12 まであった。保有期間は 2007-01〜2026-01（2025-12 に決めた 2026-01 まで）。French の 2026-02〜08 は予言変数が無いので使っていない（引き延ばさない）',
        'FRED の UNRATE は 2025-10 が欠測（政府閉鎖で未公表）。0 や補間で埋めず、その値が要る 2025-12・2026-01 の G1・X2a・X6a の判断は作らなかった（窓は 2025-11 まで）',
        'X3c・X3d（公定歩合）は GW の変数を使わないので 2026-08 まで',
        'JST（C5）の P 族は、年次の分散で割ると株の割合の平均が 0.3〜0.5 に縮み、15か国すべて負けた＝登録どおりの規則の結果だが、C5 としては予測の力より『株を持たない』ことを測っている',
        '1981-10 の G1 の判断は、失業率がちょうど12か月平均と同じ（丸めの同点）で、計算の仕方で入れ替わる。登録どおり厳密な平均（statistics.mean）を使った',
        'mw_common.py に誤りは見つからなかった（French の CAGR 10.38%・2007〜 11.13% を再現）',
    ]
    x6 = next(e for e in tested if e['name'] == 'X6a_tilt_x_G1')
    v1s = ver['V1_international']['summary']['X6a']
    v3 = ver['V3_concentration']
    pp = x6['post_pub']['post_GrowthTrend_2016']['net']
    out['headline_ja'] = (
        f"登録どおりの主の族（予言変数の組み合わせ P・単独 I・成長×トレンド G・参照 R の21本）はすべて C。"
        f"探索を5回足して計 {summary['n_graded']} 本を格付けし、S は1本だけ: X6a（失業率が上向き かつ 10か月線の下なら現金、ほかは予測÷過去の平均で 0〜1.5倍）"
        f"＝訓練 +{x6['train']['ex_ann']}%/年 t{x6['train']['t']}・保有 +{x6['hold']['ex_ann']}%/年 t{x6['hold']['t']}・全期間 t{x6['full']['t']}・"
        f"シャープ 訓練 {x6['sharpe']['train'][0]} vs {x6['sharpe']['train'][1]}／保有 {x6['sharpe']['hold'][0]} vs {x6['sharpe']['hold'][1]}。"
        f"ただし二つの部品を結果を見た後に組み合わせた事後に近い設計で、保有の上乗せは 2008年に集中（2008〜09年を除くと +{v3['hold_ex_2008_09_gross']['ex_ann']}%/年 t{v3['hold_ex_2008_09_gross']['t']}）、"
        f"公表後（2017〜）は {pp['ex_ann']}%/年、米国以外16か国では正が {v1s['positive']}/{v1s['regions']}（2/3 に届かない）。独立の答え合わせでは再現しない＝信用できる勝ちではない")
    p = M.save(OUT, out)
    log('保存', p)


if __name__ == '__main__':
    main()
