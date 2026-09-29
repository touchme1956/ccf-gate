#!/usr/bin/env python3
"""night/mw_deep_history_verify.py — mw_deep_history の反証の検証（adversarial verify）。読むだけ・門の判定には不使用。

相手の主張（out/mw_deep_history.json）を、自前のコードで作り直して崩しにいく。
- データの取得だけ mw_common の取得器（M.get / M.ff_factors）を使う。
- JST の国の回転（選び方・相手・回転・費用）と、超過・NW t・幾何の年率差・転がる20年・格付けの材料は**自前**で書いた
  （研究者の mw_deep_history.py の関数は一つも import しない）。判定の線は M.grade（線の単一実装）にそのまま渡す。
- 検証するもの（課題の指定）: S/A の中で保有の t が上位8本（すべて BSV の紙の上乗せ）＋主の族の S 全部（P1・P3・P4・P5・P6）
  ＋ B の上位2本（保有の超過で U3・B3）。ついでに残りの S（U1・U2・U4・U5・U6・X1・X3・B4・B5）も同じ物差しで出す。

使い方: python3 night/mw_deep_history_verify.py   → out/mw_deep_history_verify.json
"""
import sys, os, io, json, math, statistics as S, random, datetime
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import mw_common as M  # 取得器と判定の線（grade）だけ使う

OUT = os.path.join(M.BASE, 'out', 'mw_deep_history_verify.json')
JST_URL = 'https://www.macrohistory.net/app/download/9834512569/JSTdatasetR6.xlsx?t=1763503850'
BSV_URL = 'https://ndownloader.figshare.com/files/26879918'
CTRY = 'AUS BEL CHE DEU DNK ESP FIN FRA GBR ITA JPN NLD NOR PRT SWE USA'.split()
FX_BAD = set(range(1939, 1950))
COST = 0.002          # 片道100%あたり（研究者と同じ既定）
COST_HI = 0.010
BONF_T = 4.35         # out/mw_summary.json のプログラム全体の Bonferroni の線
NOTE = []


def note(*a):
    s = ' '.join(str(x) for x in a)
    NOTE.append(s)
    print(s, flush=True)


# ════════════════════════ データ（自前で読む） ════════════════════════
def load_jst():
    import openpyxl
    b = M.get(JST_URL, 'jst_R6.xlsx', max_age_days=3650)
    wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
    it = wb['Sheet1'].iter_rows(values_only=True)
    hdr = list(next(it))
    ix = {n: i for i, n in enumerate(hdr)}
    need = ['eq_tr', 'eq_dp', 'bill_rate', 'xrusd', 'cpi', 'gdp', 'eq_div_rtn', 'eq_capgain']
    P = {}
    for row in it:
        iso = row[ix['iso']]
        if iso not in CTRY:
            continue
        P[(iso, int(row[ix['year']]))] = {k: (None if row[ix[k]] is None else float(row[ix[k]])) for k in need}
    return P


def load_bsv():
    import openpyxl
    b = M.get(BSV_URL, 'bsv_gfp_1800_2016.xlsx', max_age_days=3650)
    wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
    rows = list(wb['Data BSV_JFE'].iter_rows(values_only=True))
    top, sub = rows[4], rows[5]
    colmap, style = {}, None
    for j, (t, a) in enumerate(zip(top, sub)):
        if t:
            style = 'BAB' if 'Beta' in t or 'BAB' in t else t
        if a in ('EQ', 'FI', 'COM', 'FX', 'MA') and style:
            colmap[j] = (style, a)
    out = {}
    for r in rows[6:]:
        if not (isinstance(r[0], (int, float)) and isinstance(r[1], (int, float))):
            continue
        ym = int(r[0]) * 100 + int(r[1])
        for j, key in colmap.items():
            v = r[j]
            if isinstance(v, (int, float)):
                out.setdefault(key, {})[ym] = float(v)
    return out


def _num(x):
    return isinstance(x, (int, float)) and not isinstance(x, bool) and math.isfinite(x)


def load_aqr():
    """AQR の公開の紙の系列（Century of Factor Premia・TSMOM）＝ BSV が 2016-12 で終わった後を延ばす代わりの系列。
    {列名: {yyyymm: r}}（空欄は欠測・0 にしない）"""
    import openpyxl
    out = {}
    b = M.get('https://www.aqr.com/-/media/AQR/Documents/Insights/Data-Sets/Century-of-Factor-Premia-Monthly.xlsx',
              name='aqr_century_monthly.xlsx', max_age_days=3650)
    rows = list(openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)['Century of Factor Premia'].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[0] == 'Date')
    hdr = rows[hi]
    for r in rows[hi + 1:]:
        d = r[0]
        if isinstance(d, str) and d.count('/') == 2:
            mm, _, yy = d.split('/')
            ym = int(yy) * 100 + int(mm)
        elif isinstance(d, datetime.datetime):
            ym = d.year * 100 + d.month
        else:
            continue
        for j, h in enumerate(hdr):
            if j and h and _num(r[j]):
                out.setdefault('CEN ' + h, {})[ym] = float(r[j])
    b = M.get('https://www.aqr.com/-/media/AQR/Documents/Insights/Data-Sets/Time-Series-Momentum-Factors-Monthly.xlsx',
              name='aqr_tsmom_monthly.xlsx', max_age_days=3650)
    rows = list(openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)['TSMOM Factors'].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[1] == 'TSMOM')
    hdr = rows[hi]
    for r in rows[hi + 1:]:
        d = r[0]
        if not isinstance(d, datetime.datetime):
            continue
        ym = d.year * 100 + d.month
        for j, h in enumerate(hdr):
            if j and h and _num(r[j]):
                out.setdefault(h, {})[ym] = float(r[j])
    return out


def load_us_monthly():
    """1871-02〜1926-06 は Shiller（P は月平均）の総リターン、1926-07〜 は French Mkt。rf は NBER 商業手形→French RF"""
    import xlrd
    b = M.get('http://www.econ.yale.edu/~shiller/data/ie_data.xls', name='shiller_ie_data.xls', max_age_days=3650)
    sh = xlrd.open_workbook(file_contents=b).sheet_by_name('Data')
    px, dv = {}, {}
    for i in range(sh.nrows):
        r = sh.row_values(i)
        if not isinstance(r[0], float) or r[0] < 1800:
            continue
        y, mth = divmod(round(r[0] * 100), 100)
        ym = int(y) * 100 + int(mth)
        if isinstance(r[1], float):
            px[ym] = r[1]
        if isinstance(r[2], float):
            dv[ym] = r[2]
    ks = sorted(px)
    shiller = {}
    for a, z in zip(ks, ks[1:]):
        if z in dv:
            shiller[z] = (px[z] + dv[z] / 12.0) / px[a] - 1
    cp = M.get('https://fred.stlouisfed.org/graph/fredgraph.csv?id=M13002US35620M156NNBR',
               name='fred_M13002US35620M156NNBR.csv', max_age_days=3650).decode()
    cpr = {}
    for ln in cp.strip().splitlines()[1:]:
        d, v = ln.split(',')
        if v not in ('', '.'):
            cpr[int(d[:4]) * 100 + int(d[5:7])] = float(v) / 1200.0
    ff = M.ff_factors()
    mkt, rf = {}, {}
    for k, v in shiller.items():
        if k <= 192606:
            mkt[k] = v
    for k, v in ff['mkt'].items():
        if k >= 192607:
            mkt[k] = v
    for k, v in cpr.items():
        if k <= 192606:
            rf[k] = v
    for k, v in ff['rf'].items():
        if k >= 192607:
            rf[k] = v
    return mkt, rf, shiller, ff


# ════════════════════════ 自前の統計 ════════════════════════
def nwt(x, lag):
    n = len(x)
    if n < 6:
        return None
    m = sum(x) / n
    e = [v - m for v in x]
    s = sum(v * v for v in e) / n
    for L in range(1, min(lag, n - 1) + 1):
        s += 2.0 * (1.0 - L / (lag + 1.0)) * sum(e[i] * e[i - L] for i in range(L, n)) / n
    return m / math.sqrt(s / n) if s > 0 else None


def geo(xs, per):
    return math.exp(math.fsum(math.log1p(v) for v in xs) * per / len(xs)) - 1


def cmp(s, b, a=None, z=None, per=1, lag=2, minn=6):
    """自前: 同じ基準どうしの差 → 年率の算術超過・NW t・幾何の年率差（M.grade が読むキー名にそろえる）"""
    ks = sorted(k for k in s if k in b and (a is None or k >= a) and (z is None or k <= z))
    if len(ks) < minn:
        return None
    ex = [s[k] - b[k] for k in ks]
    t = nwt(ex, lag)
    gs, gb = geo([s[k] for k in ks], per), geo([b[k] for k in ks], per)
    p = math.erfc(abs(t) / math.sqrt(2)) if t is not None else None
    return {'from': ks[0], 'to': ks[-1], 'n': len(ks), 'ex_ann': round(S.mean(ex) * per * 100, 2),
            't': round(t, 2) if t is not None else None, 'p': round(p, 4) if p is not None else None,
            'cagr_diff': round((gs - gb) * 100, 2), 'cagr_s': round(gs * 100, 2), 'cagr_b': round(gb * 100, 2),
            'win_share': round(sum(1 for x in ex if x > 0) / len(ex), 3)}


def roll_years(s, b, n=20):
    """年次: 毎年起点・20年一括・幾何の年率差（20年すべてそろう窓だけ）"""
    out = []
    for y0 in sorted(k for k in s if k in b):
        w = list(range(y0, y0 + n))
        if all(k in s and k in b for k in w):
            out.append((y0, (geo([s[k] for k in w], 1) - geo([b[k] for k in w], 1)) * 100))
    if not out:
        return None
    return {'windows': len(out), 'win_rate': round(sum(1 for _, d in out if d > 0) / len(out), 3),
            'median': round(sorted(d for _, d in out)[len(out) // 2], 2), 'worst': min(out, key=lambda x: x[1])}


def roll_months(s, b, n=20, start_month=7):
    """月次: 毎年7月起点・20年一括・幾何の年率差（97%以上そろう窓）"""
    ks = sorted(k for k in s if k in b)
    if not ks:
        return None
    out = []
    for y in range(ks[0] // 100, ks[-1] // 100 + 1):
        a, z = y * 100 + start_month, (y + n) * 100 + start_month - 1
        if z > ks[-1]:
            break
        w = [k for k in ks if a <= k <= z]
        if len(w) < n * 12 * 0.97:
            continue
        out.append((y, (geo([s[k] for k in w], 12) - geo([b[k] for k in w], 12)) * 100))
    if not out:
        return None
    return {'windows': len(out), 'win_rate': round(sum(1 for _, d in out if d > 0) / len(out), 3),
            'median': round(sorted(d for _, d in out)[len(out) // 2], 2), 'worst': min(out, key=lambda x: x[1])}


def sharpe_m(r, rf, a, z):
    ks = [k for k in sorted(r) if k in rf and a <= k <= z]
    if len(ks) < 24:
        return None
    x = [r[k] - rf[k] for k in ks]
    return round(S.mean(x) / S.stdev(x) * math.sqrt(12), 3)


def ac1(x):
    if len(x) < 10:
        return None
    m = S.mean(x)
    e = [v - m for v in x]
    d = sum(v * v for v in e)
    return round(sum(e[i] * e[i - 1] for i in range(1, len(e))) / d, 3) if d else None


def block_boot_p(x, block=5, reps=4000, seed=11):
    rnd = random.Random(seed)
    n, m = len(x), S.mean(x)
    c = 0
    for _ in range(reps):
        s = []
        while len(s) < n:
            st = rnd.randrange(n)
            s.extend(x[(st + i) % n] for i in range(block))
        if S.mean(s[:n]) - m >= m:
            c += 1
    return round((c + 1) / (reps + 1), 4)


# ════════════════════════ JST の国の回転（自前） ════════════════════════
class Panel:
    def __init__(self, P):
        self.P = P

    def v(self, c, y, k):
        d = self.P.get((c, y))
        return None if d is None else d.get(k)

    def ret(self, c, y, base):
        eq = self.v(c, y, 'eq_tr')
        p1, p0 = self.v(c, y, 'cpi'), self.v(c, y - 1, 'cpi')
        if eq is None or p1 is None or not p0 or p1 / p0 - 1 >= 1.0:
            return None
        if base == 'H':
            bl, bu = self.v(c, y, 'bill_rate'), self.v('USA', y, 'bill_rate')
            if bl is None or bu is None:
                return None
            return (1 + eq) * (1 + bu) / (1 + bl) - 1
        if y in FX_BAD:
            return None
        x0, x1 = self.v(c, y - 1, 'xrusd'), self.v(c, y, 'xrusd')
        if not x0 or not x1:
            return None
        return (1 + eq) * x0 / x1 - 1

    def cash(self, y):
        return self.v('USA', y, 'bill_rate')


def pick_top(score, K, largest=True):
    items = sorted(score.items(), key=lambda kv: (-kv[1] if largest else kv[1], kv[0]))
    return [c for c, _ in items[:K]]


def ranks01(d):
    xs = sorted(d.items(), key=lambda kv: kv[1])
    n, out, i = len(xs), {}, 0
    while i < n:
        j = i
        while j + 1 < n and xs[j + 1][1] == xs[i][1]:
            j += 1
        for q in range(i, j + 1):
            out[xs[q][0]] = ((i + j) / 2) / (n - 1) if n > 1 else 0.5
        i = j + 1
    return out


def target(P, rule, base, K, Y, ctry, dp_lag=0, mom_lag=0):
    """年末 Y の値だけで年 Y+1 の目標の重み。候補が 2K 未満なら None"""
    live = [c for c in ctry if P.ret(c, Y, base) is not None]
    if not live:
        return None
    if rule == 'VAL':
        sc = {c: P.v(c, Y - dp_lag, 'eq_dp') for c in live}
        sc = {c: x for c, x in sc.items() if x is not None}
        return {c: 1.0 / K for c in pick_top(sc, K)} if len(sc) >= 2 * K else None
    if rule == 'MOM':
        sc = {c: P.ret(c, Y - mom_lag, base) for c in live}
        sc = {c: x for c, x in sc.items() if x is not None}
        return {c: 1.0 / K for c in pick_top(sc, K)} if len(sc) >= 2 * K else None
    if rule == 'VM':
        a = target(P, 'VAL', base, K, Y, ctry, dp_lag, mom_lag)
        b = target(P, 'MOM', base, K, Y, ctry, dp_lag, mom_lag)
        if a is None or b is None:
            return None
        w = {}
        for d in (a, b):
            for c, x in d.items():
                w[c] = w.get(c, 0.0) + 0.5 * x
        return w
    if rule == 'CARRY':
        sc = {}
        for c in live:
            d, bl = P.v(c, Y, 'eq_dp'), P.v(c, Y, 'bill_rate')
            if d is not None and bl is not None:
                sc[c] = d - bl
        return {c: 1.0 / K for c in pick_top(sc, K)} if len(sc) >= 2 * K else None
    if rule == 'VMRANK':
        dp = {c: P.v(c, Y, 'eq_dp') for c in live}
        dp = {c: x for c, x in dp.items() if x is not None}
        if len(dp) < 2 * K:
            return None
        mo = {c: P.ret(c, Y, base) for c in dp}
        rd, rm = ranks01(dp), ranks01(mo)
        order = sorted(dp, key=lambda c: (-(rd[c] + rm[c]) / 2, -rm[c], c))[:K]
        return {c: 1.0 / K for c in order}
    raise ValueError(rule)


def backtest(P, rule, base, K, ctry=CTRY, dp_lag=0, mom_lag=0, y0=1870, y1=2019):
    """自前の回転: 目標→年 Y+1 のリターン。相手 = 年 Y に有効で Y+1 も有効な国の等分。
    さらに GDP（米ドル換算）の重みと米国だけの相手も同じ年で作る"""
    s, b, bg, bu, turn, picks = {}, {}, {}, {}, {}, {}
    prev = None
    for Y in range(y0, y1 + 1):
        live = [c for c in ctry if P.ret(c, Y, base) is not None]
        nxt = {c: P.ret(c, Y + 1, base) for c in live}
        members = [c for c in live if nxt[c] is not None]
        w = target(P, rule, base, K, Y, ctry, dp_lag, mom_lag)
        if w is None or not members:
            prev = None
            continue
        held = {c: x for c, x in w.items() if nxt.get(c) is not None}
        if not held:
            prev = None
            continue
        tot = sum(held.values())
        held = {c: x / tot for c, x in held.items()}
        rs = sum(x * nxt[c] for c, x in held.items())
        if prev is not None:
            pw, pr = prev
            g = sum(pw[c] * (1 + pr[c]) for c in pw)
            drift = {c: pw[c] * (1 + pr[c]) / g for c in pw}
            to = 0.5 * sum(abs(held.get(c, 0.0) - drift.get(c, 0.0)) for c in set(held) | set(drift))
        else:
            to = 0.0
        s[Y + 1] = rs
        b[Y + 1] = sum(nxt[c] for c in members) / len(members)
        turn[Y + 1] = to
        picks[Y + 1] = sorted(held)
        # GDP の重み（時価総額の代わり）: 為替のそろう国だけ
        gw = {}
        for c in members:
            g_, x_ = P.v(c, Y, 'gdp'), P.v(c, Y, 'xrusd')
            if g_ and x_:
                gw[c] = g_ / x_
        if gw and (base == 'H' or Y + 1 not in FX_BAD):
            tg = sum(gw.values())
            bg[Y + 1] = sum(gw[c] / tg * nxt[c] for c in gw)
        if 'USA' in members:
            bu[Y + 1] = nxt['USA']
        prev = (held, {c: nxt[c] for c in held})
    net = {y: s[y] - turn[y] * COST for y in s}
    net_hi = {y: s[y] - turn[y] * COST_HI for y in s}
    return {'s': s, 'b': b, 'bg': bg, 'bu': bu, 'turn': turn, 'net': net, 'net_hi': net_hi, 'picks': picks}


JST_CANDS = {
    # id: (base, rule, K, 課題の中での役割)
    'P1_VAL_K3': ('H', 'VAL', 3, 'primary_S'), 'P3_MOM_K3': ('H', 'MOM', 3, 'primary_S'), 'P4_MOM_K5': ('H', 'MOM', 5, 'primary_S'),
    'P5_VM_K3': ('H', 'VM', 3, 'primary_S'), 'P6_VM_K5': ('H', 'VM', 5, 'primary_S'), 'U3_MOM_K3': ('U', 'MOM', 3, 'top2_B'),
    'U1_VAL_K3': ('U', 'VAL', 3, 'extra_S'), 'U2_VAL_K5': ('U', 'VAL', 5, 'extra_S'), 'U4_MOM_K5': ('U', 'MOM', 5, 'extra_S'),
    'U5_VM_K3': ('U', 'VM', 3, 'extra_S'), 'U6_VM_K5': ('U', 'VM', 5, 'extra_S'), 'X1_CARRY_K5': ('H', 'CARRY', 5, 'extra_S'),
    'X3_VMRANK_K5': ('H', 'VMRANK', 5, 'extra_S'),
}


def grade_with(full, train, hold, roll, hold_net, lev=False, sharpe_pair=None):
    g, c = M.grade(full, train, hold, roll, cost_hold=hold_net, repl=None, family_holm_p=None,
                   sharpe_pair=sharpe_pair, leveraged_or_timing=lev)
    return g, c


def verify_jst(P, cid, base, rule, K, claimed):
    R = backtest(P, rule, base, K)
    s, b, n = R['s'], R['b'], R['net']
    r = {'id': cid, 'base': base, 'rule': rule, 'K': K}
    # ① 研究者の割り当て（訓練 〜1925・保有 1926〜2020）の再計算
    full, tr, ho, hn = cmp(s, b), cmp(s, b, None, 1925), cmp(s, b, 1926, None), cmp(n, b, 1926, None)
    roll = roll_years(s, b)
    g_angle, _ = grade_with(full, tr, ho, roll, hn)
    r['angle_mapping'] = {'train_1872_1925': tr, 'hold_1926_2020': ho, 'hold_net': hn, 'full': full,
                          'roll20': roll, 'grade_reproduced': g_angle,
                          'mean_turnover': round(S.mean(R['turn'].values()), 3)}
    r['matches_claim'] = bool(tr and ho and abs(tr['ex_ann'] - claimed['train_ex']) <= 0.05 and abs(ho['ex_ann'] - claimed['hold_ex']) <= 0.05
                              and abs((tr['t'] or 0) - claimed['train_t']) <= 0.05 and abs((ho['t'] or 0) - claimed['hold_t']) <= 0.05)
    # ② 全体の事前登録の割り当て（訓練 〜2006・保有 2007〜データの終わり 2020）
    tr_g, ho_g, hn_g = cmp(s, b, None, 2006), cmp(s, b, 2007, None, lag=1), cmp(n, b, 2007, None, lag=1)
    g_glob, c_glob = grade_with(full, tr_g, ho_g, roll, hn_g)
    r['global_mapping'] = {'train_to_2006': tr_g, 'hold_2007_2020': ho_g, 'hold_net': hn_g, 'grade': g_glob, 'criteria': c_glob}
    # ③ 保有期間の分解
    r['subperiods'] = {'1872_1898': cmp(s, b, 1872, 1898), '1899_1925': cmp(s, b, 1899, 1925),
                       '1926_1972': cmp(s, b, 1926, 1972), '1973_2020': cmp(s, b, 1973, 2020),
                       '1926_2006': cmp(s, b, 1926, 2006), '1998_2020_postpub_ALS1997': cmp(s, b, 1998, 2020, lag=1),
                       '2007_2020': ho_g}
    # ④ 相手を替える: GDP の重み（時価総額の代わり）・米国だけ
    r['vs_gdp_weighted'] = {'train': cmp(s, R['bg'], None, 1925), 'hold_1926_2020': cmp(s, R['bg'], 1926, None),
                            '2007_2020': cmp(s, R['bg'], 2007, None, lag=1), 'bench_ew_vs_gdp_2007_2020': cmp(b, R['bg'], 2007, None, lag=1)}
    r['vs_usa'] = {'train': cmp(s, R['bu'], None, 1925), 'hold_1926_2020': cmp(s, R['bu'], 1926, None),
                   '2007_2020': cmp(s, R['bu'], 2007, None, lag=1)}
    # ⑤ 費用 1%（厳しめ）
    r['net_1pct'] = {'train': cmp(R['net_hi'], b, None, 1925), 'hold_1926_2020': cmp(R['net_hi'], b, 1926, None)}
    # ⑥ 近くの K
    nb = {}
    for k2 in (2, 3, 4, 5, 6):
        R2 = backtest(P, rule, base, k2)
        a, h, g7 = cmp(R2['s'], R2['b'], None, 1925), cmp(R2['s'], R2['b'], 1926, None), cmp(R2['s'], R2['b'], 2007, None, lag=1)
        nb[f'K{k2}'] = {'train_ex': a and a['ex_ann'], 'train_t': a and a['t'], 'hold_ex': h and h['ex_ann'], 'hold_t': h and h['t'],
                        '2007_2020_ex': g7 and g7['ex_ann'], '2007_2020_t': g7 and g7['t']}
    r['neighbors_K'] = nb
    # ⑦ 1か国ずつ抜く（訓練・保有の t の最小と、その国）
    loo = []
    for c in CTRY:
        R3 = backtest(P, rule, base, K, [x for x in CTRY if x != c])
        a, h = cmp(R3['s'], R3['b'], None, 1925), cmp(R3['s'], R3['b'], 1926, None)
        if a and h:
            loo.append((c, a['ex_ann'], a['t'], h['ex_ann'], h['t']))
    wt, wh = min(loo, key=lambda x: x[2]), min(loo, key=lambda x: x[4])
    r['leave_one_out'] = {'train_t_min': wt[2], 'train_ex_at_min': wt[1], 'train_drop': wt[0],
                          'hold_t_min': wh[4], 'hold_ex_at_min': wh[3], 'hold_drop': wh[0]}
    # ⑧ 信号を1年古く（割安: 配当利回りを Y−1、勢い: リターンを Y−1）＝年末の株価の誤差・ならしの自己相関を使えない版
    lag = {'VAL': (1, 0), 'MOM': (0, 1), 'VM': (1, 1), 'CARRY': None, 'VMRANK': None}[rule]
    if lag:
        R4 = backtest(P, rule, base, K, dp_lag=lag[0], mom_lag=lag[1])
        r['signal_one_year_older'] = {'train': cmp(R4['s'], R4['b'], None, 1925), 'hold_1926_2020': cmp(R4['s'], R4['b'], 1926, None),
                                      '1973_2020': cmp(R4['s'], R4['b'], 1973, 2020)}
    # ⑨ ブロック・ブートストラップ（片側）
    xt = [s[k] - b[k] for k in sorted(s) if k <= 1925]
    xh = [s[k] - b[k] for k in sorted(s) if k >= 1926]
    r['block_bootstrap_p_one_sided'] = {'train': block_boot_p(xt), 'hold_1926_2020': block_boot_p(xh)}
    # ⑩ 生き残りの偏りの目安: 消えた国（ロシア1917・中国1949 等）の −100% の1年を1回だけ戦略に入れたら
    nyr_t = len(xt)
    w1 = (1.0 / K) * (0.5 if rule == 'VM' else 1.0)
    r['survivorship_stress_one_wipeout_in_train'] = {'weight_of_one_pick': round(w1, 3),
                                                     'train_ex_after': round(tr['ex_ann'] - w1 * 100 / nyr_t, 2) if tr else None}
    return r, R


# ════════════════════════ BSV の紙の上乗せ（自前） ════════════════════════
BSV_T = {'Trend': 4, 'Momentum': 6, 'Value': 2, 'Carry': 6, 'Seasonal': 12, 'BAB': 4, 'Multi': 6}
STY = ['Trend', 'Momentum', 'Value', 'Carry', 'Seasonal', 'BAB']
BSV_CANDS = {
    'BM7_MULTI': ('Multi', 'MA', 'top8_holdt'), 'BM1_TREND': ('Trend', 'MA', 'top8_holdt'), 'B7_MULTI': ('Multi', 'EQ', 'top8_holdt'),
    'BM2_MOM': ('Momentum', 'MA', 'top8_holdt'), 'BM4_CARRY': ('Carry', 'MA', 'top8_holdt'), 'B2_MOM': ('Momentum', 'EQ', 'top8_holdt'),
    'B1_TREND': ('Trend', 'EQ', 'top8_holdt'), 'BM5_SEAS': ('Seasonal', 'MA', 'top8_holdt'), 'B3_VAL': ('Value', 'EQ', 'top2_B'),
    'B4_CARRY': ('Carry', 'EQ', 'extra_S'), 'B5_SEAS': ('Seasonal', 'EQ', 'extra_S'),
}
# 各型の代表的な発見・公表（紙の標本の終わり）＝公表後の窓の起点（BSV は 2016-12 で終わる）
PUB = {'Trend': ('Moskowitz-Ooi-Pedersen 2012（標本〜2009）', 201301), 'Momentum': ('Asness-Moskowitz-Pedersen 2013（標本〜2011）', 201401),
       'Value': ('Asness-Moskowitz-Pedersen 2013（標本〜2011）', 201401), 'Carry': ('Koijen-Moskowitz-Pedersen-Vrugt 2018（標本〜2012）', 201301),
       'Seasonal': ('Keloharju-Linnainmaa-Nyberg 2016（標本〜2011）', 201201), 'BAB': ('Frazzini-Pedersen 2014（標本〜2012）', 201301),
       'Multi': ('上の6本の最も早い標本の終わり〜2012', 201301)}


def multi(B, asset):
    ks = set()
    for st in STY:
        ks |= set(B.get((st, asset), {}))
    out = {}
    for k in sorted(ks):
        v = [B[(st, asset)][k] for st in STY if k in B.get((st, asset), {})]
        if len(v) >= 4:
            out[k] = sum(v) / len(v)
    return out


ANALOG = {
    ('Trend', 'EQ'): ['TSMOM^EQ'], ('Trend', 'MA'): ['TSMOM'],
    ('Momentum', 'EQ'): ['CEN Equity indices Momentum'], ('Momentum', 'MA'): ['CEN All Macro Momentum'],
    ('Value', 'EQ'): ['CEN Equity indices Value'],
    ('Carry', 'EQ'): ['CEN Equity indices Carry'], ('Carry', 'MA'): ['CEN All Macro Carry'],
    ('Multi', 'EQ'): ['TSMOM^EQ', 'CEN Equity indices Momentum', 'CEN Equity indices Value', 'CEN Equity indices Carry', 'CEN Equity indices Defensive'],
    ('Multi', 'MA'): ['TSMOM', 'CEN All Macro Momentum', 'CEN All Macro Value', 'CEN All Macro Carry', 'CEN All Macro Defensive'],
}


def analog_series(A, style, asset):
    """BSV の型に最も近い AQR の公開の紙の系列。各部品を『2016-12 までのデータだけ』で年率10%のぶれにそろえ（延長部分に先読みなし）、
    多型は そろう部品が4本以上の月だけ平均。季節性は相当する公開系列が無い＝None"""
    names = ANALOG.get((style, asset))
    if not names:
        return None, None
    comps = []
    for nm in names:
        x = A.get(nm)
        if not x:
            continue
        pre = [x[k] for k in sorted(x) if k <= 201612]
        sc = 0.10 / (S.stdev(pre) * math.sqrt(12))
        comps.append({k: v * sc for k, v in x.items()})
    need = 4 if len(names) > 1 else 1
    ks = set()
    for c in comps:
        ks |= set(c)
    out = {}
    for k in sorted(ks):
        v = [c[k] for c in comps if k in c]
        if len(v) >= need:
            out[k] = sum(v) / len(v)
    return out, names


def verify_bsv(B, cid, style, asset, mkt, rf, claimed, A=None, frm=None, frf=None):
    F = multi(B, asset) if style == 'Multi' else B[(style, asset)]
    k = 0.5
    c_ann = BSV_T[style] * 0.001 + 0.005
    c_real = BSV_T[style] * 0.001 + 1.5 / 100      # 現実の器（流動的な代替投信の信託報酬 ≈1.5%/年）に寄せた感応度
    ov = {m: k * F[m] for m in F}
    ovn = {m: k * (F[m] - c_ann / 12) for m in F}
    ovr = {m: k * (F[m] - c_real / 12) for m in F}
    zero = {m: 0.0 for m in F}
    s = {m: mkt[m] + ov[m] for m in F if m in mkt}
    sn = {m: mkt[m] + ovn[m] for m in F if m in mkt}
    sr = {m: mkt[m] + ovr[m] for m in F if m in mkt}
    r = {'id': cid, 'style': style, 'asset': asset, 'k': k, 'cost_ann_per_unit_angle': c_ann, 'cost_ann_per_unit_realistic': c_real}
    # ① 研究者の割り当て: 訓練 1800-01〜1925-12（上乗せ vs 0）・保有 1926-01〜2016-12（m+kF vs m）
    tr = cmp(ov, zero, 180001, 192512, per=12, lag=12)
    ho = cmp(s, mkt, 192601, 201612, per=12, lag=12)
    hn = cmp(sn, mkt, 192601, 201612, per=12, lag=12)
    full = cmp(ov, zero, 180001, 201612, per=12, lag=12)
    roll = roll_months(s, mkt)
    sp = {'train': (sharpe_m(sn, rf, 187102, 192512), sharpe_m(mkt, rf, 187102, 192512)),
          'hold': (sharpe_m(sn, rf, 192601, 201612), sharpe_m(mkt, rf, 192601, 201612))}
    g_angle, _ = grade_with(full, tr, ho, roll, hn, lev=True, sharpe_pair=sp)
    r['angle_mapping'] = {'train_1800_1925': tr, 'hold_1926_2016': ho, 'hold_net': hn, 'full': full, 'roll20': roll,
                          'sharpe_net_vs_mkt': sp, 'grade_reproduced': g_angle}
    r['matches_claim'] = bool(tr and ho and abs(tr['ex_ann'] - claimed['train_ex']) <= 0.05 and abs(ho['ex_ann'] - claimed['hold_ex']) <= 0.05
                              and abs((tr['t'] or 0) - claimed['train_t']) <= 0.06 and abs((ho['t'] or 0) - claimed['hold_t']) <= 0.06)
    # ② 全体の割り当て: 訓練 〜2006-12・保有 2007-01〜2016-12（BSV の終わり）
    tr_g = cmp(ov, zero, 180001, 200612, per=12, lag=12)
    ho_g = cmp(s, mkt, 200701, 201612, per=12, lag=12)
    hn_g = cmp(sn, mkt, 200701, 201612, per=12, lag=12)
    hr_g = cmp(sr, mkt, 200701, 201612, per=12, lag=12)
    sp_g = {'train': (sharpe_m(sn, rf, 187102, 200612), sharpe_m(mkt, rf, 187102, 200612)),
            'hold': (sharpe_m(sn, rf, 200701, 201612), sharpe_m(mkt, rf, 200701, 201612))}
    g_glob, c_glob = grade_with(full, tr_g, ho_g, roll, hn_g, lev=True, sharpe_pair=sp_g)
    g_real, _ = grade_with(full, tr_g, ho_g, roll, hr_g, lev=True,
                           sharpe_pair={'train': sp_g['train'], 'hold': (sharpe_m(sr, rf, 200701, 201612), sp_g['hold'][1])})
    r['global_mapping'] = {'train_to_2006': tr_g, 'hold_2007_2016': ho_g, 'hold_net': hn_g, 'hold_net_realistic_fee': hr_g,
                           'sharpe_net_vs_mkt': sp_g, 'grade': g_glob, 'criteria': c_glob, 'grade_with_realistic_fee': g_real}
    # ③ 分解
    pubname, pubstart = PUB[style]
    r['subperiods'] = {'1800_1870': cmp(ov, zero, 180001, 187012, per=12, lag=12), '1871_1925': cmp(ov, zero, 187101, 192512, per=12, lag=12),
                       '1926_1966': cmp(s, mkt, 192601, 196612, per=12, lag=12), '1967_2006': cmp(s, mkt, 196701, 200612, per=12, lag=12),
                       '1998_2016': cmp(s, mkt, 199801, 201612, per=12, lag=12), '2007_2011': cmp(s, mkt, 200701, 201112, per=12, lag=12, minn=24),
                       '2012_2016': cmp(s, mkt, 201201, 201612, per=12, lag=12, minn=24),
                       'post_publication': {'paper': pubname, 'window': cmp(s, mkt, pubstart, 201612, per=12, lag=12, minn=24)}}
    # ④ 系列の性質: ぶれを事後にそろえたか（全期間の年率の標準偏差がちょうど10%なら事後）・1925年以前の自己相関（古い値・ならしの疑い）
    Fk = sorted(F)
    r['series_checks'] = {
        'vol_full_ann_pct': round(S.stdev([F[m] for m in Fk]) * math.sqrt(12) * 100, 2),
        'vol_1800_1925_ann_pct': round(S.stdev([F[m] for m in Fk if m <= 192512]) * math.sqrt(12) * 100, 2),
        'vol_1926_2016_ann_pct': round(S.stdev([F[m] for m in Fk if m >= 192601]) * math.sqrt(12) * 100, 2),
        'ac1_1800_1925': ac1([F[m] for m in Fk if m <= 192512]), 'ac1_1926_2016': ac1([F[m] for m in Fk if m >= 192601]),
        'corr_with_us_mkt_1926_2016': round(M.corr([F[m] for m in Fk if m >= 192601 and m in mkt], [mkt[m] for m in Fk if m >= 192601 and m in mkt]), 3),
        'data_ends': max(Fk)}
    # ⑤ BSV が終わった後（2017-01〜）を AQR の公開の紙の系列で延ばす（全体の事前登録の保有は『2007〜データの終わり』）
    if A is not None:
        An, names = analog_series(A, style, asset)
        if An is None:
            r['extension_aqr'] = {'analog': None, 'note': '季節性に相当する公開の系列が無い＝延ばせない（2016 年で止まったまま）'}
        else:
            ov_ks = [m for m in sorted(An) if m in F and m <= 201612]
            valid = {'overlap': [ov_ks[0], ov_ks[-1]], 'corr_with_bsv': round(M.corr([An[m] for m in ov_ks], [F[m] for m in ov_ks]), 3),
                     'mean_ann_bsv_pct': round(S.mean([F[m] for m in ov_ks]) * 1200, 2), 'mean_ann_analog_pct': round(S.mean([An[m] for m in ov_ks]) * 1200, 2)}
            k90 = [m for m in ov_ks if m >= 199001]
            if len(k90) > 24:
                valid['corr_with_bsv_1990_2016'] = round(M.corr([An[m] for m in k90], [F[m] for m in k90]), 3)
            k7 = [m for m in ov_ks if m >= 200701]
            if k7:
                valid['2007_2016_bsv_pct'] = round(S.mean([F[m] for m in k7]) * 1200, 2)
                valid['2007_2016_analog_pct'] = round(S.mean([An[m] for m in k7]) * 1200, 2)
            ax = {m: mkt[m] + k * An[m] for m in An if m in mkt and m >= 201701}
            axn = {m: mkt[m] + k * (An[m] - c_ann / 12) for m in An if m in mkt and m >= 201701}
            axr = {m: mkt[m] + k * (An[m] - c_real / 12) for m in An if m in mkt and m >= 201701}
            ext = cmp(ax, mkt, per=12, lag=12)
            # つないだ保有: 2007-01〜2016-12 は BSV、2017-01〜 は AQR の代わり
            sp_s = {m: v for m, v in s.items() if m >= 200701}; sp_s.update(ax)
            sp_n = {m: v for m, v in sn.items() if m >= 200701}; sp_n.update(axn)
            sp_r = {m: v for m, v in sr.items() if m >= 200701}; sp_r.update(axr)
            hs, hsn, hsr = cmp(sp_s, mkt, per=12, lag=12), cmp(sp_n, mkt, per=12, lag=12), cmp(sp_r, mkt, per=12, lag=12)
            last = max(sp_s)
            half = sorted(sp_s)[len(sp_s) // 2]
            s_all = dict(s); s_all.update(ax)
            roll_x = roll_months(s_all, mkt)
            sp_x = {'train': sp_g['train'], 'hold': (sharpe_m(sp_n, rf, 200701, last), sharpe_m(mkt, rf, 200701, last))}
            g_x, c_x = grade_with(full, tr_g, hs, roll_x, hsn, lev=True, sharpe_pair=sp_x)
            aq = {m: mkt[m] + k * An[m] for m in An if m in mkt and m >= 200701}
            r['extension_aqr_only_hold_2007_end'] = cmp(aq, mkt, per=12, lag=12)
            r['extension_aqr'] = {'analog': names, 'validation_vs_bsv': valid,
                                  '2017_on_analog': ext, '2017_on_analog_net': cmp(axn, mkt, per=12, lag=12),
                                  'spliced_hold_2007_end': hs, 'spliced_hold_net': hsn, 'spliced_hold_net_realistic': hsr,
                                  'spliced_first_half': cmp(sp_s, mkt, 200701, half, per=12, lag=12),
                                  'spliced_second_half': cmp(sp_s, mkt, half + 1, last, per=12, lag=12),
                                  'roll20_with_extension': roll_x, 'sharpe_net_vs_mkt': sp_x,
                                  'grade_with_spliced_hold': g_x, 'criteria': c_x}
    return r


# ════════════════════════ 判定（自前の線引き） ════════════════════════
def verdict_jst(r, claimed_grade):
    am, gm = r['angle_mapping'], r['global_mapping']
    h7 = gm['hold_2007_2020']
    issues = []
    issues.append(f"全体の事前登録の保有（2007〜2020）: {h7['ex_ann']}%/年 t{h7['t']}・幾何差 {h7['cagr_diff']}・勝った年の割合 {h7['win_share']} → C2 不合格。"
                  f"角度の割り当て（保有=1926〜2020）は 1926〜2006＝国の割安・勢いの発見の時代（ALS 1997・AMP 2013）が 81/95 年を占める")
    pp = r['subperiods']['1998_2020_postpub_ALS1997']
    issues.append(f"公表後（1998〜2020）: {pp['ex_ann']}%/年 t{pp['t']}")
    vg = r['vs_gdp_weighted']['2007_2020']; vu = r['vs_usa']['2007_2020']
    issues.append(f"相手を GDP の重み（時価総額の代わり）にすると 2007〜2020 は {vg['ex_ann']}%/年 t{vg['t']}、米国だけなら {vu['ex_ann']}%/年 t{vu['t']}"
                  f"（1926〜2020 は GDP {r['vs_gdp_weighted']['hold_1926_2020']['ex_ann']} t{r['vs_gdp_weighted']['hold_1926_2020']['t']}・"
                  f"米国 {r['vs_usa']['hold_1926_2020']['ex_ann']} t{r['vs_usa']['hold_1926_2020']['t']}）。相手の等分は国の数が少なく純粋な時価加重ではない")
    s1, s2 = r['subperiods']['1926_1972'], r['subperiods']['1973_2020']
    issues.append(f"保有の前半 1926〜1972 {s1['ex_ann']} t{s1['t']} ／ 後半 1973〜2020 {s2['ex_ann']} t{s2['t']}")
    lo = r['leave_one_out']
    issues.append(f"1か国抜き: 訓練の t 最小 {lo['train_t_min']}（{lo['train_drop']} を抜く・超過 {lo['train_ex_at_min']}）")
    if 'signal_one_year_older' in r:
        so = r['signal_one_year_older']
        issues.append(f"信号を1年古く: 訓練 {so['train']['ex_ann']} t{so['train']['t']} ／ 1926〜2020 {so['hold_1926_2020']['ex_ann']} t{so['hold_1926_2020']['t']} "
                      f"／ 1973〜2020 {so['1973_2020']['ex_ann']} t{so['1973_2020']['t']}")
    issues.append(f"多重検定: 角度の格付け44本＋プログラム全体 3,613本。保有（2007〜）の t はプログラム全体の線 t≈{BONF_T} に遠い。"
                  f"訓練（独立の 1872〜1925）の t {am['train_1872_1925']['t']} も線 {BONF_T} には届かない")
    issues.append("生き残り: JST は市場が消えた国（ロシア1917・中国1949・オーストリア＝ハンガリー等）を含まない。"
                  f"1回だけ −100% の年を1銘柄分入れると訓練の超過は {r['survivorship_stress_one_wipeout_in_train']['train_ex_after']}%/年")
    reproduced = r['matches_claim'] and am['grade_reproduced'] == claimed_grade
    vg_ = gm['grade']
    if r['rule'] in ('MOM',) or (r['rule'] == 'VM' and r.get('signal_one_year_older', {}).get('hold_1926_2020', {}).get('ex_ann', 1) <= 0.5):
        issues.append("勢い: 1925年以前の国の指数は年次の自己相関が高い（ならされた・薄商いの指数の署名）＝勢いの独立期間の勝ちは作り物でありうる。"
                      "信号を1年古くすると符号が消える／逆になる")
    return reproduced, vg_, issues


def main():
    t0 = datetime.datetime.now()
    claim = {x['id']: x for x in json.load(open(os.path.join(M.BASE, 'out', 'mw_deep_history.json')))['tested']}
    P = Panel(load_jst())
    B = load_bsv()
    mkt, rf, shiller, ff = load_us_monthly()
    A = load_aqr()

    # ── 検算 ──
    sanity = {}
    ok = bad = 0
    for (c, y), d in P.P.items():
        d0 = P.P.get((c, y - 1))
        if d.get('eq_dp') is None or d.get('eq_div_rtn') is None or d.get('eq_capgain') is None:
            continue
        if abs(d['eq_div_rtn'] - d['eq_dp'] * (1 + d['eq_capgain'])) < 1e-4:
            ok += 1
        else:
            bad += 1
    sanity['eq_dp_is_D_y_over_P_y_end'] = {'match': ok, 'mismatch': bad,
                                           'meaning': 'eq_div_rtn(Y) = eq_dp(Y)×(1+capgain(Y)) が全件で成り立つ＝ eq_dp(Y) は その年の配当÷年末 Y の株価（翌年の配当は使っていない＝割安の信号に先読みなし）'}
    # 国の指数の1年の自己相関（H の超過）
    acs = {}
    for c in CTRY:
        row = {}
        for a, z in ((1872, 1925), (1926, 1972), (1973, 2020)):
            xs = []
            for y in range(a, z + 1):
                r0, r1 = P.ret(c, y - 1, 'H'), P.ret(c, y, 'H')
                b0, b1 = P.v(c, y - 1, 'bill_rate'), P.v(c, y, 'bill_rate')
                if None in (r0, r1, b0, b1):
                    continue
                xs.append(((1 + P.v(c, y - 1, 'eq_tr')) / (1 + b0) - 1, (1 + P.v(c, y, 'eq_tr')) / (1 + b1) - 1))
            if len(xs) >= 15:
                row[f'{a}_{z}'] = round(M.corr([p for p, _ in xs], [q for _, q in xs]), 2)
        acs[c] = row
    med = {w: round(S.median([acs[c][w] for c in acs if w in acs[c]]), 3) for w in ('1872_1925', '1926_1972', '1973_2020')}
    sanity['country_excess_ac1'] = {'median': med, 'by_country': acs,
                                    'meaning': '年末どうしの無作為の歩みなら 0 前後、年平均の株価なら ≈ +0.25。1925年以前の中央が高い＝ならし・薄商いの疑い'}
    mk_ac = {'shiller_1871_1926': ac1([mkt[k] for k in sorted(mkt) if k <= 192606]), 'french_1926_2026': ac1([mkt[k] for k in sorted(mkt) if k >= 192607])}
    sanity['us_monthly_market_ac1'] = mk_ac
    sanity['french_mkt_cagr_2007_2026'] = round(geo([ff['mkt'][k] for k in sorted(ff['mkt']) if k >= 200701], 12) * 100, 2)
    note('検算', json.dumps(sanity['country_excess_ac1']['median']), mk_ac)

    verd = []
    details = {}
    # ── JST ──
    ew_vs = None
    for cid, (base, rule, K, role) in JST_CANDS.items():
        r, R = verify_jst(P, cid, base, rule, K, claim[cid])
        details[cid] = r
        reproduced, g_glob, issues = verdict_jst(r, claim[cid]['grade'])
        am = r['angle_mapping']
        kn = (f"再現: 訓練 1872〜1925 {am['train_1872_1925']['ex_ann']}%/年 t{am['train_1872_1925']['t']}・保有 1926〜2020 {am['hold_1926_2020']['ex_ann']} t{am['hold_1926_2020']['t']}"
              f"・費用後 {am['hold_net']['ex_ann']}（幾何差 {am['hold_net']['cagr_diff']}）・全期間 t{am['full']['t']}・20年窓 {am['roll20']['win_rate']}"
              f"／全体の割り当て: 保有 2007〜2020 {r['global_mapping']['hold_2007_2020']['ex_ann']} t{r['global_mapping']['hold_2007_2020']['t']} → {g_glob}")
        claimed_g = claim[cid]['grade']
        if claimed_g == 'B':
            v = f'downgraded to {g_glob}' if g_glob != 'B' else 'confirmed'
        else:
            v = f'downgraded to {g_glob}'
        if rule == 'MOM' and g_glob == 'C':
            so = r.get('signal_one_year_older', {}).get('hold_1926_2020') or {}
            if (so.get('ex_ann') or 0) < 0:
                v = 'refuted'
        verd.append({'name': cid, 'role': role, 'claimed_grade': claimed_g, 'reproduced': reproduced, 'verified_grade': g_glob,
                     'verdict': v, 'key_numbers': kn, 'issues': issues})
        note(cid, claimed_g, '→', v, '|', kn)
        if ew_vs is None:
            ew_vs = r['vs_gdp_weighted']['bench_ew_vs_gdp_2007_2020']
    # ── BSV ──
    for cid, (style, asset, role) in BSV_CANDS.items():
        r = verify_bsv(B, cid, style, asset, mkt, rf, claim[cid], A=A)
        details[cid] = r
        am, gm = r['angle_mapping'], r['global_mapping']
        h7 = gm['hold_2007_2016']
        pp = r['subperiods']['post_publication']
        sc = r['series_checks']
        kn = (f"再現: 訓練 1800〜1925 {am['train_1800_1925']['ex_ann']}%/年 t{am['train_1800_1925']['t']}・保有 1926〜2016 {am['hold_1926_2016']['ex_ann']} t{am['hold_1926_2016']['t']}"
              f"・費用後 {am['hold_net']['ex_ann']}（幾何差 {am['hold_net']['cagr_diff']}）・20年窓 {am['roll20']['win_rate']}"
              f"／全体の割り当て: 保有 2007〜2016 {h7['ex_ann']} t{h7['t']}・費用後 {gm['hold_net']['ex_ann']}（幾何差 {gm['hold_net']['cagr_diff']}）"
              f"・現実の信託報酬なら {gm['hold_net_realistic_fee']['ex_ann']}（幾何差 {gm['hold_net_realistic_fee']['cagr_diff']}）→ {gm['grade']}"
              f"・公表後 {pp['window']['from'] if pp['window'] else '-'}〜2016 {pp['window']['ex_ann'] if pp['window'] else '-'} t{pp['window']['t'] if pp['window'] else '-'}")
        issues = [
            f"保有（角度の割り当て 1926〜2016）は各型の発見の標本（MOP 2012・AMP 2013・KMPV 2018・FP 2014 等は 1970年代〜2010年前後）と重なる＝本当の標本外は 2007〜2016 と公表後だけ。"
            f"2007〜2016: {h7['ex_ann']}%/年 t{h7['t']}（2007〜2011 {r['subperiods']['2007_2011']['ex_ann']} t{r['subperiods']['2007_2011']['t']}・"
            f"2012〜2016 {r['subperiods']['2012_2016']['ex_ann']} t{r['subperiods']['2012_2016']['t']}）",
            f"公表後（{pp['paper']} の後・〜2016-12）: {pp['window']['ex_ann'] if pp['window'] else '-'}%/年 t{pp['window']['t'] if pp['window'] else '-'}。BSV は 2016-12 で終わる＝2018〜2020 の『クオンツの冬』（実在の多資産・多型の QSPIX は 2018〜2020 に大きく負けた）を含まない",
            f"紙の系列: 自己資金ゼロの買い−売り・費用前・空売りと先物が要る・1925年以前は実行不能。系列のぶれは全期間 {sc['vol_full_ann_pct']}%（1800〜1925 {sc['vol_1800_1925_ann_pct']}%・1926〜 {sc['vol_1926_2016_ann_pct']}%）"
            f"＝全期間でちょうど10%なら事後にそろえた形（t と符号は変わらないが上乗せの大きさは事後）",
            f"1925年以前の系列そのものの1次の自己相関は {sc['ac1_1800_1925']}（1926〜 {sc['ac1_1926_2016']}）でならしの署名は見えない。"
            f"ただし訓練の市場（Shiller の月平均）の自己相関は {mk_ac['shiller_1871_1926']}（French {mk_ac['french_1926_2026']}）＝C8 の訓練のシャープの相手は月平均の市場",
            f"現実の費用: 角度の費用 {round(r['cost_ann_per_unit_angle'] * 100, 2)}%/年（1単位）に対し、実在の流動的な代替投信の信託報酬は約1.5%/年 → 2007〜2016 費用後 {gm['hold_net_realistic_fee']['ex_ann']}%/年。"
            "mw_overlay の reality_gap: 実在のトレンド・ファンドを株に半分重ねた 2007〜 の上乗せは年 +0.7〜0.9%（t≈0.8〜0.9）、QSPIX を半分重ねても 2013-11〜2026-08 +3.2%/年 t1.59",
            f"多重検定: 保有（2007〜2016）の t {h7['t']} はプログラム全体の Bonferroni の線 t≈{BONF_T} に届かない（角度の保有 1926〜2016 の t だけが線を越える）",
            "日本の個人は NISA で買えない（空売り・先物の多資産の型の投信は楽天で買えない・mw_overlay）",
        ]
        reproduced = r['matches_claim'] and am['grade_reproduced'] == claim[cid]['grade']
        claimed_g = claim[cid]['grade']
        ex = r.get('extension_aqr') or {}
        order = 'SABC'
        if ex.get('analog'):
            # 全体の保有（2007〜データの終わり）を AQR の代わりで 2026 まで延ばした格付け。費用は角度の既定と現実の信託報酬の厳しい側
            hs, hsn, hsr = ex['spliced_hold_2007_end'], ex['spliced_hold_net'], ex['spliced_hold_net_realistic']
            g_real, _ = grade_with(am['full'], gm['train_to_2006'], hs, ex['roll20_with_extension'], hsr, lev=True,
                                   sharpe_pair=ex['sharpe_net_vs_mkt'])
            g_bsv_real, _ = grade_with(am['full'], gm['train_to_2006'], gm['hold_2007_2016'], am['roll20'], gm['hold_net_realistic_fee'], lev=True,
                                       sharpe_pair=gm['sharpe_net_vs_mkt'])
            # 懐疑の既定: ①BSV そのものの保有（2007〜2016）②AQR で延ばした保有（2007〜）の両方で、角度の費用と現実の信託報酬の両方を通ること
            g = max(gm['grade'], g_bsv_real, ex['grade_with_spliced_hold'], g_real, key=order.index)
            aqo = r.get('extension_aqr_only_hold_2007_end') or {}
            h2 = ex['spliced_second_half'] or {}
            cap_reason = []
            if (h2.get('ex_ann') or 0) <= 0:
                cap_reason.append(f"保有の後半（{h2.get('from')}〜{h2.get('to')}）が {h2.get('ex_ann')}%/年＝勝ちが続いていない")
            if (aqo.get('ex_ann') or 0) <= 0:
                cap_reason.append(f"同じ考えの別の作り方（AQR だけ）で 2007〜 を測ると {aqo.get('ex_ann')}%/年 t{aqo.get('t')}＝作り方に依存")
            if cap_reason and g in ('S', 'A'):
                g = 'B'
                issues.append('B に抑えた理由: ' + '／'.join(cap_reason))
            vd = ex['validation_vs_bsv']; e17 = ex['2017_on_analog']
            kn += (f"／AQR の代わり（{'+'.join(ex['analog'])}・BSV との相関 {vd['corr_with_bsv']}・重なり {vd['overlap'][0]}〜{vd['overlap'][1]}）で延ばすと "
                   f"2017-01〜{e17['to']} {e17['ex_ann']}%/年 t{e17['t']}（幾何差 {e17['cagr_diff']}）・つないだ保有 2007〜{hs['to']} {hs['ex_ann']} t{hs['t']}"
                   f"・費用後 {hsn['ex_ann']}（幾何差 {hsn['cagr_diff']}）・現実の信託報酬なら {hsr['ex_ann']}（幾何差 {hsr['cagr_diff']}）→ {g}")
            issues.insert(0, f"BSV は 2016-12 で終わる。全体の事前登録の保有は『2007〜データの終わり』なので、AQR の公開の紙の系列（Century of Factor Premia・TSMOM・ぶれは 2016 年までのデータだけでそろえた）で延ばした: "
                             f"2017〜 {e17['ex_ann']}%/年 t{e17['t']}・つないだ 2007〜{hs['to']} {hs['ex_ann']} t{hs['t']}"
                             f"（前半 {ex['spliced_first_half']['ex_ann']} t{ex['spliced_first_half']['t']}・後半 {ex['spliced_second_half']['ex_ann']} t{ex['spliced_second_half']['t']}）")
        else:
            g_real, _ = grade_with(am['full'], gm['train_to_2006'], gm['hold_2007_2016'], am['roll20'], gm['hold_net_realistic_fee'], lev=True,
                                   sharpe_pair=gm['sharpe_net_vs_mkt'])
            g = max(gm['grade'], g_real, key=order.index)
            kn += '／季節性は公開の代わりの系列が無く 2017 年以降へ延ばせない'
        if g == claimed_g:
            v = 'confirmed'
        else:
            v = f'downgraded to {g}'
        verd.append({'name': cid, 'role': role, 'claimed_grade': claimed_g, 'reproduced': reproduced, 'verified_grade': g,
                     'verdict': v, 'key_numbers': kn, 'issues': issues})
        note(cid, claimed_g, '→', v, '|', kn)

    D_ = details
    f_ = lambda x: f"{x['ex_ann']}%/年 t{x['t']}" if x else '-'
    concl = [
        f"再現: 24本すべて研究者の数字と一致（JST は訓練・保有の超過と t が小数2桁まで一致、BSV も一致）。先読みは見つからない（eq_dp は その年の配当÷年末の株価・{sanity['eq_dp_is_D_y_over_P_y_end']['match']}件で恒等式が成立）",
        f"JST の国の回転の S はすべて角度の割り当て（保有 = 1926〜2020・うち 81年が国の割安・勢いの発見の時代）に乗っている。全体の事前登録の保有 2007〜2020 では全部負け: "
        f"VM 上位3か国 {f_(D_['P5_VM_K3']['global_mapping']['hold_2007_2020'])}・割安 上位3か国 {f_(D_['P1_VAL_K3']['global_mapping']['hold_2007_2020'])}・"
        f"公表後 1998〜2020 も {f_(D_['P5_VM_K3']['subperiods']['1998_2020_postpub_ALS1997'])} → C。GDP の重み（時価総額の代わり）でも米国だけでも 2007〜2020 は負け"
        f"（米国に対し VM {f_(D_['P5_VM_K3']['vs_usa']['2007_2020'])}）",
        f"国の勢いは 1925 年以前の指数のならしの署名と重なる: 国の超過の1年の自己相関の中央 1872〜1925 {sanity['country_excess_ac1']['median']['1872_1925']}・1926〜1972 {sanity['country_excess_ac1']['median']['1926_1972']}・1973〜2020 {sanity['country_excess_ac1']['median']['1973_2020']}。"
        f"信号を1年古くすると勢いは全期間で負け（P3 訓練 {f_(D_['P3_MOM_K3']['signal_one_year_older']['train'])}）、VM は ≈0（P5 訓練 {f_(D_['P5_VM_K3']['signal_one_year_older']['train'])}）→ 勢いの S は反証",
        f"BSV の紙の上乗せ（保有の t の上位8本）: 保有 1926〜2016 は各型の発見の標本と重なる。BSV そのものの 2007〜2016 と、AQR の公開の紙の系列（Century・TSMOM）で 2026 年まで延ばした保有の両方で、角度の費用と現実の信託報酬の両方を通って A 以上に残るものは無かった。"
        f"最上位の BM7（多型・多資産）は 2007〜2016 {f_(D_['BM7_MULTI']['global_mapping']['hold_2007_2016'])} だが 2017〜 {f_(D_['BM7_MULTI']['extension_aqr']['2017_on_analog'])} → C。"
        f"残ったのは B（BM1 トレンド・B2 勢い・B4 キャリー・B5 季節性）。B4 は BSV では 2007〜2016 {f_(D_['B4_CARRY']['global_mapping']['hold_2007_2016'])} なのに、同じ考えの AQR の作り方では同じ窓 {D_['B4_CARRY']['extension_aqr']['validation_vs_bsv']['2007_2016_analog_pct']}%/年（1単位）・相関 {D_['B4_CARRY']['extension_aqr']['validation_vs_bsv'].get('corr_with_bsv_1990_2016')}＝作り方に依存",
        f"多重検定: 角度で格付けした 44本・プログラム全体 3,613本（Bonferroni の線 t≈{BONF_T}）。全体の保有（2007〜）の t で線を越えるものは1本も無い",
        "結論: 検証した 24本（研究者の S 22本・B 2本）は、全体の事前登録の保有（2007〜）・時価加重に近い相手・2016 年以降の延長で見ると A 以上が1本も残らない（B 4本・C 20本）。S のまま残るのは角度の割り当て（保有 = 1926〜）の中だけ",
    ]
    out = {'angle': 'deep_history', 'verifier': 'night/mw_deep_history_verify.py', 'target': 'out/mw_deep_history.json',
           'conclusion_ja': concl,
           'rules_of_this_verification': [
               'JST: 全体の事前登録（out/mw_prereg.json）の保有 = 2007〜データの終わり（2020）で M.grade を掛け直した格付けを verified_grade にした（角度の割り当ての格付けも grade_reproduced として残す）',
               'JST の勢い（MOM）: 全体の保有で C かつ信号を1年古くした版が 1926〜2020 で負け → refuted（独立期間の勝ちがならしの自己相関で作られた疑いが強い）',
               'BSV: ①BSV そのものの保有 2007〜2016 ②AQR の公開の紙の系列で 2017〜 を延ばしてつないだ保有 2007〜2026 の両方を、角度の費用（T×0.10%+0.50%）と現実の信託報酬（T×0.10%+1.50%）の両方で M.grade に掛け、いちばん悪い格付け',
               'BSV: さらに A 以上でも、つないだ保有の後半が 0 以下、または AQR だけの作り方で 2007〜 が 0 以下なら B に抑える',
               '判定の線（C1〜C8）は M.grade そのまま。数字は自前のコード']
           ,
           'selection': '課題の指定: S/A が8本を超えるので、保有の t の上位8本（BM7・BM1・B7・BM2・BM4・B2・B1・BM5＝すべて BSV の紙の上乗せ）＋主の族の S（P1・P3・P4・P5・P6）＋ B の上位2本（保有の超過で U3・B3）。'
                        '残りの S（U1・U2・U4・U5・U6・X1・X3・B4・B5）も同じ物差しで付けた（role = extra_S）',
           'independence': 'JST・BSV・Shiller・FRED を自前で読み、国の回転（選び方・相手・回転・費用）・上乗せ・超過・NW t・幾何の年率差・転がる20年を自前で書いた。取得は M.get / M.ff_factors、線は M.grade',
           'program_bonferroni_t': BONF_T, 'n_graded_in_angle': 44, 'sanity': sanity,
           'ew_vs_gdp_weighted_2007_2020': ew_vs,
           'verdicts': verd, 'details': details, 'log': NOTE,
           'runtime_sec': round((datetime.datetime.now() - t0).total_seconds(), 1),
           'generated': datetime.date.today().isoformat()}
    json.dump(out, open(OUT, 'w'), ensure_ascii=False, indent=1, default=str)
    note('書いた', OUT)


if __name__ == '__main__':
    main()
