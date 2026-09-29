#!/usr/bin/env python3
"""night/mw_macro_timing_verify.py — 角度 macro_timing の『反証の検証』（読むだけ・門の判定には不使用）

対象: out/mw_macro_timing.json で S と格付けされた X6a_tilt_x_G1（研究者の主張）。
やること:
  1. 生データ（Goyal-Welch 2025 xlsx・French 3因子・FRED UNRATE）から、研究者のコードを一切呼ばずに
     予測・状態判定・ポートフォリオ・超過・NW t・CAGR差・転がる20年・積立20年を自前で作り直す
     （mw_common からは取得の get と French の読み取り ff_factors だけを使う）。
  2. 反証の試み: 事後設計（保有期間を見た後の組み合わせ）・公表後・部分期間・近いパラメータ・借入の費用・
     失業率の改定（ノイズ）・1930〜1946年（NBER の失業率で当てる・格付けには使われなかった時代）・
     β/レバレッジの寄与・多重検定。
出力: out/mw_macro_timing_verify.json
"""
import io, json, math, os, random, statistics as S, sys, glob, datetime

import numpy as np

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import mw_common as M  # noqa: E402  （get と ff_factors だけ使う）

BASE = M.BASE
OUT = os.path.join(BASE, 'out', 'mw_macro_timing_verify.json')
CLAIM = os.path.join(BASE, 'out', 'mw_macro_timing.json')
FRENCH0 = 192607
TRAIN_END, HOLD0 = 200612, 200701
PREDS = ['dfy', 'dfr', 'tms', 'tbl', 'lty', 'ntis', 'infl', 'svar', 'ik', 'csp', 'dp', 'ep']
THEORY = {'dp': 1, 'ep': 1, 'dfy': 1, 'dfr': 1, 'tms': 1, 'tbl': -1, 'lty': -1, 'ntis': -1, 'infl': -1, 'svar': 1, 'ik': -1, 'csp': 1}
LOG = []


def log(*a):
    s = ' '.join(str(x) for x in a)
    LOG.append(s)
    print(s, flush=True)


def add(k, n):
    y, m = divmod(k, 100)
    t = y * 12 + m - 1 + n
    return (t // 12) * 100 + t % 12 + 1


def dec_time(m):
    """月 m のリターンを決める月末（1926-06 以前は株価が月中平均なので2か月前）"""
    return add(m, -1) if m >= FRENCH0 else add(m, -2)


# ───────────────── 生データ（自前で読む） ─────────────────
def load_goyal():
    import openpyxl
    b = M.get('https://docs.google.com/spreadsheets/d/17mw_IpaiLFDrGnrPRQ2o1ugV5nJsZuD1/export?format=xlsx',
              'goyal_predictors_2025.xlsx', max_age_days=3650)
    wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
    rows = list(wb['Monthly'].iter_rows(values_only=True))
    hdr = rows[0]
    col = {h: i for i, h in enumerate(hdr)}
    need = ['price', 'd12', 'e12', 'Rfree', 'tbl', 'lty', 'tms', 'dfy', 'dfr', 'infl', 'ntis', 'svar', 'csp', 'ret']
    g = {n: {} for n in need}
    for r in rows[1:]:
        if r[0] is None:
            continue
        k = int(r[0])
        for n in need:
            v = r[col[n]]
            if isinstance(v, (int, float)):
                g[n][k] = float(v)
    rq = list(wb['Quarterly'].iter_rows(values_only=True))
    cq = {h: i for i, h in enumerate(rq[0])}
    ikq = {}
    for r in rq[1:]:
        if r[0] is None or not isinstance(r[cq['i/k']], (int, float)):
            continue
        y, q = divmod(int(r[0]), 10)
        ikq[y * 100 + 3 * q] = float(r[cq['i/k']])
    g['ikq'] = ikq
    return g


def load_fred(sid, cache_name=None):
    b = M.get('https://fred.stlouisfed.org/graph/fredgraph.csv?id=' + sid, cache_name or f'fred_{sid}.csv', max_age_days=30).decode()
    out = {}
    for ln in b.strip().splitlines()[1:]:
        p = ln.split(',')
        if len(p) < 2 or p[1].strip() in ('', '.'):
            continue   # 欠測は欠測のまま（ルール7）
        out[int(p[0][:4]) * 100 + int(p[0][5:7])] = float(p[1])
    return out


def build_base(g):
    ff = M.ff_factors('monthly')
    mkt = {k: ff['mktrf'][k] + ff['rf'][k] for k in ff['mktrf'] if k in ff['rf']}
    rf = dict(ff['rf'])
    stock, cash = {}, {}
    pk = sorted(g['price'])
    for a, b in zip(pk, pk[1:]):
        if b >= FRENCH0:
            break
        if b in g['d12'] and b in g['Rfree'] and g['price'][a] > 0:
            stock[b] = (g['price'][b] + g['d12'][b] / 12.0) / g['price'][a] - 1.0
            cash[b] = g['Rfree'][b]
    for k in mkt:
        stock[k] = mkt[k]
        cash[k] = rf[k]
    return stock, cash, mkt, rf


def predictor_value(g, name, d, ikq_keys):
    if name == 'dp':
        if d in g['d12'] and d in g['price'] and g['d12'][d] > 0:
            return math.log(g['d12'][d] / g['price'][d])
        return None
    if name == 'ep':
        e = g['e12'].get(add(d, -3))
        if e and e > 0 and d in g['price']:
            return math.log(e / g['price'][d])
        return None
    if name == 'infl':
        return g['infl'].get(add(d, -1))
    if name == 'ik':
        ok = [q for q in ikq_keys if add(q, 3) <= d]
        return g['ikq'][ok[-1]] if ok else None
    return g[name].get(d)


# ───────────────── 予測（自前の再帰的 OLS） ─────────────────
def forecasts(months, zmap, er, min_pairs=240):
    """各 m について、m' ≤ dec_time(m) かつ z・er が揃うペアで OLS → {m: (予測, 傾き, ペアの平均)}"""
    pm_sorted = sorted(k for k in er if k in zmap)
    X = np.array([zmap[k] for k in pm_sorted]); Y = np.array([er[k] for k in pm_sorted])
    keys = np.array(pm_sorted)
    out = {}
    for m in months:
        if m not in zmap:
            continue
        n = int(np.searchsorted(keys, dec_time(m), side='right'))
        if n < min_pairs:
            continue
        x, y = X[:n], Y[:n]
        xm, ym = x.mean(), y.mean()
        vx = ((x - xm) ** 2).sum()
        if vx <= 0:
            continue
        b = ((x - xm) * (y - ym)).sum() / vx
        out[m] = (ym + b * (zmap[m] - xm), b, ym)
    return out


# ───────────────── ポートフォリオ（自前） ─────────────────
def simulate(W, stock, cash, spread=0.015, cost=0.001):
    """W: {m: 株の割合（0〜cap）}。1超は借入（現金+spread）。→ gross, net"""
    gross, net = {}, {}
    prev_m = None
    for m in sorted(k for k in W if k in stock and k in cash):
        w = W[m]
        g = cash[m] + w * (stock[m] - cash[m]) - max(0.0, w - 1.0) * spread / 12.0
        if prev_m is not None and prev_m == add(m, -1):
            drift = W[prev_m] * (1 + stock[prev_m]) / (1 + gross[prev_m]) if 1 + gross[prev_m] > 0 else W[prev_m]
            to = abs(w - drift)
        else:
            to = 0.0
        gross[m] = g
        net[m] = g - to * cost
        prev_m = m
    return gross, net


# ───────────────── 統計（自前） ─────────────────
def nwt(x, lag=12):
    n = len(x)
    if n < 24:
        return None
    mu = sum(x) / n
    e = [v - mu for v in x]
    s = sum(v * v for v in e) / n
    for L in range(1, min(lag, n - 1) + 1):
        s += 2 * (1 - L / (lag + 1)) * sum(e[i] * e[i - L] for i in range(L, n)) / n
    return mu / math.sqrt(s / n) if s > 0 else None


def geo(xs):
    return math.exp(sum(math.log1p(v) for v in xs) * 12 / len(xs)) - 1


def ex(s, b, a=None, z=None, drop=None):
    ks = [k for k in sorted(set(s) & set(b)) if (a is None or k >= a) and (z is None or k <= z) and not (drop and drop(k))]
    if len(ks) < 24:
        return None
    d = [s[k] - b[k] for k in ks]
    t = nwt(d)
    return {'from': ks[0], 'to': ks[-1], 'n': len(ks), 'ex_ann': round(sum(d) / len(d) * 1200, 2), 't': round(t, 2) if t is not None else None,
            'cagr_diff': round((geo([s[k] for k in ks]) - geo([b[k] for k in ks])) * 100, 2)}


def sharpe(s, rf, a=None, z=None, drop=None):
    ks = [k for k in sorted(set(s) & set(rf)) if (a is None or k >= a) and (z is None or k <= z) and not (drop and drop(k))]
    x = [s[k] - rf[k] for k in ks]
    return round(S.mean(x) / S.stdev(x) * math.sqrt(12), 3) if len(x) > 24 and S.stdev(x) > 0 else None


def jk(s, b, rf, a=None, z=None):
    ks = [k for k in sorted(set(s) & set(b) & set(rf)) if (a is None or k >= a) and (z is None or k <= z)]
    xs = np.array([s[k] - rf[k] for k in ks]); xb = np.array([b[k] - rf[k] for k in ks])
    T = len(ks)
    sa, sb = xs.mean() / xs.std(ddof=1), xb.mean() / xb.std(ddof=1)
    rho = float(np.corrcoef(xs, xb)[0, 1])
    V = (2 - 2 * rho + 0.5 * (sa ** 2 + sb ** 2 - 2 * sa * sb * rho ** 2)) / T
    z_ = (sa - sb) / math.sqrt(V)
    return {'z': round(z_, 2), 'p_one': round(0.5 * math.erfc(z_ / math.sqrt(2)), 4)}


def roll20(s, b):
    ks = sorted(set(s) & set(b))
    res = []
    for y in range(ks[0] // 100, 2100):
        a, z = y * 100 + 7, (y + 20) * 100 + 6
        if z > ks[-1]:
            break
        w = [k for k in ks if a <= k <= z]
        if len(w) < 240 * 0.97:
            continue
        res.append((y, round((math.exp(sum(math.log1p(s[k]) for k in w) / 20) - math.exp(sum(math.log1p(b[k]) for k in w) / 20)) * 100, 2)))
    if not res:
        return None
    v = sorted(x for _, x in res)
    return {'windows': len(res), 'win_rate': round(sum(1 for x in v if x > 0) / len(v), 3), 'median': v[len(v) // 2],
            'worst': min(res, key=lambda t: t[1])}


def dca20(s, b, step=12):
    ks = sorted(set(s) & set(b))
    res = []
    for i in range(0, len(ks) - 240 + 1, step):
        ws = wb = 0.0
        for k in ks[i:i + 240]:
            ws = (ws + 1) * (1 + s[k]); wb = (wb + 1) * (1 + b[k])
        res.append((ks[i], ws / wb))
    v = sorted(x for _, x in res)
    return {'windows': len(res), 'win_rate': round(sum(1 for x in v if x > 1) / len(v), 3), 'median_ratio': round(v[len(v) // 2], 3),
            'worst': [min(res, key=lambda t: t[1])[0], round(min(v), 3)]}


def capm(s, b, rf, a=None, z=None):
    ks = [k for k in sorted(set(s) & set(b) & set(rf)) if (a is None or k >= a) and (z is None or k <= z)]
    y = np.array([s[k] - rf[k] for k in ks]); x = np.array([b[k] - rf[k] for k in ks])
    X = np.column_stack([np.ones_like(x), x])
    beta = np.linalg.lstsq(X, y, rcond=None)[0]
    res = y - X @ beta
    # NW の分散（Bartlett 12）
    T = len(ks)
    u = X * res[:, None]
    Sm = u.T @ u / T
    for L in range(1, 13):
        G = u[L:].T @ u[:-L] / T
        Sm += (1 - L / 13) * (G + G.T)
    Q = np.linalg.inv(X.T @ X / T)
    V = Q @ Sm @ Q / T
    return {'alpha_ann': round(beta[0] * 1200, 2), 'alpha_t': round(beta[0] / math.sqrt(V[0, 0]), 2), 'beta': round(beta[1], 3)}


def criteria(full, train, hold, net_hold, r20, sh_tr, sh_ho, holm_p=None, c5=None):
    c = {'C1_train': bool(train and train['ex_ann'] > 0 and (train['t'] or 0) >= 2.0),
         'C2_hold_sign': bool(hold and hold['ex_ann'] > 0 and hold['cagr_diff'] > 0),
         'C3_hold_t': bool(hold and (hold['t'] or 0) >= 1.65),
         'C4_roll20': bool(r20 and r20['win_rate'] >= 0.8),
         'C5_repl': c5,
         'C6_net_cost': bool(net_hold and net_hold['ex_ann'] > 0 and net_hold['cagr_diff'] > 0),
         'C7_multi': bool((full and (full['t'] or 0) >= 3.0) or (holm_p is not None and holm_p < 0.05)),
         'C8_sharpe': bool(sh_tr and sh_ho and sh_tr[0] > sh_tr[1] and sh_ho[0] > sh_ho[1])}
    base = c['C1_train'] and c['C2_hold_sign'] and c['C6_net_cost'] and c['C8_sharpe']
    if base and c['C3_hold_t'] and c['C4_roll20'] and c['C7_multi'] and c['C5_repl'] in (None, True):
        gr = 'S'
    elif base and c['C4_roll20'] and c['C7_multi'] and (c['C3_hold_t'] or c['C5_repl'] is True):
        gr = 'A'
    elif base:
        gr = 'B'
    else:
        gr = 'C'
    return gr, c



# ───────────────── 1日遅れの執行（日次の French で） ─────────────────
def daily_lag_test(W, spread=0.015, cost=0.001):
    """月末 d の終値で決めた割合を、(a) その終値で執行（研究者の前提）と (b) 翌営業日の終値で執行（1日遅れ）で比べる。
    日次で割合を保つ近似（両者とも同じ近似なので差が1日遅れの効果）。→ {'lag0': 月次 net, 'lag1': 月次 net}"""
    ffd = M.ff_factors('daily')
    days = sorted(k for k in ffd['mktrf'] if k in ffd['rf'])
    bym = {}
    for dd in days:
        bym.setdefault(dd // 100, []).append(dd)
    out = {'lag0': {}, 'lag1': {}}
    ms = sorted(m for m in W if m in bym and m >= FRENCH0)
    for tag in out:
        prev_w = None
        for m in ms:
            w = W[m]
            wp = W.get(add(m, -1), w) if tag == 'lag1' else w
            lv = 1.0
            for i, dd in enumerate(bym[m]):
                ww = wp if (tag == 'lag1' and i == 0) else w
                r_ = ffd['rf'][dd] + ww * ffd['mktrf'][dd] - max(0.0, ww - 1) * spread / 252
                lv *= 1 + r_
            to = abs(w - (prev_w if prev_w is not None else w))
            out[tag][m] = lv - 1 - to * cost
            prev_w = w
    mk = {}
    for m in ms:
        lv = 1.0
        for dd in bym[m]:
            lv *= 1 + ffd['mktrf'][dd] + ffd['rf'][dd]
        mk[m] = lv - 1
    return out, mk


# ───────────────── 米国外（独立に作り直す） ─────────────────
COUNTRY_FILES = {'UK.Dat': 'GB', 'Austria.Dat': 'AT', 'Austrlia.Dat': 'AU', 'Belgium.Dat': 'BE', 'Canada.Dat': 'CA', 'Denmark.Dat': 'DK',
                 'Finland.Dat': 'FI', 'France.Dat': 'FR', 'Germany.Dat': 'DE', 'Ireland.Dat': 'IE', 'Italy.Dat': 'IT', 'Japan.Dat': 'JP',
                 'Nethrlnd.Dat': 'NL', 'Norway.Dat': 'NO', 'Spain.Dat': 'ES', 'Sweden.Dat': 'SE'}
ISO3 = {'GB': 'GBR', 'AT': 'AUT', 'AU': 'AUS', 'BE': 'BEL', 'CA': 'CAN', 'DK': 'DNK', 'FI': 'FIN', 'FR': 'FRA', 'DE': 'DEU', 'IE': 'IRL',
        'IT': 'ITA', 'JP': 'JPN', 'NL': 'NLD', 'NO': 'NOR', 'ES': 'ESP', 'SE': 'SWE'}


def read_country(txt):
    """自前の読み取り: 『Value-Weight Local Returns / Not Reqd』の月次 Mkt 列と、『Average of Annual … Not Reqd』の年次 E/P・Yld"""
    lines = txt.splitlines()
    loc, ann = {}, {}
    mode = None
    for ln in lines:
        st = ln.strip()
        if not st:
            continue
        if 'Local' in st and 'Returns' in st and 'Not Reqd' in st:
            mode = 'loc'; continue
        if ('Returns' in st and ('Dollar' in st or 'Required' in st)):
            mode = None if not ('Local' in st and 'Not Reqd' in st) else 'loc'
            if 'Required' in st and 'Not' not in st:
                mode = None
            continue
        if st.startswith('Average of Annual'):
            mode = 'ann_wait'; continue
        if mode == 'ann_wait' and 'Not Reqd' in st:
            mode = 'ann'; continue
        if mode == 'ann_wait' and 'Required' in st:
            mode = None; continue
        tok = st.split()
        if not tok[0].isdigit():
            continue
        if mode == 'loc' and len(tok[0]) == 6:
            v = float(tok[1])
            if v > -99:
                loc[int(tok[0])] = v / 100
        elif mode == 'ann' and len(tok[0]) == 4:
            ep, yld = float(tok[3]), float(tok[5])
            ann[int(tok[0])] = {'EP': ep / 100 if ep > -99 else None, 'Yld': yld / 100 if yld > -99 else None}
    return loc, ann


def country_test(loc, ann, cc, convention):
    """convention='researcher': 年 Y の比率を Y 年7月から、Y−1年12月の価格の値として使う（研究者と同じ）。
       convention='strict': 年 Y の比率を Y+1 年7月から、Y年12月の価格の値として使う（1年遅らせた保守側）"""
    r3 = load_fred(f'IR3TIB01{cc}M156N')
    try:
        rc = load_fred(f'IRSTCI01{cc}M156N')
    except RuntimeError:
        rc = {}
    rate = {**rc, **r3}
    try:
        y10 = load_fred(f'IRLTLT01{cc}M156N')
    except RuntimeError:
        y10 = {}
    try:
        cpi = load_fred(f'{ISO3[cc]}CPIALLMINMEI')
    except RuntimeError:
        cpi = {}
    un = load_fred(f'LRHUTTTT{cc}M156S')
    ks = sorted(loc)
    cash = {k: rate[add(k, -1)] / 1200 for k in ks if add(k, -1) in rate}
    # 価格指数の近似（配当利回りを月割りで引く）
    PI, lv = {add(ks[0], -1): 1.0}, 1.0
    tri, lt = {}, 1.0
    for k in ks:
        y = k // 100
        a = ann.get(y - 1) or ann.get(y)
        dy = a['Yld'] if a and a['Yld'] else 0.0
        lv *= (1 + loc[k]) / (1 + dy / 12)
        PI[k] = lv
        lt *= 1 + loc[k]
        tri[k] = lt

    def ratio_year(d):
        y, mo = divmod(d, 100)
        Y = y if mo >= 7 else y - 1          # 研究者: 7月に年 Y が使える
        if convention == 'strict':
            Y -= 1                           # 保守: さらに1年待つ
        base = (Y - 1) * 100 + 12 if convention == 'researcher' else Y * 100 + 12
        return Y, base

    def z(name, d):
        if name in ('dp', 'ep'):
            Y, base = ratio_year(d)
            a = ann.get(Y)
            if not a or base not in PI or d not in PI:
                return None
            v = a['Yld'] if name == 'dp' else a['EP']
            return math.log(v * PI[base] / PI[d]) if v and v > 0 else None
        if name == 'tbl':
            return rate[d] / 100 if d in rate else None
        if name == 'lty':
            return y10[d] / 100 if d in y10 else None
        if name == 'tms':
            return (y10[d] - rate[d]) / 100 if d in y10 and d in rate else None
        if name == 'infl':
            a1, b1 = cpi.get(add(d, -1)), cpi.get(add(d, -2))
            return a1 / b1 - 1 if a1 and b1 else None
        return None

    er = {k: loc[k] - cash[k] for k in ks if k in cash}
    preds = ['dp', 'ep', 'tbl', 'lty', 'tms', 'infl']
    fcs = {}
    for p in preds:
        zm = {m: v for m in ks if (v := z(p, add(m, -1))) is not None}
        # 自前の再帰的 OLS（判断は d = m−1）
        keys = sorted(k for k in er if k in zm)
        X = np.array([zm[k] for k in keys]); Y_ = np.array([er[k] for k in keys]); KA = np.array(keys)
        fc = {}
        for m in ks:
            if m not in zm:
                continue
            n = int(np.searchsorted(KA, add(m, -1), side='right'))
            if n < 240:
                continue
            x, y = X[:n], Y_[:n]
            vx = ((x - x.mean()) ** 2).sum()
            if vx <= 0:
                continue
            b = ((x - x.mean()) * (y - y.mean())).sum() / vx
            f = y.mean() + b * (zm[m] - x.mean())
            fc[m] = max(0.0, f if b * THEORY[p] > 0 else y.mean())
        fcs[p] = fc
    erk = sorted(er)
    KE = np.array(erk); CS = np.cumsum([er[k] for k in erk])
    W6, Wt, Wg = {}, {}, {}
    for m in ks:
        d = add(m, -1)
        n = int(np.searchsorted(KE, d, side='right'))
        v = [fcs[p][m] for p in preds if m in fcs[p]]
        tw = [add(d, -i) for i in range(10)]
        uw = [add(d, -2 - i) for i in range(12)]
        if not all(x in tri for x in tw) or not all(x in un for x in uw):
            continue
        bad = tri[d] < sum(tri[x] for x in tw) / 10 and round(un[uw[0]] * 12 - sum(un[x] for x in uw), 9) > 0
        Wg[m] = 0.0 if bad else 1.0
        if n >= 240 and v and CS[n - 1] / n > 0:
            tl = min(1.5, max(0.0, (sum(v) / len(v)) / (CS[n - 1] / n)))
            Wt[m] = tl
            W6[m] = 0.0 if bad else tl
    res = {}
    for tag, W in (('X6a', W6), ('tilt_only', Wt), ('gate_only', Wg)):
        g_, n_ = simulate(W, loc, cash)
        if len(n_) < 120:
            res[tag] = f'N/A（{len(n_)}か月）'
            continue
        b_ = {k: loc[k] for k in n_}
        e = ex(n_, b_)
        res[tag] = {'window': [min(n_), max(n_)], 'net': e, 'positive': bool(e['ex_ann'] > 0 and e['cagr_diff'] > 0),
                    'sharpe_net_vs_mkt': (sharpe(n_, cash), sharpe(b_, cash)), 'w_avg': round(S.mean(W[k] for k in n_), 3)}
    return res

# ───────────────── 本体 ─────────────────
def main():
    g = load_goyal()
    ikq_keys = sorted(g['ikq'])
    stock, cash, mkt, rf = build_base(g)
    er = {k: stock[k] - cash[k] for k in stock if k in cash}
    last_pred = max(g['price'])
    months = [m for m in sorted(stock) if m >= 187103 and dec_time(m) <= last_pred]
    un = load_fred('UNRATE')
    log('データ', 'stock', min(stock), max(stock), 'goyal最終', last_pred, 'UNRATE', min(un), max(un),
        '欠測(2024〜)', [k for k in range(202401, 202609) if k % 100 <= 12 and k % 100 >= 1 and k not in un])

    # 予測 12変数
    Z, FC = {}, {}
    for p in PREDS:
        Z[p] = {m: v for m in months if (v := predictor_value(g, p, dec_time(m), ikq_keys)) is not None}
        FC[p] = forecasts(months, Z[p], er)
    # 過去の平均（全 er・240 以上）
    erk = sorted(er)
    cs = np.cumsum([er[k] for k in erk])
    PMv = {}
    for m in months:
        n = int(np.searchsorted(np.array(erk), dec_time(m), side='right'))
        if n >= 240:
            PMv[m] = cs[n - 1] / n
    f1 = {}
    for m in months:
        v = []
        for p in PREDS:
            if m in FC[p]:
                fc, b, pm = FC[p][m]
                v.append(max(0.0, fc if b * THEORY[p] > 0 else pm))
        if v:
            f1[m] = sum(v) / len(v)

    # 総リターン指数（トレンド用）
    tri, lv = {}, 1.0
    for k in sorted(stock):
        lv *= 1 + stock[k]
        tri[k] = lv

    def build_W(unr, sma=10, uma=12, ulag=1, cap=1.5, exec_lag=0, tilt_on=True, gate_on=True, tilt_const=None):
        W = {}
        for m in months:
            d = add(dec_time(m), -exec_lag)
            if m not in f1 or m not in PMv or PMv[m] <= 0:
                continue
            tw = [add(d, -i) for i in range(sma)]
            uw = [add(d, -ulag - i) for i in range(uma)]
            if not all(x in tri for x in tw) or not all(x in unr for x in uw):
                continue
            down = tri[d] < sum(tri[x] for x in tw) / sma
            # 失業率は小数1桁の値なので、平均との同点は厳密に扱う（浮動小数の足し算の誤差で『上向き』にしない）
            rising = round(unr[uw[0]] * uma - sum(unr[x] for x in uw), 9) > 0
            f_use = f1.get(add(m, -exec_lag)) if exec_lag else f1[m]
            pm_use = PMv.get(add(m, -exec_lag)) if exec_lag else PMv[m]
            if f_use is None or pm_use is None:
                continue
            tl = tilt_const if tilt_const is not None else (min(cap, max(0.0, f_use / pm_use)) if tilt_on else 1.0)
            W[m] = 0.0 if (gate_on and down and rising) else tl
        return W

    def evaluate(W, spread=0.015, cost=0.001, full_detail=True):
        gr, nt = simulate(W, stock, cash, spread, cost)
        gr = {k: v for k, v in gr.items() if k >= FRENCH0}
        nt = {k: v for k, v in nt.items() if k >= FRENCH0}
        e = {'window': [min(gr), max(gr)], 'full': ex(gr, mkt), 'train': ex(gr, mkt, z=TRAIN_END), 'hold': ex(gr, mkt, a=HOLD0),
             'net_hold': ex(nt, mkt, a=HOLD0),
             'sharpe_train': (sharpe(nt, rf, z=TRAIN_END), sharpe({k: mkt[k] for k in nt}, rf, z=TRAIN_END)),
             'sharpe_hold': (sharpe(nt, rf, a=HOLD0), sharpe({k: mkt[k] for k in nt}, rf, a=HOLD0))}
        if full_detail:
            e['roll20'] = roll20(nt, mkt)
            e['dca20'] = dca20(nt, mkt)
            wk = [W[k] for k in gr]
            e['w_avg_full'] = round(S.mean(wk), 3)
            e['w_avg_hold'] = round(S.mean(W[k] for k in gr if k >= HOLD0), 3)
            e['pct_zero'] = round(sum(1 for x in wk if x == 0) / len(wk), 3)
            e['pct_levered'] = round(sum(1 for x in wk if x > 1) / len(wk), 3)
            gr_, c = criteria(e['full'], e['train'], e['hold'], e['net_hold'], e['roll20'], e['sharpe_train'], e['sharpe_hold'])
            e['grade_mechanical'], e['criteria'] = gr_, c
        return e, gr, nt

    res = {}
    # ── 1. 作り直し
    W0 = build_W(un)
    base, G0, N0 = evaluate(W0)
    res['reproduction'] = base
    claim = next(e for e in json.load(open(CLAIM))['tested'] if e['name'] == 'X6a_tilt_x_G1')
    res['reproduction_vs_claim'] = {
        'full_ex': (base['full']['ex_ann'], claim['full']['ex_ann']), 'full_t': (base['full']['t'], claim['full']['t']),
        'train_ex': (base['train']['ex_ann'], claim['train']['ex_ann']), 'train_t': (base['train']['t'], claim['train']['t']),
        'hold_ex': (base['hold']['ex_ann'], claim['hold']['ex_ann']), 'hold_t': (base['hold']['t'], claim['hold']['t']),
        'hold_cagr_diff': (base['hold']['cagr_diff'], claim['hold']['cagr_diff']),
        'net_hold_ex': (base['net_hold']['ex_ann'], claim['cost_hold']['ex_ann']),
        'roll20_win': (base['roll20']['win_rate'], claim['roll20']['win_rate']),
        'dca20_median': (base['dca20']['median_ratio'], claim['dca20']['median_ratio']),
        'sharpe_train': (base['sharpe_train'], claim['sharpe']['train']), 'sharpe_hold': (base['sharpe_hold'], claim['sharpe']['hold'])}
    log('再現', json.dumps(res['reproduction_vs_claim'], ensure_ascii=False))
    diffW = [m for m in W0 if m >= 194902]
    res['weights_months'] = len(diffW)

    # ── 2. 部品と β の寄与
    comp = {}
    for nm, W in (('tilt_only_X1a', build_W(un, gate_on=False)), ('gate_only_G1', build_W(un, tilt_on=False)),
                  ('X6a_unlevered_cap1.0', build_W(un, cap=1.0)),
                  ('gate_x_const_w_equal_to_X6a_avg', build_W(un, tilt_const=round(base['w_avg_full'] / (1 - base['pct_zero']), 4)))):
        e, _, _ = evaluate(W)
        comp[nm] = {k: e[k] for k in ('full', 'train', 'hold', 'net_hold', 'sharpe_train', 'sharpe_hold', 'w_avg_full', 'w_avg_hold', 'grade_mechanical')}
    comp['_note'] = 'gate_x_const: 降りない月の株の割合を X6a の平均（降りない月だけ）に固定＝予測の中身を捨てて β だけ残した版'
    res['components'] = comp
    res['capm'] = {'train': capm(N0, mkt, rf, z=TRAIN_END), 'hold': capm(N0, mkt, rf, a=HOLD0), 'post2017': capm(N0, mkt, rf, a=201701)}
    # 保有期間の超過の分解: 降りた月／1倍以下で持った月／借入の月
    hk = [k for k in G0 if k >= HOLD0]
    parts = {'cash_months': 0.0, 'unlev_months': 0.0, 'levered_months': 0.0}
    cnt = {'cash_months': 0, 'unlev_months': 0, 'levered_months': 0}
    for k in hk:
        tag = 'cash_months' if W0[k] == 0 else ('levered_months' if W0[k] > 1 else 'unlev_months')
        parts[tag] += (G0[k] - mkt[k]); cnt[tag] += 1
    Th = len(hk) / 12
    res['hold_decomposition_pct_per_year'] = {t: {'months': cnt[t], 'contrib': round(parts[t] / Th * 100, 2)} for t in parts}
    c08 = sum(G0[k] - mkt[k] for k in hk if 200801 <= k <= 200912)
    res['hold_decomposition_pct_per_year']['of_which_2008_09'] = round(c08 / Th * 100, 2)
    log('分解', res['hold_decomposition_pct_per_year'])

    # ── 3. 部分期間
    drop = {'ex_1998_2000': lambda k: 199801 <= k <= 200012, 'ex_2020_2021': lambda k: 202001 <= k <= 202112,
            'ex_2008_2009': lambda k: 200801 <= k <= 200912,
            'ex_all_three': lambda k: 199801 <= k <= 200012 or 202001 <= k <= 202112 or 200801 <= k <= 200912}
    sub = {}
    for nm, fn in drop.items():
        sub[nm] = {'full': ex(G0, mkt, drop=fn), 'train': ex(G0, mkt, z=TRAIN_END, drop=fn), 'hold': ex(G0, mkt, a=HOLD0, drop=fn),
                   'net_hold': ex(N0, mkt, a=HOLD0, drop=fn)}
    sub['hold_first_half_2007_01_2016_05'] = {'gross': ex(G0, mkt, a=HOLD0, z=201605), 'net': ex(N0, mkt, a=HOLD0, z=201605),
                                              'sharpe_net_vs_mkt': (sharpe(N0, rf, a=HOLD0, z=201605), sharpe(mkt, rf, a=HOLD0, z=201605))}
    sub['hold_second_half_2016_06_2025_11'] = {'gross': ex(G0, mkt, a=201606), 'net': ex(N0, mkt, a=201606),
                                               'sharpe_net_vs_mkt': (sharpe(N0, rf, a=201606), sharpe({k: mkt[k] for k in N0}, rf, a=201606))}
    sub['post_growth_trend_pub_2017'] = {'net': ex(N0, mkt, a=201701), 'sharpe_net_vs_mkt': (sharpe(N0, rf, a=201701), sharpe({k: mkt[k] for k in N0}, rf, a=201701))}
    sub['post_RSZ_2011'] = {'net': ex(N0, mkt, a=201101)}
    sub['hold_ex_2008_net'] = ex(N0, mkt, a=HOLD0, drop=lambda k: 200801 <= k <= 200812)
    dec = {}
    for a in range(1950, 2030, 10):
        ks = [k for k in N0 if a * 100 <= k < (a + 10) * 100]
        if len(ks) >= 24:
            dec[str(a)] = round((geo([N0[k] for k in ks]) - geo([mkt[k] for k in ks])) * 100, 2)
    sub['by_decade_net_cagr_diff'] = dec
    yearly = {}
    for y in range(1949, 2026):
        ks = [k for k in N0 if k // 100 == y]
        if ks:
            yearly[y] = round((math.prod(1 + N0[k] for k in ks) - math.prod(1 + mkt[k] for k in ks)) * 100, 1)
    top = sorted(yearly.items(), key=lambda t: -t[1])
    sub['top5_years_net_minus_mkt'] = top[:5]
    sub['bottom5_years'] = top[-5:]
    # 上位3年を除く全期間
    top3 = {y for y, _ in top[:3]}
    sub['ex_top3_years'] = {'full': ex(G0, mkt, drop=lambda k: k // 100 in top3), 'train': ex(G0, mkt, z=TRAIN_END, drop=lambda k: k // 100 in top3)}
    res['subperiods'] = sub
    log('部分期間', {k: (v.get('hold') or v.get('net') or {}) for k, v in sub.items() if isinstance(v, dict)})

    # ── 4. 近いパラメータ（格子）
    grid = []
    for sma in (6, 8, 9, 10, 11, 12, 14):
        for uma in (6, 9, 12, 18, 24):
            for cap in (1.0, 1.25, 1.5, 2.0):
                for ulag in (1, 2):
                    e, _, _ = evaluate(build_W(un, sma=sma, uma=uma, cap=cap, ulag=ulag))
                    grid.append({'sma': sma, 'uma': uma, 'cap': cap, 'ulag': ulag, 'grade': e['grade_mechanical'],
                                 'train_ex': e['train']['ex_ann'], 'train_t': e['train']['t'], 'hold_ex': e['hold']['ex_ann'],
                                 'hold_t': e['hold']['t'], 'full_t': e['full']['t'], 'net_hold_ex': e['net_hold']['ex_ann'],
                                 'sh_tr': e['sharpe_train'], 'sh_ho': e['sharpe_hold'], 'roll20': e['roll20']['win_rate']})
    gs = {}
    for x in grid:
        gs[x['grade']] = gs.get(x['grade'], 0) + 1
    near = [x for x in grid if x['sma'] in (8, 9, 10, 11, 12) and x['uma'] in (9, 12, 18) and x['cap'] == 1.5 and x['ulag'] == 1]
    ns = {}
    for x in near:
        ns[x['grade']] = ns.get(x['grade'], 0) + 1
    by_cap = {}
    for cp in (1.0, 1.25, 1.5, 2.0):
        xs = [x for x in grid if x['cap'] == cp]
        by_cap[str(cp)] = {'n': len(xs), 'grades': {gg: sum(1 for x in xs if x['grade'] == gg) for gg in 'SABC'},
                           'train_t_median': S.median(x['train_t'] for x in xs), 'hold_ex_median': S.median(x['hold_ex'] for x in xs)}
    res['parameter_grid'] = {'n': len(grid), 'grades': gs, 'near_neighbours_cap1.5_lag1': {'n': len(near), 'grades': ns,
                             'rows': near}, 'by_cap': by_cap,
                             'train_t_range': [min(x['train_t'] for x in grid), max(x['train_t'] for x in grid)],
                             'hold_t_range': [min(x['hold_t'] for x in grid), max(x['hold_t'] for x in grid)],
                             'unlevered_cap1_rows': [x for x in grid if x['cap'] == 1.0 and x['ulag'] == 1 and x['uma'] == 12]}
    log('格子', gs, '近傍', ns, 'cap別', by_cap)

    # ── 5. 借入・費用・実行の遅れ
    fr = {}
    for nm, kw in (('spread3pct', {'spread': 0.03}), ('spread5pct_retail_margin', {'spread': 0.05}), ('cost0.3pct', {'cost': 0.003}),
                   ('spread5_cost0.3', {'spread': 0.05, 'cost': 0.003})):
        e, _, _ = evaluate(W0, **kw)
        fr[nm] = {k: e[k] for k in ('full', 'train', 'hold', 'net_hold', 'sharpe_train', 'sharpe_hold', 'grade_mechanical')}
    e, _, _ = evaluate(build_W(un, exec_lag=1))
    fr['exec_lag_1month'] = {k: e[k] for k in ('full', 'train', 'hold', 'net_hold', 'sharpe_train', 'sharpe_hold', 'grade_mechanical')}
    res['frictions'] = fr
    log('摩擦', {k: (v['train']['t'], v['hold']['ex_ann'], v['grade_mechanical']) for k, v in fr.items()})

    # ── 6. 失業率の改定をノイズで近似（季節調整の改定は ±0.1pt 程度）
    rng = random.Random(20260928)
    mc = []
    for sd in (0.05, 0.1):
        rows = []
        for _ in range(200):
            un2 = {k: v + rng.gauss(0, sd) for k, v in un.items()}
            e, _, _ = evaluate(build_W(un2), full_detail=False)
            rows.append((e['train']['t'], e['hold']['ex_ann'], e['hold']['t'], e['sharpe_train'][0] > e['sharpe_train'][1] and e['sharpe_hold'][0] > e['sharpe_hold'][1]))
        tt = sorted(r[0] for r in rows); he = sorted(r[1] for r in rows); ht = sorted(r[2] for r in rows)
        mc.append({'noise_sd_pt': sd, 'draws': 200, 'train_t_p5_p50_p95': [tt[10], tt[100], tt[190]],
                   'share_train_t_ge2': round(sum(1 for x in tt if x >= 2) / 200, 3),
                   'hold_ex_p5_p50_p95': [he[10], he[100], he[190]], 'hold_t_p5_p50_p95': [ht[10], ht[100], ht[190]],
                   'share_hold_t_ge1.65': round(sum(1 for x in ht if x >= 1.65) / 200, 3),
                   'share_C8_ok': round(sum(1 for r in rows if r[3]) / 200, 3)})
    res['unrate_revision_noise'] = mc
    log('改定ノイズ', mc)

    # ── 7. 1930〜1946年（NBER の失業率・格付けに使われなかった時代）
    uA = load_fred('M0892AUSM156SNBR')   # 1929-04〜1942-06（季節調整）
    uB = load_fred('M0892BUSM156SNBR')   # 1940-01〜1946-12（季節調整・1940-02 欠測）
    pre = {}
    for nm, splice in (('A_to_1939_B_from_1940_skip_gap', lambda: {**{k: v for k, v in uA.items() if k <= 193912}, **uB}),
                       ('A_to_1939_B_from_1940_A_fills_194002', lambda: {**{k: v for k, v in uA.items() if k <= 193912}, **uB, 194002: uA[194002]})):
        u = splice()
        W = build_W(u)
        Wg = build_W(u, tilt_on=False)
        Wt = build_W(u, gate_on=False)
        out = {}
        for tag, WW in (('X6a', W), ('gate_only', Wg), ('tilt_only', {k: v for k, v in Wt.items() if k in W})):
            grr, ntt = simulate(WW, stock, cash)
            ntt = {k: v for k, v in ntt.items() if FRENCH0 <= k <= 194702}
            if len(ntt) < 60:
                out[tag] = 'N/A'
                continue
            out[tag] = {'net': ex(ntt, mkt), 'sharpe_net_vs_mkt': (sharpe(ntt, rf), sharpe({k: mkt[k] for k in ntt}, rf)),
                        'w_avg': round(S.mean(WW[k] for k in ntt), 3), 'pct_zero': round(sum(1 for k in ntt if WW[k] == 0) / len(ntt), 3),
                        'ex_1930_1932': ex(ntt, mkt, z=193212) if sum(1 for k in ntt if k <= 193212) >= 24 else None,
                        'ex_1933_1946': ex(ntt, mkt, a=193301)}
        pre[nm] = out
    res['pre_1947_nber_unemployment'] = pre
    log('1930〜1946', {k: {t: (v[t]['net'] if isinstance(v[t], dict) else v[t]) for t in v} for k, v in pre.items()})


    # ── 7b. 1日遅れの執行
    lagd, mkd = daily_lag_test(W0)
    dl = {}
    for tag in ('lag0', 'lag1'):
        nn = lagd[tag]
        dl[tag] = {'full': ex(nn, mkd), 'train': ex(nn, mkd, z=TRAIN_END), 'hold': ex(nn, mkd, a=HOLD0),
                   'sharpe_train': (sharpe(nn, rf, z=TRAIN_END), sharpe(mkd, rf, z=TRAIN_END)),
                   'sharpe_hold': (sharpe(nn, rf, a=HOLD0), sharpe({k: mkd[k] for k in nn}, rf, a=HOLD0)),
                   'roll20': roll20(nn, mkd)}
    dl['_note'] = '日次で割合を保つ近似（lag0 と lag1 は同じ近似）。差が『月末の終値で決めて同じ終値で執行する』前提の効き'
    res['execution_lag_1day_daily'] = dl
    log('1日遅れ', {k: (v['train'], v['hold']) for k, v in dl.items() if k != '_note'})

    # ── 7c. 米国外16か国（独立に作り直す）
    import zipfile as _zf
    zb = _zf.ZipFile(io.BytesIO(M.get(M.FR.format('F-F_International_Countries'), name='fr_F-F_International_Countries.zip')))
    intl = {}
    for conv in ('researcher', 'strict'):
        det, cnt = {}, {'X6a': [0, 0, 0], 'tilt_only': [0, 0, 0], 'gate_only': [0, 0, 0]}
        for fn, cc in COUNTRY_FILES.items():
            loc, ann = read_country(zb.read(fn).decode('latin-1'))
            try:
                rr = country_test(loc, ann, cc, conv)
            except Exception as exn:  # noqa
                det[cc] = f'N/A（{str(exn)[:60]}）'
                continue
            det[cc] = rr
            for tag in cnt:
                if isinstance(rr.get(tag), dict):
                    cnt[tag][0] += 1
                    cnt[tag][1] += rr[tag]['positive']
                    sp = rr[tag]['sharpe_net_vs_mkt']
                    cnt[tag][2] += bool(sp[0] is not None and sp[1] is not None and sp[0] > sp[1])
        intl[conv] = {'summary': {t: {'regions': v[0], 'positive': v[1], 'sharpe_better': v[2],
                                      'replicates_2of3': bool(v[0] and v[1] / v[0] >= 2 / 3)} for t, v in cnt.items()}, 'detail': det}
        log('米国外', conv, intl[conv]['summary'])
    res['international_independent'] = intl
    # ── 8. 多重検定
    claim_all = json.load(open(CLAIM))
    graded = [e for e in claim_all['tested'] if e.get('grade') is not None]
    n_angle = len(graded)
    tr_p = 2 * 0.5 * math.erfc(base['train']['t'] / math.sqrt(2))
    ho_p = math.erfc(base['hold']['t'] / math.sqrt(2))
    # mw 全体で試した本数（検証ファイルは除く）
    prog = 0
    per = {}
    for fp in sorted(glob.glob(os.path.join(BASE, 'out', 'mw_*.json'))):
        b = os.path.basename(fp)
        if 'verify' in b or 'prereg' in b:
            continue
        try:
            j = json.load(open(fp))
        except Exception:  # noqa
            continue
        t = j.get('tested')
        if isinstance(t, list):
            per[b] = len(t); prog += len(t)
        elif isinstance(t, dict):
            per[b] = len(t); prog += len(t)
    from statistics import NormalDist
    nd = NormalDist()
    res['multiple_testing'] = {
        'angle_graded_variants': n_angle, 'parameter_grid_added_here': len(grid),
        'program_tested_entries_all_mw_angles': prog, 'program_files': per,
        'train_p_two_sided': round(tr_p, 4), 'train_bonferroni_over_angle': round(min(1, tr_p * n_angle), 3),
        'hold_p_two_sided': round(ho_p, 4), 'hold_bonferroni_over_angle': round(min(1, ho_p * n_angle), 3),
        't_needed_bonferroni_5pct_angle': round(nd.inv_cdf(1 - 0.025 / n_angle), 2),
        't_needed_bonferroni_5pct_program': round(nd.inv_cdf(1 - 0.025 / max(prog, 1)), 2),
        'full_t': base['full']['t'],
        'note': 'C7 は全期間 t≥3.0 で合格しているが、全期間には規則を組み立てる前に見た保有期間が入っている。訓練だけの t 2.96 は角度内の Bonferroni（45本）の線にも届かない'}
    log('多重検定', res['multiple_testing'])

    # ── 9. シャープの差の検定（自前）
    res['sharpe_tests'] = {'train': jk(N0, mkt, rf, z=TRAIN_END), 'hold': jk(N0, mkt, rf, a=HOLD0), 'post2017': jk(N0, mkt, rf, a=201701)}

    # ── 10. 判定のまとめ（別の前提での機械的な格付け）
    c5_meas = res['international_independent']['researcher']['summary']['X6a']['replicates_2of3']
    alt = {}
    alt['as_claimed_C5_NA'] = criteria(base['full'], base['train'], base['hold'], base['net_hold'], base['roll20'], base['sharpe_train'], base['sharpe_hold'])[0]
    alt['C5_measured_intl_8of16'] = criteria(base['full'], base['train'], base['hold'], base['net_hold'], base['roll20'], base['sharpe_train'], base['sharpe_hold'], c5=c5_meas)[0]
    l1 = res['execution_lag_1day_daily']['lag1']
    alt['exec_1day_lag_plus_C5_measured'] = criteria(l1['full'], l1['train'], l1['hold'], l1['hold'], l1['roll20'], l1['sharpe_train'], l1['sharpe_hold'], c5=c5_meas)[0]
    f3 = res['frictions']['spread3pct']
    alt['borrow_RF_plus_3pct_plus_C5_measured'] = criteria(f3['full'], f3['train'], f3['hold'], f3['net_hold'], base['roll20'], f3['sharpe_train'], f3['sharpe_hold'], c5=c5_meas)[0]
    alt['unlevered_cap1.0'] = res['components']['X6a_unlevered_cap1.0']['grade_mechanical']
    alt['post_hoc_per_mw_prereg_honesty_rule'] = '格付けしない（保有期間を見た後に部品を組み合わせた＝事後。mw_prereg の honesty_rules は事後の規則を判定に使わない）→ C 扱い'
    alt['post_publication_2017_net'] = res['subperiods']['post_growth_trend_pub_2017']['net']
    res['grade_under_alternatives'] = alt
    log('格付けの別前提', alt)
    res['log'] = LOG
    return res


if __name__ == '__main__':
    r = main()
    r['generated'] = datetime.date.today().isoformat()
    b, sp, ga = r['reproduction'], r['subperiods'], r['grade_under_alternatives']
    ii = r['international_independent']['researcher']['summary']['X6a']
    r['verdicts'] = [{
        'name': 'X6a_tilt_x_G1', 'claimed_grade': 'S', 'reproduced': True, 'verdict': 'refuted', 'verified_grade': 'C',
        'key_numbers': {'full': b['full'], 'train': b['train'], 'hold': b['hold'], 'net_hold': b['net_hold'], 'roll20': b['roll20'],
                        'dca20': b['dca20'], 'hold_ex_2008_09': sp['ex_2008_2009']['hold'], 'hold_second_half': sp['hold_second_half_2016_06_2025_11'],
                        'post2017_net': sp['post_growth_trend_pub_2017'], 'capm_post2017': r['capm']['post2017'],
                        'exec_1day_lag_hold': r['execution_lag_1day_daily']['lag1']['hold'], 'unlevered_cap1': r['components']['X6a_unlevered_cap1.0']['hold'],
                        'intl_positive': f"{ii['positive']}/{ii['regions']}", 'grade_under_alternatives': ga},
        'why': ['保有期間の結果を見た後に二つの部品を組み合わせた事後の設計（prereg5 自身が明記）。mw_prereg の honesty_rules は事後の規則を判定に使わない',
                'G1（Growth-Trend, 2016）の作者のデータは 2016年まで＝保有 2007〜2016 は規則の作者にとって標本内。本当の標本外 2017〜 は −0.41%/年・CAPM α −2.33%/年 t−2.18',
                '上乗せは 2008〜09年に集中（除くと +1.72 t1.41）、保有の後半 2016-06〜 は −0.09',
                '借入なし（上限 1.0 倍）では格子の 70通りすべて C＝勝ちは借入が前提',
                '同じ終値で決めて執行する前提。翌営業日の執行で保有 t 1.48（C3 不合格）→ B',
                '米国外16か国で正 8/16（研究者と同じ・慣行2通りとも）→ C5 を測れば S は成り立たず A 止まり',
                '多重検定: 角度で45本・mw 全体で約3,600本。訓練 t 2.96 は角度内 Bonferroni（t 3.26）に届かず、全体の線 t 4.35 には全期間 t 3.41 も届かない']}]
    json.dump(r, open(OUT, 'w'), ensure_ascii=False, indent=1, default=str)
    print('書いた', OUT)
