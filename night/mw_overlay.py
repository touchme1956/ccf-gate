#!/usr/bin/env python3
"""night/mw_overlay.py — 『市場に勝てる歴史検証』角度 overlay（読むだけ・門の判定には不使用）

株式市場（French Mkt）を100%持ったまま、先物の証拠金で『株とほぼ無相関のリターン源』を k だけ上乗せする
（リターン・スタッキング）。AQR の時系列モメンタム（TSMOM）と Century of Factor Premia を
運用報酬・成功報酬・売買費用・証拠金の目減りを引いて重ね、リターンとシャープレシオの両方で市場に勝つかを
事前登録 out/mw_overlay_prereg.json の規則どおりに測り、mw_common.grade（C1〜C8）で裁く。
さらに紙の上乗せを実在のマネージド・フューチャーズ・ファンドに置き換えて答え合わせをする。

  python3 night/mw_overlay.py            → out/mw_overlay.json
"""
import datetime, json, math, os, subprocess, sys, time
import numpy as np

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import mw_common as M  # noqa: E402

PREREG = 'out/mw_overlay_prereg.json'
OUT = 'mw_overlay.json'
TSMOM_URL = 'https://www.aqr.com/-/media/AQR/Documents/Insights/Data-Sets/Time-Series-Momentum-Factors-Monthly.xlsx'
CEN_URL = 'https://www.aqr.com/-/media/AQR/Documents/Insights/Data-Sets/Century-of-Factor-Premia-Monthly.xlsx'
TE, HS, RS = M.TRAIN_END, M.HOLD_START, M.RECENT_START
DRAG = 0.002
PERF = 0.20
LOG = []


def log(*a):
    s = ' '.join(str(x) for x in a)
    LOG.append(s)
    print(s, flush=True)


# ───────────────────────── データ ─────────────────────────
def _isnum(x):
    return isinstance(x, (int, float)) and not isinstance(x, bool) and math.isfinite(x)


def load_tsmom():
    import openpyxl
    M.get(TSMOM_URL, name='aqr_tsmom_monthly.xlsx')
    wb = openpyxl.load_workbook(os.path.join(M.CACHE, 'aqr_tsmom_monthly.xlsx'), read_only=True, data_only=True)
    rows = list(wb['TSMOM Factors'].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[1] == 'TSMOM')
    hdr = rows[hi]
    out = {h: {} for h in hdr[1:] if h}
    for r in rows[hi + 1:]:
        d = r[0]
        if not isinstance(d, datetime.datetime):
            continue
        ym = d.year * 100 + d.month
        for j, h in enumerate(hdr):
            if j == 0 or not h:
                continue
            if _isnum(r[j]):
                out[h][ym] = float(r[j])
    return out


def load_century():
    import openpyxl
    M.get(CEN_URL, name='aqr_century_monthly.xlsx')
    wb = openpyxl.load_workbook(os.path.join(M.CACHE, 'aqr_century_monthly.xlsx'), read_only=True, data_only=True)
    rows = list(wb['Century of Factor Premia'].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[0] == 'Date')
    hdr = rows[hi]
    out = {h: {} for h in hdr[1:] if h}
    for r in rows[hi + 1:]:
        d = r[0]
        if isinstance(d, str) and d.count('/') == 2:
            mm, dd, yy = d.split('/')
            ym = int(yy) * 100 + int(mm)
        elif isinstance(d, datetime.datetime):
            ym = d.year * 100 + d.month
        else:
            continue
        for j, h in enumerate(hdr):
            if j == 0 or not h:
                continue
            if _isnum(r[j]):   # 空文字・None は欠測（0 にしない）
                out[h][ym] = float(r[j])
    return out


def french_region(name):
    for t, v in M.french_tables(name).items():
        if v['freq'] == 'monthly':
            cols = v['cols']
            i_m, i_rf = cols.index('Mkt-RF'), cols.index('RF')
            mk, rf = {}, {}
            for d, row in v['data'].items():
                if row[i_m] is not None and row[i_rf] is not None:
                    mk[d] = (row[i_m] + row[i_rf]) / 100
                    rf[d] = row[i_rf] / 100
            return mk, rf
    raise KeyError(name)


def av_dead(t):
    p = os.path.join(M.CACHE, f'av_dead_mf_{t}.csv')
    if not os.path.exists(p):
        return None
    import csv
    rows = sorted(csv.DictReader(open(p)), key=lambda r: r['date'])
    px = {}
    for r in rows:
        try:
            a = float(r['adj'])
        except (ValueError, KeyError):
            continue
        y, m = int(r['date'][:4]), int(r['date'][5:7])
        px[y * 100 + m] = a
    ks = sorted(px)
    return {k: px[k] / px[p] - 1 for p, k in zip(ks, ks[1:])}


def cboe_monthly(name):
    b = M.get(f'https://cdn.cboe.com/api/global/us_indices/daily_prices/{name}_History.csv', name=f'cboe_{name}_History.csv')
    last = {}
    for line in b.decode('latin-1').splitlines()[1:]:
        c = line.split(',')
        if len(c) < 2:
            continue
        try:
            mm, dd, yy = c[0].split('/')
            v = float(c[1])
        except ValueError:
            continue
        ym, d = int(yy) * 100 + int(mm), int(yy) * 10000 + int(mm) * 100 + int(dd)
        if ym not in last or d > last[ym][0]:
            last[ym] = (d, v)
    ks = sorted(last)
    out = {}
    for p, k in zip(ks, ks[1:]):
        # 連続した月だけ（間が空いた月はリターンにしない）
        py, pm_ = divmod(p, 100)
        ny, nm_ = divmod(k, 100)
        if (ny * 12 + nm_) - (py * 12 + pm_) == 1:
            out[k] = last[k][1] / last[p][1] - 1
    return out


# ───────────────────────── 上乗せの組み立て ─────────────────────────
def sd_ann(d, a=None, z=None):
    v = [x for k, x in d.items() if (a is None or k >= a) and (z is None or k <= z)]
    return float(np.std(v, ddof=1) * math.sqrt(12))


def sleeve(a, k, mgmt, tcy, perf=PERF):
    """ファンドの持ち分 k の月次の寄与（費用後）。成功報酬は暦年の最後の月に 20%×max(0, 年の合計)"""
    ks = sorted(a)
    sl = {t: k * (a[t] - mgmt / 12 - tcy / 12) for t in ks}
    if perf:
        by = {}
        for t in ks:
            by.setdefault(t // 100, []).append(t)
        for y, ms in by.items():
            tot = math.fsum(sl[t] for t in ms)
            if tot > 0:
                sl[max(ms)] -= perf * tot
    return sl


def stack(m, sl, drag=DRAG):
    return {t: m[t] + sl[t] - drag / 12 for t in sorted(set(m) & set(sl))}


# ───────────────────────── 統計 ─────────────────────────
def keys(s, b, a=None, z=None):
    return sorted(k for k in set(s) & set(b) if (a is None or k >= a) and (z is None or k <= z))


def jk_memmel(s, b, rf, a=None, z=None):
    ks = [k for k in keys(s, b, a, z) if k in rf]
    if len(ks) < 24:
        return None
    y1 = np.array([s[k] - rf[k] for k in ks]); y2 = np.array([b[k] - rf[k] for k in ks])
    s1, s2 = y1.mean() / y1.std(ddof=1), y2.mean() / y2.std(ddof=1)
    rho = float(np.corrcoef(y1, y2)[0, 1]); T = len(ks)
    var = (2 - 2 * rho + 0.5 * (s1 ** 2 + s2 ** 2 - 2 * s1 * s2 * rho ** 2)) / T
    zz = float((s1 - s2) / math.sqrt(var))
    return {'z': round(zz, 2), 'p_two': round(math.erfc(abs(zz) / math.sqrt(2)), 4),
            'p_one': round(0.5 * math.erfc(zz / math.sqrt(2)), 4), 'rho': round(rho, 3), 'n': T,
            'sharpe_diff_ann': round((s1 - s2) * math.sqrt(12), 3)}


def boot_sharpe(s, b, rf, a=None, z=None, reps=2000, block=12, seed=7):
    ks = [k for k in keys(s, b, a, z) if k in rf]
    if len(ks) < 24:
        return None
    y1 = np.array([s[k] - rf[k] for k in ks]); y2 = np.array([b[k] - rf[k] for k in ks])
    T = len(ks); rng = np.random.default_rng(seed)
    nb = int(math.ceil(T / block))
    st = rng.integers(0, T, (reps, nb))
    idx = ((st[:, :, None] + np.arange(block)[None, None, :]).reshape(reps, -1)[:, :T]) % T
    a1, a2 = y1[idx], y2[idx]
    d = a1.mean(1) / a1.std(1, ddof=1) - a2.mean(1) / a2.std(1, ddof=1)
    cnt = int((d <= 0).sum())
    ds = np.sort(d)
    return {'p_one_sided': round((cnt + 1) / (reps + 1), 4),
            'ci90_ann': [round(float(ds[int(0.05 * reps)]) * math.sqrt(12), 3), round(float(ds[int(0.95 * reps)]) * math.sqrt(12), 3)]}


def corr_w(x, y, a=None, z=None):
    ks = keys(x, y, a, z)
    if len(ks) < 24:
        return None
    return round(M.corr([x[k] for k in ks], [y[k] for k in ks]), 3)


def maxdd_w(r, ks):
    return round(M.maxdd({k: r[k] for k in ks}) * 100, 1) if ks else None


def ann_stats(x, a=None, z=None):
    v = [x[k] for k in sorted(x) if (a is None or k >= a) and (z is None or k <= z)]
    if len(v) < 12:
        return None
    mu, sd = float(np.mean(v)) * 12, float(np.std(v, ddof=1)) * math.sqrt(12)
    return {'months': len(v), 'mean_ann': round(mu * 100, 2), 'vol_ann': round(sd * 100, 2),
            'sharpe': round(mu / sd, 3) if sd else None, 't_nw': (lambda t: round(t, 2) if t is not None else None)(M.nw_t(v))}


def drop_years(d, years):
    return {k: v for k, v in d.items() if k // 100 not in years}


def vol_matched(s, m, rf, spread=0.005):
    """訓練期間の戦略のぶれに合わせて市場を L 倍（借入 RF+spread）"""
    ks = [k for k in keys(s, m, None, TE) if k in rf]
    if len(ks) < 24:
        return None, None
    ss = np.std([s[k] - rf[k] for k in ks], ddof=1); sm = np.std([m[k] - rf[k] for k in ks], ddof=1)
    L = float(ss / sm)
    mL = {k: rf[k] + L * (m[k] - rf[k]) - max(L - 1, 0) * spread / 12 for k in m if k in rf}
    return round(L, 3), mL


# ───────────────────────── 日本の税の近似（報告） ─────────────────────────
def jp_tax_annual(m, sl, drag=DRAG, rate=0.20315, carry=3):
    """株 NISA（非課税）＋上乗せを課税口座: 上乗せの暦年の損益（費用後・drag込み）が正なら12月に課税。損失は3年繰越"""
    ks = sorted(set(m) & set(sl))
    by = {}
    for t in ks:
        by.setdefault(t // 100, []).append(t)
    losses = []  # (年, 残り)
    out = {}
    for y in sorted(by):
        ms = by[y]
        pl = math.fsum(sl[t] - drag / 12 for t in ms)
        tax = 0.0
        losses = [(yy, v) for yy, v in losses if y - yy <= carry]
        if pl > 0:
            rem = pl
            nl = []
            for yy, v in losses:
                use = min(v, rem); rem -= use
                if v - use > 1e-12:
                    nl.append((yy, v - use))
            losses = nl
            tax = rate * rem
        elif pl < 0:
            losses.append((y, -pl))
        for t in ms:
            out[t] = m[t] + sl[t] - drag / 12 - (tax if t == max(ms) else 0.0)
    return out


def jp_tax_deferred(m, sl, a, z, drag=DRAG, rate=0.20315):
    """上乗せの累積の利益に最後に一度だけ課税（分配しない投信を持ち続ける形）→ 年率差（戦略の税後 − 市場）"""
    ks = keys(m, sl, a, z)
    if not ks:
        return None
    w = 1.0; gain = 0.0
    for t in ks:
        contrib = w * (sl[t] - drag / 12)
        gain += contrib
        w *= 1 + m[t] + sl[t] - drag / 12
    w_after = w - rate * max(gain, 0.0)
    wm = math.prod(1 + m[t] for t in ks)
    yrs = len(ks) / 12
    return round(((w_after ** (1 / yrs)) - (wm ** (1 / yrs))) * 100, 2)


# ───────────────────────── 日本で買えるか ─────────────────────────
def toushin_search(keywords):
    import http.cookiejar, urllib.request
    BASE = 'https://toushin-lib.fwg.ne.jp'
    ARR = ['s_investAssetKindCd', 's_investArea3kindCd', 's_instCd', 's_fdsInstCd', 's_dcFundCD', 't_investArea10kindCd',
           't_investAssetKindCd', 't_instCd', 't_fdsInstCd', 's_investArea10kindCd', 's_setlFqcy', 's_dividend1y',
           's_totalNetAssets', 's_nowToRedemptionDate', 's_establishedDateToNow', 's_isinCd']
    KEEP = ['isinCd', 'fundNm', 'entrustCmpNm', 'trustReward', 'establishedDate', 'totalNetAssets', 'standardDate',
            'nisaFlg', 'nisaGrowthFlg', 'standardPriceRa1y', 'standardPriceRa3y', 'standardPriceRa5y', 'standardPriceRa10y',
            'riskRa5y', 'sharpRa5y']
    ua = M.UA['User-Agent']
    cj = http.cookiejar.CookieJar()
    op = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cj))
    op.addheaders = [('User-Agent', ua)]
    op.open(BASE + '/FdsWeb/FDST000000', timeout=40).read()
    res, seen = {}, {}
    for kw in keywords:
        rows, start = [], 0
        try:
            while start < 200:
                body = {f: [] for f in ARR}
                body.update({'t_keyword': kw, 't_kensakuKbn': '1', 't_searchInfoFlag': '1', 'startNo': start, 'draw': 1, 'searchBtnClickFlg': True})
                req = urllib.request.Request(BASE + '/FdsWeb/FDST999900/fundDataSearch', data=json.dumps(body).encode(),
                                             headers={'Content-Type': 'application/json', 'X-Requested-With': 'XMLHttpRequest',
                                                      'Referer': BASE + '/FdsWeb/FDST999900', 'User-Agent': ua})
                d = json.loads(op.open(req, timeout=40).read().decode())
                info = d.get('searchResultInfo') or {}
                got = info.get('resultInfoMapList') or []
                rows += got
                total = int(info.get('recordsTotal') or 0)
                start += len(got)
                if not got or start >= total:
                    break
                time.sleep(0.4)
            res[kw] = len(rows)
            for h in rows:
                seen[h.get('isinCd') or h.get('fundNm')] = {k: h.get(k) for k in KEEP}
        except Exception as e:  # noqa
            res[kw] = f'取得失敗: {str(e)[:80]}'
        time.sleep(0.5)
    return {'hits_per_keyword': res, 'funds': sorted(seen.values(), key=lambda x: str(x.get('fundNm')))}


# ───────────────────────── 本体 ─────────────────────────
def git_sha(path):
    try:
        return subprocess.check_output(['git', 'log', '-1', '--format=%h', '--', path], cwd=M.BASE, text=True).strip() or None
    except Exception:  # noqa
        return None


def r2(x):
    return None if x is None else round(x, 3)


def evaluate(name, spec, m, rf, regions, fam):
    """spec: {'a': dict, 'k', 'c', 'mgmt_base', 'mgmt_stress', 'T', 'perf'}"""
    a, k, c, T, perf = spec['a'], spec['k'], spec['c'], spec['T'], spec['perf']
    sl_b = sleeve(a, k, spec['mgmt_base'], T * c * 0.0005, PERF if perf else 0)
    sl_s = sleeve(a, k, spec['mgmt_stress'], T * c * 0.0010, PERF if perf else 0)
    s = stack(m, sl_b)
    s_st = stack(m, sl_s)
    s_gr = {t: m[t] + k * a[t] for t in sorted(set(m) & set(a))}
    b = m
    full = M.excess_stats(s, b); train = M.excess_stats(s, b, z=TE); hold = M.excess_stats(s, b, a=HS)
    recent = M.excess_stats(s, b, a=RS)
    cost_hold = M.excess_stats(s_st, b, a=HS)
    roll = M.rolling(s, b, 20); dc = M.dca(s, b, 20)
    sp = {'train': (M.sharpe(s, rf, z=TE), M.sharpe(b, rf, z=TE)) if train else None,
          'hold': (M.sharpe(s, rf, a=HS), M.sharpe(b, rf, a=HS)) if hold else None}
    sp_full = (M.sharpe(s, rf), M.sharpe(b, rf))
    tests = {'train': {'jk_memmel': jk_memmel(s, b, rf, z=TE), 'block_boot': boot_sharpe(s, b, rf, z=TE)},
             'hold': {'jk_memmel': jk_memmel(s, b, rf, a=HS), 'block_boot': boot_sharpe(s, b, rf, a=HS)}}
    th = tests['hold']
    c8sig = bool(th['jk_memmel'] and th['block_boot'] and th['jk_memmel']['p_one'] < 0.05 and th['block_boot']['p_one_sided'] < 0.05)
    # C5
    rep_detail, pos, nreg = {}, 0, 0
    for rn, (rm, rrf) in regions.items():
        sr = stack(rm, sleeve({t: v for t, v in a.items() if t in rm}, k, spec['mgmt_base'], T * c * 0.0005, PERF if perf else 0))
        es = M.excess_stats(sr, rm)
        esh = M.excess_stats(sr, rm, a=HS)
        if es:
            nreg += 1
            ok = es['ex_ann'] > 0 and es['cagr_diff'] > 0
            pos += ok
            rep_detail[rn] = {'full': es, 'hold': esh, 'positive': ok,
                              'sharpe_full': (M.sharpe(sr, rrf), M.sharpe(rm, rrf)), 'sharpe_hold': (M.sharpe(sr, rrf, a=HS), M.sharpe(rm, rrf, a=HS))}
    repl = {'regions': nreg, 'positive': pos}
    L, mL = vol_matched(s, m, rf)
    vm = {'L': L, 'hold': M.excess_stats(s, mL, a=HS) if mL else None, 'full': M.excess_stats(s, mL) if mL else None}
    ksh = keys(s, b, HS); ksf = keys(s, b)
    ent = {
        'name': name, 'family': fam, 'desc': spec['desc'], 'overlay': spec['ov'], 'k': k, 'c': r2(c), 'T_turnover': T,
        'leak': spec.get('leak', False),
        'overlay_stats': {'train': ann_stats(a, z=TE), 'hold': ann_stats(a, a=HS), 'full': ann_stats(a),
                          'corr_mktrf_train': corr_w(a, {t: m[t] - rf[t] for t in m if t in rf}, z=TE),
                          'corr_mktrf_hold': corr_w(a, {t: m[t] - rf[t] for t in m if t in rf}, a=HS)},
        'full': full, 'train': train, 'hold': hold, 'recent': recent,
        'postpub_2013': M.excess_stats(s, b, a=201301), 'postpub_2022': M.excess_stats(s, b, a=202201),
        'hold_ex_2008_2022': M.excess_stats(drop_years(s, {2008, 2022}), b, a=HS),
        'pre1985': M.excess_stats(s, b, z=198412),
        'sharpe_pre1985': (M.sharpe(s, rf, z=198412), M.sharpe(b, rf, z=198412)) if M.excess_stats(s, b, z=198412) else None,
        'sharpe_ex_2008_2022_hold': (M.sharpe(drop_years(s, {2008, 2022}), rf, a=HS), M.sharpe(drop_years(b, {2008, 2022}), rf, a=HS)),
        'sharpe_recent': (M.sharpe(s, rf, a=RS), M.sharpe(b, rf, a=RS)),
        'cost_hold_stress': cost_hold,
        'gross': {'full': M.excess_stats(s_gr, b), 'train': M.excess_stats(s_gr, b, z=TE), 'hold': M.excess_stats(s_gr, b, a=HS)},
        'roll20': roll, 'dca20': dc,
        'sharpe_pair': sp, 'sharpe_full': sp_full, 'sharpe_tests': tests, 'c8_significant_hold': c8sig,
        'maxdd': {'full': (maxdd_w(s, ksf), maxdd_w(b, ksf)), 'hold': (maxdd_w(s, ksh), maxdd_w(b, ksh))},
        'repl': repl, 'repl_detail': rep_detail, 'vol_matched': vm,
        '_series': s, '_sleeve': sl_b,
    }
    return ent


def main():
    t0 = time.time()
    pre = json.load(open(os.path.join(M.BASE, PREREG)))
    ff = M.ff_factors()
    m, rf = ff['mkt'], ff['rf']
    mktrf = ff['mktrf']
    log('French Mkt', min(m), max(m), 'CAGR', round(M.cagr(m) * 100, 2), '2007〜', round(M.cagr(M.window(m, HS)) * 100, 2))
    TS = load_tsmom(); CEN = load_century()
    for h, d in TS.items():
        log('TSMOM', h, min(d), max(d), len(d))
    regions = {}
    for rn in ['Developed_ex_US_3_Factors', 'Japan_3_Factors', 'Europe_3_Factors']:
        regions[rn.replace('_3_Factors', '')] = french_region(rn)

    def cc(o):
        return 0.10 / sd_ann(o, z=TE)

    specs = {}
    # 主の族
    for nm, key, src, T in [('TSMOM', 'TSMOM', TS, 5), ('CENmom', 'All asset classes Momentum', CEN, 10), ('CENms', 'All asset classes Multi-style', CEN, 10)]:
        for k in (0.25, 0.5):
            specs[f'P_{nm}_k{int(k * 100)}'] = {'fam': 'primary', 'a': src[key], 'ov': key, 'k': k, 'c': 1.0, 'T': T, 'perf': True,
                                                'mgmt_base': 0.01, 'mgmt_stress': 0.02, 'desc': f'Mkt + {k}×{key}（c=1）'}
    # 探索 E1
    for nm, key, src, T in [('TSMOM', 'TSMOM', TS, 5), ('CENmom', 'All asset classes Momentum', CEN, 10), ('CENms', 'All asset classes Multi-style', CEN, 10)]:
        c = cc(src[key])
        for k in (0.5, 1.0):
            specs[f'E1_{nm}_v10_k{int(k * 100)}'] = {'fam': 'explore', 'a': {t: c * v for t, v in src[key].items()}, 'ov': key, 'k': k, 'c': c, 'T': T, 'perf': True,
                                                     'mgmt_base': 0.01, 'mgmt_stress': 0.02, 'desc': f'Mkt + {k}×（{key} をぶれ10%に）'}
    # E2
    for nm, key, T, leak in [('ACval', 'All asset classes Value', 10, False), ('ACcarry', 'All asset classes Carry', 10, False),
                             ('ACdef', 'All asset classes Defensive', 10, False), ('MACmom', 'All Macro Momentum', 5, False),
                             ('MACms', 'All Macro Multi-style', 5, False), ('SSms', 'All Stock Selection Multi-style', 10, False),
                             ('EQIXmom', 'Equity indices Momentum', 5, True)]:
        c = cc(CEN[key])
        specs[f'E2_{nm}_v10_k50'] = {'fam': 'explore', 'a': {t: c * v for t, v in CEN[key].items()}, 'ov': key, 'k': 0.5, 'c': c, 'T': T, 'perf': True,
                                     'mgmt_base': 0.01, 'mgmt_stress': 0.02, 'leak': leak, 'desc': f'Mkt + 0.5×（{key} をぶれ10%に）' + ('【leak: 下見で保有期間を見た系列】' if leak else '')}
    # E3
    for key in ['TSMOM^CM', 'TSMOM^EQ', 'TSMOM^FI', 'TSMOM^FX']:
        c = cc(TS[key])
        specs[f'E3_{key.replace("^", "_")}_v10_k50'] = {'fam': 'explore', 'a': {t: c * v for t, v in TS[key].items()}, 'ov': key, 'k': 0.5, 'c': c, 'T': 5, 'perf': True,
                                                         'mgmt_base': 0.01, 'mgmt_stress': 0.02, 'desc': f'Mkt + 0.5×（{key} をぶれ10%に）'}
    # E4
    for nm, key, k in [('FImkt', 'Fixed income Market', 0.5), ('FImkt', 'Fixed income Market', 1.0), ('CMmkt', 'Commodities Market', 0.5)]:
        specs[f'E4_{nm}_k{int(k * 100)}'] = {'fam': 'explore', 'a': CEN[key], 'ov': key, 'k': k, 'c': 1.0, 'T': 2, 'perf': False,
                                             'mgmt_base': 0.002, 'mgmt_stress': 0.005, 'desc': f'Mkt + {k}×{key}（先物をそのまま・成功報酬なし）'}
    # E5
    c1, c2 = cc(TS['TSMOM']), cc(CEN['All asset classes Multi-style'])
    comb = {t: 0.5 * c1 * TS['TSMOM'][t] + 0.5 * c2 * CEN['All asset classes Multi-style'][t] for t in set(TS['TSMOM']) & set(CEN['All asset classes Multi-style'])}
    specs['E5_TSMOMxCENms_v10_k100'] = {'fam': 'explore', 'a': comb, 'ov': 'TSMOM(v10)/2 + CENms(v10)/2', 'k': 1.0, 'c': 1.0,
                                        'T': 0.5 * 5 * c1 + 0.5 * 10 * c2, 'perf': True, 'mgmt_base': 0.01, 'mgmt_stress': 0.02,
                                        'desc': 'Mkt + 1.0×（ぶれ10%の TSMOM と ぶれ10%の Century Multi-style を半分ずつ）'}
    assert len([s for s in specs.values() if s['fam'] == 'primary']) == 6 and len(specs) == 27, len(specs)

    res = {}
    for nm, sp in specs.items():
        res[nm] = evaluate(nm, sp, m, rf, regions, sp['fam'])
        e = res[nm]
        log(f"{nm:28s} c={e['c']} tr {e['train']['ex_ann'] if e['train'] else None} t{e['train']['t'] if e['train'] else None} | "
            f"ho {e['hold']['ex_ann']} t{e['hold']['t']} geo{e['hold']['cagr_diff']} | SR tr {e['sharpe_pair']['train']} ho {e['sharpe_pair']['hold']}")
    # Holm
    prim = {n: res[n]['hold']['p'] for n in res if res[n]['family'] == 'primary'}
    allp = {n: res[n]['hold']['p'] for n in res}
    hp, ha = M.holm(prim), M.holm(allp)
    for n, e in res.items():
        fp = hp[n] if e['family'] == 'primary' else ha[n]
        e['family_holm_p'] = fp
        e['holm_p_all27'] = ha[n]
        g, crit = M.grade(e['full'], e['train'], e['hold'], e['roll20'], cost_hold=e['cost_hold_stress'], repl=e['repl'],
                          family_holm_p=fp, sharpe_pair=e['sharpe_pair'], leveraged_or_timing=True)
        e['grade'], e['criteria'] = g, crit
        log(f"  {n:28s} grade {g} {crit} c8sig={e['c8_significant_hold']}")

    # ── 下見の再現（パーサーの検算・判定に不使用）
    ts = TS['TSMOM']
    chk = {t: m[t] + 0.5 * (ts[t] - 0.01 / 12) for t in set(m) & set(ts)}
    rep = {'train': M.excess_stats(chk, m, z=TE), 'hold': M.excess_stats(chk, m, a=HS), 'recent': M.excess_stats(chk, m, a=RS),
           'sharpe_hold': (M.sharpe(chk, rf, a=HS), M.sharpe(m, rf, a=HS)), 'corr_tsmom_mktrf': corr_w(ts, mktrf),
           'brief_says': 'train +8.4 t8.7 / hold +2.4 t1.79 / recent +2.0 t1.07 / SR hold 0.81 vs 0.655 / corr −0.08'}
    log('下見の再現', json.dumps({k: (v if not isinstance(v, dict) else {kk: v[kk] for kk in ('ex_ann', 't')}) for k, v in rep.items()}, ensure_ascii=False))

    # ── 日本の税（報告）
    jp_tax = {}
    for n in ['P_TSMOM_k50', 'P_CENmom_k50', 'P_CENms_k50', 'E1_TSMOM_v10_k100', 'E1_CENmom_v10_k100', 'E1_CENms_v10_k100']:
        e = res[n]; sl = e['_sleeve']
        sa = jp_tax_annual(m, sl)
        jp_tax[n] = {'annual_realization_hold': M.excess_stats(sa, m, a=HS), 'annual_realization_train': M.excess_stats(sa, m, z=TE),
                     'deferred_hold_cagr_diff': jp_tax_deferred(m, sl, HS, None),
                     'no_tax_hold_cagr_diff': e['hold']['cagr_diff']}

    # ── VRP（報告）
    vrp = {}
    try:
        put = cboe_monthly('PUT')
        pe = {t: put[t] - rf[t] for t in put if t in rf}
        s1 = {t: m[t] + 0.5 * pe[t] for t in pe if t in m}
        hed = {}
        ks = sorted(t for t in pe if t in mktrf)
        for i, t in enumerate(ks):
            if i < 36:
                continue
            w = ks[i - 36:i]
            x = np.array([mktrf[u] for u in w]); y = np.array([pe[u] for u in w])
            beta = float(np.cov(x, y, ddof=1)[0, 1] / np.var(x, ddof=1))
            hed[t] = pe[t] - beta * mktrf[t]
        s2 = {t: m[t] + 0.5 * hed[t] - DRAG / 12 for t in hed if t in m}
        vrp = {'months_put': [min(put), max(put), len(put)],
               'unhedged_k50': {'hold': M.excess_stats(s1, m, a=HS), 'sharpe_hold': (M.sharpe(s1, rf, a=HS), M.sharpe(m, rf, a=HS))},
               'beta_hedged_k50': {'hold': M.excess_stats(s2, m, a=HS), 'sharpe_hold': (M.sharpe(s2, rf, a=HS), M.sharpe(m, rf, a=HS)),
                                   'jk': jk_memmel(s2, m, rf, a=HS)},
               'overlay_corr_mktrf_hedged': corr_w(hed, mktrf)}
    except Exception as ex:  # noqa
        vrp = {'error': str(ex)[:200]}

    # ── 実在のファンド（答え合わせ）
    RG = pre['reality_gap']
    trend = ['RYMFX', 'AQMIX', 'ASFYX', 'WTMF', 'CSAIX', 'MFTNX', 'EBSIX', 'QMHIX', 'FMF', 'PQTIX', 'ABYIX', 'AHLIX', 'DBMF', 'KMLM', 'CTA', 'FMFFX', 'ASMF']
    multi = ['QSPIX', 'QRPIX', 'LFMIX']
    direct = ['RSST', 'NTSX', 'BLNDX']
    dead_try = ['MHFIX', 'MHFAX', 'FUTS', 'WFIIX', 'PFFTX', 'RTSIX', 'WAVEX']
    live = {}
    for t in trend + multi + direct:
        try:
            r = M.yahoo(t)
            live[t] = {k: v for k, v in r.items() if k in rf}
        except Exception as ex:  # noqa
            live[t] = None
            log('Yahoo 失敗', t, str(ex)[:80])
    dead = {}
    for t in dead_try:
        r = av_dead(t)
        dead[t] = {k: v for k, v in r.items() if k in rf} if r else None
    # 異常値の印
    anomalies = {t: sorted((k, round(v, 3)) for k, v in (r or {}).items() if abs(v) > 0.3) for t, r in {**live, **dead}.items() if r}
    anomalies = {t: v for t, v in anomalies.items() if v}
    # 紙の比較相手（ぶれ10%・基本の費用後・k=1 のファンド）
    cT, cM = cc(TS['TSMOM']), cc(CEN['All asset classes Multi-style'])
    paper_T = sleeve({t: cT * v for t, v in TS['TSMOM'].items()}, 1.0, 0.01, 5 * cT * 0.0005)
    paper_M = sleeve({t: cM * v for t, v in CEN['All asset classes Multi-style'].items()}, 1.0, 0.01, 10 * cM * 0.0005)

    def fund_entry(t, r, paper, kind):
        if not r or len(r) < 24:
            return {'ticker': t, 'kind': kind, 'status': 'データ不足・取れず', 'months': len(r) if r else 0}
        ex = {k: r[k] - rf[k] for k in r}
        ks = keys(ex, paper)
        e = {'ticker': t, 'kind': kind, 'from': min(r), 'to': max(r), 'months': len(r),
             'live_excess': ann_stats(ex), 'cagr_total': round(M.cagr(r) * 100, 2),
             'mkt_same_months': ann_stats({k: m[k] - rf[k] for k in r}),
             'corr_mktrf': corr_w(ex, mktrf)}
        if len(ks) >= 24:
            pl = {k: paper[k] for k in ks}; lv = {k: ex[k] for k in ks}
            e['paper_same_months'] = ann_stats(pl); e['live_on_paper_months'] = ann_stats(lv)
            e['corr_paper'] = corr_w(lv, pl)
            e['gap_paper_minus_live_ann'] = round(e['paper_same_months']['mean_ann'] - e['live_on_paper_months']['mean_ann'], 2)
            e['paper_window'] = [ks[0], ks[-1]]
        for k in (0.5, 1.0):
            s = {u: m[u] + k * ex[u] - DRAG / 12 for u in ex if u in m}
            e[f'stack_k{int(k * 100)}'] = {'ex': M.excess_stats(s, m), 'sharpe': (M.sharpe(s, rf), M.sharpe(m, {u: rf[u] for u in s})),
                                          'jk': jk_memmel(s, m, rf), 'maxdd': (maxdd_w(s, sorted(s)), maxdd_w(m, sorted(s)))}
        return e

    funds = {}
    for t in trend:
        funds[t] = fund_entry(t, live.get(t), paper_T, 'trend')
    for t in multi:
        funds[t] = fund_entry(t, live.get(t), paper_M, 'multi-style/macro')
    for t in dead_try:
        funds['DEAD_' + t] = fund_entry(t, dead.get(t), paper_T, 'trend(dead・AV)') if dead.get(t) else {'ticker': t, 'status': 'AV で取れず（欠測のまま・埋めない）'}
    for t in direct:
        r = live.get(t)
        if r and len(r) >= 24:
            funds['DIRECT_' + t] = {'ticker': t, 'kind': 'stacked product (direct vs Mkt)', 'from': min(r), 'to': max(r), 'months': len(r),
                                    'ex': M.excess_stats(r, m), 'sharpe': (M.sharpe(r, rf), M.sharpe({u: m[u] for u in r if u in m}, rf)),
                                    'jk': jk_memmel(r, m, rf), 'maxdd': (maxdd_w(r, sorted(r)), maxdd_w(m, sorted(r)))}
        else:
            funds['DIRECT_' + t] = {'ticker': t, 'status': 'データ不足'}

    def composite(tickers, src):
        allk = sorted(set().union(*[set(src[t]) for t in tickers if src.get(t)]))
        comp, cnt = {}, {}
        for k in allk:
            v = [src[t][k] for t in tickers if src.get(t) and k in src[t]]
            if v:
                comp[k] = sum(v) / len(v); cnt[k] = len(v)
        return comp, cnt

    comp_live, cnt_live = composite(trend, live)
    both = {**{t: live.get(t) for t in trend}, **{'DEAD_' + t: dead.get(t) for t in dead_try}}
    comp_all, cnt_all = composite(list(both), both)
    comps = {}
    for nm, comp, cnt in [('survivors_only', comp_live, cnt_live), ('with_dead', comp_all, cnt_all)]:
        e = fund_entry(nm, comp, paper_T, 'composite')
        e['funds_per_month'] = {'first': [min(cnt), cnt[min(cnt)]], 'last': [max(cnt), cnt[max(cnt)]], 'max': max(cnt.values())}
        e['hold_2007'] = {}
        for k in (0.5, 1.0):
            s = {u: m[u] + k * (comp[u] - rf[u]) - DRAG / 12 for u in comp if u in m}
            e['hold_2007'][f'k{int(k * 100)}'] = {'ex_hold': M.excess_stats(s, m, a=HS), 'sharpe_hold': (M.sharpe(s, rf, a=HS), M.sharpe(m, {u: rf[u] for u in s}, a=HS)),
                                                  'ex_ex_2008_2022': M.excess_stats(drop_years(s, {2008, 2022}), m, a=HS)}
        comps[nm] = e
    # 同じ窓の紙（2007-03〜）
    paper_same = {}
    for n in ['P_TSMOM_k50', 'E1_TSMOM_v10_k50', 'E1_TSMOM_v10_k100']:
        s = res[n]['_series']
        a0 = min(comp_live)
        paper_same[n] = {'ex': M.excess_stats(s, m, a=a0), 'sharpe': (M.sharpe(s, rf, a=a0), M.sharpe(m, rf, a=a0))}

    # ── 日本で買えるか
    try:
        jp = toushin_search(['マネージド・フューチャーズ', 'マネージドフューチャーズ', 'マン・AHL', 'AHL', 'フューチャーズ', 'トレンド・フォロー',
                             'トレンドフォロー', 'ウィントン', 'アルファシンプレックス', 'リターン・スタック'])
    except Exception as ex:  # noqa
        jp = {'error': str(ex)[:200]}
    tl = json.load(open(os.path.join(M.BASE, 'out', 'tsumitate_lineup.json')))
    tl_hits = []
    for sec, v in tl.items():
        if isinstance(v, dict) and 'rows' in v:
            for r in v['rows']:
                if any(w in (r.get('nm') or '') for w in ['フューチャーズ', 'AHL', 'トレンド', 'ウィントン', 'マネージド']):
                    tl_hits.append({'sec': sec, **r})
    bl = json.load(open(os.path.join(M.BASE, 'out', 'broker_lineup.json')))
    bl_check = {t: (t in bl['etfs']) for t in ['DBMF', 'KMLM', 'CTA', 'RSST', 'RSBT', 'NTSX', 'WTMF', 'FMF', 'ASMF', 'BLNDX']}

    # ── 保存
    tested = []
    for n, e in res.items():
        tested.append({'name': n, 'family': e['family'], 'graded': True, 'grade': e['grade'], 'leak': e['leak'],
                       'hold_ex': e['hold']['ex_ann'], 'hold_t': e['hold']['t'], 'hold_cagr_diff': e['hold']['cagr_diff']})
    for n in jp_tax:
        tested.append({'name': 'R04_jp_tax_' + n, 'family': 'report', 'graded': False})
    tested += [{'name': 'R05_vrp_unhedged_k50', 'family': 'report', 'graded': False}, {'name': 'R05_vrp_beta_hedged_k50', 'family': 'report', 'graded': False}]
    for n in funds:
        tested.append({'name': 'RG_' + n, 'family': 'reality_gap', 'graded': False})
    for n in comps:
        tested.append({'name': 'RG_composite_' + n, 'family': 'reality_gap', 'graded': False})

    def strip(e):
        return {k: v for k, v in e.items() if not k.startswith('_')}

    out = {
        'angle': 'overlay', 'prereg': PREREG, 'prereg_commit': git_sha(PREREG),
        'question': pre['question'],
        'sanity': {'french_mkt_cagr_1926': round(M.cagr(m) * 100, 2), 'french_mkt_cagr_2007': round(M.cagr(M.window(m, HS)) * 100, 2),
                   'tsmom_range': [min(TS['TSMOM']), max(TS['TSMOM']), len(TS['TSMOM'])],
                   'century_range': {k: [min(CEN[k]), max(CEN[k]), len(CEN[k])] for k in ['All asset classes Momentum', 'All asset classes Multi-style', 'Fixed income Market']},
                   'brief_reproduction': rep, 'vol_target_c': {n: res[n]['c'] for n in res},
                   'live_fund_monthly_abs_gt_30pct': anomalies},
        'n_graded': len(res), 'tested': tested,
        'strategies': {n: strip(e) for n, e in res.items()},
        'jp_tax': jp_tax, 'vrp': vrp,
        'reality_gap': {'funds': funds, 'composite': comps, 'paper_same_window_as_composite': paper_same},
        'japan_investability': {'toushin_lib': jp, 'tsumitate_lineup_hits': tl_hits, 'rakuten_us_etf_lineup': bl_check},
        'runtime_sec': round(time.time() - t0, 1), 'log_tail': LOG[-80:],
    }
    p = M.save(OUT, out)
    log('saved', p, os.path.getsize(p))


if __name__ == '__main__':
    main()
