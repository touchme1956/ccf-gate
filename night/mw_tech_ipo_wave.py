#!/usr/bin/env python3
"""night/mw_tech_ipo_wave.py — 『市場に勝てる歴史検証』角度 tech_ipo_wave（読むだけ・門の判定には不使用）

問い: テック・VC 出資の株の新しい供給（Ritter の IPO 件数）が熱い間は、ハイテク側（NASDAQ100・半導体の器）の
      新規資金（R3・売らない＝NISA 向け）や持ち分（R1 全部・R2 半分）を市場全体へ逃がすと、
      今の規則（ハイテクへ積立・持ち続ける）と市場への積立に勝つか。集中の頂点の後の反転を避けられるか。

事前登録: out/mw_tech_ipo_wave_prereg.json（規則・線は測る前に固定。ここで動かさない）
出力    : out/mw_tech_ipo_wave.json

約束: 月次リターンは小数。総リターンどうしで比べる（ハイテクの器・French Mkt とも総リターン）。
      読み H_s は月 s の終わりまでに分かる値だけ。月 t の持ち方は H_{t−12}〜H_{t−1}（t の値は使わない）。
      欠測を0で埋めない（リターンの欠けた月は評価から外す。Internet の印は 2021-12 で打ち切る）。
"""
import bisect, collections, datetime, io, json, math, os, re, subprocess, sys, statistics as S

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import mw_common as M  # noqa: E402

PREREG = 'mw_tech_ipo_wave_prereg.json'
PREREG2 = 'mw_tech_ipo_wave_prereg2.json'   # 探索の族 X（上場の波 × 相対トレンド）
OUT = 'mw_tech_ipo_wave.json'
END = 202608                  # French の終わり
INT_SIG_END = 202112          # Internet の印を使う最後の月（Ritter が近年更新していない）
INT_POS_END = 202201          # その読みで決まる持ち方の最後の月
COST = 0.001                  # 持ち替え1回（資産の100%）あたり 0.10%
COST_HI = 0.002               # 報告: 0.20%
TAX = 0.20315
NDX_DIV = 0.005               # 1999-03 までの ^NDX に足す推定配当（mw_trend と同じ）
Q5, T3 = 0.80, 2 / 3
MIN_N = 60
DCA_N = 240
LOG = []


def log(*a):
    s = ' '.join(str(x) for x in a)
    LOG.append(s)
    print(s, flush=True)


def ym_add(k, n):
    y, m = divmod(k, 100)
    t = y * 12 + (m - 1) + n
    return (t // 12) * 100 + t % 12 + 1


def months(a, z):
    out, k = [], a
    while k <= z:
        out.append(k)
        k = ym_add(k, 1)
    return out


def git_sha(path):
    try:
        return subprocess.check_output(['git', '-C', M.BASE, 'log', '-1', '--format=%H', '--', path], text=True).strip() or None
    except Exception:  # noqa
        return None


# ───────────────────────── データ ─────────────────────────
def load_ipo_age():
    """Ritter IPO-age.xlsx → 月ごとの VC 件数・Internet 件数・不明の件数（完全な一覧なので IPO の無い月の 0 は本当の 0）"""
    import openpyxl
    b = M.get('https://site.warrington.ufl.edu/ritter/files/IPO-age.xlsx', 'ritter_IPO-age.xlsx', max_age_days=365)
    rows = list(openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True))
    assert rows[0][0] == 'offer date' and rows[0][5] == 'VC' and rows[0][8] == 'Internet', rows[0]
    ms = months(197501, 202512)
    vc = {k: 0 for k in ms}; it = {k: 0 for k in ms}; vcu = {k: 0 for k in ms}; itu = {k: 0 for k in ms}; n = {k: 0 for k in ms}
    for r in rows[1:]:
        d = int(r[0]); k = d // 100
        n[k] += 1
        if r[5] in (1, 2):
            vc[k] += 1
        elif r[5] != 0:
            vcu[k] += 1
        if r[8] == 1:
            it[k] += 1
        elif r[8] != 0:
            itu[k] += 1
    return {'vc': vc, 'int': it, 'vc_unknown': vcu, 'int_unknown': itu, 'n': n, 'rows': len(rows) - 1}


def load_ipoall():
    import openpyxl
    b = M.get('https://site.warrington.ufl.edu/ritter/files/IPOALL.xlsx', 'ritter_IPOALL.xlsx', max_age_days=365)
    rows = list(openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True))
    gross, net, fdr = {}, {}, {}
    prev = None
    for r in rows:
        if not (isinstance(r[0], (int, float)) and isinstance(r[1], (int, float))):
            continue
        mo, yy = int(r[0]), int(r[1])
        y = 1900 + yy if yy >= 60 else 2000 + yy
        k = y * 100 + mo
        assert prev is None or k == ym_add(prev, 1), (prev, k)
        prev = k
        gross[k] = int(r[3])
        if y >= 1975 and isinstance(r[4], (int, float)):
            net[k] = int(r[4])
        if isinstance(r[2], (int, float)):
            fdr[k] = float(r[2])   # '.' = その月に IPO が無い → 値なし（0 と読まない）
    return {'gross': gross, 'net': net, 'fdr': fdr}


def load_table4b(pre):
    t = {int(y): v for y, v in pre['data']['ritter_table4b']['year: [tech, life_sci, other]'].items()}
    sums = [sum(v[i] for v in t.values()) for i in range(3)]
    chk = {'sum_tech': sums[0], 'sum_life': sums[1], 'sum_other': sums[2], 'expected': [3365, 1031, 4947], 'ok': sums == [3365, 1031, 4947]}
    # pypdf があれば PDF から読み直して一致を確かめる（無ければ合計の検算だけ）
    try:
        sys.modules.setdefault('cryptography', None)
        sys.path.insert(0, os.path.join(M.CACHE, 'tiw_pylib'))
        import logging
        logging.disable(logging.CRITICAL)
        from pypdf import PdfReader
        b = M.get('https://site.warrington.ufl.edu/ritter/files/IPOs-LifeScience.pdf', 'ritter_IPOs-LifeScience.pdf', max_age_days=365)
        txt = PdfReader(io.BytesIO(b)).pages[2].extract_text()
        re_t = {}
        for ln in txt.splitlines():
            m = re.match(r'\s*(19[89]\d|20[0-2]\d)\s+(\d+)\s+(\d+)\s+(\d+)\s', ln)
            if m:
                re_t[int(m.group(1))] = [int(m.group(2)), int(m.group(3)), int(m.group(4))]
        chk['pdf_reparse_equal'] = re_t == t
    except Exception as e:  # noqa
        chk['pdf_reparse_equal'] = f'未確認（{type(e).__name__}）'
    return t, chk


def load_french():
    t = M.french_tables('49_Industry_Portfolios')
    vw = t['Average Value Weighted Returns -- Monthly']
    cols = vw['cols']
    ret = {c: {} for c in cols}
    for d, row in vw['data'].items():
        for c, x in zip(cols, row):
            if x is not None:
                ret[c][d] = x / 100
    cnt_t = t['Number of Firms in Portfolios']; size_t = t['Average Firm Size']
    assert cnt_t['cols'] == cols and size_t['cols'] == cols
    cnt = {c: {} for c in cols}; size = {c: {} for c in cols}
    for d, row in cnt_t['data'].items():
        for c, x in zip(cols, row):
            if x is not None:
                cnt[c][d] = x
    for d, row in size_t['data'].items():
        for c, x in zip(cols, row):
            if x is not None:
                size[c][d] = x
    return cols, ret, cnt, size


def make_t3(ret, cnt):
    """T3 = Hardw・Softw・Chips のうち前月の銘柄数 ≥ 5 の業種を等分（総リターン）"""
    out, used = {}, {}
    for k in sorted(ret['Hardw']):
        p = ym_add(k, -1)
        inds = [c for c in ('Hardw', 'Softw', 'Chips') if cnt[c].get(p, 0) >= 5 and k in ret[c]]
        if len(inds) >= 2:
            out[k] = sum(ret[c][k] for c in inds) / len(inds)
            used[k] = len(inds)
    return out, used


def make_ndx():
    px = M.yahoo('^NDX', '1mo'); q = M.yahoo('QQQ', '1mo')
    q0 = min(q)
    dd = (1 + NDX_DIV) ** (1 / 12) - 1
    out = {}
    for k, v in px.items():
        if k < q0:
            out[k] = (1 + v) * (1 + dd) - 1
    for k, v in q.items():
        out[k] = v
    out = {k: v for k, v in out.items() if k <= END}
    chk = {'qqq_from': q0, 'ndx_from': min(px), 'cagr_qqq_minus_ndx_price_1999_04+': round((M.cagr(M.window(q, q0, END)) - M.cagr(M.window(px, q0, END))) * 100, 2)}
    return out, chk


# ───────────────────────── 信号 ─────────────────────────
def trailing_sum(series, start, end, n=12):
    out = {}
    for s in months(ym_add(start, n - 1), end):
        w = [series.get(ym_add(s, -j)) for j in range(n)]
        if any(x is None for x in w):
            continue
        out[s] = sum(w)
    return out


def trailing_mean_avail(series, start, end, n=12, need=6):
    out = {}
    for s in months(ym_add(start, n - 1), end):
        w = [series[ym_add(s, -j)] for j in range(n) if ym_add(s, -j) in series]
        if len(w) >= need:
            out[s] = sum(w) / len(w)
    return out


def expanding_pct(X, min_n=MIN_N):
    """広がる窓の中位順位（s を含む・s 以前の値だけ）。N ≥ min_n の月だけ"""
    srt, out = [], {}
    for s in sorted(X):
        bisect.insort(srt, X[s])
        n = len(srt)
        if n < min_n:
            continue
        lo = bisect.bisect_left(srt, X[s]); hi = bisect.bisect_right(srt, X[s])
        out[s] = (lo + 0.5 * (hi - lo)) / n
    return out


def readings(pct, thr):
    return {s: p >= thr for s, p in pct.items()}


def state_of(H, t0, t1, dur=12, lag=0):
    """hot_state(t) = H_{t−lag−dur}〜H_{t−lag−1} のどれかが熱い。t0 は最初の読みの翌月以降。返り値 {t: bool} と使った読みの月の最大（検算用）"""
    st, used_max = {}, {}
    first = min(H)
    for t in months(max(t0, ym_add(first, 1 + lag)), t1):
        src = [ym_add(t, -lag - j) for j in range(1, dur + 1)]
        st[t] = any(H.get(s, False) for s in src)
        used_max[t] = max(s for s in src if s in H) if any(s in H for s in src) else None
    return st, used_max


def runs(state):
    out, cur = [], None
    for t in sorted(state):
        if state[t]:
            if cur and ym_add(cur[1], 1) == t:
                cur[1] = t
            else:
                cur = [t, t]; out.append(cur)
    return [tuple(x) for x in out]


# ───────────────────────── 戦略 ─────────────────────────
def timing(tech, mkt, state, wh, cost=COST):
    """R1（wh=0）・R2（wh=0.5）: 月初に目標の割合へ。費用 = 動かした割合 × cost。最初の月は『ハイテクを全部持っている』から始める"""
    g, n, w_of, turn = {}, {}, {}, {}
    wd = 1.0
    prev = None
    miss = 0
    for t in sorted(state):
        if t not in tech or t not in mkt:
            miss += 1
            continue
        w = wh if state[t] else 1.0
        tv = abs(w - wd)
        rt = w * tech[t] + (1 - w) * mkt[t]
        g[t] = rt; n[t] = rt - tv * cost; w_of[t] = w; turn[t] = tv
        den = w * (1 + tech[t]) + (1 - w) * (1 + mkt[t])
        wd = w * (1 + tech[t]) / den if den else w
        prev = t
    return g, n, w_of, turn, miss


def sub(d, keys):
    return {k: d[k] for k in keys if k in d}


def eval_timing(name, family, desc, tech, mkt, rf, state, wh, sig_meta):
    g, n, w, tv, miss = timing(tech, mkt, state, wh)
    ks = sorted(g)
    res = {'name': name, 'family': family, 'desc': desc, 'rule': 'R1' if wh == 0 else 'R2', 'from': ks[0], 'to': ks[-1],
           'months': len(ks), 'missing_months_skipped': miss, 'hot_share': round(sum(1 for k in ks if state[k]) / len(ks), 3),
           'switches_per_year': round(sum(tv.values()) / (len(ks) / 12), 2), 'signal': sig_meta, 'cmp': {}}
    for bn, b in (('tech', tech), ('mkt', mkt)):
        bb = sub(b, ks)
        full = M.excess_stats(g, bb); train = M.excess_stats(g, bb, z=M.TRAIN_END); hold = M.excess_stats(g, bb, a=M.HOLD_START)
        recent = M.excess_stats(g, bb, a=M.RECENT_START)
        full_net = M.excess_stats(n, bb); hold_net = M.excess_stats(n, bb, a=M.HOLD_START); train_net = M.excess_stats(n, bb, z=M.TRAIN_END)
        roll20 = M.rolling(n, bb, 20); dca20 = M.dca(n, bb, 20)
        sp = {'train': (M.sharpe(n, rf, z=M.TRAIN_END), M.sharpe(bb, rf, z=M.TRAIN_END)),
              'hold': (M.sharpe(n, rf, a=M.HOLD_START), M.sharpe(bb, rf, a=M.HOLD_START))}
        post = {'BW2000_2001+': M.excess_stats(n, bb, a=200101), 'PV2005_2006+': M.excess_stats(n, bb, a=200601)}
        hi = timing(tech, mkt, state, wh, cost=COST_HI)[1]
        res['cmp'][bn] = {'full': full, 'train': train, 'hold': hold, 'recent': recent, 'full_net': full_net, 'train_net': train_net,
                          'hold_net': hold_net, 'hold_net_cost0.20': M.excess_stats(hi, bb, a=M.HOLD_START),
                          'roll20_net': roll20, 'dca20_net': dca20, 'sharpe_pair_net': sp, 'post_publication_net': post,
                          'maxdd_s': round(M.maxdd(n) * 100, 1), 'maxdd_b': round(M.maxdd(bb) * 100, 1)}
    res['_series'] = (g, n, w, tv)
    return res


def grade_family(strats, comp):
    ps = {s['name']: (s['cmp'][comp]['hold_net'] or {}).get('p') for s in strats}
    hp = M.holm(ps)
    for s in strats:
        c = s['cmp'][comp]
        gr, cr = M.grade(c['full'], c['train'], c['hold'], c['roll20_net'], cost_hold=c['hold_net'], repl=None,
                         family_holm_p=hp.get(s['name']), sharpe_pair=c['sharpe_pair_net'], leveraged_or_timing=True)
        c['family_holm_p'] = hp.get(s['name'])
        c['grade'] = gr
        c['criteria'] = cr
        s['grade_vs_' + comp] = gr


# ───────────────────────── 積立（R3） ─────────────────────────
def pctl(v, q):
    v = sorted(v)
    return v[int(q * (len(v) - 1))]


def irr_m(W, n):
    """毎月初に 1 を n か月入れて最終額 W になる月利 m（年率で返す）"""
    f = lambda m: (n if abs(m) < 1e-12 else (1 + m) * ((1 + m) ** n - 1) / m) - W
    lo, hi = -0.05, 0.08
    for _ in range(80):
        mid = (lo + hi) / 2
        if f(mid) > 0:
            hi = mid
        else:
            lo = mid
    return (1 + (lo + hi) / 2) ** 12 - 1


def dca_run(tech, mkt, state, ms):
    c = 1 - COST
    bT = bM = wT = wM = 0.0
    for k in ms:
        if state[k]:
            bM += c
        else:
            bT += c
        bT *= 1 + tech[k]; bM *= 1 + mkt[k]
        wT = (wT + c) * (1 + tech[k]); wM = (wM + c) * (1 + mkt[k])
    return bT + bM, wT, wM


def eval_dca(name, family, desc, tech, mkt, state, t_end, sig_meta):
    ms_all = [t for t in sorted(state) if t <= t_end]
    ms_all = [t for t in ms_all if t in tech and t in mkt]
    # 連続した区間だけ（欠けがあればそこで切る）
    ok = all(ym_add(a, 1) == b for a, b in zip(ms_all, ms_all[1:]))
    assert ok, f'{name}: 月が連続していない'
    W = []
    for i in range(len(ms_all) - DCA_N + 1):
        w = ms_all[i:i + DCA_N]
        r3, wt, wm = dca_run(tech, mkt, state, w)
        W.append({'start': w[0], 'end': w[-1], 'r3': r3, 'tech': wt, 'mkt': wm, 'hot_months': sum(1 for k in w if state[k])})
    res = {'name': name, 'family': family, 'desc': desc, 'rule': 'R3', 'windows': len(W), 'from': ms_all[0], 'to': ms_all[-1], 'signal': sig_meta, 'cmp': {}}
    for bn in ('tech', 'mkt'):
        tr = [x['r3'] / x[bn] for x in W if x['end'] <= M.TRAIN_END]
        ho = [x['r3'] / x[bn] for x in W if x['end'] >= M.HOLD_START]
        allr = [x['r3'] / x[bn] for x in W]
        d1 = bool(tr) and S.median(tr) > 1 and sum(1 for r in tr if r > 1) / len(tr) >= 0.8
        d2 = bool(ho) and S.median(ho) > 1 and sum(1 for r in ho if r > 1) / len(ho) >= 0.8
        p5s, p5b = pctl([x['r3'] / DCA_N for x in W], 0.05), pctl([x[bn] / DCA_N for x in W], 0.05)
        d3 = p5s >= p5b
        irr_d = sorted(round((irr_m(x['r3'], DCA_N) - irr_m(x[bn], DCA_N)) * 100, 2) for x in W)
        summ = lambda v: None if not v else {'n': len(v), 'median_ratio': round(S.median(v), 3), 'win_rate': round(sum(1 for r in v if r > 1) / len(v), 3),
                                             'worst': round(min(v), 3), 'best': round(max(v), 3)}
        res['cmp'][bn] = {'train_windows': summ(tr), 'hold_windows': summ(ho), 'all_windows': summ(allr),
                          'worst_window': min(((x['start'], round(x['r3'] / x[bn], 3)) for x in W), key=lambda z: z[1]) if W else None,
                          'p5_final_over_contrib': {'r3': round(p5s, 3), 'bench': round(p5b, 3)},
                          'irr_diff_pct': {'median': irr_d[len(irr_d) // 2], 'worst': irr_d[0], 'best': irr_d[-1]} if irr_d else None,
                          'D1_train': d1, 'D2_hold': d2, 'D3_risk': d3, 'D4_repl': None,
                          'judgement': '積立で勝ち（再現なし・格付け外）' if (d1 and d2 and d3) else '不合格'}
        # 報告: 10年窓（起点 2007-01 以降）
        ms10 = [t for t in ms_all if t >= M.HOLD_START]
        r10 = []
        for i in range(len(ms10) - 120 + 1):
            a, b, c = dca_run(tech, mkt, state, ms10[i:i + 120])
            r10.append(a / (b if bn == 'tech' else c))
        res['cmp'][bn]['report_10y_windows_from_2007'] = summ(r10)
    res['judgement_vs_tech'] = res['cmp']['tech']['judgement']
    res['judgement_vs_mkt'] = res['cmp']['mkt']['judgement']
    return res


# ───────────────────────── 報告 ─────────────────────────
def nw_slope(y, h, lag):
    n = len(y)
    if n < 24 or sum(h) in (0, n):
        return None, None
    hb = sum(h) / n; yb = sum(y) / n
    sxx = sum((x - hb) ** 2 for x in h)
    b = sum((x - hb) * (v - yb) for x, v in zip(h, y)) / sxx
    a = yb - b * hb
    u = [(x - hb) * (v - a - b * x) for x, v in zip(h, y)]
    g0 = sum(x * x for x in u) / n
    s = g0
    for L in range(1, min(lag, n - 1) + 1):
        s += 2 * (1 - L / (lag + 1)) * sum(u[i] * u[i - L] for i in range(L, n)) / n
    se = math.sqrt(n * s) / sxx if s > 0 else None
    return b, (b / se if se else None)


def predictive(H, tech, mkt, k, a=None, z=None, train=False):
    ys, hs = [], []
    for s in sorted(H):
        if a is not None and s < a:
            continue
        fw = [ym_add(s, j) for j in range(1, k + 1)]
        if fw[-1] > END or any(f not in tech or f not in mkt for f in fw):
            continue
        if train and fw[-1] > M.TRAIN_END:
            continue
        if z is not None and s > z:
            continue
        ys.append(sum(math.log1p(tech[f]) - math.log1p(mkt[f]) for f in fw))
        hs.append(1.0 if H[s] else 0.0)
    b, t = nw_slope(ys, hs, k)
    if b is None:
        return {'n': len(ys), 'hot_n': int(sum(hs)), 'diff_pct_per_yr': None, 't_nw': None}
    mh = S.mean([y for y, h in zip(ys, hs) if h]) * 12 / k * 100
    mc = S.mean([y for y, h in zip(ys, hs) if not h]) * 12 / k * 100
    return {'n': len(ys), 'hot_n': int(sum(hs)), 'hot_mean_pct_per_yr': round(mh, 2), 'cold_mean_pct_per_yr': round(mc, 2),
            'diff_pct_per_yr': round(b * 12 / k * 100, 2), 't_nw': round(t, 2) if t is not None else None}


def binom_ge(k, n):
    return round(sum(math.comb(n, i) for i in range(k, n + 1)) / 2 ** n, 4) if n else None


def episodes(state, tech, mkt):
    eps = []
    for a, b in runs(state):
        ms = [t for t in months(a, b) if t in tech and t in mkt]
        during = sum(math.log1p(tech[t]) - math.log1p(mkt[t]) for t in ms)
        f36 = [ym_add(a, j) for j in range(36)]
        f36v = sum(math.log1p(tech[t]) - math.log1p(mkt[t]) for t in f36) if all(t in tech and t in mkt for t in f36) else None
        eps.append({'start': a, 'end': b, 'months': len(ms), 'tech_minus_mkt_during_logpct': round(during * 100, 1),
                    'tech_minus_mkt_36m_from_start_logpct': round(f36v * 100, 1) if f36v is not None else None,
                    'hit_market_won_during': during < 0})
    out = {}
    for lab, sel in (('train', [e for e in eps if e['start'] <= M.TRAIN_END]), ('hold', [e for e in eps if e['start'] >= M.HOLD_START]), ('all', eps)):
        h = sum(1 for e in sel if e['hit_market_won_during'])
        out[lab] = {'episodes': len(sel), 'hits': h, 'binom_p_one_sided': binom_ge(h, len(sel))}
    out['list'] = eps
    return out


def tax_sim(tech, mkt, state, wh, a=M.HOLD_START, z=END):
    """報告: 課税口座（総平均法・同じ暦年の中だけ通算・繰越なし・年末に税を資産から払う）。最後に全部売る"""
    ks = [t for t in sorted(state) if a <= t <= z and t in tech and t in mkt]
    vT, vM, bT, bM = 1.0, 0.0, 1.0, 0.0   # 2006-12 末にハイテクを 1（取得原価 1）持っている
    realized = 0.0
    hv, hb = 1.0, 1.0                     # 買って持つだけ
    for i, t in enumerate(ks):
        tot = vT + vM
        w = wh if state[t] else 1.0
        tgtT = w * tot
        if tgtT < vT - 1e-15:   # ハイテクを売る
            s = vT - tgtT
            realized += s - bT * s / vT
            bT *= (vT - s) / vT; vT -= s
            buy = s * (1 - COST); vM += buy; bM += buy
        elif tgtT > vT + 1e-15 and vM > 0:
            s = min(vM, tgtT - vT)
            realized += s - bM * s / vM
            bM *= (vM - s) / vM; vM -= s
            buy = s * (1 - COST); vT += buy; bT += buy
        vT *= 1 + tech[t]; vM *= 1 + mkt[t]; hv *= 1 + tech[t]
        last = (i == len(ks) - 1)
        if t % 100 == 12 or last:
            if realized > 0:
                tax = realized * TAX
                tot = vT + vM
                fT = vT / tot
                vT -= tax * fT; vM -= tax * (1 - fT)   # 税を持ち分の比で払う（原価の調整は省く＝近似）
            realized = 0.0
    pre = vT + vM
    post = pre - max(0.0, (vT - bT) + (vM - bM)) * TAX
    hpost = hv - max(0.0, hv - hb) * TAX
    return {'from': ks[0], 'to': ks[-1], 'strategy_after_tax': round(post, 3), 'buyhold_after_tax': round(hpost, 3),
            'ratio_after_tax': round(post / hpost, 3)}


def main():
    pre = json.load(open(os.path.join(M.BASE, 'out', PREREG)))
    sanity = {}
    ff = M.ff_factors(); mkt = M.window(ff['mkt'], None, END); rf = ff['rf']
    sanity['us_mkt_cagr_full'] = round(M.cagr(mkt) * 100, 2)
    sanity['us_mkt_cagr_2007+'] = round(M.cagr(M.window(mkt, M.HOLD_START)) * 100, 2)
    log('French Mkt CAGR', sanity['us_mkt_cagr_full'], '2007〜', sanity['us_mkt_cagr_2007+'])

    ipo = load_ipo_age(); allx = load_ipoall(); t4b, t4chk = load_table4b(pre)
    sanity['table4b'] = t4chk
    sanity['ipo_age_rows'] = ipo['rows']
    cols, ret, cnt, size = load_french()
    t3, t3n = make_t3(ret, cnt); t3 = M.window(t3, None, END)
    ndx, ndxchk = make_ndx()
    sanity['ndx'] = ndxchk
    sanity['t3_industries_used'] = {'2_until': max(k for k, v in t3n.items() if v == 2), '3_from': min(k for k, v in t3n.items() if v == 3)}
    sanity['t3_missing_after_1965'] = [k for k in months(196512, END) if k not in t3]
    sanity['mkt_missing'] = [k for k in months(196001, END) if k not in mkt]
    log('T3', min(t3), max(t3), 'NDX', min(ndx), max(ndx), ndxchk)

    # 不明の VC / Internet の件数（直近12か月の最大）
    vcu12 = trailing_sum(ipo['vc_unknown'], 197501, 202512); itu12 = trailing_sum(ipo['int_unknown'], 197501, 202512)
    sanity['unknown_flags_max_per_12m'] = {'vc': max(vcu12.values()), 'int': max(itu12.values())}
    sanity['unknown_flags_by_year'] = {y: [sum(ipo['vc_unknown'][k] for k in months(y * 100 + 1, y * 100 + 12)),
                                           sum(ipo['int_unknown'][k] for k in months(y * 100 + 1, y * 100 + 12))] for y in range(1975, 2026)
                                       if sum(ipo['vc_unknown'][k] + ipo['int_unknown'][k] for k in months(y * 100 + 1, y * 100 + 12))}

    # ── 信号
    X = {}
    X['VC'] = trailing_sum(ipo['vc'], 197501, 202512)
    X['INT'] = trailing_sum({k: v for k, v in ipo['int'].items() if k <= INT_SIG_END}, 197501, INT_SIG_END)
    X['NET'] = trailing_sum(allx['net'], 197501, 202512)
    X['GROSS'] = trailing_sum(allx['gross'], 196001, 202512)
    X['FDR'] = trailing_mean_avail(allx['fdr'], 196001, 202512)
    tech_n = {k: sum(cnt[c].get(k, 0) for c in ('Hardw', 'Softw', 'Chips')) for k in cnt['Hardw']}
    X['TNL'] = {s: tech_n[s] / tech_n[ym_add(s, -12)] - 1 for s in months(197407, END) if s in tech_n and ym_add(s, -12) in tech_n and tech_n[ym_add(s, -12)] > 0}
    pct = {k: expanding_pct(v) for k, v in X.items()}
    # 年次（Table 4b）: 読みは12月だけ
    def annual_pct(col):
        srt, out = [], {}
        for y in sorted(t4b):
            v = t4b[y][col]
            bisect.insort(srt, v)
            if len(srt) >= 5:
                lo = bisect.bisect_left(srt, v); hi = bisect.bisect_right(srt, v)
                out[y * 100 + 12] = (lo + 0.5 * (hi - lo)) / len(srt)
        return out
    pct['TECHY'] = annual_pct(0)
    pct['LIFEY'] = annual_pct(1)
    H = {}
    for k in pct:
        H[k + '_Q5'] = readings(pct[k], Q5)
        H[k + '_T3'] = readings(pct[k], T3)
    sig_meta = {}
    for k, h in H.items():
        rr = runs({s: v for s, v in h.items()})
        sig_meta[k] = {'first_reading': min(h), 'last_reading': max(h), 'readings': len(h), 'hot_readings': sum(1 for v in h.values() if v),
                       'hot_runs_of_readings': [f'{a}-{b}' for a, b in rr]}

    # ── 状態（持ち方）
    def st(sig, t_end=END, dur=12, lag=0):
        h = H[sig]
        t_end = min(t_end, INT_POS_END) if sig.startswith('INT') else t_end
        s, used = state_of(h, 196001, t_end, dur, lag)
        # 検算: 使った読みの月 < t
        assert all(u is None or u < t for t, u in used.items())
        return s

    tested = []
    strats = {'P': [], 'N': [], 'E': []}
    for fam, tech, tname in (('P', t3, 'T3'), ('N', ndx, 'NDX')):
        i = 0
        for thr in ('Q5', 'T3'):
            for sig in ('VC', 'INT'):
                for wh in (0.0, 0.5):
                    i += 1
                    nm = f'{fam}{i}_{"R1" if wh == 0 else "R2"}_{sig}_{thr}'
                    s = st(f'{sig}_{thr}')
                    s = {t: v for t, v in s.items() if t in tech}
                    r = eval_timing(nm, {'P': '主（T3）', 'N': '副（NDX）'}[fam], f'{"R1 全部" if wh == 0 else "R2 半分"}を市場へ・信号 {sig}・閾値 {thr}・器 {tname}',
                                    tech, mkt, rf, s, wh, sig_meta[f'{sig}_{thr}'])
                    r['primary'] = (fam == 'P' and thr == 'Q5')
                    strats[fam].append(r)
    # 名前を事前登録の並びに合わせる（R1 VC Q5, R2 VC Q5, R1 INT Q5, R2 INT Q5, … ）
    order = [('VC', 'Q5', 'R1'), ('VC', 'Q5', 'R2'), ('INT', 'Q5', 'R1'), ('INT', 'Q5', 'R2'), ('VC', 'T3', 'R1'), ('VC', 'T3', 'R2'), ('INT', 'T3', 'R1'), ('INT', 'T3', 'R2')]
    for fam in ('P', 'N'):
        byk = {(r['name'].split('_')[2], r['name'].split('_')[3], r['name'].split('_')[1]): r for r in strats[fam]}
        strats[fam] = []
        for j, (sg, th, rl) in enumerate(order, 1):
            r = byk[(sg, th, rl)]
            r['name'] = f'{fam}{j}_{rl}_{sg}_{th}'
            strats[fam].append(r)
    for j, sig in enumerate(('TECHY', 'NET', 'GROSS', 'FDR', 'TNL'), 1):
        s = st(f'{sig}_Q5')
        s = {t: v for t, v in s.items() if t in t3}
        r = eval_timing(f'E{j}_R1_{sig}_Q5', '探索（T3）', f'R1 全部を市場へ・信号 {sig}・Q5・器 T3', t3, mkt, rf, s, 0.0, sig_meta[f'{sig}_Q5'])
        r['primary'] = False
        strats['E'].append(r)
    # ── 探索族 X（事前登録2）: 上場の波 × ハイテクの相対トレンドの崩れ
    def rel_trend(tech):
        ks = sorted(k for k in tech if k in mkt)
        assert all(ym_add(a, 1) == b for a, b in zip(ks, ks[1:])), '相対指数の月が連続していない'
        ri, v = {}, 1.0
        for k in ks:
            v *= (1 + tech[k]) / (1 + mkt[k]); ri[k] = v
        return {k: ri[k] < sum(ri[ks[j]] for j in range(i - 9, i + 1)) / 10 for i, k in enumerate(ks) if i >= 9}
    td = {'T3': rel_trend(t3), 'NDX': rel_trend(ndx)}

    def xstate(sig, tn, t0):
        s = st(sig) if sig else None
        out = {}
        for t in months(t0, END):
            p_ = ym_add(t, -1)
            if p_ not in td[tn] or (s is not None and t not in s):
                continue
            out[t] = (s[t] and td[tn][p_]) if s is not None else td[tn][p_]
        return out
    xspec = [('X1_R1_VC_Q5_TD_T3', 'VC_Q5', 'T3', 198012), ('X2_R1_TD_T3', None, 'T3', 196512), ('X3_R1_GROSS_Q5_TD_T3', 'GROSS_Q5', 'T3', 196512),
             ('X4_R1_TNL_Q5_TD_T3', 'TNL_Q5', 'T3', 197907), ('X5_R1_VC_Q5_TD_NDX', 'VC_Q5', 'NDX', 198609), ('X6_R1_TD_NDX', None, 'NDX', 198609)]
    strats['X'] = []
    xstates = {}
    for nm, sig, tn, t0 in xspec:
        tech = t3 if tn == 'T3' else ndx
        s = xstate(sig, tn, t0)
        assert min(s) == t0, (nm, min(s), t0)
        xstates[nm] = s
        r = eval_timing(nm, '探索 X（上場の波×相対トレンド・事前登録2）', f'R1・{"信号 " + sig + " かつ " if sig else "【対照】"}ハイテクの相対指数が10か月平均の下・器 {tn}',
                        tech, mkt, rf, s, 0.0, sig_meta[sig] if sig else {'trend_only': True})
        r['primary'] = False
        strats['X'].append(r)
    for fam in strats:
        for comp in ('tech', 'mkt'):
            grade_family(strats[fam], comp)
    for fam in strats:
        for r in strats[fam]:
            log(r['name'], r['from'], r['to'], 'hot', r['hot_share'], 'vsTech', r['grade_vs_tech'], (r['cmp']['tech']['hold_net'] or {}).get('ex_ann'),
                'vsMkt', r['grade_vs_mkt'], (r['cmp']['mkt']['hold_net'] or {}).get('ex_ann'), (r['cmp']['mkt']['train'] or {}).get('ex_ann'), (r['cmp']['mkt']['train'] or {}).get('t'))

    # ── 積立（R3）
    dres = []
    j = 0
    for tech, tname in ((t3, 'T3'), (ndx, 'NDX')):
        for thr in ('Q5', 'T3'):
            for sig in ('VC', 'INT'):
                j += 1
                s = st(f'{sig}_{thr}')
                s = {t: v for t, v in s.items() if t in tech}
                t_end = INT_POS_END if sig == 'INT' else END
                dres.append(eval_dca(f'D{j}_R3_{sig}_{thr}_{tname}', '積立（D判定）', f'R3 新規資金だけ市場へ・信号 {sig}・閾値 {thr}・器 {tname}', tech, mkt, s, t_end, sig_meta[f'{sig}_{thr}']))
    for j, sig in enumerate(('TECHY', 'NET', 'GROSS', 'FDR', 'TNL'), 1):
        s = st(f'{sig}_Q5'); s = {t: v for t, v in s.items() if t in t3}
        dres.append(eval_dca(f'E{j}D_R3_{sig}_Q5', '探索の積立（D判定）', f'R3・信号 {sig}・Q5・器 T3', t3, mkt, s, END, sig_meta[f'{sig}_Q5']))
    for (nm, sig, tn, t0) in xspec:
        tech = t3 if tn == 'T3' else ndx
        nmd = nm.replace('_R1_', 'D_R3_', 1)
        dres.append(eval_dca(nmd, '探索 X の積立（D判定・事前登録2）', f'R3・{"信号 " + sig + " かつ " if sig else "【対照】"}相対トレンドの下・器 {tn}', tech, mkt, xstates[nm], END,
                             sig_meta[sig] if sig else {'trend_only': True}))
    for d in dres:
        log(d['name'], d['windows'], 'vsTech', d['judgement_vs_tech'], d['cmp']['tech']['train_windows'], d['cmp']['tech']['hold_windows'],
            'vsMkt', d['judgement_vs_mkt'], d['cmp']['mkt']['hold_windows'])

    # ── 検算: 常に熱くない状態なら R1 の超過 0・R3 の比 1
    s0 = {t: False for t in months(198101, END)}
    g0 = timing(t3, mkt, s0, 0.0)[0]
    sanity['never_hot_R1_excess_vs_tech'] = M.excess_stats(g0, sub(t3, g0))['ex_ann']
    r3, wt, _ = dca_run(t3, mkt, s0, months(198101, 200012))
    sanity['never_hot_R3_ratio'] = round(r3 / wt, 6)

    # ── 報告
    reports = {}
    # 予言の検定
    pr = {}
    for sig in ('VC_Q5', 'INT_Q5', 'VC_T3', 'INT_T3', 'TECHY_Q5', 'NET_Q5', 'GROSS_Q5', 'FDR_Q5', 'TNL_Q5', 'LIFEY_Q5'):
        h = H[sig]
        if sig.startswith('INT'):
            h = {s: v for s, v in h.items() if s <= INT_SIG_END}
        for tech, tname in ((t3, 'T3'), (ndx, 'NDX')):
            for k in (12, 36):
                pr[f'{sig}|{tname}|k{k}'] = {'full': predictive(h, tech, mkt, k), 'train': predictive(h, tech, mkt, k, train=True),
                                             'hold': predictive(h, tech, mkt, k, a=M.HOLD_START)}
    reports['F_predictive'] = pr
    # 独立の波
    ep = {}
    for sig in ('VC_Q5', 'INT_Q5', 'VC_T3', 'INT_T3', 'TECHY_Q5', 'NET_Q5', 'GROSS_Q5', 'FDR_Q5', 'TNL_Q5'):
        for tech, tname in ((t3, 'T3'), (ndx, 'NDX')):
            s = st(sig); s = {t: v for t, v in s.items() if t in tech}
            ep[f'{sig}|{tname}'] = episodes(s, tech, mkt)
    reports['episodes'] = ep
    # 集中の頂点
    cap = {}
    for k in cnt['Hardw']:
        tot = 0.0; tc = 0.0; ok = True
        for c in cols:
            if k in cnt[c] and k in size[c]:
                v = cnt[c][k] * size[c][k]
                tot += v
                if c in ('Hardw', 'Softw', 'Chips'):
                    tc += v
        if tot > 0:
            cap[k] = tc / tot
    ks = sorted(cap); rec = []
    mx = -1
    for i, k in enumerate(ks):
        if i >= 120 and cap[k] > mx:
            rec.append(k)
        mx = max(mx, cap[k])
    conc = {'tech_weight_first': ks[0], 'records': len(rec), 'record_months_since_1965': [k for k in rec if k >= 196501]}
    for sig in ('VC_Q5', 'INT_Q5', 'GROSS_Q5', 'FDR_Q5', 'TECHY_Q5', 'TNL_Q5'):
        s = st(sig)
        rr = [k for k in rec if k in s]
        f60 = {}
        for k in rr:
            fw = [ym_add(k, j) for j in range(1, 61)]
            if all(f in t3 and f in mkt for f in fw):
                f60[k] = sum(math.log1p(t3[f]) - math.log1p(mkt[f]) for f in fw) / 5 * 100
        hot = [f60[k] for k in f60 if s[k]]; cold = [f60[k] for k in f60 if not s[k]]
        conc[sig] = {'record_months_in_period': len(rr), 'share_hot': round(sum(1 for k in rr if s[k]) / len(rr), 3) if rr else None,
                     'fwd60_tech_minus_mkt_pct_per_yr_hot': {'n': len(hot), 'mean': round(S.mean(hot), 2) if hot else None},
                     'fwd60_tech_minus_mkt_pct_per_yr_cold': {'n': len(cold), 'mean': round(S.mean(cold), 2) if cold else None}}
    reports['concentration_reversal'] = conc
    # 業種での再現
    srep = {}
    s = st('LIFEY_Q5'); drugs = M.window(ret['Drugs'], None, END); s = {t: v for t, v in s.items() if t in drugs}
    rl = eval_timing('SR_a_R1_LIFEY_Drugs', '再現の報告', 'R1・ライフサイエンス IPO（年）・Q5・器 French Drugs', drugs, mkt, rf, s, 0.0, sig_meta['LIFEY_Q5'])
    for comp in ('tech', 'mkt'):
        c = rl['cmp'][comp]
        g, cr = M.grade(c['full'], c['train'], c['hold'], c['roll20_net'], cost_hold=c['hold_net'], repl=None, family_holm_p=None,
                        sharpe_pair=c['sharpe_pair_net'], leveraged_or_timing=True)
        c['grade_report_only'] = g; c['criteria'] = cr
    rl.pop('_series')
    srep['a_lifesci_drugs'] = rl
    ind = {}
    for c in cols:
        n_c = cnt[c]
        Xc = {s_: n_c[s_] / n_c[ym_add(s_, -12)] - 1 for s_ in months(197407, END) if s_ in n_c and ym_add(s_, -12) in n_c and n_c[ym_add(s_, -12)] > 0}
        if len(Xc) < MIN_N + 12:
            continue
        Hc = readings(expanding_pct(Xc), Q5)
        sc, _ = state_of(Hc, 196001, END)
        rc = M.window(ret[c], None, END)
        sc = {t: v for t, v in sc.items() if t in rc and t in mkt}
        g, n, _, _, _ = timing(rc, mkt, sc, 0.0)
        bb = sub(rc, g)
        tr = M.excess_stats(g, bb, z=M.TRAIN_END); ho = M.excess_stats(n, bb, a=M.HOLD_START)
        ind[c] = {'train_ex': tr['ex_ann'] if tr else None, 'train_t': tr['t'] if tr else None, 'hold_net_ex': ho['ex_ann'] if ho else None,
                  'hold_net_cagr_diff': ho['cagr_diff'] if ho else None, 'hot_share': round(sum(sc.values()) / len(sc), 3)}
    trp = [v for v in ind.values() if v['train_ex'] is not None]; hop = [v for v in ind.values() if v['hold_net_ex'] is not None]
    kt = sum(1 for v in trp if v['train_ex'] > 0); kh = sum(1 for v in hop if v['hold_net_ex'] > 0 and v['hold_net_cagr_diff'] > 0)
    srep['b_net_listings_49'] = {'industries': len(ind), 'train_positive': kt, 'train_binom_p': binom_ge(kt, len(trp)),
                                 'hold_positive_net': kh, 'hold_binom_p': binom_ge(kh, len(hop)), 'by_industry': ind}
    reports['sector_replication'] = srep
    # 感度（報告のみ）
    sens = {}
    for base in [x for x in strats['P'] + strats['N'] if x['name'].split('_')[1] == 'R1' and x['name'].split('_')[3] == 'Q5'] + strats['E']:
        nm = base['name']; sig = nm.split('_')[2] + '_Q5'
        tech = ndx if nm.startswith('N') else t3
        out = {}
        for lab, kw in (('lag+1', {'lag': 1}), ('dur24', {'dur': 24})):
            s = st(sig, **kw); s = {t: v for t, v in s.items() if t in tech}
            g, n, _, _, _ = timing(tech, mkt, s, 0.0)
            out[lab] = {cn: {'train': M.excess_stats(g, sub(b, g), z=M.TRAIN_END), 'hold_net': M.excess_stats(n, sub(b, g), a=M.HOLD_START)}
                        for cn, b in (('tech', tech), ('mkt', mkt))}
        sens[nm] = out
    reports['sensitivity'] = sens
    # 税
    tx = {}
    for base in strats['P'] + strats['N']:
        nm = base['name']; parts = nm.split('_'); sig = parts[2] + '_' + parts[3]
        tech = ndx if nm.startswith('N') else t3
        s = st(sig); s = {t: v for t, v in s.items() if t in tech}
        tx[nm] = tax_sim(tech, mkt, s, 0.0 if parts[1] == 'R1' else 0.5, z=INT_POS_END if parts[2] == 'INT' else END)
    reports['tax_taxable_account_2007+'] = tx

    # 事前登録2の報告: IPO の信号が対照に足した分・基準の積立・持ち替え
    xs = {r['name']: r for r in strats['X']}
    av = {}
    for a_, b_ in (('X1_R1_VC_Q5_TD_T3', 'X2_R1_TD_T3'), ('X3_R1_GROSS_Q5_TD_T3', 'X2_R1_TD_T3'), ('X4_R1_TNL_Q5_TD_T3', 'X2_R1_TD_T3'), ('X5_R1_VC_Q5_TD_NDX', 'X6_R1_TD_NDX')):
        na = xs[a_]['_series'][1]; nb = sub(xs[b_]['_series'][1], na)
        av[f'{a_} − {b_}'] = {'full': M.excess_stats(na, nb), 'train': M.excess_stats(na, nb, z=M.TRAIN_END), 'hold': M.excess_stats(na, nb, a=M.HOLD_START)}
    reports['X_ipo_added_value'] = av
    base = {}
    for tn, tech, t0 in (('T3', t3, 198012), ('NDX', ndx, 198511)):
        s0_ = {t: False for t in months(t0, END) if t in tech}
        bd = eval_dca(f'BASE_{tn}', '基準', f'今の規則: 毎月 1 を {tn} へ（vs 市場へ）', tech, mkt, s0_, END, {})
        base[tn] = bd['cmp']['mkt']
    reports['X_baseline_dca_tech_vs_mkt'] = base
    reports['X_switches'] = {r['name']: {'switches_per_year': r['switches_per_year'], 'hot_share(=市場にいた割合)': r['hot_share']} for r in strats['X']}

    # ── 出力
    all_timing = strats['P'] + strats['N'] + strats['E'] + strats['X']
    for r in all_timing:
        r.pop('_series', None)
    for r in all_timing:
        tested.append({'name': r['name'], 'family': r['family'], 'kind': 'timing', 'primary': r.get('primary', False),
                       'grade_vs_tech': r['grade_vs_tech'], 'grade_vs_mkt': r['grade_vs_mkt']})
    for d in dres:
        tested.append({'name': d['name'], 'family': d['family'], 'kind': 'dca_D', 'judgement_vs_tech': d['judgement_vs_tech'], 'judgement_vs_mkt': d['judgement_vs_mkt']})
    tested.append({'name': 'SR_a_R1_LIFEY_Drugs', 'family': '再現の報告', 'kind': 'report'})
    tested.append({'name': f'SR_b_net_listings_x{len(ind)}', 'family': '再現の報告', 'kind': 'report', 'count': len(ind)})
    tested.append({'name': 'BASE_T3・BASE_NDX（今の規則の積立 vs 市場の積立）', 'family': '基準の報告（事前登録2）', 'kind': 'report'})
    out = {'angle': 'tech_ipo_wave', 'prereg': [PREREG, PREREG2], 'prereg_commit': {PREREG: git_sha(f'out/{PREREG}'), PREREG2: git_sha(f'out/{PREREG2}')},
           'benchmark_note': 'grade_vs_tech = 同じハイテクの器を買って持つだけ（全体の事前登録のタイミング型の相手）／grade_vs_mkt = French Mkt（市場に勝つか）',
           'tested_count': len(tested), 'tested': tested,
           'families': {'P': strats['P'], 'N': strats['N'], 'E': strats['E'], 'X': strats['X'], 'D': dres},
           'signals': sig_meta, 'reports': reports, 'sanity': sanity, 'log': LOG}
    p = M.save(OUT, out)
    log('saved', p, os.path.getsize(p))


if __name__ == '__main__':
    main()
