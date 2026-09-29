#!/usr/bin/env python3
"""night/nx_regime_ind_data.py — 角度 nx_regime_ind（マクロの局面で『どの業種を持つか』を替える）の
データの取得と整形だけ（読むだけ・門の判定には不使用・**規則の成績は計算しない**）。

事前登録 out/nx_regime_ind_prereg.json の道具。業種・市場のリターンは『形』（始まり・終わり・欠け）だけを数え、
平均・t・シャープ・累積など成績に当たる数字は一切計算しない（局面の信号＝マクロの値だけを数える）。

取るもの
- 米国の局面の信号（月末 t に分かる値だけで作る。持つのは t+1 月）
    CPI: FRED CPIAUCNS（季節調整なし＝季節の係数の改定が無い・1913-01〜）の前年比 π。公表は翌月の中旬＝t 月末には t−1 月まで
    逆イールド: FRED GS10 − TB3MS（どちらも月平均・1953-04〜）。t 月の平均は t 月末に分かる。10年−1年（GS10 − GS1）は探索
                1953-03 以前は Goyal-Welch の lty − tbl（探索の早い延長だけ。Shiller の GS10 は 1953 年より前は年次の直線補間で、
                Shiller には月次の短期金利が無い＝q07leu の fam_riskparity が記録した事実）
    実時点の Sahm ルール: FRED SAHMREALTIME（その月に手に入った失業率の系列で作った値・1959-12〜）。雇用統計と同時に翌月初に公表
    SPF: フィラデルフィア連銀の専門家調査の『翌四半期に実質GDPが減る確率』の平均（RECESS2＝anxious index・1968Q4〜）。
         1990Q2 以降は公表日の月から、それより前（公表日が不明）は次の四半期の最初の月から使う
- 後知恵の対照（報告だけ・信号ではない）: Neville ほか (2021) の事後の8局面（持つ月そのもの）・NBER の景気後退の月（FRED USREC）
- C5: BIS の長期 CPI（WS_LONG_CPI・前年比・7か国）と FRED の OECD MEI 10年金利（IRLTLT01）・3か月金利（IR3TIB01）、
      JKP の国別 GICS セクターと国の市場（形だけ）。OECD の CPI（FRED の CPALTT01・OECD SDMX）は日本が 2021-06 で止まり、
      FRED の CPALTT01 は各国 2025-03 前後で止まっているので使わない（2026-09-28 に確かめた）
- 実在の答え合わせ（形だけ）: Select Sector SPDR（XLE・XLB・XLP・XLV・XLU）と Fidelity Select（FSENX・FSDPX・FDFAX・FSPHX・FSUTX）・SPY

出力 out/_nx_cache/nx_regime_ind_signals.json（gitignore）。'frozen'（業種の組・信号・後知恵の月・期間）の sha256 を事前登録に書き、
測る道具は最初にこの値と一致するかを確かめる（信号を結果を見てから作り直していないことの証明）。
業種の組は nx_leadlag_data.FR2GICS（French 49 → GICS 11 の対応・2026-09-28 に leadlag の事前登録で固定済み）から機械的に作る。

使い方: python3 night/nx_regime_ind_data.py          → 取得・信号・形の要約を書く
        python3 night/nx_regime_ind_data.py --show   → 保存済みの形の要約を表示
"""
import sys, os, io, re, csv, json, hashlib, zipfile, datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import nx_common as N  # noqa: E402
import nx_leadlag_data as LL  # noqa: E402  （FR2GICS の単一の正本。写さない）

OUT = os.path.join(N.CACHE, 'nx_regime_ind_signals.json')
FRED = 'https://fred.stlouisfed.org/graph/fredgraph.csv?id={}'
GW_URL = 'https://docs.google.com/spreadsheets/d/17mw_IpaiLFDrGnrPRQ2o1ugV5nJsZuD1/export?format=xlsx'  # eknzbh の mw_macro_timing と同じ原本
SPF_BASE = 'https://www.philadelphiafed.org/-/media/frbp/assets/surveys-and-data/survey-of-professional-forecasters/'
BIS_CPI = 'https://stats.bis.org/api/v2/data/dataflow/BIS/WS_LONG_CPI/1.0/M.{}.771?format=csv'
JKP_IND = 'https://jkpfactors-data.s3.amazonaws.com/public/industry/%5B{c}%5D_%5Bgics%5D_%5Bmonthly%5D_%5Bvw%5D.zip'

TRAIN_END, HOLD_START = N.TRAIN_END, N.HOLD_START
MAX_STALE = 3           # 公表されなかった月（2025-10 の CPI・失業率は米政府の閉鎖で出なかった）は、その前の最新の公表値を使う。3か月より古ければ不明
EPISODE_GAP = 6          # 局面の『塊』: 局面でない月が6か月未満しか挟まらない続きは同じ塊（Neville ほかの『6か月未満は局面と数えない』に合わせた）

# ── 業種の組（French 49 の名前。FR2GICS から機械的に。測る前に固定）
FR2GICS = LL.FR2GICS
GICS_REAL = ['10', '15']              # エネルギー・素材
GICS_DEF = ['30', '35', '55']         # 生活必需品・ヘルスケア・公益
REAL = sorted(k for k, v in FR2GICS.items() if v in GICS_REAL)
DEF = sorted(k for k, v in FR2GICS.items() if v in GICS_DEF)
SETS = {
    'REAL': REAL,                                           # 主: GICS 10・15 に当たる French 業種（10業種）
    'DEF': DEF,                                             # 主: GICS 30・35・55 に当たる French 業種（10業種）
    'REAL_NARROW': ['Coal', 'Gold', 'Mines', 'Oil'],        # 探索: 掘る業種だけ（Neville ほかの本文を読んだ後の選び方＝汚れあり）
    'REAL_UTIL': sorted(REAL + ['Util']),                   # 探索: 課題文の例（公益を実物に近い側へ）
    'DEF_EKNZBH': ['Drugs', 'Food', 'Hlth', 'Hshld', 'MedEq', 'Util'],  # 探索: eknzbh の mw_industry の DEF（無条件の A3 と比べられるように）
    'GICS_REAL': GICS_REAL, 'GICS_DEF': GICS_DEF,
}

# ── 後知恵の対照: Neville, Draaisma, Funnell, Harvey & Van Hemert (2021, JPM 47(8)) の Exhibit の米国8局面（事後に谷と山で日付を付けたもの）
NEVILLE = [(194104, 194205), (194603, 194703), (195008, 195102), (196602, 197001),
           (197207, 197412), (197702, 198003), (198702, 199011), (200709, 200807)]

C5 = {'jpn': 'JP', 'gbr': 'GB', 'can': 'CA', 'fra': 'FR', 'deu': 'DE', 'aus': 'AU', 'che': 'CH'}   # leadlag・eknzbh mw_industry と同じ7か国
C5_CPI_LAG = {'aus': 3}   # 豪州の CPI は四半期（月次の指標は 2018 年から）で、四半期の終わりの約1か月後に出る＝3か月遅れで使う。他は1か月
C5_SPAN = (199908, 202512)
ETFS = ['SPY', 'VFINX', 'XLE', 'XLB', 'XLP', 'XLV', 'XLU', 'FSENX', 'FSDPX', 'FDFAX', 'FSPHX', 'FSUTX']


# ───────────────────────── 月の計算 ─────────────────────────
def add(ym, k):
    y, m = divmod(ym, 100)
    n = y * 12 + (m - 1) + k
    return (n // 12) * 100 + n % 12 + 1


def mrange(a, z):
    out, m = [], a
    while m <= z:
        out.append(m)
        m = add(m, 1)
    return out


# ───────────────────────── 取得 ─────────────────────────
def fred(sid, max_age=7):
    b = N.get(FRED.format(sid), name=f'fred_{sid}.csv', max_age_days=max_age)
    out = {}
    for line in b.decode().strip().splitlines()[1:]:
        d, v = line.split(',')[:2]
        if v.strip() in ('', '.'):
            continue
        out[int(d[:4]) * 100 + int(d[5:7])] = float(v)
    return out


def goyal():
    import openpyxl
    b = N.get(GW_URL, name='goyal_predictors_2025.xlsx', max_age_days=3650)
    wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
    rows = list(wb['Monthly'].iter_rows(values_only=True))
    h = list(rows[0])
    out = {'tbl': {}, 'lty': {}}
    for r in rows[1:]:
        if r[0] is None:
            continue
        for k in out:
            v = r[h.index(k)]
            if isinstance(v, (int, float)):
                out[k][int(r[0])] = float(v)
    return out


def spf_release():
    b = N.get(SPF_BASE + 'spf-release-dates.txt', name='spf_release_dates.txt', max_age_days=30)
    rel, year = {}, None
    for line in b.decode('latin-1').splitlines():
        m = re.match(r'^\s*(\d{4})?\s*Q([1-4])\s+(\S+)\s+(\S+)', line)
        if not m:
            continue
        if m.group(1):
            year = int(m.group(1))
        mo, _, yy = m.group(4).rstrip('*').split('/')
        yy = int(yy)
        yy = 1900 + yy if yy >= 50 else 2000 + yy
        rel[(year, int(m.group(2)))] = yy * 100 + int(mo)
    return rel


def spf():
    """[{'survey': (年, 四半期), 'recess2': 確率%, 'avail': 使える最初の信号の月}]（調査の順）"""
    import openpyxl
    b = N.get(SPF_BASE + 'data-files/files/mean_recess_level.xlsx', name='spf_mean_recess_level.xlsx', max_age_days=30)
    wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
    rows = list(wb['Mean_Level'].iter_rows(values_only=True))
    h = list(rows[0])
    iy, iq, ir = h.index('YEAR'), h.index('QUARTER'), h.index('RECESS2')
    rel = spf_release()
    out = []
    for r in rows[1:]:
        if r[iy] is None or not isinstance(r[ir], (int, float)):
            continue
        y, q = int(r[iy]), int(r[iq])
        if (y, q) in rel:
            avail, how = rel[(y, q)], 'release'
        else:
            assert (y, q) < (1990, 2), f'1990Q2 以降なのに公表日が無い: {y}Q{q}'
            avail, how = (y * 100 + 3 * q + 1) if q < 4 else (y + 1) * 100 + 1, 'assumed_next_quarter_first_month'
        out.append({'survey': [y, q], 'recess2': float(r[ir]), 'avail': avail, 'how': how})
    out.sort(key=lambda x: tuple(x['survey']))
    return out


def bis_cpi():
    ccs = sorted(set(C5.values()) | {'US'})
    b = N.get(BIS_CPI.format('+'.join(ccs)), name='bis_long_cpi_M_771.csv', max_age_days=30)
    out = {}
    for r in csv.DictReader(io.StringIO(b.decode())):
        v = r.get('OBS_VALUE')
        if v in (None, '', 'NaN'):
            continue
        out.setdefault(r['REF_AREA'], {})[int(r['TIME_PERIOD'][:4]) * 100 + int(r['TIME_PERIOD'][5:7])] = float(v)
    return out


def jkp_gics_shape(country):
    b = N.get(JKP_IND.format(c=country), name=f'jkp_industry_{country}_gics_vw_monthly.zip')
    z = zipfile.ZipFile(io.BytesIO(b))
    d = {}
    for x in csv.DictReader(io.StringIO(z.read(z.namelist()[0]).decode())):
        if x['ret'] in ('', 'NA', 'na'):
            continue
        d.setdefault(str(int(float(x['gics']))), []).append(N._ym(x['date']))
    return {g: {'from': min(v), 'to': max(v), 'n': len(v), 'gaps': len(mrange(min(v), max(v))) - len(set(v))} for g, v in sorted(d.items())}


# ───────────────────────── 信号（月末 t に分かる値だけ。返り値は (旗, 使った最後のデータの月)） ─────────────────────────
def last_known(s, m, max_stale=MAX_STALE):
    """m 以前で公表された最新の値（公表されなかった月は飛ばす＝その時点の投資家が見た最新）。max_stale か月より古ければ不明"""
    for k in range(max_stale + 1):
        x = add(m, -k)
        if x in s:
            return s[x], x
    return None, None


def sig_inf(pi, t, thr=5.0, rise=6, lag=1):
    p, m = last_known(pi, add(t, -lag))
    if p is None:
        return None, None
    if rise:
        p0, _ = last_known(pi, add(m, -rise))
        if p0 is None:
            return None, None
        return int(p >= thr and p > p0), m
    return int(p >= thr), m


def sig_yc(sp, t, win=12):
    ms = [add(t, -k) for k in range(win)]
    if any(m not in sp for m in ms):
        return None, None
    return int(min(sp[m] for m in ms) < 0), t


def sig_level(s, t, thr, lag=1):
    v, m = last_known(s, add(t, -lag))
    if v is None:
        return None, None
    return int(v >= thr), m


def quantile7(vals, q):
    v = sorted(vals)
    h = (len(v) - 1) * q
    lo = int(h)
    return v[lo] if lo + 1 >= len(v) else v[lo] + (h - lo) * (v[lo + 1] - v[lo])


def sig_spf(sv, t, q=0.8, min_n=40):
    av = [x for x in sv if x['avail'] <= t]
    if len(av) < min_n:
        return None, None
    cur = av[-1]   # 調査の順に並んでいる＝使える最新の調査
    assert all(tuple(x['survey']) <= tuple(cur['survey']) for x in av)
    return int(cur['recess2'] >= quantile7([x['recess2'] for x in av], q)), cur['avail']


def build_signal(fn, months, lag_check):
    """{t: 旗} と先読みの検査。lag_check = 使ってよい最後の月 = t − lag_check"""
    out = {}
    for t in months:
        f, used = fn(t)
        if f is not None:
            assert used <= add(t, -lag_check), f'先読み: t={t} used={used}'
        out[t] = f
    return out


# ───────────────────────── 形の要約（成績は数えない） ─────────────────────────
def episodes(flag_by_hold, a, z, gap=EPISODE_GAP):
    ms = mrange(a, z)
    on = [m for m in ms if flag_by_hold.get(m) == 1]
    unk = [m for m in ms if flag_by_hold.get(m) is None]
    runs = []
    for m in on:
        if runs and add(runs[-1][1], 1) == m:
            runs[-1][1] = m
        else:
            runs.append([m, m])
    eps = []
    for r in runs:
        if eps and len(mrange(add(eps[-1][1], 1), add(r[0], -1))) < gap:
            eps[-1][1] = r[1]
            eps[-1][2] += len(mrange(r[0], r[1]))
        else:
            eps.append([r[0], r[1], len(mrange(r[0], r[1]))])
    return {'span': [a, z], 'months': len(ms), 'on': len(on), 'unknown': len(unk), 'runs': len(runs),
            'episodes_n': len(eps), 'episodes': eps}


def split_shape(flag_by_hold, a, z):
    d = {'full': episodes(flag_by_hold, a, z)}
    if a <= TRAIN_END:
        d['train'] = episodes(flag_by_hold, a, min(z, TRAIN_END))
    if z >= HOLD_START:
        d['hold'] = episodes(flag_by_hold, max(a, HOLD_START), z)
    return d


def to_hold(sig):
    """信号 {t: 旗} → 持つ月 {t+1: 旗}"""
    return {add(t, 1): f for t, f in sig.items()}


def first_full(hold, start, z):
    """start 以降で、そこから z まで欠けの無い最初の持つ月"""
    ms = mrange(start, z)
    for i, m in enumerate(ms):
        if all(hold.get(x) is not None for x in ms[i:]):
            return m
    return None


# ───────────────────────── 本体 ─────────────────────────
def build():
    # French 49（形だけ）
    T = N.french_tables('49_Industry_Portfolios')
    vw = T['Average Value Weighted Returns -- Monthly']
    cols = [c.strip() for c in vw['cols']]
    fr_months = sorted(vw['data'])
    fr_first = {c: None for c in cols}
    fr_last = {c: None for c in cols}
    fr_gaps = {c: 0 for c in cols}
    for m in fr_months:
        for c, v in zip(cols, vw['data'][m]):
            if v is not None:
                fr_first[c] = fr_first[c] or m
                fr_last[c] = m
    for c in cols:
        if fr_first[c]:
            fr_gaps[c] = sum(1 for m in fr_months if fr_first[c] <= m <= fr_last[c]
                             and vw['data'][m][cols.index(c)] is None)
    Z = fr_months[-1]
    for nm in ('REAL', 'DEF', 'REAL_NARROW', 'REAL_UTIL', 'DEF_EKNZBH'):
        assert all(x in cols for x in SETS[nm]), (nm, [x for x in SETS[nm] if x not in cols])

    # 米国の系列
    cpi = fred('CPIAUCNS')
    pi = {m: (cpi[m] / cpi[add(m, -12)] - 1) * 100 for m in cpi if add(m, -12) in cpi}
    gs10, tb3, gs1 = fred('GS10'), fred('TB3MS'), fred('GS1')
    gw = goyal()
    sp = {m: gs10[m] - tb3[m] for m in gs10 if m in tb3}
    sp_gw = {m: (gw['lty'][m] - gw['tbl'][m]) * 100 for m in gw['lty'] if m in gw['tbl']}
    sp_splice = {**{m: v for m, v in sp_gw.items() if m < 195304}, **sp}
    sp1 = {m: gs10[m] - gs1[m] for m in gs10 if m in gs1}
    sahm = fred('SAHMREALTIME')
    usrec = fred('USREC')
    sv = spf()

    ts = mrange(191401, add(Z, -1))   # 信号の月 t（持つのは t+1 ≤ Z）
    S = {}
    S['INF5_r6'] = build_signal(lambda t: sig_inf(pi, t, 5.0, 6), ts, 1)
    S['INF4_r6'] = build_signal(lambda t: sig_inf(pi, t, 4.0, 6), ts, 1)
    S['INF5_lvl'] = build_signal(lambda t: sig_inf(pi, t, 5.0, 0), ts, 1)
    S['INF5_r3'] = build_signal(lambda t: sig_inf(pi, t, 5.0, 3), ts, 1)
    S['INF5_r12'] = build_signal(lambda t: sig_inf(pi, t, 5.0, 12), ts, 1)
    S['YC12'] = build_signal(lambda t: sig_yc(sp, t, 12), ts, 0)
    S['YC12_splice'] = build_signal(lambda t: sig_yc(sp_splice, t, 12), ts, 0)
    S['YC1'] = build_signal(lambda t: sig_yc(sp, t, 1), ts, 0)
    S['YC24'] = build_signal(lambda t: sig_yc(sp, t, 24), ts, 0)
    S['YC12_10y1y'] = build_signal(lambda t: sig_yc(sp1, t, 12), ts, 0)
    S['SAHM'] = build_signal(lambda t: sig_level(sahm, t, 0.5, 1), ts, 1)
    S['SPF80'] = build_signal(lambda t: sig_spf(sv, t, 0.8, 40), ts, 0)

    # Sahm の点いた月から12か月（点いた月 u: SAHM(u)=1 ∧ SAHM(u−1)=0。t の旗 = [t−11, t] に点いた月がある）
    trig = {t: (1 if S['SAHM'].get(t) == 1 and S['SAHM'].get(add(t, -1)) == 0 else
                (None if S['SAHM'].get(t) is None or S['SAHM'].get(add(t, -1)) is None else 0)) for t in ts}
    S['SAHM_12m'] = {}
    for t in ts:
        w = [trig.get(add(t, -k)) for k in range(12)]
        S['SAHM_12m'][t] = None if any(x is None for x in w) else int(any(w))

    # 組み合わせ（E20）: インフレ → REAL、そうでなく 逆イールド12 か Sahm → DEF、どれでもなければ市場
    S['COMBO'] = {}
    for t in ts:
        a, b, c = S['INF5_r6'].get(t), S['YC12'].get(t), S['SAHM'].get(t)
        S['COMBO'][t] = None if None in (a, b, c) else ('REAL' if a == 1 else ('DEF' if (b == 1 or c == 1) else 'MKT'))

    # 後知恵（持つ月そのもの）
    H = {'NEVILLE_expost': {m: int(any(a <= m <= z for a, z in NEVILLE)) for m in mrange(192607, Z)},
         'NBER_USREC': {m: int(usrec[m] == 1) for m in mrange(192607, Z) if m in usrec}}

    # 期間（持つ月）。主の族はデータの質で始まりを決めた（事前登録の本文を見よ）
    want = {'P1_INF_REAL': ('INF5_r6', 194702), 'P2_YC_DEF': ('YC12', 195404), 'P3_SAHM_DEF': ('SAHM', 196002),
            'P4_SPF_DEF': ('SPF80', 196901),
            'E1_INF_REAL_early': ('INF5_r6', 192608), 'E2_YC_DEF_early': ('YC12_splice', 192608),
            'E7_INF4': ('INF4_r6', 194702), 'E8_INF5_lvl': ('INF5_lvl', 194702), 'E9_INF5_r3': ('INF5_r3', 194702),
            'E10_INF5_r12': ('INF5_r12', 194702), 'E13_YC1': ('YC1', 195404), 'E14_YC24': ('YC24', 195404),
            'E15_YC12_10y1y': ('YC12_10y1y', 195404), 'E17_SAHM_12m': ('SAHM_12m', 196002), 'E20_COMBO': ('COMBO', 196002)}
    spans, shape = {}, {}
    for rule, (sig, start) in want.items():
        hold = to_hold(S[sig])
        a = first_full(hold, start, Z)
        spans[rule] = {'signal': sig, 'hold_from': a, 'hold_to': Z}
        if sig == 'COMBO':
            for st in ('REAL', 'DEF'):
                shape[f'{rule}:{st}'] = split_shape({m: (None if f is None else int(f == st)) for m, f in hold.items()}, a, Z)
        else:
            shape[rule] = split_shape(hold, a, Z)
    for k, v in H.items():
        shape[f'HINDSIGHT_{k}'] = split_shape(v, 192607, Z)

    # C5
    bis = bis_cpi()
    c5_sig, c5_shape, c5_data = {}, {}, {}
    for c, cc in C5.items():
        cpi_c = bis.get(cc, {})
        lt, st = fred(f'IRLTLT01{cc}M156N'), fred(f'IR3TIB01{cc}M156N')
        sp_c = {m: lt[m] - st[m] for m in lt if m in st}
        lag = C5_CPI_LAG.get(c, 1)
        tt = mrange(add(C5_SPAN[0], -1), add(C5_SPAN[1], -1))
        c5_sig[c] = {'INF5_r6': build_signal(lambda t: sig_inf(cpi_c, t, 5.0, 6, lag), tt, lag),
                     'YC12': build_signal(lambda t: sig_yc(sp_c, t, 12), tt, 0)}
        c5_shape[c] = {k: episodes(to_hold(v), C5_SPAN[0], C5_SPAN[1]) for k, v in c5_sig[c].items()}
        au_q = None
        if c == 'aus':   # 月次の値が四半期の中で同じ値の繰り返しか（＝四半期の値）を数える（形）
            qs = [m for m in cpi_c if 199901 <= m <= 201712]
            same = sum(1 for m in qs if (m % 100) % 3 != 1 and cpi_c.get(add(m, -1)) == cpi_c[m])
            au_q = {'months_1999_2017': len(qs), 'same_as_prev_within_quarter': same}
        c5_data[c] = {'cpi_from': min(cpi_c) if cpi_c else None, 'cpi_to': max(cpi_c) if cpi_c else None,
                      'lt_from': min(lt) if lt else None, 'lt_to': max(lt) if lt else None,
                      'st_from': min(st) if st else None, 'st_to': max(st) if st else None,
                      'cpi_lag': lag, 'aus_quarterly_check': au_q,
                      'jkp_gics': jkp_gics_shape(c), 'jkp_mkt_vw': None}
        try:
            mk = N.jkp_mkt(c, 'vw')
            c5_data[c]['jkp_mkt_vw'] = [min(mk), max(mk), len(mk)]
        except Exception as e:  # noqa
            c5_data[c]['jkp_mkt_vw'] = f'取得失敗 {e}'

    # 実在の器（形だけ）
    etf = {}
    for tk in ETFS:
        try:
            r = N.yahoo(tk)
            etf[tk] = [min(r), max(r), len(r)]
        except Exception as e:  # noqa
            etf[tk] = f'取得失敗 {str(e)[:80]}'

    # US の CPI の二つの原本の一致（形）: BIS の US 前年比 と CPIAUCNS から作った前年比
    us_b = bis.get('US', {})
    common = [m for m in pi if m in us_b and m >= 195001]
    cpi_agree = {'months': len(common), 'max_abs_diff_pt': round(max(abs(pi[m] - us_b[m]) for m in common), 3) if common else None}

    frozen = {'sets': SETS, 'neville': NEVILLE, 'episode_gap': EPISODE_GAP, 'spans': spans,
              'us_signals': {k: {str(t): f for t, f in v.items() if f is not None} for k, v in S.items()},
              'hindsight': {k: {str(m): f for m, f in v.items()} for k, v in H.items()},
              'c5_signals': {c: {k: {str(t): f for t, f in v.items() if f is not None} for k, v in d.items()} for c, d in c5_sig.items()},
              'c5_span': list(C5_SPAN), 'c5_cpi_lag': {c: C5_CPI_LAG.get(c, 1) for c in C5}}
    sha = hashlib.sha256(json.dumps(frozen, sort_keys=True, separators=(',', ':'), ensure_ascii=False).encode()).hexdigest()

    series = {'pi_us': {str(m): round(v, 4) for m, v in pi.items()}, 'spread_10y3m': {str(m): round(v, 4) for m, v in sp.items()},
              'spread_gw_lty_tbl': {str(m): round(v, 4) for m, v in sp_gw.items()},
              'spread_10y1y': {str(m): round(v, 4) for m, v in sp1.items()}, 'sahm_realtime': {str(m): v for m, v in sahm.items()},
              'spf_recess2': sv}
    meta = {
        'generated': datetime.datetime.utcnow().strftime('%Y-%m-%dT%H:%MZ'),
        'french49': {'from': fr_months[0], 'to': Z, 'first': fr_first, 'last': fr_last, 'gaps': fr_gaps},
        'us_series_span': {'CPIAUCNS': [min(cpi), max(cpi)], 'GS10': [min(gs10), max(gs10)], 'TB3MS': [min(tb3), max(tb3)],
                           'GS1': [min(gs1), max(gs1)], 'GW_lty': [min(gw['lty']), max(gw['lty'])], 'GW_tbl': [min(gw['tbl']), max(gw['tbl'])],
                           'SAHMREALTIME': [min(sahm), max(sahm)], 'USREC': [min(usrec), max(usrec)],
                           'SPF_RECESS2': [sv[0]['survey'], sv[-1]['survey'], len(sv)],
                           'SPF_release_known_from': min(tuple(x['survey']) for x in sv if x['how'] == 'release')},
        'splice_check_1953': {'GW_minus_FRED_spread_195304_195312_maxabs': round(max(abs(sp_gw[m] - sp[m]) for m in mrange(195304, 195312)), 3)},
        'cpi_us_bis_vs_cpiaucns': cpi_agree,
        'c5': c5_data, 'etf': etf,
    }
    out = {'frozen': frozen, 'frozen_sha256': sha, 'shape': shape, 'c5_shape': c5_shape, 'series': series, 'meta': meta}
    tmp = OUT + '.tmp'
    json.dump(out, open(tmp, 'w'), ensure_ascii=False, indent=0, default=list)
    os.replace(tmp, OUT)
    # 事前登録に書いた版を上書きしない写し（データが新しい月へ伸びると sha が変わる。測る道具は登録した sha の写しを読む）
    frz = OUT.replace('.json', f'_{sha[:12]}.json')
    if not os.path.exists(frz):
        json.dump(out, open(frz, 'w'), ensure_ascii=False, indent=0, default=list)
    return out


def fmt(m):
    return f'{m // 100}-{m % 100:02d}' if isinstance(m, int) else str(m)


def show(o):
    print('frozen sha256:', o['frozen_sha256'])
    print('sets:', {k: v for k, v in o['frozen']['sets'].items()})
    print('spans:')
    for k, v in o['frozen']['spans'].items():
        print(f'  {k:22s} {v["signal"]:12s} {fmt(v["hold_from"])}〜{fmt(v["hold_to"])}')
    print('局面の形（持つ月・塊は間が6か月未満ならつなぐ）:')
    for k, d in o['shape'].items():
        line = []
        for part in ('train', 'hold'):
            if part in d:
                e = d[part]
                line.append(f'{part} {e["on"]}/{e["months"]}か月 塊{e["episodes_n"]}（連{e["runs"]}） 不明{e["unknown"]}')
        print(f'  {k:28s} ' + ' | '.join(line))
        print('      塊: ' + ', '.join(f'{fmt(a)}〜{fmt(b)}({n})' for a, b, n in d['full']['episodes']))
    print('C5（1999-08〜2025-12 の持つ月）:')
    for c, d in o['c5_shape'].items():
        print(f'  {c}: ' + ' | '.join(f'{k} {e["on"]}/{e["months"]} 塊{e["episodes_n"]} 不明{e["unknown"]} ['
                                     + ', '.join(f'{fmt(a)}〜{fmt(b)}' for a, b, _ in e['episodes']) + ']' for k, e in d.items()))
    m = o['meta']
    print('US series:', m['us_series_span'])
    print('splice:', m['splice_check_1953'], ' cpi:', m['cpi_us_bis_vs_cpiaucns'])
    for c, d in m['c5'].items():
        print(f'  {c}: cpi {d["cpi_from"]}〜{d["cpi_to"]} lag{d["cpi_lag"]} | 10y {d["lt_from"]}〜{d["lt_to"]} | 3m {d["st_from"]}〜{d["st_to"]} | mkt {d["jkp_mkt_vw"]} | aus {d["aus_quarterly_check"]}')
        print('      gics:', {g: (x['from'], x['to'], x['gaps']) for g, x in d['jkp_gics'].items()})
    print('ETF:', m['etf'])
    fr = m['french49']
    print('French49:', fr['from'], fr['to'], {k: fr['first'][k] for k in o['frozen']['sets']['REAL'] + o['frozen']['sets']['DEF']},
          'gaps', {k: v for k, v in fr['gaps'].items() if v})


if __name__ == '__main__':
    if '--show' in sys.argv:
        show(json.load(open(OUT)))
    else:
        show(build())
