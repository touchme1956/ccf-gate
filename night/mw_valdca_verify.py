#!/usr/bin/env python3
"""night/mw_valdca_verify.py — 反証の検証: mw_valdca の A 2本（E3c・E3d）と B の保有期間上位2本（F1・E3b）
（＋文脈として E2c・F2）。読むだけ・門の判定には不使用。

独立性:
- 取得は mw_common の fetcher（M.get・M.ff_factors・M.yahoo）だけを使う。Shiller xls・Goyal xlsx・FRED・World Bank・
  French の国別 zip・JST の解析、CAPE/ECY の信号、債券の総リターン（半年利払いの額面債＝研究者の年1回利払いと別の式）、
  売買の模擬、超過・NW t・CAGR・転がる20年・積立20年・回帰（HAC）はすべてこのファイルの自作。
- 研究者のコード（night/mw_valdca.py）は import しない。

攻め筋（事前に決めた順）:
 1 再現（全期間・訓練・保有・直近・費用後・20年窓・積立20年・シャープ）
 2 β を増やしただけか（α の回帰・同じ平均倍率の一定レバレッジ・シャープの差の検定）
 3 小区間（2008 を除く・1998-2000 と 2020-2021 を除く・保有の前半/後半・2021〜・2013-07〜）
 4 パラメータの脆さ（ECY の線の分位 30〜70%・移動平均 6〜14・信号の遅れ）
 5 実行できたか（Reg T の証拠金規制で 1945〜1974 の2倍は違法だった・歴史の借入金利＝コールレート/プライム・往復の費用・日本の課税）
 6 米国外の再現（French 国別を自作で組み直す・リスク調整・債券の月中平均による後知恵の点検）
 7 JST 1870〜1974（自作で組み直す・ドイツのハイパーインフレ・対数差）
 8 Shiller 時代（月中平均の株価によるトレンドの後知恵）
 9 多重検定（この角度の試行数）
"""
import io, json, math, os, re, sys, zipfile, datetime, subprocess, statistics as S

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import mw_common as M  # noqa: E402  （fetcher だけを使う）

OUT = 'mw_valdca_verify.json'
END = 202608
TRAIN_END, HOLD_START, RECENT_START = 200612, 200701, 201307
THR_ECY = 0.03308           # 事前登録の値（訓練の中央値）＝規則はこれで裁く（自作の中央値は検算として別に出す）
SPREAD = 0.015
COST = 0.001
LOG = []


def log(*a):
    s = ' '.join(str(x) for x in a)
    LOG.append(s)
    print(s, flush=True)


def ymadd(k, n):
    y, m = divmod(k, 100)
    t = y * 12 + m - 1 + n
    return (t // 12) * 100 + t % 12 + 1


def mrange(a, z):
    out = []
    while a <= z:
        out.append(a)
        a = ymadd(a, 1)
    return out


# ═════════════════════════ 取得（自作の解析） ═════════════════════════
def fred(sid):
    b = M.get('https://fred.stlouisfed.org/graph/fredgraph.csv?id=' + sid, f'fred_{sid}.csv', max_age_days=3650).decode()
    out = {}
    for ln in b.strip().splitlines()[1:]:
        p = ln.split(',')
        if len(p) < 2 or p[1].strip() in ('', '.'):
            continue
        out[p[0].strip()] = float(p[1])
    return out


def fred_m(sid):
    return {int(d[:4]) * 100 + int(d[5:7]): v for d, v in fred(sid).items()}


def fred_month_end(sid):
    out = {}
    d0 = fred(sid)
    for d in sorted(d0):
        out[int(d[:4]) * 100 + int(d[5:7])] = d0[d]   # 月の最後の観測で上書き
    return out


def load_shiller():
    import xlrd
    b = M.get('http://www.econ.yale.edu/~shiller/data/ie_data.xls', 'shiller_ie_data.xls', max_age_days=3650)
    sh = xlrd.open_workbook(file_contents=b).sheet_by_name('Data')
    hr = next(i for i in range(sh.nrows) if str(sh.row_values(i)[0]).strip() == 'Date')
    hdr = [str(x).strip() for x in sh.row_values(hr)]
    col = {'P': hdr.index('P'), 'D': hdr.index('D'), 'E': hdr.index('E'), 'CPI': hdr.index('CPI'),
           'GS10': hdr.index('Rate GS10'), 'CAPE': hdr.index('CAPE'), 'ECY': hdr.index('Yield')}
    # 月次の債券の総リターン（行 t = t→t+1）: 上の行に 'Bond' があり見出しが 'Returns' の最初の列
    up1 = [str(x).strip() for x in sh.row_values(hr - 1)]
    col['BOND'] = next(j for j, h in enumerate(hdr) if h == 'Returns' and up1[j] == 'Bond')
    out = {c: {} for c in col}
    for i in range(hr + 1, sh.nrows):
        r = sh.row_values(i)
        d = r[0]
        if not isinstance(d, float):
            continue
        y = int(d + 1e-9)
        m = int(round((d - y) * 100))
        if not 1 <= m <= 12:
            continue
        k = y * 100 + m
        for c, j in col.items():
            if j < len(r) and isinstance(r[j], float):
                out[c][k] = r[j]
    return out


def load_goyal():
    import openpyxl
    b = M.get('https://docs.google.com/spreadsheets/d/17mw_IpaiLFDrGnrPRQ2o1ugV5nJsZuD1/export?format=xlsx',
              'goyal_predictors_2025.xlsx', max_age_days=3650)
    wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
    rows = list(wb['Monthly'].iter_rows(values_only=True))
    h = list(rows[0])
    I = {c: h.index(c) for c in ('yyyymm', 'e12', 'd12', 'ltr', 'Rfree', 'price')}
    out = {c: {} for c in I if c != 'yyyymm'}
    for r in rows[1:]:
        if r[I['yyyymm']] is None:
            continue
        k = int(r[I['yyyymm']])
        for c, j in I.items():
            if c != 'yyyymm' and r[j] is not None:
                out[c][k] = float(r[j])
    return out


def gspc_month_avg():
    u = 'https://query1.finance.yahoo.com/v8/finance/chart/%5EGSPC?period1=-1400000000&period2=1790000000&interval=1d'
    j = json.loads(M.get(u, 'valdca_yh_GSPC_1d.json', max_age_days=3650))
    r = j['chart']['result'][0]
    acc = {}
    for ts, c in zip(r['timestamp'], r['indicators']['quote'][0]['close']):
        if c is None:
            continue
        d = datetime.datetime.utcfromtimestamp(ts)
        acc.setdefault(d.year * 100 + d.month, []).append(c)
    return {k: sum(v) / len(v) for k, v in acc.items()}


def carry_forward(d, a, z):
    out, last = {}, None
    for k in mrange(a, z):
        if k in d:
            last = d[k]
        if last is not None:
            out[k] = last
    return out


# ═════════════════════════ 債券（自作の式: 半年利払いの10年額面債） ═════════════════════════
def par_bond_semi(y0, y1, years=10.0):
    """利回り y0 の新発10年債（半年ごとに y0/2 の利札）を1か月後に利回り y1（半年複利）で評価した総リターン。
    研究者（年1回利払いの Shiller 式）とは別の式で独立に作る"""
    y1 = max(y1, 1e-6)
    v = 0.0
    n = int(round(years * 2))
    for j in range(1, n + 1):
        tj = 0.5 * j - 1 / 12
        cf = y0 / 2 + (1.0 if j == n else 0.0)
        v += cf * (1 + y1 / 2) ** (-2 * tj)
    return v - 1


# ═════════════════════════ 米国のデータと信号（自作） ═════════════════════════
def build_us():
    ff = M.ff_factors('monthly')
    mkt = {k: v for k, v in ff['mkt'].items() if k <= END}
    rf = {k: v for k, v in ff['rf'].items() if k <= END}
    sh = load_shiller()
    gy = load_goyal()
    cpi_f = fred_m('CPIAUCNS')
    gs10_f = fred_m('GS10')
    # 物価: 〜1912 Shiller・1913〜 FRED（未公表の月は直前の値）
    cpi = {k: v for k, v in sh['CPI'].items() if k < 191301}
    cpi.update(cpi_f)
    cpi = carry_forward(cpi, 187101, END)
    # 10年金利（月中平均）
    gs10 = {k: v for k, v in sh['GS10'].items() if k < 195304}
    gs10.update(gs10_f)
    # 株価（月中平均）: Shiller は 2023-08 まで（2023-09 は 9/1 の終値なので使わない）、以降 Yahoo の月中平均
    P = {k: v for k, v in sh['P'].items() if k <= 202308}
    for k, v in gspc_month_avg().items():
        if 202309 <= k <= END:
            P[k] = v
    # 利益: Shiller E（〜2023-06）→ Goyal e12（〜2025-12）→ それ以降は最後に分かった値を持ち越す（研究者の multpl 延長とは別）
    E = dict(sh['E'])
    for k, v in gy['e12'].items():
        if k > max(E):
            E[k] = v
    E = carry_forward(E, min(E), END)
    # 信号（月末 t に分かる値）
    sig = {'CAPE': {}, 'ECY': {}, 'ECY_shiller_col': {}}
    for t in sorted(P):
        c1 = cpi.get(ymadd(t, -1))
        c121 = cpi.get(ymadd(t, -121))
        if c1 is None:
            continue
        ks = [ymadd(t, -3 - i) for i in range(120)]
        if any(k not in E or k not in cpi for k in ks):
            continue
        e10 = sum(E[k] / cpi[k] for k in ks) / 120
        if e10 <= 0:
            continue
        cape = (P[t] / c1) / e10
        sig['CAPE'][t] = cape
        if t in gs10 and c121:
            pi10 = (c1 / c121) ** 0.1 - 1
            sig['ECY'][t] = 1 / cape - (gs10[t] / 100 - pi10)
    for k, v in sh['ECY'].items():
        sig['ECY_shiller_col'][k] = v
    # 10年国債の総リターン（t 月 = t−1 月末→t 月末）
    bond = {}
    for k in mrange(192607, 196112):
        if k in gy['ltr']:
            bond[k] = gy['ltr'][k]
    me = fred_month_end('DGS10')
    for k in mrange(196202, END):
        a = ymadd(k, -1)
        if a in me and k in me:
            bond[k] = par_bond_semi(me[a] / 100, me[k] / 100)
    bond[196201] = par_bond_semi(gs10[196112] / 100, gs10[196201] / 100)
    # 別の債券系列（頑健さ）: Goyal ltr（長期国債）を 2025-12 まで、2026 は DGS10
    bond_ltr = {k: v for k, v in gy['ltr'].items() if 192607 <= k <= END}
    for k in mrange(202601, END):
        if k not in bond_ltr and k in bond:
            bond_ltr[k] = bond[k]
    return {'mkt': mkt, 'rf': rf, 'sig': sig, 'bond': bond, 'bond_ltr': bond_ltr, 'cpi': cpi, 'gs10': gs10, 'P': P, 'E': E,
            'sh': sh, 'gy': gy}


def tr_index(r):
    ks = sorted(r)
    idx = {ymadd(ks[0], -1): 1.0}
    lv = 1.0
    for k in ks:
        lv *= 1 + r[k]
        idx[k] = lv
    return idx


def sma_up(idx, n=10):
    ks = sorted(idx)
    out = {}
    for i in range(n - 1, len(ks)):
        w = ks[i - n + 1:i + 1]
        if ymadd(w[0], n - 1) != w[-1]:
            continue
        out[ks[i]] = idx[ks[i]] >= sum(idx[k] for k in w) / n
    return out


def tsmom_up(r, cash, n=12):
    out = {}
    for t in sorted(r):
        w = [ymadd(t, -i) for i in range(n)]
        if any(k not in r or k not in cash for k in w):
            continue
        a = b = 1.0
        for k in w:
            a *= 1 + r[k]
            b *= 1 + cash[k]
        out[t] = a > b
    return out


# ═════════════════════════ 規則と模擬（自作） ═════════════════════════
def agree_state(ecy, trend, thr, lev, lag=1):
    """t 月の目標倍率（t−lag 月末の信号）。割安∧上→lev・割高∧下→0（逃げ先）・ほか 1"""
    def f(t):
        k = ymadd(t, -lag)
        e, u = ecy.get(k), trend.get(k)
        if e is None or u is None:
            return None
        if e > thr and u:
            return lev
        if e < thr and not u:
            return 0.0
        return 1.0
    return f


def simulate(months_, wfun, stock, alt, rf, spread=SPREAD, cost=COST, two_leg=False, borrow=None, cap=None):
    """月次の模擬。w>1 は借入（borrow[t] があればその月利、無ければ rf+spread/12）。w<1 は残りを alt。
    費用 = 株の売買（前月末に値動きで変わった後の割合との差）×cost（two_leg なら逃げ先の売買も数える）"""
    g, n, W = {}, {}, {}
    ws_d = wa_d = None
    for t in months_:
        w = wfun(t)
        if cap is not None and w > 1:
            w = min(w, cap(t))
        r = stock[t]
        if w > 1:
            b = borrow[t] if borrow is not None else rf[t] + spread / 12
            rp = w * r - (w - 1) * b
            wa, ra = 0.0, 0.0
        else:
            ra = alt[t] if w < 1 else 0.0
            rp = w * r + (1 - w) * ra
            wa = 1 - w
        to = 0.0 if ws_d is None else abs(w - ws_d) + (abs(wa - wa_d) if two_leg else 0.0)
        g[t] = rp
        n[t] = rp - to * cost
        W[t] = w
        ws_d = w * (1 + r) / (1 + rp) if 1 + rp > 1e-9 else w
        wa_d = wa * (1 + ra) / (1 + rp) if 1 + rp > 1e-9 else wa
    return g, n, W


def valid_months(wfun, need, a, z):
    ok = [t for t in mrange(a, z) if wfun(t) is not None and all(t in d for d in need)]
    # 最後の月で終わる最長の連続区間
    if not ok:
        return []
    run = [ok[-1]]
    for k in reversed(ok[:-1]):
        if ymadd(k, 1) == run[-1]:
            run.append(k)
        else:
            break
    return sorted(run)


# ═════════════════════════ 統計（自作） ═════════════════════════
def nwt(x, L=12):
    n = len(x)
    if n < 24:
        return None
    m = sum(x) / n
    e = [v - m for v in x]
    s = sum(v * v for v in e) / n
    for l in range(1, min(L, n - 1) + 1):
        s += 2 * (1 - l / (L + 1)) * sum(e[i] * e[i - l] for i in range(l, n)) / n
    return m / math.sqrt(s / n) if s > 0 else None


def geo(x, py=12):
    return math.exp(math.fsum(math.log1p(v) for v in x) * py / len(x)) - 1


def ex(s, b, a=None, z=None, py=12, L=12, drop=None):
    ks = [k for k in sorted(s) if k in b and (a is None or k >= a) and (z is None or k <= z) and not (drop and drop(k))]
    if len(ks) < 24:
        return None
    d = [s[k] - b[k] for k in ks]
    t = nwt(d, L)
    gs, gb = geo([s[k] for k in ks], py), geo([b[k] for k in ks], py)
    return {'from': ks[0], 'to': ks[-1], 'n': len(ks), 'ex_ann': round(sum(d) / len(d) * py * 100, 2), 't': round(t, 2) if t is not None else None,
            'cagr_s': round(gs * 100, 2), 'cagr_b': round(gb * 100, 2), 'cagr_diff': round((gs - gb) * 100, 2)}


def sharpe(r, rf, a=None, z=None, drop=None):
    ks = [k for k in sorted(r) if k in rf and (a is None or k >= a) and (z is None or k <= z) and not (drop and drop(k))]
    x = [r[k] - rf[k] for k in ks]
    sd = S.stdev(x)
    return sum(x) / len(x) / sd * math.sqrt(12) if sd else None


def sharpe_diff_test(r1, r2, rf, a=None, z=None):
    """Jobson-Korkie（Memmel 2003 の補正）: 月次シャープの差の z（iid の近似・楽観側）"""
    ks = [k for k in sorted(r1) if k in r2 and k in rf and (a is None or k >= a) and (z is None or k <= z)]
    x = [r1[k] - rf[k] for k in ks]
    y = [r2[k] - rf[k] for k in ks]
    T = len(ks)
    s1, s2 = S.mean(x) / S.stdev(x), S.mean(y) / S.stdev(y)
    rho = M.corr(x, y)
    var = (2 * (1 - rho) + 0.5 * (s1 ** 2 + s2 ** 2 - 2 * s1 * s2 * rho ** 2)) / T
    return {'sr_s': round(s1 * math.sqrt(12), 3), 'sr_b': round(s2 * math.sqrt(12), 3), 'rho': round(rho, 3),
            'z': round((s1 - s2) / math.sqrt(var), 2) if var > 0 else None}


def hac_reg(y, x, L=12):
    """y = a + b x の OLS と Newey-West（Bartlett）の標準誤差 → (a, t_a, b, t_b)"""
    n = len(y)
    mx, my = sum(x) / n, sum(y) / n
    sxx = sum((v - mx) ** 2 for v in x)
    b = sum((xi - mx) * (yi - my) for xi, yi in zip(x, y)) / sxx
    a = my - b * mx
    u = [yi - a - b * xi for xi, yi in zip(x, y)]
    # X'X
    Sx, Sxx = sum(x), sum(v * v for v in x)
    det = n * Sxx - Sx * Sx
    inv = [[Sxx / det, -Sx / det], [-Sx / det, n / det]]
    # S = Σ w_l Σ_t u_t u_{t-l} x_t x_{t-l}'
    Sm = [[0.0, 0.0], [0.0, 0.0]]
    for l in range(0, L + 1):
        wl = 1.0 if l == 0 else 2 * (1 - l / (L + 1))
        a00 = a01 = a10 = a11 = 0.0
        for t in range(l, n):
            p = u[t] * u[t - l]
            a00 += p
            a01 += p * x[t - l]
            a10 += p * x[t]
            a11 += p * x[t] * x[t - l]
        if l == 0:
            Sm[0][0] += a00; Sm[0][1] += a01; Sm[1][0] += a10; Sm[1][1] += a11
        else:
            # 対称化（l と −l を足す）
            Sm[0][0] += wl * a00; Sm[0][1] += wl * (a01 + a10) / 2; Sm[1][0] += wl * (a01 + a10) / 2; Sm[1][1] += wl * a11
    V = [[sum(inv[i][k] * sum(Sm[k][j2] * inv[j2][j] for j2 in range(2)) for k in range(2)) for j in range(2)] for i in range(2)]
    ta = a / math.sqrt(V[0][0]) if V[0][0] > 0 else None
    tb = b / math.sqrt(V[1][1]) if V[1][1] > 0 else None
    return a, ta, b, tb


def alpha(s, m, rf, a=None, z=None):
    ks = [k for k in sorted(s) if k in m and k in rf and (a is None or k >= a) and (z is None or k <= z)]
    y = [s[k] - rf[k] for k in ks]
    x = [m[k] - rf[k] for k in ks]
    al, ta, be, tb = hac_reg(y, x)
    return {'alpha_ann': round(al * 1200, 2), 't_alpha': round(ta, 2) if ta else None, 'beta': round(be, 3)}


def roll20(s, b, years=20):
    ks = sorted(k for k in s if k in b)
    res = []
    for y in range(ks[0] // 100, 2100):
        a0, z0 = y * 100 + 7, (y + years) * 100 + 6
        if z0 > ks[-1]:
            break
        w = [k for k in ks if a0 <= k <= z0]
        if len(w) < years * 12 * 0.97:
            continue
        gs = math.exp(math.fsum(math.log1p(s[k]) for k in w) / years) - 1
        gb = math.exp(math.fsum(math.log1p(b[k]) for k in w) / years) - 1
        res.append((y, round((gs - gb) * 100, 2)))
    if not res:
        return None
    v = sorted(c for _, c in res)
    return {'windows': len(res), 'wins': sum(c > 0 for _, c in res), 'win_rate': round(sum(c > 0 for _, c in res) / len(res), 3),
            'median': v[len(v) // 2], 'worst': min(res, key=lambda x: x[1])}


def dca20(s, b, n=240):
    ks = sorted(k for k in s if k in b)
    res = []
    for i in range(len(ks) - n + 1):
        w = ks[i:i + n]
        if ymadd(w[0], n - 1) != w[-1]:
            continue
        a = c = 0.0
        for k in w:
            a = (a + 1) * (1 + s[k])
            c = (c + 1) * (1 + b[k])
        res.append((w[0], w[-1], a / c))
    def summ(rr):
        if not rr:
            return None
        v = sorted(r for *_, r in rr)
        return {'windows': len(v), 'win_rate': round(sum(r > 1 for r in v) / len(v), 3), 'median': round(v[len(v) // 2], 3),
                'worst': round(v[0], 3)}
    return {'all': summ(res), 'hold_end': summ([x for x in res if x[1] >= HOLD_START])}


def maxdd(r, a=None, z=None):
    w = pk = 1.0
    dd = 0.0
    for k in sorted(r):
        if (a and k < a) or (z and k > z):
            continue
        w *= 1 + r[k]
        pk = max(pk, w)
        dd = min(dd, w / pk - 1)
    return round(dd * 100, 1)


def grade_mine(full, train, hold, costh, rl, repl_ok, sh_tr, sh_ho, lev=True):
    """out/mw_prereg.json の線を自分で当てる（研究者の grade 関数は使わない）"""
    c = {'C1': bool(train and train['ex_ann'] > 0 and (train['t'] or 0) >= 2.0),
         'C2': bool(hold and hold['ex_ann'] > 0 and hold['cagr_diff'] > 0),
         'C3': bool(hold and (hold['t'] or 0) >= 1.65),
         'C4': bool(rl and rl['win_rate'] >= 0.8),
         'C5': repl_ok,
         'C6': bool(costh and costh['ex_ann'] > 0 and costh['cagr_diff'] > 0),
         'C7': bool(full and (full['t'] or 0) >= 3.0),
         'C8': bool(sh_tr and sh_ho and sh_tr[0] > sh_tr[1] and sh_ho[0] > sh_ho[1]) if lev else None}
    base = c['C1'] and c['C2'] and c['C6'] and (c['C8'] in (None, True))
    if base and c['C3'] and c['C4'] and c['C7'] and c['C5'] in (None, True):
        g = 'S'
    elif base and c['C4'] and c['C7'] and (c['C3'] or c['C5'] is True):
        g = 'A'
    elif base:
        g = 'B'
    else:
        g = 'C'
    return g, c


# ═════════════════════════ 米国の評価（1規則ぶん） ═════════════════════════
def evaluate(nm_, g, n, W, mkt, rf, extra=True):
    e = {'full': ex(g, mkt), 'train': ex(g, mkt, z=TRAIN_END), 'hold': ex(g, mkt, a=HOLD_START),
         'recent_2013_07': ex(g, mkt, a=RECENT_START), 'from_2021': ex(g, mkt, a=202101),
         'net_full': ex(n, mkt), 'net_train': ex(n, mkt, z=TRAIN_END), 'net_hold': ex(n, mkt, a=HOLD_START),
         'roll20_net': roll20(n, mkt),
         'sharpe_net': {'train': [round(sharpe(n, rf, z=TRAIN_END), 3), round(sharpe(mkt, rf, a=min(n), z=TRAIN_END), 3)],
                        'hold': [round(sharpe(n, rf, a=HOLD_START), 3), round(sharpe(mkt, rf, a=HOLD_START), 3)]},
         'avg_w': {'full': round(S.mean(W.values()), 3), 'hold': round(S.mean(v for k, v in W.items() if k >= HOLD_START), 3)},
         'months_by_state': {str(s): sum(1 for v in W.values() if v == s) for s in sorted(set(W.values()))},
         'window': [min(W), max(W)]}
    if extra:
        e['dca20_net'] = dca20(n, mkt)
        e['maxdd_net'] = {'full': [maxdd(n), maxdd({k: mkt[k] for k in n})], 'hold': [maxdd(n, a=HOLD_START), maxdd(mkt, a=HOLD_START)]}
    return e


def const_lev(mkt, rf, L, keys, spread=SPREAD):
    """同じ平均倍率の一定レバレッジ（毎月 L 倍に戻す・借入 rf+spread）"""
    out = {}
    for t in keys:
        out[t] = L * mkt[t] - (L - 1) * (rf[t] + spread / 12) if L > 1 else L * mkt[t] + (1 - L) * rf[t]
    return out


# Reg T（連邦準備制度の当初証拠金率）の歴史。1934-10 以前は規制なし（2倍は可能とみなす）。
# 出典: Federal Reserve の Regulation T 当初証拠金率の改定表（記憶からの転記・日付は月単位に丸めた。要確認と明記）
REGT = [(193410, 0.45), (193602, 0.55), (193711, 0.40), (194502, 0.50), (194507, 0.75), (194601, 1.00), (194702, 0.75),
        (194903, 0.50), (195101, 0.75), (195302, 0.50), (195501, 0.60), (195504, 0.70), (195801, 0.50), (195808, 0.70),
        (195810, 0.90), (196007, 0.70), (196207, 0.50), (196311, 0.70), (196806, 0.80), (197005, 0.65), (197112, 0.55),
        (197211, 0.65), (197401, 0.50)]


def regt_cap(t):
    m = None
    for k, v in REGT:
        if t >= k:
            m = v
    return 99.0 if m is None else 1 / m


def us_block(us):
    mkt, rf, sig = us['mkt'], us['rf'], us['sig']
    bond = us['bond']
    idx = tr_index(mkt)
    up10 = sma_up(idx, 10)
    tsm = tsmom_up(mkt, rf)
    ecy = sig['ECY']
    rules = {
        'E3c_ECY_agree_lev2': dict(lev=2.0, trend=up10, alt=bond),
        'E3d_ECY_agree_tsmom': dict(lev=1.5, trend=tsm, alt=bond),
        'E3b_ECY_agree_lev_cash': dict(lev=1.5, trend=up10, alt=rf),
        'E2c_agree_lev': dict(lev=1.5, trend=up10, alt=bond),
    }
    R = {}
    series = {}
    for nm_, p in rules.items():
        wf = agree_state(ecy, p['trend'], THR_ECY, p['lev'])
        keys = valid_months(wf, [mkt, rf, p['alt']], 192607, END)
        g, n, W = simulate(keys, wf, mkt, p['alt'], rf)
        e = evaluate(nm_, g, n, W, mkt, rf)
        series[nm_] = (g, n, W, keys, wf, p)
        R[nm_] = e
        log(f"{nm_:24s} 全期間 {e['full']['ex_ann']:+.2f}(t{e['full']['t']}) 訓練 {e['train']['ex_ann']:+.2f}(t{e['train']['t']}) "
            f"保有 {e['hold']['ex_ann']:+.2f}(t{e['hold']['t']}) 幾何 {e['hold']['cagr_diff']:+.2f} 費用後保有 {e['net_hold']['ex_ann']:+.2f} "
            f"20年 {e['roll20_net']['win_rate']} シャープ {e['sharpe_net']} 窓 {e['window']}")
    return R, series, {'idx': idx, 'up10': up10, 'tsm': tsm}


# ═════════════════════════ 攻め 2: β を増やしただけか ═════════════════════════
def attack_beta(R, series, us):
    mkt, rf = us['mkt'], us['rf']
    out = {}
    for nm_, (g, n, W, keys, wf, p) in series.items():
        o = {}
        for per, (a, z) in {'full': (None, None), 'train': (None, TRAIN_END), 'hold': (HOLD_START, None), 'from_2013_07': (RECENT_START, None)}.items():
            o[f'alpha_net_{per}'] = alpha(n, mkt, rf, a, z)
        # 同じ平均倍率（全期間・保有期間それぞれ）の一定レバレッジと比べる
        Lf = S.mean(W.values())
        Lh = S.mean(v for k, v in W.items() if k >= HOLD_START)
        cf = const_lev(mkt, rf, Lf, keys)
        ch = const_lev(mkt, rf, Lh, keys)
        o['const_lev_same_avg_w'] = {'L_full': round(Lf, 3), 'L_hold': round(Lh, 3),
                                     'vs_const_full': ex(n, cf), 'vs_const_hold': ex(n, ch, a=HOLD_START),
                                     'const_vs_mkt_hold': ex(ch, mkt, a=HOLD_START), 'const_roll20': roll20(cf, mkt)}
        # 同じ変動率（保有期間）に揃えた市場（一定レバレッジ）と幾何で比べる
        ks_h = [k for k in keys if k >= HOLD_START]
        vs = S.stdev(n[k] for k in ks_h)
        vm = S.stdev(mkt[k] for k in ks_h)
        Lv = vs / vm
        cv = const_lev(mkt, rf, Lv, keys)
        o['vol_matched_hold'] = {'L': round(Lv, 3), 'vs_volmatched': ex(n, cv, a=HOLD_START)}
        o['sharpe_diff_hold'] = sharpe_diff_test(n, mkt, rf, a=HOLD_START)
        o['sharpe_diff_full'] = sharpe_diff_test(n, mkt, rf)
        out[nm_] = o
        log(f"β点検 {nm_:24s} α保有 {o['alpha_net_hold']} α全期間 {o['alpha_net_full']} 一定{Lh:.2f}倍との差(保有) {o['const_lev_same_avg_w']['vs_const_hold']['ex_ann']:+.2f}"
            f"(t{o['const_lev_same_avg_w']['vs_const_hold']['t']}) 同じぶれの市場との幾何差 {o['vol_matched_hold']['vs_volmatched']['cagr_diff']:+.2f} シャープ差z {o['sharpe_diff_hold']['z']}")
    return out


# ═════════════════════════ 攻め 3: 小区間 ═════════════════════════
def attack_subperiods(series, us):
    mkt = us['mkt']
    out = {}
    halves = {'hold_1st_2007_01_2016_10': (200701, 201610), 'hold_2nd_2016_11_2026_08': (201611, 202608)}
    for nm_, (g, n, W, keys, wf, p) in series.items():
        o = {}
        o['hold_net_drop_2008'] = ex(n, mkt, a=HOLD_START, drop=lambda k: 200801 <= k <= 200812)
        o['hold_net_drop_2008_2009'] = ex(n, mkt, a=HOLD_START, drop=lambda k: 200801 <= k <= 200912)
        o['hold_net_drop_2020_2021'] = ex(n, mkt, a=HOLD_START, drop=lambda k: 202001 <= k <= 202112)
        o['full_net_drop_1998_2000_2020_2021'] = ex(n, mkt, drop=lambda k: 199801 <= k <= 200012 or 202001 <= k <= 202112)
        o['full_net_drop_1927_1945'] = ex(n, mkt, a=194601)
        o['net_1946_2006'] = ex(n, mkt, a=194601, z=TRAIN_END)
        o['net_1975_2006'] = ex(n, mkt, a=197501, z=TRAIN_END)
        for h, (a, z) in halves.items():
            o[h] = ex(n, mkt, a=a, z=z)
        o['net_2009_2016'] = ex(n, mkt, a=200901, z=201612)
        o['net_2017_2026'] = ex(n, mkt, a=201701)
        o['net_from_2021'] = ex(n, mkt, a=202101)
        o['net_from_2013_07'] = ex(n, mkt, a=RECENT_START)
        # 保有期間の超過（算術・年率）を状態ごとに分ける
        ks = [k for k in keys if k >= HOLD_START]
        yrs = len(ks) / 12
        by = {}
        for k in ks:
            s = str(W[k])
            by.setdefault(s, [0, 0.0])
            by[s][0] += 1
            by[s][1] += n[k] - mkt[k]
        o['hold_excess_by_state_pp_per_year'] = {s: {'months': c, 'contrib_pp_per_yr': round(v / yrs * 100, 2)} for s, (c, v) in by.items()}
        # 年ごとの超過（幾何）
        yr = {}
        for k in ks:
            y = k // 100
            a_, b_ = yr.get(y, (1.0, 1.0))
            yr[y] = (a_ * (1 + n[k]), b_ * (1 + mkt[k]))
        o['hold_year_by_year_pp'] = {str(y): round((a_ - b_) * 100, 1) for y, (a_, b_) in sorted(yr.items())}
        # 2008 年の寄与を外すと費用後の幾何差はどうなるか → 上の hold_net_drop_2008
        out[nm_] = o
        log(f"小区間 {nm_:24s} 2008除く {o['hold_net_drop_2008']['ex_ann']:+.2f}(t{o['hold_net_drop_2008']['t']}) 幾何{o['hold_net_drop_2008']['cagr_diff']:+.2f} "
            f"前半 {o['hold_1st_2007_01_2016_10']['ex_ann']:+.2f}(t{o['hold_1st_2007_01_2016_10']['t']}) 後半 {o['hold_2nd_2016_11_2026_08']['ex_ann']:+.2f}(t{o['hold_2nd_2016_11_2026_08']['t']}) "
            f"2021〜 {o['net_from_2021']['ex_ann']:+.2f}(t{o['net_from_2021']['t']}) 1946-2006 {o['net_1946_2006']['ex_ann']:+.2f}(t{o['net_1946_2006']['t']}) "
            f"1975-2006 {o['net_1975_2006']['ex_ann']:+.2f}(t{o['net_1975_2006']['t']}) 状態別 {o['hold_excess_by_state_pp_per_year']}")
    return out


# ═════════════════════════ 攻め 4: パラメータの脆さ ═════════════════════════
def qt(v, p):
    v = sorted(v)
    return v[int(p * (len(v) - 1))]


def attack_grid(us, tr):
    mkt, rf, bond, ecy = us['mkt'], us['rf'], us['bond'], us['sig']['ECY']
    idx = tr['idx']
    ecy_tr = [v for k, v in ecy.items() if k <= TRAIN_END]
    qs = [0.30, 0.35, 0.40, 0.45, 0.50, 0.55, 0.60, 0.65, 0.70]
    ns = [6, 8, 10, 12, 14]
    trends = {n_: sma_up(idx, n_) for n_ in ns}
    trends['tsmom12'] = tr['tsm']
    out = {}
    for lev, alt_name in ((2.0, 'bond'), (1.5, 'bond'), (1.5, 'cash')):
        alt = bond if alt_name == 'bond' else rf
        rows = []
        for q in qs:
            thr = qt(ecy_tr, q)
            for tn, trend in trends.items():
                for lag in (1, 2):
                    wf = agree_state(ecy, trend, thr, lev, lag)
                    keys = valid_months(wf, [mkt, rf, alt], 192607, END)
                    g, n, W = simulate(keys, wf, mkt, alt, rf)
                    h = ex(n, mkt, a=HOLD_START)
                    f = ex(g, mkt)
                    t_ = ex(g, mkt, z=TRAIN_END)
                    hg = ex(g, mkt, a=HOLD_START)
                    rows.append({'q': q, 'thr': round(thr, 5), 'trend': str(tn), 'lag': lag, 'train_ex': t_['ex_ann'], 'train_t': t_['t'],
                                 'hold_ex_gross': hg['ex_ann'], 'hold_t_gross': hg['t'], 'hold_net_cagr_diff': h['cagr_diff'],
                                 'full_t': f['t'], 'from_2013_07_net_cagr': ex(n, mkt, a=RECENT_START)['cagr_diff']})
        # 要約
        def summ(rr):
            v = sorted(r['hold_net_cagr_diff'] for r in rr)
            return {'n': len(rr), 'hold_net_cagr_pos': sum(x > 0 for x in v), 'hold_t_ge_1.65': sum((r['hold_t_gross'] or 0) >= 1.65 for r in rr),
                    'full_t_ge_3': sum((r['full_t'] or 0) >= 3 for r in rr), 'train_t_ge_2': sum((r['train_t'] or 0) >= 2 for r in rr),
                    'hold_net_cagr_median': v[len(v) // 2], 'hold_net_cagr_min': v[0], 'hold_net_cagr_max': v[-1]}
        key = f'lev{lev}_{alt_name}'
        out[key] = {'all': summ(rows), 'lag1_only': summ([r for r in rows if r['lag'] == 1]),
                    'lag1_sma10_by_q': {str(r['q']): [r['hold_net_cagr_diff'], r['hold_t_gross'], r['full_t']] for r in rows if r['lag'] == 1 and r['trend'] == '10'},
                    'lag1_q50_by_trend': {r['trend']: [r['hold_net_cagr_diff'], r['hold_t_gross'], r['full_t']] for r in rows if r['lag'] == 1 and r['q'] == 0.5},
                    'rows': rows}
        log(f"格子 {key}: {out[key]['all']} ／ SMA10・遅れ1の分位ごと(保有費用後の幾何差, 保有t, 全期間t) {out[key]['lag1_sma10_by_q']}")
    return out


# ═════════════════════════ 攻め 5: 実行できたか ═════════════════════════
def hist_margin_rate(us):
    """個人の信用取引の金利の近似（月利）: 1926-07〜1970-11 は NY のコールレート＋1%（下限 RF+1.5%）、
    1970-12〜2006-12 はプライムレート（下限 RF+1.5%）、2007〜 は RF+1.5%（今のネット証券並み）"""
    call = fred_m('M13001USM156NNBR')
    prime = fred_m('MPRIME')
    rf = us['rf']
    out = {}
    for t in rf:
        base = rf[t] + SPREAD / 12
        a = ymadd(t, -1)
        if t <= 197011 and a in call:
            out[t] = max(base, (call[a] + 1.0) / 1200)
        elif t <= TRAIN_END and a in prime:
            out[t] = max(base, prime[a] / 1200)
        else:
            out[t] = base
    return out


def tax_japan(keys, W, stock, alt, rf, spread=SPREAD, rate=0.20315, cost=COST):
    """日本の課税口座の近似: 平均取得価額で売った分の損益を実現・年末に課税（3年の繰越控除つき）・借入の利息は経費で差し引く。
    最後に全部売って課税。相手は買って持って最後に売る"""
    Sv = Bs = 0.0   # 株の時価・取得価額
    Av = Ba = 0.0   # 逃げ先
    Lo = 0.0        # 借入
    nav = 1.0
    real = 0.0
    carry = []      # (年, 繰越損失)
    first = True
    for t in keys:
        w = W[t]
        # 月初に目標へ
        tgt_s = w * nav
        tgt_a = max(0.0, 1 - w) * nav if w < 1 else 0.0
        tgt_l = max(0.0, w - 1) * nav
        if first:
            Sv, Bs, Av, Ba, Lo = tgt_s, tgt_s, tgt_a, tgt_a, tgt_l
            nav -= (tgt_s + tgt_a) * cost
            first = False
        else:
            tr_ = abs(tgt_s - Sv) + abs(tgt_a - Av)
            if tgt_s < Sv and Sv > 0:
                q = (Sv - tgt_s) / Sv
                real += (Sv - tgt_s) - Bs * q
                Bs *= 1 - q
            elif tgt_s > Sv:
                Bs += tgt_s - Sv
            if tgt_a < Av and Av > 0:
                q = (Av - tgt_a) / Av
                real += (Av - tgt_a) - Ba * q
                Ba *= 1 - q
            elif tgt_a > Av:
                Ba += tgt_a - Av
            Sv, Av, Lo = tgt_s, tgt_a, tgt_l
            c = tr_ * cost
            real -= c
            Sv -= c * (Sv / (Sv + Av)) if Sv + Av > 0 else 0
        # 月の値動き
        rs = stock[t]
        ra = alt[t] if Av > 0 else 0.0
        interest = Lo * (rf[t] + spread / 12)
        Sv *= 1 + rs
        Av *= 1 + ra
        real -= interest
        nav = Sv + Av - Lo - interest
        # 利息は借入に上乗せせず、その月に払った（株を少し売った）とみなす: 近似として nav に反映済み
        Sv_adj = Sv - interest if Sv > interest else Sv
        Sv = Sv_adj
        if t % 100 == 12 or t == keys[-1]:
            if t == keys[-1]:
                real += (Sv - Bs) + (Av - Ba)
            y = t // 100
            carry = [(yy, v) for yy, v in carry if y - yy <= 3]
            if real > 0:
                for i, (yy, v) in enumerate(carry):
                    use = min(real, v)
                    real -= use
                    carry[i] = (yy, v - use)
                tax = real * rate
                if t == keys[-1]:
                    nav -= tax
                else:
                    # 税は株と逃げ先から比例で払う（その売りで生じる小さな実現は無視）
                    tot = Sv + Av
                    if tot > 0:
                        Sv -= tax * Sv / tot
                        Av -= tax * Av / tot
                        Bs -= tax * Bs / max(tot, 1e-12) * 0  # 取得価額は変えない近似
                    nav -= tax
            elif real < 0:
                carry.append((y, -real))
            real = 0.0
    bh = 1.0
    for t in keys:
        bh *= 1 + stock[t]
    bh_after = bh - max(0.0, bh - 1) * rate
    return {'strategy_after_tax': round(nav, 4), 'bh_after_tax': round(bh_after, 4), 'ratio': round(nav / bh_after, 4),
            'years': round(len(keys) / 12, 1),
            'cagr_diff_after_tax_pp': round(((nav) ** (12 / len(keys)) - (bh_after) ** (12 / len(keys))) * 100, 2)}


def attack_implementable(series, us, tr):
    mkt, rf, bond = us['mkt'], us['rf'], us['bond']
    hm = hist_margin_rate(us)
    out = {}
    for nm_, (g0, n0, W0, keys, wf, p) in series.items():
        o = {}
        alt = p['alt']
        # (a) Reg T の当初証拠金率で倍率に上限（1934-10〜）
        g, n, W = simulate(keys, wf, mkt, alt, rf, cap=regt_cap)
        o['regT_cap'] = {'full': ex(g, mkt), 'train': ex(g, mkt, z=TRAIN_END), 'hold': ex(g, mkt, a=HOLD_START), 'net_hold': ex(n, mkt, a=HOLD_START),
                         'roll20_net': roll20(n, mkt), 'sharpe_train': [round(sharpe(n, rf, z=TRAIN_END), 3), round(sharpe(mkt, rf, a=min(n), z=TRAIN_END), 3)],
                         'months_capped': sum(1 for k in keys if wf(k) > 1 and regt_cap(k) < wf(k)),
                         'net_1946_1974': ex(n, mkt, a=194601, z=197412)}
        # (b) 歴史の個人の信用金利（コール＋1% → プライム → RF+1.5%）
        g, n, W = simulate(keys, wf, mkt, alt, rf, borrow=hm)
        o['hist_margin_rate'] = {'full': ex(g, mkt), 'train': ex(g, mkt, z=TRAIN_END), 'net_hold': ex(n, mkt, a=HOLD_START), 'roll20_net': roll20(n, mkt)}
        # (c) 両方（Reg T ＋ 歴史の金利）
        g, n, W = simulate(keys, wf, mkt, alt, rf, borrow=hm, cap=regt_cap)
        o['regT_and_hist_rate'] = {'full': ex(g, mkt), 'train': ex(g, mkt, z=TRAIN_END), 'net_full': ex(n, mkt), 'net_hold': ex(n, mkt, a=HOLD_START),
                                   'roll20_net': roll20(n, mkt)}
        # (d) 借入 RF+3%
        g, n, W = simulate(keys, wf, mkt, alt, rf, spread=0.03)
        o['spread_3pct'] = {'full': ex(g, mkt), 'net_hold': ex(n, mkt, a=HOLD_START)}
        # (e) 往復（逃げ先の売買も数える）・片道 0.10% と 0.30%
        for c in (0.001, 0.003):
            g, n, W = simulate(keys, wf, mkt, alt, rf, cost=c, two_leg=True)
            o[f'two_leg_cost_{c}'] = {'net_full': ex(n, mkt), 'net_hold': ex(n, mkt, a=HOLD_START), 'turnover_per_yr_note': 'stock+alt legs'}
        # (f) 債券の系列を Goyal ltr（長期国債）に替える
        if alt is bond:
            g, n, W = simulate(keys, wf, mkt, us['bond_ltr'], rf)
            o['bond_goyal_ltr'] = {'full': ex(g, mkt), 'net_hold': ex(n, mkt, a=HOLD_START)}
        # (g) 日本の課税口座（保有期間に一括1）
        kh = [k for k in keys if k >= HOLD_START]
        o['tax_japan_hold'] = tax_japan(kh, W0, mkt, alt, rf)
        out[nm_] = o
        log(f"実行 {nm_:24s} RegT: 全期間 {o['regT_cap']['full']['ex_ann']:+.2f}(t{o['regT_cap']['full']['t']}) 訓練t {o['regT_cap']['train']['t']} "
            f"20年 {o['regT_cap']['roll20_net']['win_rate']} 上限にかかった月 {o['regT_cap']['months_capped']} ／ 歴史の金利: 全期間t {o['hist_margin_rate']['full']['t']} "
            f"／ 両方: 全期間 {o['regT_and_hist_rate']['full']['ex_ann']:+.2f}(t{o['regT_and_hist_rate']['full']['t']}) 訓練t {o['regT_and_hist_rate']['train']['t']} 20年 {o['regT_and_hist_rate']['roll20_net']['win_rate']} "
            f"／ RF+3% 保有費用後 {o['spread_3pct']['net_hold']['cagr_diff']:+.2f} ／ 往復0.3% 保有 {o['two_leg_cost_0.003']['net_hold']['cagr_diff']:+.2f} ／ 課税 {o['tax_japan_hold']}")
    return out


# ═════════════════════════ F1・F2: 日次レバレッジ投信/ETF で持つ形 ═════════════════════════
def lev_daily_monthly(rd, rfd, L, fee=0.009, spread=0.005):
    """日次 L 倍（借入 rf_日 + spread/252・信託報酬 fee/252）を月へ複利 → {yyyymm: r}"""
    out = {}
    for d in sorted(rd):
        f = rfd.get(d)
        if f is None:
            continue
        v = L * rd[d] - (L - 1) * (f + spread / 252) - fee / 252
        m = d // 100
        out[m] = (1 + out.get(m, 0.0)) * (1 + v) - 1
    return out


def daily_to_month(rd):
    out = {}
    for d in sorted(rd):
        m = d // 100
        out[m] = (1 + out.get(m, 0.0)) * (1 + rd[d]) - 1
    return out


def attack_F(us, tr):
    mkt, rf, bond, ecy = us['mkt'], us['rf'], us['bond'], us['sig']['ECY']
    up = tr['up10']
    ffd = M.ff_factors('daily')
    md = {k: v for k, v in ffd['mkt'].items() if k <= 20260831}
    rfd = {k: v for k, v in ffd['rf'].items() if k <= 20260831}
    chk = daily_to_month(md)
    out = {'french_daily_vs_monthly_max_abs_pp': round(max(abs(chk[k] - mkt[k]) for k in chk if k in mkt) * 100, 3)}
    # 実在の SSO（2006-06 設定・日次2倍 S&P500）とモデルの差
    try:
        sso = M.yahoo('SSO', '1mo')
        sptr_d = M.yahoo('^SP500TR', '1d')
        model_sp = lev_daily_monthly(sptr_d, {d: rfd[d] for d in rfd}, 2.0)
        model_fr = lev_daily_monthly(md, rfd, 2.0)
        ks = [k for k in sorted(sso) if 200607 <= k <= END and k in model_sp and k in model_fr]
        out['SSO_reality'] = {'from': ks[0], 'to': ks[-1], 'months': len(ks),
                              'cagr_SSO': round(geo([sso[k] for k in ks]) * 100, 2),
                              'cagr_model_on_SP500TR': round(geo([model_sp[k] for k in ks]) * 100, 2),
                              'cagr_model_on_French_mkt': round(geo([model_fr[k] for k in ks]) * 100, 2),
                              'gap_SSO_minus_model_SP_pp': round((geo([sso[k] for k in ks]) - geo([model_sp[k] for k in ks])) * 100, 2)}
        log('SSO の実績とモデル', out['SSO_reality'])
    except Exception as e:  # noqa
        out['SSO_reality'] = f'取得失敗: {e}'
    for name, L in (('F1_SP_ETF2x_agree', 2.0), ('F2_SP_ETF15x_agree', 1.5)):
        levm = {k: v for k, v in lev_daily_monthly(md, rfd, L).items() if k <= END}
        ser = {L: levm, 1.0: mkt, 0.0: bond}
        wf = agree_state(ecy, up, THR_ECY, L)
        keys = valid_months(wf, [mkt, rf, bond, levm], 192607, END)
        g, n1, n2, W = {}, {}, {}, {}
        prev = None
        for t in keys:
            w = wf(t)
            r = ser[w][t]
            g[t] = r
            # 研究者の費用: |Δw|×0.10%。自分の費用: 状態が変わるたびに全部売って別の商品を全部買う（往復 2×0.10%）
            n1[t] = r - (0.0 if prev is None else abs(w - prev) * COST)
            n2[t] = r - (0.0 if prev is None or w == prev else 2 * COST)
            W[t] = w
            prev = w
        e = evaluate(name, g, n1, W, mkt, rf)
        e['net_hold_two_leg'] = ex(n2, mkt, a=HOLD_START)
        e['alpha_net_hold'] = alpha(n2, mkt, rf, HOLD_START)
        e['sharpe_diff_hold'] = sharpe_diff_test(n2, mkt, rf, a=HOLD_START)
        e['vs_const_same_avg_w_hold'] = ex(n2, const_lev(mkt, rf, S.mean(v for k, v in W.items() if k >= HOLD_START), keys), a=HOLD_START)
        e['hold_net_drop_2008'] = ex(n2, mkt, a=HOLD_START, drop=lambda k: 200801 <= k <= 200812)
        e['hold_1st'] = ex(n2, mkt, a=200701, z=201610)
        e['hold_2nd'] = ex(n2, mkt, a=201611)
        e['from_2021'] = ex(n2, mkt, a=202101)
        e['product_existed_from'] = '2006-06（ProShares SSO）。1940年投資会社法の借入制限（資産の1/3）で、デリバティブ以前の2倍の投信は作れなかった＝訓練期間の2倍は仮想'
        # 実在の商品がある期間だけ（2006-07〜）の2倍の月は SSO の実績に置き換える
        if isinstance(out.get('SSO_reality'), dict):
            n3 = {}
            prev = None
            for t in keys:
                w = W[t]
                r = sso[t] if (w == 2.0 and L == 2.0 and t in sso and t >= 200607) else g[t]
                n3[t] = r - (0.0 if prev is None or w == prev else 2 * COST)
                prev = w
            e['net_hold_with_real_SSO'] = ex(n3, mkt, a=HOLD_START) if L == 2.0 else None
        # 実在の商品だけで組んだ保有期間（2x=SSO・1x=SPY・0=IEF・相手=SPY）。信号は同じ（ECY＋French 指数の10か月線）と、S&P の指数の10か月線の2通り
        if L == 2.0 and isinstance(out.get('SSO_reality'), dict):
            try:
                spy = M.yahoo('SPY', '1mo')
                ief = M.yahoo('IEF', '1mo')
                sp_idx = tr_index({k: v for k, v in spy.items() if k <= END})
                up_sp = sma_up(sp_idx, 10)
                real = {}
                for tag, trend in (('trend_French', up), ('trend_SPY', up_sp)):
                    wf2 = agree_state(ecy, trend, THR_ECY, 2.0)
                    kk = [t for t in mrange(HOLD_START, END) if wf2(t) is not None and t in sso and t in spy and t in ief]
                    nr, prev = {}, None
                    for t in kk:
                        w = wf2(t)
                        r = {2.0: sso, 1.0: spy, 0.0: ief}[w][t]
                        nr[t] = r - (0.0 if prev is None or w == prev else 2 * COST)
                        prev = w
                    real[tag] = {'vs_SPY': ex(nr, spy), 'alpha_vs_SPY': alpha(nr, spy, rf), 'from_2021': ex(nr, spy, a=202101),
                                 'sharpe': [round(sharpe(nr, rf), 3), round(sharpe({k: spy[k] for k in nr}, rf), 3)]}
                e['real_products_hold'] = real
            except Exception as ee:  # noqa
                e['real_products_hold'] = f'取得失敗: {ee}'
        out[name] = e
        log(f"{name}: 実在の商品（SSO/SPY/IEF）保有期間 {e.get('real_products_hold')}")
        log(f"{name}: 全期間 {e['full']['ex_ann']:+.2f}(t{e['full']['t']}) 訓練 {e['train']['ex_ann']:+.2f}(t{e['train']['t']}) 保有 {e['hold']['ex_ann']:+.2f}(t{e['hold']['t']}) "
            f"費用後(研究者) {e['net_hold']['cagr_diff']:+.2f} 費用後(往復) {e['net_hold_two_leg']['cagr_diff']:+.2f} 実SSO {e.get('net_hold_with_real_SSO')} α保有 {e['alpha_net_hold']} "
            f"一定倍率との差 {e['vs_const_same_avg_w_hold']['ex_ann']:+.2f}(t{e['vs_const_same_avg_w_hold']['t']}) 2008除く {e['hold_net_drop_2008']['ex_ann']:+.2f}(t{e['hold_net_drop_2008']['t']}) "
            f"前半 {e['hold_1st']['ex_ann']:+.2f} 後半 {e['hold_2nd']['ex_ann']:+.2f} 2021〜 {e['from_2021']['ex_ann']:+.2f}(t{e['from_2021']['t']}) 20年 {e['roll20_net']['win_rate']} シャープ {e['sharpe_net']}")
    return out


# ═════════════════════════ 攻め 6: 米国外 18か国（自作で組み直す） ═════════════════════════
CTRY = {  # French のファイル名: (ISO2, ISO3)
    'UK.Dat': ('GB', 'GBR'), 'Austria.Dat': ('AT', 'AUT'), 'Austrlia.Dat': ('AU', 'AUS'), 'Belgium.Dat': ('BE', 'BEL'),
    'Canada.Dat': ('CA', 'CAN'), 'Denmark.Dat': ('DK', 'DNK'), 'Finland.Dat': ('FI', 'FIN'), 'France.Dat': ('FR', 'FRA'),
    'Germany.Dat': ('DE', 'DEU'), 'Ireland.Dat': ('IE', 'IRL'), 'Italy.Dat': ('IT', 'ITA'), 'Japan.Dat': ('JP', 'JPN'),
    'Nethrlnd.Dat': ('NL', 'NLD'), 'NewZland.Dat': ('NZ', 'NZL'), 'Norway.Dat': ('NO', 'NOR'), 'Spain.Dat': ('ES', 'ESP'),
    'Sweden.Dat': ('SE', 'SWE'), 'Swtzrlnd.Dat': ('CH', 'CHE')}


def parse_country_file(txt):
    """French の国別 .Dat（自作の解析）→ 月次の現地通貨 Mkt・年次の比率（B/M・E/P・Yld、小数）"""
    lines = txt.splitlines()
    loc, ann = {}, {}
    mode = None
    seen_local = 0
    seen_avg = False
    for ln in lines:
        s = ln.strip()
        if 'Local' in s and 'Returns' in s and 'Not Reqd' in s:
            seen_local += 1
            mode = 'loc' if seen_local == 1 else None
            continue
        if ('Returns' in s and 'Value-Weight' in s):
            mode = None
            continue
        if s.startswith('Average of Annual'):
            mode = 'avg_hdr'
            continue
        if mode == 'avg_hdr' and s.startswith('All 4 Data Items'):
            mode = 'ann' if ('Not Reqd' in s and not seen_avg) else None
            seen_avg = True
            continue
        p = s.split()
        if not p or not re.match(r'^\d{4,6}$', p[0]):
            continue
        if mode == 'loc' and len(p[0]) == 6:
            v = float(p[1])
            if v > -99:
                loc[int(p[0])] = v / 100
        elif mode == 'ann' and len(p[0]) == 4:
            vals = [float(x) for x in p[1:6]]
            ann[int(p[0])] = {'BM': vals[1] / 100 if vals[1] > -99 else None, 'EP': vals[2] / 100 if vals[2] > -99 else None,
                              'Yld': vals[4] / 100 if vals[4] > -99 else None}
    return loc, ann


def wb_cpi_all():
    iso3 = ';'.join(v[1] for v in CTRY.values())
    b = M.get(f'https://api.worldbank.org/v2/country/{iso3}/indicator/FP.CPI.TOTL?format=json&per_page=2000&date=1955:2026',
              'wb_cpi_valdca.json', max_age_days=3650)
    out = {}
    for r in json.loads(b)[1]:
        if r['value'] is not None:
            out.setdefault(r['countryiso3code'], {})[int(r['date'])] = float(r['value'])
    return out


def build_country(fname, txt, cpi_all):
    iso2, iso3 = CTRY[fname]
    tr, ann = parse_country_file(txt)
    cpi = cpi_all.get(iso3, {})
    y10 = fred_m(f'IRLTLT01{iso2}M156N')
    r3 = fred_m(f'IR3TIB01{iso2}M156N')
    rc = fred_m(f'IRSTCI01{iso2}M156N')
    ks = sorted(tr)

    def ky(t):  # 年 Y の比率は Y 年7月から
        y, m = divmod(t, 100)
        return y if m >= 7 else y - 1
    PI = {ymadd(ks[0], -1): 1.0}
    lv = 1.0
    for t in ks:
        a = ann.get(ky(t)) or ann.get(t // 100)
        dy = (a['Yld'] if a and a['Yld'] is not None else 0.0) / 12
        lv *= (1 + tr[t]) / (1 + dy)
        PI[t] = lv
    Efy = {}
    for Y, a in ann.items():
        dec = (Y - 1) * 100 + 12
        if a['EP'] is not None and dec in PI:
            Efy[Y - 1] = a['EP'] * PI[dec]
    ecy = {}
    for t in ks:
        F = ky(t) - 1
        fs = []
        for i in range(10):
            if (F - i) in Efy and (F - i) in cpi:
                fs.append(F - i)
            else:
                break
        cy = t // 100 - 1
        if len(fs) < 5 or cy not in cpi or (cy - 10) not in cpi or t not in y10:
            continue
        er = sum(Efy[f] / cpi[f] for f in fs) / len(fs)
        if er <= 0:
            continue
        cape = (PI[t] / cpi[cy]) / er
        pi10 = (cpi[cy] / cpi[cy - 10]) ** 0.1 - 1
        ecy[t] = 1 / cape - (y10[t] / 100 - pi10)
    cash, bond, bond_shift = {}, {}, {}
    for t in ks:
        a = ymadd(t, -1)
        rate = r3.get(a, rc.get(a))
        if rate is not None:
            cash[t] = rate / 1200
        if a in y10 and t in y10:
            bond[t] = par_bond_semi(y10[a] / 100, y10[t] / 100)
        nx = ymadd(t, 1)
        if t in y10 and nx in y10:
            bond_shift[t] = par_bond_semi(y10[t] / 100, y10[nx] / 100)
    ev = [v for k, v in ecy.items() if k <= TRAIN_END]
    thr = qt(ev, 0.5) if len(ev) >= 60 else None
    idx = tr_index(tr)
    return {'mkt': tr, 'cash': cash, 'bond': bond, 'bond_shift': bond_shift, 'ecy': ecy, 'thr': thr,
            'up': sma_up(idx, 10), 'tsm': tsmom_up(tr, cash), 'n_train_ecy': len(ev)}


def attack_countries():
    z = zipfile.ZipFile(io.BytesIO(M.get(M.FR.format('F-F_International_Countries'), name='fr_F-F_International_Countries.zip', max_age_days=3650)))
    cpi_all = wb_cpi_all()
    C = {}
    for fn in CTRY:
        C[fn] = build_country(fn, z.read(fn).decode('latin-1'), cpi_all)
    rules = {'E3c_ECY_agree_lev2': (2.0, 'up', 'bond'), 'E3d_ECY_agree_tsmom': (1.5, 'tsm', 'bond'),
             'E3b_ECY_agree_lev_cash': (1.5, 'up', 'cash'), 'E2c_agree_lev': (1.5, 'up', 'bond'),
             'E3c_cash_version': (2.0, 'up', 'cash'), 'E3c_bond_shifted_no_lookahead': (2.0, 'up', 'bond_shift')}
    out = {}
    for rn, (L, tk, ak) in rules.items():
        per, pooled_m, pooled_c, pooled_m_h, pooled_c_h = {}, {}, {}, {}, {}
        cnt = {'regions': 0, 'pos_vs_mkt': 0, 'const_lev_pos_vs_mkt': 0, 'pos_vs_const_lev': 0, 'sharpe_gt_mkt': 0, 'alpha_t_ge_2': 0,
               'hold_pos_vs_mkt': 0, 'hold_pos_vs_const': 0}
        for fn, c in C.items():
            if c['thr'] is None:
                per[fn] = 'N/A 訓練の信号<60か月'
                continue
            alt = c[ak]
            wf = agree_state(c['ecy'], c[tk], c['thr'], L)
            keys = [k for k in valid_months(wf, [c['mkt'], c['cash'], alt], 197501, 202512)]
            if len(keys) < 120:
                per[fn] = f'N/A {len(keys)}か月'
                continue
            g, n, W = simulate(keys, wf, c['mkt'], alt, c['cash'])
            Lbar = S.mean(W.values())
            cl = const_lev(c['mkt'], c['cash'], Lbar, keys)
            xm = ex(n, c['mkt'])
            xc = ex(n, cl)
            cm = ex(cl, c['mkt'])
            xh = ex(n, c['mkt'], a=HOLD_START)
            xch = ex(n, cl, a=HOLD_START)
            al = alpha(n, c['mkt'], c['cash'])
            shs, shm = sharpe(n, c['cash']), sharpe({k: c['mkt'][k] for k in n}, c['cash'])
            cnt['regions'] += 1
            cnt['pos_vs_mkt'] += xm['ex_ann'] > 0 and xm['cagr_diff'] > 0
            cnt['const_lev_pos_vs_mkt'] += cm['ex_ann'] > 0 and cm['cagr_diff'] > 0
            cnt['pos_vs_const_lev'] += xc['ex_ann'] > 0 and xc['cagr_diff'] > 0
            cnt['sharpe_gt_mkt'] += shs > shm
            cnt['alpha_t_ge_2'] += (al['t_alpha'] or 0) >= 2
            if xh:
                cnt['hold_pos_vs_mkt'] += xh['ex_ann'] > 0 and xh['cagr_diff'] > 0
                cnt['hold_pos_vs_const'] += xch['ex_ann'] > 0 and xch['cagr_diff'] > 0
            per[fn] = {'window': [keys[0], keys[-1]], 'thr': round(c['thr'], 4), 'avg_w': round(Lbar, 3), 'vs_mkt': xm, 'vs_const_lev': xc,
                       'const_lev_vs_mkt': cm, 'hold_vs_mkt': xh, 'hold_vs_const': xch, 'alpha': al, 'sharpe': [round(shs, 3), round(shm, 3)]}
            for k in n:
                pooled_m.setdefault(k, []).append(n[k] - c['mkt'][k])
                pooled_c.setdefault(k, []).append(n[k] - cl[k])
        pm = {k: sum(v) / len(v) for k, v in pooled_m.items() if len(v) >= 5}
        pc = {k: sum(v) / len(v) for k, v in pooled_c.items() if len(v) >= 5}
        zero = {k: 0.0 for k in pm}
        out[rn] = {'counts': cnt,
                   'pooled_vs_mkt': {'full': ex(pm, zero), 'train': ex(pm, zero, z=TRAIN_END), 'hold': ex(pm, zero, a=HOLD_START), 'from_2013_07': ex(pm, zero, a=RECENT_START)},
                   'pooled_vs_const_lev': {'full': ex(pc, zero), 'train': ex(pc, zero, z=TRAIN_END), 'hold': ex(pc, zero, a=HOLD_START), 'from_2013_07': ex(pc, zero, a=RECENT_START)},
                   'pooled_subperiods_vs_mkt': {'2007_2016': ex(pm, zero, a=200701, z=201612), '2017_2025': ex(pm, zero, a=201701), 'from_2021': ex(pm, zero, a=202101),
                                                'hold_drop_2008': ex(pm, zero, a=HOLD_START, drop=lambda k: 200801 <= k <= 200812),
                                                '1979_1994': ex(pm, zero, z=199412), '1995_2006': ex(pm, zero, a=199501, z=TRAIN_END)},
                   'pooled_subperiods_vs_const': {'2007_2016': ex(pc, zero, a=200701, z=201612), '2017_2025': ex(pc, zero, a=201701), 'from_2021': ex(pc, zero, a=202101)},
                   'per_country': per}
        sub_m = {k: (v['ex_ann'], v['t']) for k, v in out[rn]['pooled_subperiods_vs_mkt'].items()}
        sub_c = {k: (v['ex_ann'], v['t']) for k, v in out[rn]['pooled_subperiods_vs_const'].items()}
        log(f"国 {rn:30s} {cnt} 束ね(対市場) 保有 {out[rn]['pooled_vs_mkt']['hold']['ex_ann']:+.2f}(t{out[rn]['pooled_vs_mkt']['hold']['t']}) "
            f"束ね(対 同じ平均倍率の一定レバレッジ) 全期間 {out[rn]['pooled_vs_const_lev']['full']['ex_ann']:+.2f}(t{out[rn]['pooled_vs_const_lev']['full']['t']}) "
            f"保有 {out[rn]['pooled_vs_const_lev']['hold']['ex_ann']:+.2f}(t{out[rn]['pooled_vs_const_lev']['hold']['t']}) 小区間(対市場) {sub_m} (対一定) {sub_c}")
    return out


# ═════════════════════════ 攻め 7: JST 1870〜1974（自作で組み直す） ═════════════════════════
def load_jst():
    import openpyxl
    b = M.get('https://www.macrohistory.net/app/download/9834512569/JSTdatasetR6.xlsx?t=1763503850', 'jst_R6.xlsx', max_age_days=3650)
    wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
    rows = list(wb['Sheet1'].iter_rows(values_only=True))
    h = list(rows[0])
    cols = ('eq_tr', 'eq_dp', 'bond_tr', 'bill_rate', 'cpi', 'ltrate')
    I = {c: h.index(c) for c in cols + ('year', 'iso')}
    out = {}
    for r in rows[1:]:
        out.setdefault(r[I['iso']], {})[int(r[I['year']])] = {c: (float(r[I[c]]) if r[I[c]] is not None else None) for c in cols}
    return out


def jst_run(dat, L, alt_key='bond_tr', skip_hyper=False, spread=SPREAD):
    """年次（研究者と同じ規則を自作）: Y 年末の EDY（配当利回り − (長期金利 − 10年の物価上昇率)）が拡大窓の中央値より上＝割安、
    eq_tr_Y > bill_rate_Y＝上。Y+1 年: 割安∧上→株 L 倍（借入 bill+1.5%）・割高∧下→長期国債（alt）・ほか株100%。
    skip_hyper: その年の物価上昇率が 50% を超える年（持つ年 Y+1 か信号の年 Y）を外す"""
    ys = sorted(dat)
    edy, up = {}, {}
    for Y in ys:
        d, d10 = dat[Y], dat.get(Y - 10, {})
        if d['eq_dp'] is not None and d['ltrate'] is not None and d['cpi'] and d10.get('cpi'):
            edy[Y] = d['eq_dp'] - (d['ltrate'] / 100 - ((d['cpi'] / d10['cpi']) ** 0.1 - 1))
        if d['eq_tr'] is not None and d['bill_rate'] is not None:
            up[Y] = d['eq_tr'] > d['bill_rate']
    hist = []
    med = {}
    for Y in sorted(edy):
        hist.append(edy[Y])
        if len(hist) >= 10:
            med[Y] = S.median(hist)

    def infl(Y):
        a, b = dat.get(Y, {}).get('cpi'), dat.get(Y - 1, {}).get('cpi')
        return (a / b - 1) if a and b else None
    st, bm, W = {}, {}, {}
    prev = None
    ruined = False
    for Y in ys:
        Y1 = Y + 1
        d1 = dat.get(Y1)
        if ruined or d1 is None or d1['eq_tr'] is None or d1['bill_rate'] is None or Y not in med or Y not in up:
            prev = None
            continue
        if skip_hyper and ((infl(Y1) or 0) > 0.5 or (infl(Y) or 0) > 0.5):
            prev = None
            continue
        cheap, dear = edy[Y] > med[Y], edy[Y] < med[Y]
        w = L if (cheap and up[Y]) else (0.0 if (dear and not up[Y]) else 1.0)
        alt = d1['bill_rate'] if alt_key == 'bill_rate' else d1['bond_tr']
        if w < 1 and alt is None:
            prev = None
            continue
        r = d1['eq_tr']
        if w > 1:
            rp = w * r - (w - 1) * (d1['bill_rate'] + spread)
        elif w == 1:
            rp = r
        else:
            rp = alt
        rp -= 0.0 if prev is None else abs(w - prev) * COST
        if rp <= -0.999:
            rp = -0.999
            ruined = True
        st[Y1], bm[Y1], W[Y1] = rp, r, w
        prev = w
    return st, bm, W


def attack_jst():
    data = load_jst()
    out = {}
    for rn, L, alt in (('J2_agree_2.0', 2.0, 'bond_tr'), ('J1_agree_1.5', 1.5, 'bond_tr'), ('J3_agree_1.5_cash', 1.5, 'bill_rate')):
        for hyper in (False, True):
            per, pool_a, pool_l, pool_c = {}, {}, {}, {}
            pos = reg = 0
            for iso, dat in sorted(data.items()):
                st, bm, W = jst_run(dat, L, alt, skip_hyper=hyper)
                pre = [y for y in sorted(st) if y <= 1974]
                if len(pre) < 24:
                    continue
                s_pre = {y: st[y] for y in pre}
                b_pre = {y: bm[y] for y in pre}
                x = ex(s_pre, b_pre, py=1, L=2)
                # 同じ平均倍率の一定レバレッジ（年次）
                Lb = S.mean(W[y] for y in pre)
                cl = {y: (Lb * bm[y] - (Lb - 1) * (data[iso][y]['bill_rate'] + SPREAD)) if Lb > 1 else bm[y] for y in pre}
                xc = ex(s_pre, cl, py=1, L=2)
                per[iso] = {'years': len(pre), 'vs_mkt': x, 'vs_const': xc, 'avg_w': round(Lb, 3)}
                if iso != 'USA' and len(pre) >= 30:
                    reg += 1
                    pos += x['ex_ann'] > 0 and x['cagr_diff'] > 0
                if iso != 'USA':
                    for y in pre:
                        pool_a.setdefault(y, []).append(st[y] - bm[y])
                        pool_l.setdefault(y, []).append(math.log1p(st[y]) - math.log1p(bm[y]))
                        pool_c.setdefault(y, []).append(st[y] - cl[y])
            def pooled(pl, drop_iso=None):
                p = {y: sum(v) / len(v) for y, v in pl.items() if len(v) >= 5}
                return ex(p, {y: 0.0 for y in p}, py=1, L=2)
            key = f'{rn}{"_skip_hyperinflation" if hyper else ""}'
            out[key] = {'countries_30y': reg, 'positive': pos, 'pooled_arith': pooled(pool_a), 'pooled_logdiff': pooled(pool_l),
                        'pooled_vs_const_lev': pooled(pool_c), 'per_country': per}
            # ドイツを除いた束ね
            if not hyper:
                pa2, pl2 = {}, {}
                for iso, dat in sorted(data.items()):
                    if iso in ('USA', 'DEU'):
                        continue
                    st, bm, W = jst_run(dat, L, alt)
                    for y in st:
                        if y <= 1974:
                            pa2.setdefault(y, []).append(st[y] - bm[y])
                            pl2.setdefault(y, []).append(math.log1p(st[y]) - math.log1p(bm[y]))
                out[key]['pooled_arith_ex_DEU'] = pooled(pa2)
                out[key]['pooled_logdiff_ex_DEU'] = pooled(pl2)
            o = out[key]
            log(f"JST {key:40s} 正 {pos}/{reg} 束ね算術 {o['pooled_arith']['ex_ann']:+.2f}(t{o['pooled_arith']['t']}) 束ね対数差 {o['pooled_logdiff']['ex_ann']:+.2f}(t{o['pooled_logdiff']['t']}) "
                f"対一定 {o['pooled_vs_const_lev']['ex_ann']:+.2f}(t{o['pooled_vs_const_lev']['t']}) "
                + (f"ドイツ除く 算術 {o['pooled_arith_ex_DEU']['ex_ann']:+.2f}(t{o['pooled_arith_ex_DEU']['t']}) 対数差 {o['pooled_logdiff_ex_DEU']['ex_ann']:+.2f}(t{o['pooled_logdiff_ex_DEU']['t']}) " if 'pooled_arith_ex_DEU' in o else '')
                + f"米国 {o['per_country'].get('USA', {}).get('vs_mkt')}")
    return out


# ═════════════════════════ 攻め 8: Shiller 時代（1881〜1926・月中平均の株価） ═════════════════════════
def attack_shiller_era(us):
    sh, gy, sig = us['sh'], us['gy'], us['sig']
    P, D = sh['P'], sh['D']
    ks = sorted(k for k in P if k in D and k <= 192606)
    tr = {}
    for a, k in zip(ks, ks[1:]):
        if ymadd(a, 1) == k:
            tr[k] = (P[k] + D[k] / 12) / P[a] - 1
    cash = {k: gy['Rfree'][k] for k in tr if k in gy['Rfree']}
    bond = {k: sh['BOND'][ymadd(k, -1)] - 1 for k in tr if ymadd(k, -1) in sh['BOND']}
    up = sma_up(tr_index(tr), 10)
    out = {}
    for lag in (1, 2):
        # lag=2: 月中平均の株価の重なり（t 月のリターンの前半は t−1 月の後半）を避けるため信号を1か月さらに遅らせる
        wf = agree_state(sig['ECY'], up, THR_ECY, 2.0, lag)
        keys = valid_months(wf, [tr, cash, bond], 187102, 192606)
        g, n, W = simulate(keys, wf, tr, bond, cash)
        # 月中平均の株価のリターンの自己相関（1次）
        rr = [tr[k] for k in keys]
        ac1 = M.corr(rr[1:], rr[:-1])
        out[f'E3c_lag{lag}'] = {'net': ex(n, tr), 'alpha': alpha(n, tr, cash), 'ac1_stock_returns': round(ac1, 3)}
        log(f"Shiller 時代 E3c 信号の遅れ{lag}: {out[f'E3c_lag{lag}']}")
    # 比較のため French 時代の月末データでも遅れ2を出す
    return out


# ═════════════════════════ 攻め 9: 倍率を振る（保有期間の超過は倍率に比例するか） ═════════════════════════
def attack_leverage_sweep(us, tr):
    mkt, rf, bond, ecy = us['mkt'], us['rf'], us['bond'], us['sig']['ECY']
    out = {}
    for L in (1.0, 1.25, 1.5, 1.75, 2.0, 2.5, 3.0):
        wf = agree_state(ecy, tr['up10'], THR_ECY, L)
        keys = valid_months(wf, [mkt, rf, bond], 192607, END)
        g, n, W = simulate(keys, wf, mkt, bond, rf)
        Lh = S.mean(v for k, v in W.items() if k >= HOLD_START)
        out[str(L)] = {'hold_gross': ex(g, mkt, a=HOLD_START), 'full_gross_t': ex(g, mkt)['t'],
                       'hold_net_vs_const': ex(n, const_lev(mkt, rf, Lh, keys), a=HOLD_START),
                       'sharpe_hold': round(sharpe(n, rf, a=HOLD_START), 3), 'maxdd_full': maxdd(n)}
        log(f"倍率 {L}: 保有 {out[str(L)]['hold_gross']['ex_ann']:+.2f}(t{out[str(L)]['hold_gross']['t']}) 全期間t {out[str(L)]['full_gross_t']} "
            f"対一定 {out[str(L)]['hold_net_vs_const']['ex_ann']:+.2f}(t{out[str(L)]['hold_net_vs_const']['t']}) シャープ保有 {out[str(L)]['sharpe_hold']} 最大下落 {out[str(L)]['maxdd_full']}")
    return out


# ═════════════════════════ 攻め 10: 「一定の借入だけ」で線を越えるか・リスク調整した格付け ═════════════════════════
def attack_const_lev_grade(series, us, countries):
    """相手を『同じ平均倍率で常に借りて持つ市場』に替えて同じ線（C1〜C8）を当てる。
    ついでに、その一定レバレッジ自体が市場相手の線（C1・C2・C4・C6・C7）をどこまで越えるかも出す"""
    mkt, rf = us['mkt'], us['rf']
    out = {}
    for nm_, (g, n, W, keys, wf, p) in series.items():
        Lf = S.mean(W.values())
        cl = const_lev(mkt, rf, Lf, keys)
        cl_net = cl  # 一定倍率は毎月わずかに戻すだけ（費用はほぼ0とみなす＝戦略に有利ではなく不利側）
        o = {'L_const': round(Lf, 3)}
        # (1) 一定レバレッジそのものの市場相手の成績
        o['const_itself_vs_mkt'] = {'full': ex(cl, mkt), 'train': ex(cl, mkt, z=TRAIN_END), 'hold': ex(cl, mkt, a=HOLD_START), 'roll20': roll20(cl, mkt),
                                    'sharpe_train_hold': [round(sharpe(cl, rf, z=TRAIN_END), 3), round(sharpe(cl, rf, a=HOLD_START), 3)]}
        # (2) 戦略を一定レバレッジ相手に
        full, train, hold = ex(g, cl), ex(g, cl, z=TRAIN_END), ex(g, cl, a=HOLD_START)
        costh = ex(n, cl_net, a=HOLD_START)
        rl = roll20(n, cl_net)
        c5 = None
        if countries and nm_ in countries:
            cc = countries[nm_]['counts']
            c5 = cc['pos_vs_const_lev'] / cc['regions'] >= 2 / 3 if cc['regions'] else None
        sh_tr = [sharpe(n, rf, z=TRAIN_END), sharpe(cl, rf, z=TRAIN_END)]
        sh_ho = [sharpe(n, rf, a=HOLD_START), sharpe(cl, rf, a=HOLD_START)]
        gr, cr = grade_mine(full, train, hold, costh, rl, c5, sh_tr, sh_ho)
        o['strategy_vs_const'] = {'full': full, 'train': train, 'hold': hold, 'net_hold': costh, 'roll20_net': rl, 'grade_vs_const': gr, 'criteria': cr}
        out[nm_] = o
        ci = o['const_itself_vs_mkt']
        log(f"一定{Lf:.2f}倍そのもの vs 市場: 全期間 {ci['full']['ex_ann']:+.2f}(t{ci['full']['t']}) 訓練t {ci['train']['t']} 保有 {ci['hold']['ex_ann']:+.2f} 20年 {ci['roll20']['win_rate']} ／ "
            f"{nm_} vs 一定: 全期間 {full['ex_ann']:+.2f}(t{full['t']}) 訓練 {train['ex_ann']:+.2f}(t{train['t']}) 保有 {hold['ex_ann']:+.2f}(t{hold['t']}) 20年 {rl['win_rate']} → 格付け {gr} {cr}")
    return out


# ═════════════════════════ 本体 ═════════════════════════
CLAIMED = {
    'E3c_ECY_agree_lev2': {'grade': 'A', 'full_ex': 5.47, 'full_t': 3.87, 'train_ex': 5.74, 'train_t': 3.53, 'hold_ex': 4.38, 'hold_t': 1.61,
                           'hold_cagr_diff': 4.25, 'net_cost_hold_ex': 4.15, 'roll20_win': 0.988},
    'E3d_ECY_agree_tsmom': {'grade': 'A', 'full_ex': 3.49, 'full_t': 3.64, 'train_ex': 3.91, 'train_t': 3.49, 'hold_ex': 1.79, 'hold_t': 1.09,
                            'hold_cagr_diff': 1.59, 'net_cost_hold_ex': 1.69, 'roll20_win': 1.0},
    'F1_SP_ETF2x_agree': {'grade': 'B', 'full_ex': 5.56, 'full_t': 3.9, 'train_ex': 5.92, 'train_t': 3.6, 'hold_ex': 4.13, 'hold_t': 1.52,
                          'hold_cagr_diff': 3.92, 'net_cost_hold_ex': 3.93, 'roll20_win': 0.988},
    'E3b_ECY_agree_lev_cash': {'grade': 'B', 'full_ex': 2.96, 'full_t': 2.78, 'train_ex': 3.29, 'train_t': 2.67, 'hold_ex': 1.63, 'hold_t': 0.82,
                               'hold_cagr_diff': 1.71, 'net_cost_hold_ex': 1.47, 'roll20_win': 0.912},
}


def git_sha():
    try:
        return subprocess.check_output(['git', 'rev-parse', '--short', 'HEAD'], cwd=M.BASE).decode().strip()
    except Exception:  # noqa
        return None


def close(a, b, tol):
    return a is not None and b is not None and abs(a - b) <= tol


def main():
    t0 = datetime.datetime.now()
    us = build_us()
    sig = us['sig']
    ev = [v for k, v in sig['ECY'].items() if k <= TRAIN_END]
    cv = [v for k, v in sig['CAPE'].items() if k <= TRAIN_END]
    ov = [k for k in sig['ECY'] if k in sig['ECY_shiller_col']]
    sanity = {'us_mkt_cagr_full': round(geo([us['mkt'][k] for k in sorted(us['mkt'])]) * 100, 2),
              'us_mkt_cagr_2007': round(geo([us['mkt'][k] for k in sorted(us['mkt']) if k >= HOLD_START]) * 100, 2),
              'own_ECY_train_median': round(qt(ev, 0.5), 5), 'prereg_ECY_median': THR_ECY, 'own_CAPE_train_median': round(qt(cv, 0.5), 3),
              'own_ECY_vs_Shiller_ECY_column_corr': round(M.corr([sig['ECY'][k] for k in ov], [sig['ECY_shiller_col'][k] for k in ov]), 4),
              'bond_formula': '半年利払いの10年額面債（研究者は年1回利払いの Shiller 式）',
              'E_after_2025_12': '最後の Goyal e12 を持ち越し（研究者は multpl の実質利益を名目へ戻して延長）',
              'P_2023_09': 'Yahoo の月中平均（Shiller の 2023-09 は 9/1 の終値のため使わない）'}
    R, series, tr = us_block(us)
    # 自分の数字での格付け（米国外の再現は下の自作の国別で C5）
    countries = attack_countries()
    beta = attack_beta(R, series, us)
    subp = attack_subperiods(series, us)
    grid = attack_grid(us, tr)
    impl = attack_implementable(series, us, tr)
    fam_f = attack_F(us, tr)
    jst = attack_jst()
    shera = attack_shiller_era(us)
    lsweep = attack_leverage_sweep(us, tr)
    clg = attack_const_lev_grade(series, us, countries)
    # いまの信号（2026-08 月末 → 2026-09 の持ち方）
    last = max(k for k in sig['ECY'] if k <= END)
    now = {'month': last, 'ECY': round(sig['ECY'][last], 4), 'CAPE': round(sig['CAPE'][last], 2), 'above_10m_SMA': tr['up10'].get(last),
           'E3c_weight_next_month': agree_state(sig['ECY'], tr['up10'], THR_ECY, 2.0)(ymadd(last, 1)),
           'last_cheap_month': max(k for k, v in sig['ECY'].items() if v > THR_ECY),
           'months_dear_since_2017': sum(1 for k, v in sig['ECY'].items() if k >= 201701 and v < THR_ECY),
           'months_since_2017': sum(1 for k in sig['ECY'] if k >= 201701)}
    log('いまの信号', now)
    # 格付け（市場相手・自分の数字）
    mine = {}
    for nm_ in ('E3c_ECY_agree_lev2', 'E3d_ECY_agree_tsmom', 'E3b_ECY_agree_lev_cash', 'E2c_agree_lev'):
        e = R[nm_]
        c5 = countries[nm_]['counts']['pos_vs_mkt'] / countries[nm_]['counts']['regions'] >= 2 / 3
        gr, cr = grade_mine(e['full'], e['train'], e['hold'], e['net_hold'], e['roll20_net'], c5,
                            e['sharpe_net']['train'], e['sharpe_net']['hold'])
        mine[nm_] = {'grade': gr, 'criteria': cr}
    for nm_ in ('F1_SP_ETF2x_agree', 'F2_SP_ETF15x_agree'):
        e = fam_f[nm_]
        gr, cr = grade_mine(e['full'], e['train'], e['hold'], e['net_hold_two_leg'], e['roll20_net'], None,
                            e['sharpe_net']['train'], e['sharpe_net']['hold'])
        mine[nm_] = {'grade': gr, 'criteria': cr}
    log('自分の数字での形式上の格付け（市場相手）', {k: v['grade'] for k, v in mine.items()})
    # 多重検定
    from math import erfc, sqrt

    def z_for(p):
        lo, hi = 0.0, 10.0
        for _ in range(100):
            mid = (lo + hi) / 2
            if erfc(mid / sqrt(2)) > p:
                lo = mid
            else:
                hi = mid
        return round(lo, 2)
    mt = {'n_graded_by_researcher': 47, 'n_evaluated_including_reports': 73,
          'bonferroni_t_needed_two_sided_5pct': {'47': z_for(0.05 / 47), '73': z_for(0.05 / 73)},
          'E3c_full_t': R['E3c_ECY_agree_lev2']['full']['t'], 'E3c_hold_p_two_sided': round(erfc(abs(R['E3c_ECY_agree_lev2']['hold']['t']) / sqrt(2)), 4),
          'E3c_hold_p_bonferroni_73': min(1.0, round(erfc(abs(R['E3c_ECY_agree_lev2']['hold']['t']) / sqrt(2)) * 73, 3)),
          'leverage_makes_C7': 'この型では全期間 t が倍率の単調増加（倍率1→3 で t 1.11→4.30）。一定1.22倍の借入だけで全期間 t3.43（C1・C2・C4・C6・C7 を満たす）。C7 は借入の型には効かない',
          'timeline': '事前登録1 13:16 → 第1族の結果 13:29（E2c 全期間 t2.99＝C7 に 0.01 届かず B）→ 同じ時刻に事前登録2（倍率2の E3c・12か月モメンタムの E3d）→ 13:52 E3c・E3d が A',
          'posthoc_rule': 'out/mw_prereg.json honesty_rules[1]: 保有期間の結果を見て規則を変えたら『事後』と明記し判定には使わない。第2・3族は第1族の保有期間の結果を見た後の登録（研究者も『探索』と明記）'}
    verdicts = build_verdicts(R, beta, subp, grid, impl, fam_f, countries, jst, shera, lsweep, clg, mine, now)
    res = {'angle': 'valdca', 'role': '反証の検証（adversarial verifier）', 'generated': datetime.date.today().isoformat(),
           'repo_head_at_run': git_sha(), 'target_files': ['night/mw_valdca.py', 'out/mw_valdca_prereg*.json', 'out/mw_valdca.json'],
           'independence': '取得は mw_common の fetcher だけ。Shiller/Goyal/FRED/World Bank/French 国別/JST の解析・CAPE/ECY・債券の式（半年利払い）・売買の模擬・統計（NW t・CAGR・20年窓・積立・HAC 回帰・シャープ差の検定）・格付けはすべて自作。研究者のコードは import していない',
           'benchmark': 'French Mkt（Mkt-RF+RF・上限なしの時価加重・CRSP 全上場）。国は French の国別の現地通貨の時価加重市場。F1 の実在商品版は SPY',
           'sanity': sanity, 'current_signal': now, 'reproduction_us': R, 'formal_grade_own_numbers_vs_market': mine,
           'attack_beta_alpha_constlev': beta, 'attack_subperiods': subp, 'attack_param_grid': grid, 'attack_implementable': impl,
           'attack_F_etf': fam_f, 'attack_countries_rebuilt': countries, 'attack_jst_rebuilt': jst, 'attack_shiller_era_averaging': shera,
           'attack_leverage_sweep': lsweep, 'attack_grade_vs_const_leverage': clg, 'multiple_testing': mt, 'verdicts': verdicts,
           'regT_schedule_used': {'table': REGT, 'note': 'Regulation T の当初証拠金率の改定表を月単位に丸めて記憶から転記（一次資料で未照合）。最大倍率 = 1/証拠金率。既存の建玉の据え置き（grandfathering）は無視。1934-10 以前は規制なし。結論は上限を掛けても掛けなくても変わらない（全期間 t 3.87 → 3.77）'},
           'hist_margin_rate_used': '1926-07〜1970-11: FRED M13001USM156NNBR（NY コールレート）+1%（下限 RF+1.5%）／1970-12〜2006-12: FRED MPRIME（下限 RF+1.5%）／2007〜: RF+1.5%',
           'log': LOG, 'runtime_sec': round((datetime.datetime.now() - t0).total_seconds(), 1)}
    p = os.path.join(M.BASE, 'out', OUT)
    json.dump(res, open(p, 'w'), ensure_ascii=False, indent=1, default=str)
    log('書いた', p)
    return res


def build_verdicts(R, beta, subp, grid, impl, F, C, jst, shera, lsw, clg, mine, now):
    v = []
    # ── E3c
    e, b, s, i, c = R['E3c_ECY_agree_lev2'], beta['E3c_ECY_agree_lev2'], subp['E3c_ECY_agree_lev2'], impl['E3c_ECY_agree_lev2'], C['E3c_ECY_agree_lev2']
    g2 = grid['lev2.0_bond']
    cl = clg['E3c_ECY_agree_lev2']
    cm = CLAIMED['E3c_ECY_agree_lev2']
    rep = close(e['full']['t'], cm['full_t'], 0.1) and close(e['hold']['ex_ann'], cm['hold_ex'], 0.15) and close(e['net_hold']['ex_ann'], cm['net_cost_hold_ex'], 0.15)
    v.append({
        'name': 'E3c_ECY_agree_lev2', 'claimed_grade': 'A', 'verified_grade': 'B', 'reproduced': bool(rep),
        'verdict': 'downgraded to B',
        'key_numbers': (f"再現: 全期間 {e['full']['ex_ann']:+.2f}%/年 t{e['full']['t']}・訓練 {e['train']['ex_ann']:+.2f} t{e['train']['t']}・保有 {e['hold']['ex_ann']:+.2f} t{e['hold']['t']}・"
                        f"幾何差 {e['hold']['cagr_diff']:+.2f}・費用後保有 {e['net_hold']['ex_ann']:+.2f}・20年窓 {e['roll20_net']['wins']}/{e['roll20_net']['windows']}・"
                        f"シャープ 訓練 {e['sharpe_net']['train']} 保有 {e['sharpe_net']['hold']}（形式上は A を再現）。"
                        f"同じ平均倍率1.22倍の一定レバレッジ相手: 全期間 {cl['strategy_vs_const']['full']['ex_ann']:+.2f} t{cl['strategy_vs_const']['full']['t']}（C7 不合格）・保有 {cl['strategy_vs_const']['hold']['ex_ann']:+.2f} t{cl['strategy_vs_const']['hold']['t']} → {cl['strategy_vs_const']['grade_vs_const']}。"
                        f"保有の α {b['alpha_net_hold']['alpha_ann']:+.2f} t{b['alpha_net_hold']['t_alpha']}・シャープ差 z{b['sharpe_diff_hold']['z']}。"
                        f"2021〜 {s['net_from_2021']['ex_ann']:+.2f} t{s['net_from_2021']['t']}・保有の後半 {s['hold_2nd_2016_11_2026_08']['ex_ann']:+.2f}。"
                        f"国18か国（自作）: 対一定レバレッジで 18/18 正・束ねの保有 +{c['pooled_vs_const_lev']['hold']['ex_ann']} t{c['pooled_vs_const_lev']['hold']['t']}"),
        'issues': [
            f"事後の選択: 第1族の E2c（1.5倍）が全期間 t2.98〜2.99 で C7（3.0）に届かなかった直後に倍率2の E3c を登録した。この型では全期間 t は倍率の単調増加（倍率1.0/1.5/2.0/2.5/3.0 → t {lsw['1.0']['full_gross_t']}/{lsw['1.5']['full_gross_t']}/{lsw['2.0']['full_gross_t']}/{lsw['2.5']['full_gross_t']}/{lsw['3.0']['full_gross_t']}）で、保有期間の超過も倍率に比例（{lsw['1.0']['hold_gross']['ex_ann']:+.2f}→{lsw['3.0']['hold_gross']['ex_ann']:+.2f}）するが、同じ平均倍率の一定レバレッジとの差は保有期間でどの倍率でも t<1（{lsw['2.0']['hold_net_vs_const']['t']}）。全体の事前登録 honesty_rules の『保有期間の結果を見て規則を変えたら事後』に当たる",
            f"C7 は借入だけで満たせる: 一定1.22倍の借入そのものが市場相手に全期間 t{cl['const_itself_vs_mkt']['full']['t']}・訓練 t{cl['const_itself_vs_mkt']['train']['t']}・20年窓 {cl['const_itself_vs_mkt']['roll20']['win_rate']}（C1・C2・C4・C6・C7 を満たし C8 だけ落ちる）。相手を一定レバレッジに替えると全期間 t{cl['strategy_vs_const']['full']['t']} で C7 不合格＝B",
            f"米国の保有期間の勝ちは 2x の86か月（2009-06〜2016・2020-06〜2021-02＝実質金利ゼロで ECY が 3.3〜5.5% に上がった一度の局面）から: 状態別の寄与 2x {s['hold_excess_by_state_pp_per_year']['2.0']['contrib_pp_per_yr']:+.2f}pp/年・国債 {s['hold_excess_by_state_pp_per_year']['0.0']['contrib_pp_per_yr']:+.2f}pp/年（下落回避は保有期間では損）。2008 年も大半（2008-08〜2009-05）は『割安∧線の下』で株100%だった",
            f"小区間: 保有の前半 {s['hold_1st_2007_01_2016_10']['ex_ann']:+.2f}(t{s['hold_1st_2007_01_2016_10']['t']})・後半 {s['hold_2nd_2016_11_2026_08']['ex_ann']:+.2f}(t{s['hold_2nd_2016_11_2026_08']['t']})・2021〜 {s['net_from_2021']['ex_ann']:+.2f}(t{s['net_from_2021']['t']})・2008除く {s['hold_net_drop_2008']['ex_ann']:+.2f}(t{s['hold_net_drop_2008']['t']})。2017 年以降の米国は ECY が線の下の月が {now['months_dear_since_2017']}/{now['months_since_2017']}＝規則は『線を割ったら国債』しかしておらず、その部分は保有期間で負け",
            f"パラメータの崖: ECY の線を訓練の分位 0.30〜0.55 に置くと保有の費用後幾何差は +3.0〜+4.7 だが、0.60 以上（ECY 5.0% 以上）で {g2['lag1_sma10_by_q']['0.6'][0]:+.2f}〜{g2['lag1_sma10_by_q']['0.7'][0]:+.2f}。近い設定108本で保有 t≥1.65 は {g2['all']['hold_t_ge_1.65']}/108",
            f"Shiller 時代（1881〜1926）の +3.42 t2.45 は月中平均の株価によるトレンドの重なり（月次の自己相関 {shera['E3c_lag1']['ac1_stock_returns']}）を含む。信号を1か月遅らせて重なりを外すと {shera['E3c_lag2']['net']['ex_ann']:+.2f} t{shera['E3c_lag2']['net']['t']}",
            f"JST の『+4.92%/年 t2.2』は算術平均で、ドイツ1922年（ハイパーインフレ・2倍が +4131% vs 市場 +2069%）1年だけで束ねに約 +1.6%/年を足している。対数差では {jst['J2_agree_2.0']['pooled_logdiff']['ex_ann']:+.2f} t{jst['J2_agree_2.0']['pooled_logdiff']['t']}、物価上昇率50%超の年を除くと算術 {jst['J2_agree_2.0_skip_hyperinflation']['pooled_arith']['ex_ann']:+.2f} t{jst['J2_agree_2.0_skip_hyperinflation']['pooled_arith']['t']}・対数差 {jst['J2_agree_2.0_skip_hyperinflation']['pooled_logdiff']['ex_ann']:+.2f} t{jst['J2_agree_2.0_skip_hyperinflation']['pooled_logdiff']['t']}（再現は残るが大きさは1/4）。米国自身は JST で {jst['J2_agree_2.0']['per_country']['USA']['vs_mkt']['cagr_diff']:+.2f}%/年（幾何）",
            f"実行面は崩れない（反証にならない）: Reg T の証拠金規制で1945〜1974 の2倍を制限しても全期間 t{i['regT_cap']['full']['t']}、コールレート/プライムの歴史金利と両方でも t{i['regT_and_hist_rate']['full']['t']}・20年窓 {i['regT_and_hist_rate']['roll20_net']['win_rate']}。往復0.30% で保有 {i['two_leg_cost_0.003']['net_hold']['cagr_diff']:+.2f}、日本の課税口座（保有期間の一括）で幾何 {i['tax_japan_hold']['cagr_diff_after_tax_pp']:+.2f}pp/年。国の債券の月中平均による後知恵も効いていない（ずらしても {C['E3c_bond_shifted_no_lookahead']['pooled_vs_mkt']['hold']['ex_ann']:+.2f}）",
            f"B に残す理由（反証まではできない）: 自作で組み直した18か国で、同じ平均倍率の一定レバレッジ相手でも 18/18 正・シャープ 18/18 で市場超え・束ねの保有 +{c['pooled_vs_const_lev']['hold']['ex_ann']} t{c['pooled_vs_const_lev']['hold']['t']}（2017〜2025 は +{c['pooled_subperiods_vs_const']['2017_2025']['ex_ann']} t{c['pooled_subperiods_vs_const']['2017_2025']['t']}）。ただし同じ世界的な低金利の局面で、独立の試行ではない",
            f"いま（{now['month']} 月末）: ECY {now['ECY']*100:.2f}%（線 3.31% の下）・CAPE {now['CAPE']}・10か月線の {'上' if now['above_10m_SMA'] else '下'} → 規則の答えは {now['E3c_weight_next_month']} 倍。最後に割安だったのは {now['last_cheap_month']}。借入2倍は NISA で持てない",
        ]})
    # ── E3d
    e, b, s, c = R['E3d_ECY_agree_tsmom'], beta['E3d_ECY_agree_tsmom'], subp['E3d_ECY_agree_tsmom'], C['E3d_ECY_agree_tsmom']
    cl = clg['E3d_ECY_agree_tsmom']
    g15 = grid['lev1.5_bond']
    ts_by_q = {r['q']: r['hold_net_cagr_diff'] for r in g15['rows'] if r['trend'] == 'tsmom12' and r['lag'] == 1}
    cm = CLAIMED['E3d_ECY_agree_tsmom']
    rep = close(e['full']['t'], cm['full_t'], 0.1) and close(e['hold']['ex_ann'], cm['hold_ex'], 0.15)
    v.append({
        'name': 'E3d_ECY_agree_tsmom', 'claimed_grade': 'A', 'verified_grade': 'B', 'reproduced': bool(rep),
        'verdict': 'downgraded to B',
        'key_numbers': (f"再現: 全期間 {e['full']['ex_ann']:+.2f} t{e['full']['t']}・訓練 {e['train']['ex_ann']:+.2f} t{e['train']['t']}・保有 {e['hold']['ex_ann']:+.2f} t{e['hold']['t']}・幾何差 {e['hold']['cagr_diff']:+.2f}・"
                        f"費用後 {e['net_hold']['ex_ann']:+.2f}・20年窓 {e['roll20_net']['wins']}/{e['roll20_net']['windows']}（形式上は A を再現・国 {c['counts']['pos_vs_mkt']}/18）。"
                        f"保有の α {b['alpha_net_hold']['alpha_ann']:+.2f} t{b['alpha_net_hold']['t_alpha']}・一定1.03倍相手の保有 {cl['strategy_vs_const']['hold']['ex_ann']:+.2f} t{cl['strategy_vs_const']['hold']['t']}・"
                        f"2008除く {s['hold_net_drop_2008']['ex_ann']:+.2f} t{s['hold_net_drop_2008']['t']}・2021〜 {s['net_from_2021']['ex_ann']:+.2f}・国の束ね（対一定）保有 +{c['pooled_vs_const_lev']['hold']['ex_ann']} t{c['pooled_vs_const_lev']['hold']['t']}"),
        'issues': [
            '事後の選択: 第1族の結果（E2c が C7 に 0.01 届かず B）を見た後に、同じ規則のトレンドの定義だけを替えた兄弟4本（E3a〜d）の1本。全体の事前登録の honesty_rules では判定に使わない『事後』（mw_volmanaged の検証と同じ扱い）',
            f"米国の保有期間は弱い: α t{b['alpha_net_hold']['t_alpha']}・シャープ差 z{b['sharpe_diff_hold']['z']}・2008 を除くと {s['hold_net_drop_2008']['ex_ann']:+.2f}(t{s['hold_net_drop_2008']['t']})・保有の後半 {s['hold_2nd_2016_11_2026_08']['ex_ann']:+.2f}・2021〜 {s['net_from_2021']['ex_ann']:+.2f}(t{s['net_from_2021']['t']})。全期間 t の強さはほぼ訓練期間（トレンド追随の文献が見つけた時代）から",
            f"国の再現は符号では 17/18 だが、保有期間を一定レバレッジ相手で束ねると +{c['pooled_vs_const_lev']['hold']['ex_ann']} t{c['pooled_vs_const_lev']['hold']['t']}、2017〜2025 は +{c['pooled_subperiods_vs_const']['2017_2025']['ex_ann']} t{c['pooled_subperiods_vs_const']['2017_2025']['t']}＝保有期間の独立の証拠としては弱い",
            f"パラメータの崖: 12か月モメンタム版で ECY の線を訓練の分位 0.60 以上にすると保有の費用後幾何差 {ts_by_q.get(0.6)}〜{ts_by_q.get(0.7)}（0.30〜0.55 では {min(ts_by_q[q] for q in (0.3,0.35,0.4,0.45,0.5,0.55)):+.2f}〜{max(ts_by_q[q] for q in (0.3,0.35,0.4,0.45,0.5,0.55)):+.2f}）",
            f"C7 は倍率ではなく訓練期間で通る（同じ平均倍率1.03倍の一定レバレッジ相手でも全期間 t{cl['strategy_vs_const']['full']['t']}）。形式上の A はここでは崩れないので、格下げの主な根拠は事後の選択と保有期間の弱さ",
        ]})
    # ── F1
    e = F['F1_SP_ETF2x_agree']
    cm = CLAIMED['F1_SP_ETF2x_agree']
    rep = close(e['full']['t'], cm['full_t'], 0.1) and close(e['hold']['ex_ann'], cm['hold_ex'], 0.15)
    real = e.get('real_products_hold') if isinstance(e.get('real_products_hold'), dict) else {}
    v.append({
        'name': 'F1_SP_ETF2x_agree', 'claimed_grade': 'B', 'verified_grade': 'B', 'reproduced': bool(rep),
        'verdict': 'confirmed',
        'key_numbers': (f"再現: 全期間 {e['full']['ex_ann']:+.2f} t{e['full']['t']}・訓練 {e['train']['ex_ann']:+.2f} t{e['train']['t']}・保有 {e['hold']['ex_ann']:+.2f} t{e['hold']['t']}・"
                        f"費用後（研究者の |Δ倍率|×0.1%）{e['net_hold']['ex_ann']:+.2f}・往復（切替ごとに全部売って買う）幾何 {e['net_hold_two_leg']['cagr_diff']:+.2f}・20年窓 {e['roll20_net']['win_rate']}・"
                        f"シャープ 訓練 {e['sharpe_net']['train']} 保有 {e['sharpe_net']['hold']}。実在の SSO/SPY/IEF だけで組むと保有 SPY 相手 {real.get('trend_French', {}).get('vs_SPY', {}).get('cagr_diff')} pp/年（t{real.get('trend_French', {}).get('vs_SPY', {}).get('t')}）"),
        'issues': [
            f"B の線（C1・C2・C6・C8）は自作の数字でも満たす。SSO の実績とモデルの差は年 {F['SSO_reality'].get('gap_SSO_minus_model_SP_pp') if isinstance(F.get('SSO_reality'), dict) else None}pp で小さい",
            '訓練期間（1927〜2006）の日次2倍の投信は存在しない（1940年投資会社法の借入制限・SSO は 2006-06 設定）＝C1 は仮想の商品。信用取引で Reg T の証拠金規制を掛けた E3c でも訓練 t3.4 なので考え方としては成り立つ',
            f"保有期間の中身は E3c と同じ: 一定倍率相手 {e['vs_const_same_avg_w_hold']['ex_ann']:+.2f}(t{e['vs_const_same_avg_w_hold']['t']})・2008除く {e['hold_net_drop_2008']['ex_ann']:+.2f}(t{e['hold_net_drop_2008']['t']})・前半 {e['hold_1st']['ex_ann']:+.2f}／後半 {e['hold_2nd']['ex_ann']:+.2f}・2021〜 {e['from_2021']['ex_ann']:+.2f}(t{e['from_2021']['t']})",
            f"実在の商品で S&P のトレンドを使うと SPY 相手 {real.get('trend_SPY', {}).get('vs_SPY', {}).get('cagr_diff')}pp/年（t{real.get('trend_SPY', {}).get('vs_SPY', {}).get('t')}）・2021〜 {real.get('trend_SPY', {}).get('from_2021', {}).get('cagr_diff')}",
            '第3族は第2族の結果を見た後の登録（事後）。B（有望・弱い）のままが妥当で、上には上げられない',
        ]})
    # ── E3b
    e, b, s, i, c = R['E3b_ECY_agree_lev_cash'], beta['E3b_ECY_agree_lev_cash'], subp['E3b_ECY_agree_lev_cash'], impl['E3b_ECY_agree_lev_cash'], C['E3b_ECY_agree_lev_cash']
    cm = CLAIMED['E3b_ECY_agree_lev_cash']
    rep = close(e['full']['t'], cm['full_t'], 0.1) and close(e['hold']['ex_ann'], cm['hold_ex'], 0.15)
    v.append({
        'name': 'E3b_ECY_agree_lev_cash', 'claimed_grade': 'B', 'verified_grade': 'B', 'reproduced': bool(rep),
        'verdict': 'confirmed',
        'key_numbers': (f"再現: 全期間 {e['full']['ex_ann']:+.2f} t{e['full']['t']}・訓練 {e['train']['ex_ann']:+.2f} t{e['train']['t']}・保有 {e['hold']['ex_ann']:+.2f} t{e['hold']['t']}・幾何差 {e['hold']['cagr_diff']:+.2f}・"
                        f"費用後 {e['net_hold']['ex_ann']:+.2f}・シャープ 訓練 {e['sharpe_net']['train']} 保有 {e['sharpe_net']['hold']}。2008除く {s['hold_net_drop_2008']['ex_ann']:+.2f} t{s['hold_net_drop_2008']['t']}・"
                        f"国18/18・束ね（対一定）保有 +{c['pooled_vs_const_lev']['hold']['ex_ann']} t{c['pooled_vs_const_lev']['hold']['t']}"),
        'issues': [
            f"B の線は壊れない: Reg T＋歴史の金利でも訓練 t{i['regT_and_hist_rate']['train']['t']}・往復0.30% で保有 {i['two_leg_cost_0.003']['net_hold']['cagr_diff']:+.2f}・シャープは訓練・保有とも市場超え",
            f"ただし保有期間の米国の勝ちは 2008 年だけ: 2008 を除くと {s['hold_net_drop_2008']['ex_ann']:+.2f}(t{s['hold_net_drop_2008']['t']})・同じ平均倍率の一定持ち相手 t{b['const_lev_same_avg_w']['vs_const_hold']['t']}・2021〜 {s['net_from_2021']['ex_ann']:+.2f}・α t{b['alpha_net_hold']['t_alpha']}",
            f"パラメータ: 近い設定108本（1.5倍・現金）で保有の費用後幾何差が正は {grid['lev1.5_cash']['all']['hold_net_cagr_pos']}/108、ECY の線を分位0.60以上にすると負",
            '第2族（事後）の1本。B（有望・弱い）以上には上げられない',
        ]})
    return v


if __name__ == '__main__':
    main()
