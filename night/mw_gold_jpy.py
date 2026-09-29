#!/usr/bin/env python3
"""night/mw_gold_jpy.py — 『市場に勝てる歴史検証』角度 gold_jpy（読むだけ・門の判定には不使用）

問い: 円で毎月積み立てる投資家が、資産の10〜20%を円建ての金にし、その割合を『売らずに、入金を足りない側へ回すだけ』
      （台帳の gap）で保つと、20/25/30年の積立の最終資産（実質・円）の下の裾は S&P500（円建て）100% より良くなるか。
      中央ではどれだけ負けるか。投資家の NASDAQ-100/SMH（75:25）に対しても。

事前登録: out/mw_gold_jpy_prereg.json（規則・線は測る前に固定。ここで動かさない）
出力    : out/mw_gold_jpy.json

約束: 月次リターンは小数。円建てはドル建ての総リターン×為替の変化。総リターンどうしで比べる。
      欠測は0と読まない（欠けた月・年を含む窓は捨てる・絶対のルール7）。
      入金は月初（前月末の値段）→ その月のリターン。金の信号は月末 t の値だけで作り t+1 月に使う。

使い方: python3 night/mw_gold_jpy.py --check   （データの有無だけ・戦略と市場を比べた数字は出さない）
        python3 night/mw_gold_jpy.py           （全部測って out/mw_gold_jpy.json）
"""
import csv, io, json, math, os, subprocess, sys, statistics as S

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import mw_common as M  # noqa: E402

PREREG = 'mw_gold_jpy_prereg.json'
PREREG2 = 'mw_gold_jpy_prereg2.json'   # 探索2（L: G4 を長い代理で・M: 金と S&P500 の相対の勢い）
PREREG3 = 'mw_gold_jpy_prereg3.json'   # 探索3（D: 金を弾薬に・弱気相場で株へ）
OUT = 'mw_gold_jpy.json'
END_M = 202608                 # French の終わり
FEE = 0.0044                   # 金の器の信託報酬（主）
FEE_LO = 0.0020                # 変種
BUY = 0.001                    # 金を買うたびの売買の幅
SELL = 0.001                   # 金・株を売るたび（G3・毎月リバランス）
TAX = 0.20315
COST_UNIT = 0.001              # 世界の線の費用後: 片道100%あたり 0.10%
LOG = []

LBMA_URL = 'https://prices.lbma.org.uk/json/gold_pm.json'
FRED = 'https://fred.stlouisfed.org/graph/fredgraph.csv?id={}'
ESTAT_CPI = 'https://www.e-stat.go.jp/stat-search/file-download?statInfId=000032103842&fileKind=1'
WB_CMO = 'https://thedocs.worldbank.org/en/doc/5d903e848db1d1b83e0ec8f744e55570-0350012021/related/CMO-Historical-Data-Monthly.xlsx'
JST_URL = 'https://www.macrohistory.net/app/download/9834512569/JSTdatasetR6.xlsx?t=1763503850'
BOJ = 'https://www.stat-search.boj.or.jp/api/v1/getDataCode?format=json&lang=en&db=FM02&code={}&startDate={}&endDate=202612'
JST_ALL = 'AUS BEL CAN CHE DEU DNK ESP FIN FRA GBR IRL ITA JPN NLD NOR PRT SWE USA'.split()
C16 = 'AUS BEL CHE DEU DNK ESP FIN FRA GBR ITA JPN NLD NOR PRT SWE USA'.split()
WB_API = 'https://api.worldbank.org/v2/country/all/indicator/{}?format=json&per_page=20000&date={}'


def log(*a):
    s = ' '.join(str(x) for x in a)
    LOG.append(s)
    print(s, flush=True)


def ym_add(k, n):
    y, m = divmod(k, 100)
    t = y * 12 + (m - 1) + n
    return (t // 12) * 100 + t % 12 + 1


def mdiff(a, b):
    return (b // 100 * 12 + b % 100) - (a // 100 * 12 + a % 100)


# ───────────────────────── データ ─────────────────────────
def lbma_monthend():
    """LBMA 金 PM（ドル）の各月の最後の値"""
    d = json.loads(M.get(LBMA_URL, 'lbma_gold_pm.json', max_age_days=3650))
    me = {}
    for x in sorted(d, key=lambda x: x['d']):
        v = x['v'][0]
        if v is None or v <= 0:
            continue
        me[int(x['d'][:4]) * 100 + int(x['d'][5:7])] = float(v)
    return me


def lbma_monthavg():
    d = json.loads(M.get(LBMA_URL, 'lbma_gold_pm.json', max_age_days=3650))
    a = {}
    for x in d:
        v = x['v'][0]
        if v:
            a.setdefault(int(x['d'][:4]) * 100 + int(x['d'][5:7]), []).append(float(v))
    return {k: sum(v) / len(v) for k, v in a.items()}


def fred_monthend(sid):
    b = M.get(FRED.format(sid), f'fred_{sid}.csv', max_age_days=3650).decode()
    out = {}
    for ln in b.strip().splitlines()[1:]:
        dd, v = ln.split(',')[:2]
        if v.strip() in ('', '.'):
            continue
        out[int(dd[:4]) * 100 + int(dd[5:7])] = float(v)   # 日付順＝最後に書いた値が月末
    return out


def jp_cpi():
    t = M.get(ESTAT_CPI, 'jh_estat', max_age_days=3650).decode('cp932', 'replace')
    rows = list(csv.reader(io.StringIO(t)))
    assert rows[0][1] == '総合', rows[0][:3]
    return {int(r[0]): float(r[1]) for r in rows if r and len(r[0]) == 6 and r[0].isdigit() and r[1].strip()}


def wb_gold():
    import openpyxl
    b = M.get(WB_CMO, 'wb_cmo_monthly.xlsx', max_age_days=3650)
    wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
    rows = list(wb['Monthly Prices'].iter_rows(values_only=True))
    h = rows[4]
    gi = [j for j, c in enumerate(h) if c and str(c).strip() == 'Gold'][0]
    out = {}
    for r in rows[6:]:
        if r[0] and isinstance(r[gi], (int, float)):
            out[int(r[0][:4]) * 100 + int(r[0][5:])] = float(r[gi])
    return out


def boj_monthly(code, start):
    j = json.loads(M.get(BOJ.format(code, start), f'jh_boj_{code}.json', max_age_days=3650))
    r = j['RESULTSET'][0]
    return {int(d): float(v) for d, v in zip(r['VALUES']['SURVEY_DATES'], r['VALUES']['VALUES']) if v is not None}


def jpy_rf():
    col = boj_monthly('STRACLCOON', 196001)
    unc = boj_monthly('STRACLUCON', 198501)
    out = {k: v for k, v in col.items() if k < 198507}
    out.update({k: v for k, v in unc.items() if k >= 198507})
    return {k: v / 100 / 12 for k, v in out.items()}


def jst():
    import openpyxl
    b = M.get(JST_URL, 'jst_R6.xlsx', max_age_days=3650)
    wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
    rows = list(wb['Sheet1'].iter_rows(values_only=True))
    h = rows[0]
    cols = ['eq_tr', 'xrusd', 'cpi']
    I = {c: h.index(c) for c in cols + ['year', 'iso']}
    D = {}
    for r in rows[1:]:
        D.setdefault(r[I['iso']], {})[int(r[I['year']])] = {c: (float(r[I[c]]) if r[I[c]] is not None else None) for c in cols}
    return D


def price_to_ret(p):
    ks = sorted(p)
    return {k: p[k] / p[q] - 1 for q, k in zip(ks, ks[1:]) if mdiff(q, k) == 1}


def to_jpy(r, fx):
    """ドル建ての月次リターン → 円建て（その月と前月の月末の為替がそろう月だけ）"""
    out = {}
    for k, v in r.items():
        q = ym_add(k, -1)
        if k in fx and q in fx:
            out[k] = (1 + v) * fx[k] / fx[q] - 1
    return out


def load_all():
    D = {}
    D['gold_me'] = {k: v for k, v in lbma_monthend().items() if k <= END_M}
    D['fx'] = {k: v for k, v in fred_monthend('DEXJPUS').items() if k <= END_M}
    D['cpi'] = jp_cpi()
    ff = M.ff_factors()
    D['sp_usd'] = {k: v for k, v in ff['mkt'].items() if k <= END_M}
    D['rf_usd'] = ff['rf']
    w = M.jkp_mkt('world', 'vw')
    D['world_usd'] = {k: v + ff['rf'][k] for k, v in w.items() if k in ff['rf']}
    q, s = M.yahoo('QQQ'), M.yahoo('SMH')
    D['mix1_usd'] = {k: 0.75 * q[k] + 0.25 * s[k] for k in set(q) & set(s) if k <= END_M}
    n, f = M.yahoo('^NDX'), M.yahoo('FSELX')
    D['mix2_usd'] = {k: 0.75 * n[k] + 0.25 * f[k] for k in set(n) & set(f) if k <= END_M}
    D['gold_usd'] = price_to_ret(D['gold_me'])
    D['gold_jpy_px'] = {k: D['gold_me'][k] * D['fx'][k] for k in D['gold_me'] if k in D['fx']}
    D['gold_jpy'] = price_to_ret(D['gold_jpy_px'])
    D['sp_jpy'] = to_jpy(D['sp_usd'], D['fx'])
    D['world_jpy'] = to_jpy(D['world_usd'], D['fx'])
    D['mix1_jpy'] = to_jpy(D['mix1_usd'], D['fx'])
    D['mix2_jpy'] = to_jpy(D['mix2_usd'], D['fx'])
    D['rf_jpy'] = jpy_rf()
    return D


# ───────────────────────── 金の信号（X 族） ─────────────────────────
def targets_x1(D, w=0.20, minm=36):
    """RG_t = 円建ての金(t) ÷ CPI(t−1)。RG_t < 1971-01 からの累積平均（最低36か月）→ t+1 月の目標 w、そうでなければ 0。
    36か月に満たない間は w。戻り値 {月 k: その月の入金に使う目標}"""
    px, cpi = D['gold_jpy_px'], D['cpi']
    ks = sorted(k for k in px if k >= 197101 and ym_add(k, -1) in cpi)
    out, run, n = {}, 0.0, 0
    for k in ks:
        rg = px[k] / cpi[ym_add(k, -1)]
        run += rg; n += 1
        avg = run / n
        out[ym_add(k, 1)] = w if n < minm else (w if rg < avg else 0.0)
    return out


def targets_x2(D, w=0.20, L=10):
    px = D['gold_jpy_px']
    ks = sorted(k for k in px if k >= 197101)
    out = {}
    for i, k in enumerate(ks):
        if i + 1 < L:
            out[ym_add(k, 1)] = w
            continue
        win = ks[i - L + 1:i + 1]
        if mdiff(win[0], win[-1]) != L - 1:
            continue
        sma = sum(px[j] for j in win) / L
        out[ym_add(k, 1)] = w if px[k] > sma else 0.0
    return out


# ───────────────────────── 積立の模擬 ─────────────────────────
def sim(K, RE, RG, C, i0, n, w=0.0, fee=FEE, mode='gap', tgt=None, tax_annual=False, tax_end_gold=False, per_year=12, check=None,
        deploy=None, buy=BUY, sell=SELL):
    """1つの窓の最終資産（名目）。K=月（または年）の並び、RE/RG=株と金のリターンの並び、C=入金の並び（窓の頭で1に揃えたもの）。
    mode: 'gap'（不足按分・売らない）/ 'annual'（gap＋毎年1月にリバランス）/ 'monthly'（毎回リバランス）"""
    E = G = bE = bG = 0.0
    fper = fee / per_year
    for j in range(i0, i0 + n):
        wt = w if tgt is None else tgt[j]
        c = C[j]
        if deploy is not None and deploy[j] and G > 0:      # 探索3: 予備を全部売って株へ（NISA・税なし）
            E += G * (1 - sell); bE += G * (1 - sell)
            G = bG = 0.0
        if mode == 'annual' and j > i0 and (per_year == 1 or K[j] % 100 == 1) and E + G > 0:
            T0 = E + G
            gs = wt * T0
            if G > gs + 1e-12:
                y = G - gs
                gain = y * (1 - bG / G) if G > 0 else 0.0
                tax = TAX * max(0.0, gain) if tax_annual else 0.0
                bG *= (1 - y / G); G -= y
                net = y * (1 - sell) - tax
                E += net; bE += net
            elif G < gs - 1e-12 and E > 0:
                y = min(E, gs - G)
                gain = y * (1 - bE / E)
                tax = TAX * max(0.0, gain) if tax_annual else 0.0
                bE *= (1 - y / E); E -= y
                net = y * (1 - sell) - tax
                G += net * (1 - buy); bG += net
        if mode == 'monthly':
            T = E + G + c
            gs = wt * T
            if gs >= G:                      # 金を買い足す（入金と株の売りで。株の売りに費用なし）
                x = gs - G
                G += x * (1 - buy); bG += x
                E = T - gs
            else:                            # 金を売って株へ
                y = G - gs
                G = gs
                E = T - gs - y * sell
        else:
            T = E + G + c
            need = wt * T - G
            x = min(max(need, 0.0), c)
            if check is not None:
                check.append(0.0 <= x <= c + 1e-12)
            G += x * (1 - buy); bG += x
            E += c - x; bE += c - x
        E *= 1 + RE[j]
        G *= (1 + RG[j]) * (1 - fper)
    W = E + G
    if tax_end_gold:
        W -= TAX * max(0.0, G - bG)
    return W, (G / W if W > 0 else None)


def aligned(K_all, *series):
    """全ての系列がそろう月だけ（連続した区間の塊ごとに分けて返すのではなく、並びとして返す。連続性は窓ごとに確かめる）"""
    return [k for k in K_all if all(k in s for s in series)]


def windows(K, n, step_ok=None, per_year=12):
    """連続した n 期の窓の頭の番号"""
    out = []
    for i in range(0, len(K) - n + 1):
        a, z = K[i], K[i + n - 1]
        span = mdiff(a, z) if per_year == 12 else z - a
        if span == n - 1:
            out.append(i)
    return out


def q(v, p):
    s = sorted(v)
    if not s:
        return None
    x = p * (len(s) - 1)
    lo = int(math.floor(x)); hi = min(lo + 1, len(s) - 1)
    return s[lo] + (s[hi] - s[lo]) * (x - lo)


def summarize(rows):
    """rows = [(起点, 終点, 戦略の倍率, 相手の倍率)]"""
    if len(rows) < 3:
        return {'n': len(rows)}
    ms = [r[2] for r in rows]; mb = [r[3] for r in rows]
    rr = [r[2] / r[3] for r in rows]
    P = (0, 0.1, 0.25, 0.5, 0.75, 0.9, 1)
    b10 = q(mb, 0.1)
    bad = [r[2] / r[3] for r in rows if r[3] <= b10]
    wi_s = min(rows, key=lambda r: r[2]); wi_b = min(rows, key=lambda r: r[3])
    wr = min(rows, key=lambda r: r[2] / r[3]); br = max(rows, key=lambda r: r[2] / r[3])
    out = {'n': len(rows), 'from': rows[0][0], 'to': rows[-1][0],
           'strat_q': {str(p): round(q(ms, p), 3) for p in P}, 'bench_q': {str(p): round(q(mb, p), 3) for p in P},
           'tail_ratio_q10': round(q(ms, 0.1) / q(mb, 0.1), 4), 'tail_ratio_q25': round(q(ms, 0.25) / q(mb, 0.25), 4),
           'median_ratio_q50': round(q(ms, 0.5) / q(mb, 0.5), 4), 'worst_ratio_min': round(min(ms) / min(mb), 4),
           'paired_q': {str(p): round(q(rr, p), 4) for p in P}, 'paired_median': round(q(rr, 0.5), 4),
           'paired_win_rate': round(sum(1 for x in rr if x > 1) / len(rr), 3),
           'paired_worst': [wr[0], round(wr[2] / wr[3], 4)], 'paired_best': [br[0], round(br[2] / br[3], 4)],
           'strat_worst_window': [wi_s[0], round(wi_s[2], 3)], 'bench_worst_window': [wi_b[0], round(wi_b[3], 3), 'strat', round(wi_b[2], 3)],
           'bad10_paired_median': round(q(bad, 0.5), 4), 'bad10_n': len(bad)}
    out['pass'] = bool(out['tail_ratio_q10'] > 1 and out['paired_median'] >= 0.97)
    return out


def decades(rows):
    out = {}
    for d0 in (1960, 1970, 1980, 1990, 2000):
        sub = [r for r in rows if d0 * 100 <= r[0] < (d0 + 10) * 100 or d0 <= r[0] < d0 + 10]
        if sub:
            out[f'{d0}s'] = {'n': len(sub), 'bench_med': round(q([r[3] for r in sub], 0.5), 3), 'strat_med': round(q([r[2] for r in sub], 0.5), 3),
                             'paired_med': round(q([r[2] / r[3] for r in sub], 0.5), 4), 'paired_min': round(min(r[2] / r[3] for r in sub), 4),
                             'paired_max': round(max(r[2] / r[3] for r in sub), 4)}
    return out


def run_dca(D, bench_key, gold_key='gold_jpy', w=0.1, fee=FEE, mode='gap', tgt_map=None, tax_annual=False, tax_end_gold=False,
            contrib='real', horizons=(20, 25, 30), start_min=None, cpi_key='cpi', detail=False, check=None,
            deploy_map=None, buy=BUY, sell=SELL):
    """月次の積立の全ての窓 → 角度の判定。倍率は実質（CPI がある場合）"""
    RE_d, RG_d = D[bench_key], D[gold_key]
    cpi = D.get(cpi_key) if cpi_key else None
    K_all = sorted(set(RE_d) & set(RG_d))
    if start_min:
        K_all = [k for k in K_all if k >= start_min]
    if cpi is not None:                      # 実質の倍率に要る月だけ（CPI が無い系列は名目同額）
        K_all = [k for k in K_all if k in cpi and ym_add(k, -1) in cpi]
    if tgt_map is not None:
        K_all = [k for k in K_all if k in tgt_map]
    K = K_all
    RE = [RE_d[k] for k in K]; RG = [RG_d[k] for k in K]
    tgt = [tgt_map[k] for k in K] if tgt_map is not None else None
    dep = [bool(deploy_map.get(k)) for k in K] if deploy_map is not None else None
    res = {}
    for H in horizons:
        n = H * 12
        rows = []
        for i0 in windows(K, n):
            a, z = K[i0], K[i0 + n - 1]
            if cpi is not None:
                base = cpi[ym_add(a, -1)]
                if contrib == 'real':
                    C = {j: cpi[ym_add(K[j], -1)] / base for j in range(i0, i0 + n)}
                else:
                    C = {j: 1.0 for j in range(i0, i0 + n)}
                real_contrib = sum(C[j] / (cpi[ym_add(K[j], -1)] / base) for j in range(i0, i0 + n))
                defl = cpi[z] / base
            else:
                C = {j: 1.0 for j in range(i0, i0 + n)}
                real_contrib, defl = float(n), 1.0
            Cl = [0.0] * len(K)
            for j, v in C.items():
                Cl[j] = v
            ws, _ = sim(K, RE, RG, Cl, i0, n, w=w, fee=fee, mode=mode, tgt=tgt, tax_annual=tax_annual, tax_end_gold=tax_end_gold, check=check,
                        deploy=dep, buy=buy, sell=sell)
            wb, _ = sim(K, RE, RG, Cl, i0, n, w=0.0, fee=fee, mode='gap')
            rows.append((a, z, ws / defl / real_contrib, wb / defl / real_contrib))
        sub = {
            'all': summarize(rows),
            'le1990': summarize([r for r in rows if r[0] <= 199012]),
            'ge1991': summarize([r for r in rows if r[0] >= 199101]),
            'train_start_le1986': summarize([r for r in rows if r[0] <= 198612]),
            'hold_end_ge2007': summarize([r for r in rows if r[1] >= 200701]),
        }
        if rows and rows[0][0] <= 197012:
            sub['start_le1970'] = summarize([r for r in rows if r[0] <= 197012])
        sub['pass'] = sub['all'].get('pass', False)
        sub['robust'] = bool(sub['le1990'].get('pass') and sub['ge1991'].get('pass'))
        sub['decades'] = decades(rows)
        if detail and H == 20:
            sub['by_start_year_jan'] = [[r[0], round(r[3], 3), round(r[2], 3), round(r[2] / r[3], 4), r[1]] for r in rows if r[0] % 100 == 1]
            sub['worst10_bench'] = [[r[0], r[1], round(r[3], 3), round(r[2], 3), round(r[2] / r[3], 4)] for r in sorted(rows, key=lambda r: r[3])[:10]]
            sub['worst10_strat'] = [[r[0], r[1], round(r[2], 3), round(r[3], 3), round(r[2] / r[3], 4)] for r in sorted(rows, key=lambda r: r[2])[:10]]
        res[f'H{H}'] = sub
    return res


# ───────────────────────── 世界の線（C1〜C8）用の系列 ─────────────────────────
def constmix(re, rg, wmap, fee=0.0):
    """毎月リバランスの定率（w は月ごとに変えられる）。戻り値: 粗い系列・費用後・年の片道回転率"""
    ks = sorted(k for k in set(re) & set(rg) if k in wmap)
    gross, net, turn = {}, {}, []
    prev_drift = None
    for k in ks:
        w = wmap[k]
        r = (1 - w) * re[k] + w * rg[k]
        gross[k] = r
        if prev_drift is not None:
            turn.append(abs(w - prev_drift))
        else:
            turn.append(w)
        prev_drift = w * (1 + rg[k]) / (1 + r) if 1 + r > 0 else w
        net[k] = r - w * fee / 12
    tpy = 12 * S.mean(turn) if turn else 0.0
    net = M.apply_cost(net, tpy, COST_UNIT)
    return gross, net, round(tpy, 3)


def annual_rebal_series(re, rg, w, fee=0.0, taxed=False, cost=False, start=None):
    """一人の投資家が start に一括で入れて毎年1月にリバランス（課税口座なら実現益に課税）した時間加重の月次リターン"""
    ks = sorted(k for k in set(re) & set(rg) if start is None or k >= start)
    E, G = 1 - w, w
    bE, bG = E, G
    out = {}
    prev = None
    for k in ks:
        if prev is not None and mdiff(prev, k) != 1:
            break
        V0 = E + G
        if prev is not None and k % 100 == 1:
            gs = w * V0
            if G > gs:
                y = G - gs
                gain = y * (1 - bG / G)
                tax = TAX * max(0.0, gain) if taxed else 0.0
                bG *= (1 - y / G); G -= y
                net = y * (1 - (SELL if cost else 0)) - tax
                E += net; bE += net
            elif G < gs and E > 0:
                y = min(E, gs - G)
                gain = y * (1 - bE / E)
                tax = TAX * max(0.0, gain) if taxed else 0.0
                bE *= (1 - y / E); E -= y
                net = y * (1 - (SELL if cost else 0)) - tax
                G += net * (1 - (BUY if cost else 0)); bG += net
        E *= 1 + re[k]
        G *= (1 + rg[k]) * (1 - fee / 12)
        out[k] = (E + G) / V0 - 1
        prev = k
    return out


def global_eval(s_gross, s_net, b, repl=None, rf=None, timing=False):
    x = {'full': M.excess_stats(s_gross, b), 'train': M.excess_stats(s_gross, b, z=M.TRAIN_END),
         'hold': M.excess_stats(s_gross, b, a=M.HOLD_START), 'recent': M.excess_stats(s_gross, b, a=M.RECENT_START),
         'net_full': M.excess_stats(s_net, b), 'cost_hold': M.excess_stats(s_net, b, a=M.HOLD_START),
         'roll20': M.rolling(s_net, b, 20), 'roll20_gross': M.rolling(s_gross, b, 20), 'dca20': M.dca(s_net, b, 20),
         'maxdd_s': round(M.maxdd(s_net) * 100, 1), 'maxdd_b': round(M.maxdd({k: b[k] for k in s_net if k in b}) * 100, 1), 'repl': repl}
    if timing and rf is not None:
        x['sharpe'] = {'train': (M.sharpe(s_net, rf, z=M.TRAIN_END), M.sharpe(b, rf, a=min(s_net), z=M.TRAIN_END)),
                       'hold': (M.sharpe(s_net, rf, a=M.HOLD_START), M.sharpe(b, rf, a=M.HOLD_START))}
    return x


# ───────────────────────── JST（年次）R1・R2・E2 ─────────────────────────
def gold_year_end(me):
    g = {}
    for y in range(1870, 1968):
        g[y] = 20.67 if y <= 1933 else 35.0
    for y in range(1968, 2026):
        k = y * 100 + 12
        if k in me:
            g[y] = me[k]
    return g


def us_annual(sp_usd):
    out = {}
    for y in range(1927, 2026):
        ks = [y * 100 + m for m in range(1, 13)]
        if all(k in sp_usd for k in ks):
            v = 1.0
            for k in ks:
                v *= 1 + sp_usd[k]
            out[y] = v - 1
    return out


def jst_series(J, iso, gy, usa, bench, y0, y1, fx_jump=10.0):
    """その国の年次の (金, 相手, CPI) を、データのそろう年だけ。bench='us'（S&P500 現地通貨）/'home'（自国株）"""
    d = J.get(iso, {})
    rg, re, cpi, bad = {}, {}, {}, set()
    for y in range(y0, y1 + 1):
        a, b0 = d.get(y, {}), d.get(y - 1, {})
        x1, x0 = a.get('xrusd'), b0.get('xrusd')
        if x1 is None or x0 in (None, 0) or y not in gy or (y - 1) not in gy:
            continue
        fxr = x1 / x0
        if fxr > fx_jump or fxr < 1 / fx_jump:
            bad.add(y)
        rg[y] = gy[y] / gy[y - 1] * fxr - 1
        if bench == 'us':
            if y in usa:
                re[y] = (1 + usa[y]) * fxr - 1
        else:
            if a.get('eq_tr') is not None:
                re[y] = a['eq_tr']
        if b0.get('cpi') is not None:
            cpi[y - 1] = b0['cpi']
        if a.get('cpi') is not None:
            cpi[y] = a['cpi']
    return rg, re, cpi, bad


def run_annual(rg, re, cpi, w, H=20, fee=FEE, bad=frozenset(), mode='gap'):
    ks = sorted(y for y in set(rg) & set(re) if (y - 1) in cpi and y in cpi)
    RE = [re[y] for y in ks]; RG = [rg[y] for y in ks]
    rows, excl = [], 0
    for i0 in windows(ks, H, per_year=1):
        a, z = ks[i0], ks[i0 + H - 1]
        if any(y in bad for y in range(a, z + 1)):
            excl += 1
            continue
        base = cpi[a - 1]
        Cl = [0.0] * len(ks)
        for j in range(i0, i0 + H):
            Cl[j] = cpi[ks[j] - 1] / base
        defl = cpi[z] / base
        ws, _ = sim(ks, RE, RG, Cl, i0, H, w=w, fee=fee, mode=mode, per_year=1)
        wb, _ = sim(ks, RE, RG, Cl, i0, H, w=0.0, fee=fee, mode='gap', per_year=1)
        rows.append((a, z, ws / defl / H, wb / defl / H))
    s = summarize(rows)
    s['excluded_windows_fx_jump'] = excl
    return s, rows


def annual_constmix_excess(rg, re, w, y0, y1, fee=FEE):
    ys = [y for y in range(y0, y1 + 1) if y in rg and y in re]
    ex = [((1 - w) * re[y] + w * rg[y] - w * fee) - re[y] for y in ys]
    return (S.mean(ex) if ex else None), len(ys)


# ───────────────────────── 探索2（out/mw_gold_jpy_prereg2.json） ─────────────────────────
def wb_ind(ind, name, dates):
    j = json.loads(M.get(WB_API.format(ind, dates), name, max_age_days=3650))
    return {(r['countryiso3code'], int(r['date'])): float(r['value']) for r in j[1] if r['value'] is not None}


def world_proxy(J, wmap, y0, y1):
    """年次のドル建ての世界の株。重み＝wmap[(国, y−1)]。株・為替・重みのどれかが無い国はその年だけ外す。USA と JPN が両方そろう年だけ"""
    out, used = {}, {}
    for y in range(y0, y1 + 1):
        num = den = 0.0
        got = []
        for iso in C16:
            d = J.get(iso, {})
            a, b = d.get(y, {}), d.get(y - 1, {})
            eq, x1, x0 = a.get('eq_tr'), a.get('xrusd'), b.get('xrusd')
            wgt = wmap.get((iso, y - 1))
            if eq is None or not x1 or not x0 or wgt is None:
                continue
            num += wgt * ((1 + eq) * x0 / x1 - 1); den += wgt; got.append(iso)
        if den > 0 and 'USA' in got and 'JPN' in got:
            out[y] = num / den; used[y] = len(got)
    return out, used


def annual_from_monthly(r, y0, y1):
    out = {}
    for y in range(y0, y1 + 1):
        ks = [y * 100 + m for m in range(1, 13)]
        if all(k in r for k in ks):
            v = 1.0
            for k in ks:
                v *= 1 + r[k]
            out[y] = v - 1
    return out


def part2(D, tested, J, gy):
    """探索2: L（G4 を長い代理で・角度の判定のみ）と M（相対の勢い・世界の線で格付け）"""
    res = {'prereg2': PREREG2, 'prereg2_commit': git_sha(os.path.join('out', PREREG2))}
    # ── L ──
    L = {}
    fxd = {y: D['fx'][y * 100 + 12] for y in range(1971, 2026) if y * 100 + 12 in D['fx']}
    cpd = {y: D['cpi'][y * 100 + 12] for y in range(1970, 2026) if y * 100 + 12 in D['cpi']}
    wa = annual_from_monthly(D['world_usd'], 1986, 2025)
    re = {y: (1 + wa[y]) * fxd[y] / fxd[y - 1] - 1 for y in wa if y in fxd and y - 1 in fxd}
    rg = {y: gy[y] / gy[y - 1] * fxd[y] / fxd[y - 1] - 1 for y in re if y in gy and y - 1 in gy}
    L['annual_world_cagr_check'] = {'annual': round((math.prod(1 + v for v in wa.values()) ** (1 / len(wa)) - 1) * 100, 2),
                                    'monthly': round(M.cagr({k: v for k, v in D['world_usd'].items() if 198601 <= k <= 202512}) * 100, 2)}
    xj = {y: J['JPN'][y]['xrusd'] for y in range(1970, 2021) if J['JPN'].get(y, {}).get('xrusd')}
    cj = {y: J['JPN'][y]['cpi'] for y in range(1970, 2021) if J['JPN'].get(y, {}).get('cpi')}
    gdp = wb_ind('NY.GDP.MKTP.CD', 'wb_gdp_all.json', '1960:2025')
    cap = wb_ind('CM.MKT.LCAP.CD', 'wb_mktcap_all.json', '1970:2025')
    wg, ng = world_proxy(J, gdp, 1971, 2020)
    wc, nc = world_proxy(J, cap, 1976, 2020)
    ov = sorted(set(wg) & set(wa))
    L['gdp_vs_jkp_corr_1986_2020'] = round(M.corr([wg[y] for y in ov], [wa[y] for y in ov]), 3)
    ov2 = sorted(set(wc) & set(wa))
    L['cap_vs_jkp_corr_1986_2020'] = round(M.corr([wc[y] for y in ov2], [wa[y] for y in ov2]), 3)
    L['countries_used_gdp'] = {str(y): ng[y] for y in sorted(ng)}
    L['countries_used_cap'] = {str(y): nc[y] for y in sorted(nc)}
    series = {
        'G4A': (re, rg, cpd, 'JKP world vw の暦年×DEXJPUS 12月末（1986〜2025）'),
    }
    for nm, wu in (('G4L_gdp', wg), ('G4L_cap', wc)):
        rj = {y: (1 + wu[y]) * xj[y] / xj[y - 1] - 1 for y in wu if y in xj and y - 1 in xj}
        gj = {y: gy[y] / gy[y - 1] * xj[y] / xj[y - 1] - 1 for y in rj if y in gy and y - 1 in gy}
        series[nm] = (rj, gj, cj, {'G4L_gdp': 'GDP 加重の世界（JST 16か国・World Bank GDP）1971〜2020', 'G4L_cap': '時価総額加重の世界（JST 16か国・World Bank 時価総額）1976〜2020'}[nm])
    for nm, (rj, gj, cc, desc) in series.items():
        for w in (0.10, 0.20):
            sm, rows = run_annual(gj, rj, cc, w)
            sub = {'all': sm,
                   'le1990': summarize([r for r in rows if r[0] <= 1990]),
                   'ge1991': summarize([r for r in rows if r[0] >= 1991]),
                   'train_start_le1986': summarize([r for r in rows if r[0] <= 1986]),
                   'hold_end_ge2007': summarize([r for r in rows if r[1] >= 2007])}
            sub['pass'] = sm.get('pass', False)
            sub['robust'] = bool(sub['le1990'].get('pass') and sub['ge1991'].get('pass'))
            sub['by_start_year'] = [[r[0], round(r[3], 3), round(r[2], 3), round(r[2] / r[3], 4), r[1]] for r in rows]
            name = f'{nm}_w{int(w * 100)}'
            L[name] = sub
            e = {'name': name, 'family': 'L_long_world', 'description': f'探索2: 金{int(w * 100)}%・gap（年次の積立）＋{desc} vs 同じ世界100%（円建て）',
                 'graded_global': False, 'exploratory': True, 'angle': {'H20': sub}, 'angle_pass_H20': sub['pass'], 'angle_robust_H20': sub['robust']}
            tested.append(e)
            log(f"{name}: n={sm.get('n')} 裾の比={sm.get('tail_ratio_q10')} 対の比の中央={sm.get('paired_median')} 勝率={sm.get('paired_win_rate')} "
                f"相手の最悪={sm.get('bench_worst_window')} 合格={sub['pass']} 頑丈={sub['robust']} ≤1990 {sub['le1990'].get('pass')} ≥1991 {sub['ge1991'].get('pass')}")
    res['L'] = L
    # ── M ──
    sp, g = D['sp_jpy'], D['gold_jpy']
    ks = sorted(set(sp) & set(g))
    Mres = {}
    sig = {}
    for L_ in (12, 6):
        s_ = {}
        for i in range(L_ - 1, len(ks)):
            win = ks[i - L_ + 1:i + 1]
            if mdiff(win[0], win[-1]) != L_ - 1:
                continue
            rs = math.prod(1 + sp[k] for k in win) - 1
            rgm = math.prod(1 + g[k] for k in win) - 1
            s_[ym_add(ks[i], 1)] = 'g' if rgm > rs else 's'
        sig[L_] = s_
    XMG = {}
    for L_ in (12, 6):
        s_ = sig[L_]
        gross, net = {}, {}
        prev = None
        sw = 0
        for k in ks:
            if k not in s_:
                continue
            h = s_[k]
            gross[k] = g[k] if h == 'g' else sp[k]
            n_ = g[k] - FEE / 12 if h == 'g' else sp[k]
            if prev is not None and h != prev:
                n_ -= COST_UNIT; sw += 1
            net[k] = n_
            prev = h
        x = global_eval(gross, net, sp, repl=None, rf=D['rf_jpy'], timing=True)
        yrs = len(gross) / 12
        x['switches'] = sw; x['switches_per_year'] = round(sw / yrs, 2)
        x['gold_share_of_months'] = round(sum(1 for k in gross if s_[k] == 'g') / len(gross), 3)
        x['gold_share_hold'] = round(sum(1 for k in gross if s_[k] == 'g' and k >= M.HOLD_START) / max(1, sum(1 for k in gross if k >= M.HOLD_START)), 3)
        XMG[f'XM{L_}_sw'] = x
    hx = M.holm({nm: (XMG[nm]['hold'] or {}).get('p') for nm in XMG})
    for nm, x in XMG.items():
        x['holm_p'] = hx.get(nm)
        g_, c_ = M.grade(x['full'], x['train'], x['hold'], x['roll20'], cost_hold=x['cost_hold'], repl=None, family_holm_p=x['holm_p'],
                         sharpe_pair=x['sharpe'], leveraged_or_timing=True)
        x['grade'], x['criteria'] = g_, c_
        L_ = int(nm[2:-3])
        tested.append({'name': nm, 'family': 'M_momentum', 'description': f'探索2: 円建ての S&P500 と金の直近{L_}か月の総リターンを比べ、高いほうへ翌月の持ち分100%',
                       'graded_global': True, 'exploratory': True, 'global': x, 'grade': g_, 'criteria': c_})
        log(f"[格付け・探索2] {nm} {g_} 訓練 {x['train']['ex_ann']}%/年 t={x['train']['t']} 保有 {x['hold']['ex_ann']} t={x['hold']['t']} "
            f"費用後保有 {x['cost_hold']['ex_ann']} CAGR差 {x['cost_hold']['cagr_diff']} 最近 {x['recent']['ex_ann']} 転がる20年 {x['roll20']['win_rate'] if x['roll20'] else None} "
            f"シャープ {x['sharpe']} 切替/年 {x['switches_per_year']} 金の月の割合 {x['gold_share_of_months']} {c_}")
    Mres['switch'] = XMG
    s12 = sig[12]
    t20 = {k: (0.20 if v == 'g' else 0.0) for k, v in s12.items()}
    t100 = {k: (1.0 if v == 'g' else 0.0) for k, v in s12.items()}
    for nm, tm, desc in (('XM12_flow20', t20, '金の12か月が S&P500 を上回った翌月だけ金の目標20%（gap・売らない）'),
                         ('XM12_flow100', t100, '入金の全額を12か月の勝者へ（売らない）')):
        r_ = run_dca(D, 'sp_jpy', tgt_map=tm, detail=True)
        h = r_['H20']
        tested.append({'name': nm, 'family': 'M_momentum', 'description': '探索2: ' + desc, 'graded_global': False, 'exploratory': True,
                       'angle': r_, 'angle_pass_H20': h.get('pass'), 'angle_robust_H20': h.get('robust')})
        a = h['all']
        log(f"{nm}: H20 n={a.get('n')} 裾の比={a.get('tail_ratio_q10')} 対の比の中央={a.get('paired_median')} 勝率={a.get('paired_win_rate')} "
            f"悪い10%での対の比={a.get('bad10_paired_median')} 合格={h.get('pass')} 頑丈={h.get('robust')}")
    res['M'] = Mres
    return res


# ───────────────────────── 探索3（out/mw_gold_jpy_prereg3.json） ─────────────────────────
def bear_maps(D, dd=0.20):
    """円建て S&P500 の指数（1971-01=1）の最高値から dd 以上の下げで発火（まだその下げで発火していなければ）。
    戻り値: deploy{実行する月: True}・bear{金の目標0%の月: True}・episodes。
    読み方（事前登録3の文言どおり）: 発火は月末 t の信号 → t+1 月に実行。0% の期間は t+1 から『新高値の月の翌月』まで（その月を含む）"""
    sp = D['sp_jpy']
    ks = sorted(k for k in sp if k >= 197102)
    I = P = 1.0
    in_ep = False
    deploy, bear, eps = {}, {}, []
    cur = None
    for k in ks:
        I *= 1 + sp[k]
        newhigh = I > P
        if newhigh:
            P = I
        nx = ym_add(k, 1)
        if in_ep:
            bear[nx] = True
            if newhigh:
                in_ep = False
                cur['recovered_newhigh'] = k
                eps.append(cur); cur = None
            continue
        if I <= (1 - dd) * P:
            in_ep = True
            deploy[nx] = True
            bear[nx] = True
            cur = {'trigger_signal': k, 'executed': nx, 'drawdown_at_signal': round(I / P - 1, 3)}
    if cur:
        cur['recovered_newhigh'] = None
        eps.append(cur)
    return deploy, bear, eps


def part3(D, tested):
    res = {'prereg3': PREREG3, 'prereg3_commit': git_sha(os.path.join('out', PREREG3))}
    deploy, bear, eps = bear_maps(D)
    res['episodes'] = eps
    res['bear_share_of_months'] = round(sum(1 for k in D['sp_jpy'] if k >= 197102 and bear.get(k)) / sum(1 for k in D['sp_jpy'] if k >= 197102), 3)
    log('探索3 下げの期間', len(eps), '件・0%の月の割合', res['bear_share_of_months'], [(e['executed'], e['recovered_newhigh']) for e in eps])
    ks = sorted(k for k in D['sp_jpy'] if k >= 197102)
    Dc = dict(D)
    Dc['cash_jpy'] = {k: D['rf_jpy'][k] for k in ks if k in D['rf_jpy']}
    for nm, w, gk, fee, bs, desc in (('D1', 0.10, 'gold_jpy', FEE, (BUY, SELL), '金10%・gap。弱気相場（高値から−20%）で金を全部株へ、新高値まで金0%'),
                                     ('D2', 0.20, 'gold_jpy', FEE, (BUY, SELL), '金20%・同じ規則'),
                                     ('D1c', 0.10, 'cash_jpy', 0.0, (0.0, 0.0), '対照: 予備を円の現金（日銀コール）で10%・同じ規則')):
        tm = {k: (0.0 if bear.get(k) else w) for k in ks}
        r_ = run_dca(Dc, 'sp_jpy', gold_key=gk, fee=fee, tgt_map=tm, deploy_map=deploy, buy=bs[0], sell=bs[1], detail=True)
        h = r_['H20']
        tested.append({'name': nm, 'family': 'D_dry_powder', 'description': '探索3: ' + desc, 'graded_global': False, 'exploratory': True,
                       'angle': r_, 'angle_pass_H20': h.get('pass'), 'angle_robust_H20': h.get('robust')})
        a = h['all']
        log(f"{nm}: H20 n={a.get('n')} 裾の比={a.get('tail_ratio_q10')} 対の比の中央={a.get('paired_median')} 勝率={a.get('paired_win_rate')} "
            f"悪い10%での対の比={a.get('bad10_paired_median')} 合格={h.get('pass')} 頑丈={h.get('robust')} ≤1990 {h['le1990'].get('pass')} {h['le1990'].get('paired_median')} ≥1991 {h['ge1991'].get('pass')} {h['ge1991'].get('paired_median')}")
    sp, g = D['sp_jpy'], D['gold_jpy']
    DG = {}
    for nm, w in (('D1s', 0.10), ('D2s', 0.20)):
        wmap = {k: (0.0 if bear.get(k) else w) for k in ks if k in g}
        gs, ns, tpy = constmix(sp, g, wmap, fee=FEE)
        x = global_eval(gs, ns, sp, repl=None, rf=D['rf_jpy'], timing=True)
        x['turnover_per_year'] = tpy
        DG[nm] = x
    hx = M.holm({nm: (DG[nm]['hold'] or {}).get('p') for nm in DG})
    for nm, x in DG.items():
        x['holm_p'] = hx.get(nm)
        g_, c_ = M.grade(x['full'], x['train'], x['hold'], x['roll20'], cost_hold=x['cost_hold'], repl=None, family_holm_p=x['holm_p'],
                         sharpe_pair=x['sharpe'], leveraged_or_timing=True)
        x['grade'], x['criteria'] = g_, c_
        tested.append({'name': nm, 'family': 'D_dry_powder', 'description': f'探索3: 定率 金{int(nm[1]) * 10}%・弱気相場の期間だけ金0%（毎月リバランス・格付け用）',
                       'graded_global': True, 'exploratory': True, 'global': x, 'grade': g_, 'criteria': c_})
        log(f"[格付け・探索3] {nm} {g_} 訓練 {x['train']['ex_ann']} t={x['train']['t']} 保有 {x['hold']['ex_ann']} t={x['hold']['t']} 費用後保有 {x['cost_hold']['ex_ann']} "
            f"CAGR差 {x['cost_hold']['cagr_diff']} シャープ {x['sharpe']} {c_}")
    return res


def make_summary(out):
    """結果の数字から要約を組む（手で数字を書かない）"""
    T = {t['name']: t for t in out['tested']}

    def a(nm, H='H20', sub='all'):
        return T[nm]['angle'][H][sub]
    g1, g2 = a('G1'), a('G2')
    sm = {
        'answer_angle': {
            'G1_10pct': {'tail_ratio_q10': g1['tail_ratio_q10'], 'paired_median': g1['paired_median'], 'win_rate': g1['paired_win_rate'],
                         'bad10_paired_median': g1['bad10_paired_median'], 'bench_worst': g1['bench_worst_window'],
                         'pass': T['G1']['angle']['H20']['pass'], 'robust': T['G1']['angle']['H20']['robust']},
            'G2_20pct': {'tail_ratio_q10': g2['tail_ratio_q10'], 'paired_median': g2['paired_median'], 'win_rate': g2['paired_win_rate'],
                         'bad10_paired_median': g2['bad10_paired_median'], 'bench_worst': g2['bench_worst_window'],
                         'pass': T['G2']['angle']['H20']['pass'], 'robust': T['G2']['angle']['H20']['robust']},
            'decades_G1': T['G1']['angle']['H20']['decades'],
            'G4_world': {'pass': T['G4']['angle']['H20']['pass'], 'robust': T['G4']['angle']['H20']['robust'],
                         'tail': a('G4')['tail_ratio_q10'], 'median': a('G4')['paired_median'], 'n': a('G4')['n'], 'starts': [a('G4')['from'], a('G4')['to']]},
        },
        'angle_passes': sorted(n for n, t in T.items() if t.get('angle_pass_H20')),
        'angle_robust': sorted(n for n, t in T.items() if t.get('angle_robust_H20')),
        'global_grades': {n: t.get('grade') for n, t in T.items() if t.get('graded_global')},
    }
    L = out.get('part2', {}).get('L', {})
    if L:
        sm['G4_long_proxies'] = {nm: {'pass': L[nm]['pass'], 'robust': L[nm]['robust'], 'tail': L[nm]['all'].get('tail_ratio_q10'),
                                      'median': L[nm]['all'].get('paired_median'), 'le1990_median': L[nm]['le1990'].get('paired_median'), 'n': L[nm]['all'].get('n')}
                                 for nm in L if nm.startswith('G4')}
    R = out['replication']
    sm['replication'] = {'R1_w10': R['R1']['count_w10'], 'R1_w20': R['R1']['count_w20'], 'R2_w10': R['R2']['count_w10'], 'R2_w20': R['R2']['count_w20']}
    return sm


# ───────────────────────── 本体 ─────────────────────────
def check():
    D = load_all()
    for k in ('gold_me', 'fx', 'cpi', 'sp_usd', 'world_usd', 'mix1_usd', 'mix2_usd', 'gold_jpy', 'sp_jpy', 'world_jpy', 'mix1_jpy', 'mix2_jpy', 'rf_jpy'):
        ks = sorted(D[k]); print(k, ks[0], ks[-1], len(ks))
    J = jst()
    for iso in JST_ALL:
        d = J.get(iso, {})
        full = [y for y in range(1971, 2021) if d.get(y, {}).get('xrusd') is not None and d.get(y, {}).get('cpi') is not None]
        eq = [y for y in range(1971, 2021) if d.get(y, {}).get('eq_tr') is not None]
        print(iso, 'fx+cpi 1971-2020', len(full), 'eq_tr', len(eq))


def git_sha(path):
    try:
        return subprocess.check_output(['git', 'log', '-1', '--format=%h', '--', path], cwd=M.BASE).decode().strip()
    except Exception:  # noqa
        return None


def main():
    D = load_all()
    out = {'angle': 'gold_jpy', 'prereg': PREREG, 'prereg_commit': git_sha(os.path.join('out', PREREG)),
           'global_prereg': 'out/mw_prereg.json', 'sanity': {}, 'deviations': [], 'posthoc': []}
    tested = []

    # ── 検算 ──
    san = out['sanity']
    m = D['sp_usd']
    san['french_cagr_1926_2026'] = round(M.cagr(m) * 100, 2)
    san['french_cagr_2007'] = round(M.cagr(M.window(m, M.HOLD_START)) * 100, 2)
    gj = D['gold_jpy_px']
    san['gold_jpy_annual_avg'] = {y: round(S.mean(gj[y * 100 + mm] for mm in range(1, 13))) for y in (1980, 2000, 2012, 2025)}
    wbg, la = wb_gold(), lbma_monthavg()
    dif = sorted(abs(wbg[k] / la[k] - 1) for k in wbg if k in la)
    san['wb_vs_lbma_monthavg'] = {'n': len(dif), 'median_abs_diff_pct': round(q(dif, 0.5) * 100, 4), 'max_abs_diff_pct': round(dif[-1] * 100, 2)}
    # w=0 は相手と同じ・w=1 は金だけの積立と同じ・gap は売らない
    chk = []
    r0 = run_dca(D, 'sp_jpy', w=0.0, horizons=(20,), check=chk)
    san['w0_identity'] = {'paired_min': r0['H20']['all']['paired_q']['0'], 'paired_max': r0['H20']['all']['paired_q']['1']}
    K = sorted(set(D['sp_jpy']) & set(D['gold_jpy']) & {k for k in D['cpi'] if ym_add(k, -1) in D['cpi']})
    i0, n = 0, 240
    RE = [D['sp_jpy'][k] for k in K]; RG = [D['gold_jpy'][k] for k in K]
    base = D['cpi'][ym_add(K[0], -1)]
    Cl = [D['cpi'][ym_add(k, -1)] / base for k in K]
    w1, _ = sim(K, RE, RG, Cl, i0, n, w=1.0)
    direct = 0.0
    for j in range(i0, i0 + n):
        direct = direct + Cl[j] * (1 - BUY)
        direct *= (1 + RG[j]) * (1 - FEE / 12)
    san['w1_equals_gold_only_dca'] = {'sim': round(w1, 6), 'direct': round(direct, 6), 'ok': abs(w1 / direct - 1) < 1e-9}
    chk2 = []
    run_dca(D, 'sp_jpy', w=0.2, horizons=(20,), check=chk2)
    san['gap_never_sells'] = {'checks': len(chk2), 'all_ok': all(chk2)}
    log('検算', json.dumps(san, ensure_ascii=False))

    def add(name, family, desc, res, graded=False, extra=None):
        e = {'name': name, 'family': family, 'description': desc, 'graded_global': graded, 'angle': res}
        if extra:
            e.update(extra)
        h = res.get('H20', {})
        e['angle_pass_H20'] = h.get('pass'); e['angle_robust_H20'] = h.get('robust')
        tested.append(e)
        a = h.get('all', {})
        log(f"{name}: H20 n={a.get('n')} 裾の比(10%点)={a.get('tail_ratio_q10')} 対の比の中央={a.get('paired_median')} "
            f"勝率={a.get('paired_win_rate')} 相手の最悪={a.get('bench_worst_window')} 悪い10%での対の比={a.get('bad10_paired_median')} "
            f"合格={h.get('pass')} 頑丈={h.get('robust')} ≤1990 {h.get('le1990', {}).get('pass')} ≥1991 {h.get('ge1991', {}).get('pass')}")
        return e

    # ── P 主 ──
    P = {}
    P['G1'] = add('G1', 'P_primary', '金10%・gap・相手 S&P500 円建て', run_dca(D, 'sp_jpy', w=0.10, detail=True))
    P['G2'] = add('G2', 'P_primary', '金20%・gap・相手 S&P500 円建て', run_dca(D, 'sp_jpy', w=0.20, detail=True))
    P['G3'] = add('G3', 'P_primary', '金20%・gap＋毎年1月リバランス・課税口座（参考）・相手 S&P500 円建て',
                  run_dca(D, 'sp_jpy', w=0.20, mode='annual', tax_annual=True, detail=True))
    P['G4'] = add('G4', 'P_primary', '金10%・gap＋世界の時価加重90%・相手 世界の時価加重 円建て', run_dca(D, 'world_jpy', w=0.10, detail=True))

    # ── P_mix ──
    add('G1m_M1', 'P_mix', '金10%・gap＋実物の配合 QQQ75/SMH25・相手 配合100%（2000-07〜）', run_dca(D, 'mix1_jpy', w=0.10, detail=True))
    add('G2m_M1', 'P_mix', '金20%・gap＋実物の配合・相手 配合100%', run_dca(D, 'mix1_jpy', w=0.20, detail=True))
    add('G1m_M2', 'P_mix', '金10%・gap＋代理の配合 ^NDX価格75/FSELX25・相手 代理100%（1985-11〜）', run_dca(D, 'mix2_jpy', w=0.10, detail=True))
    add('G2m_M2', 'P_mix', '金20%・gap＋代理の配合・相手 代理100%', run_dca(D, 'mix2_jpy', w=0.20, detail=True))

    # ── V 変種 ──
    for w in (0.05, 0.15, 0.30):
        add(f'V_w{int(w * 100)}', 'V_variants', f'金{int(w * 100)}%・gap・相手 S&P500 円建て', run_dca(D, 'sp_jpy', w=w))
    for nm, w in (('G1', 0.10), ('G2', 0.20)):
        add(f'V_{nm}_fee020', 'V_variants', f'{nm} を信託報酬 0.20%/年で', run_dca(D, 'sp_jpy', w=w, fee=FEE_LO))
        add(f'V_{nm}_nominal', 'V_variants', f'{nm} を名目同額の入金で', run_dca(D, 'sp_jpy', w=w, contrib='nominal'))
        add(f'V_{nm}_taxT1', 'V_variants', f'{nm} で金だけ課税口座（最後に金の含み益へ20.315%）', run_dca(D, 'sp_jpy', w=w, tax_end_gold=True))
        add(f'V_{nm}_usd', 'V_variants', f'{nm} をドルの投資家で（名目同額・相手 S&P500 ドル建て）',
            run_dca(D, 'sp_usd', gold_key='gold_usd', w=w, contrib='nominal', cpi_key=None, start_min=197102))
        add(f'V_{nm}_monthly', 'V_variants', f'金{int(w * 100)}%を毎月リバランス（税なし）', run_dca(D, 'sp_jpy', w=w, mode='monthly'))
    add('V_G3_notax', 'V_variants', 'G3 を税なし（NISA の中の年1回リバランス）', run_dca(D, 'sp_jpy', w=0.20, mode='annual', tax_annual=False))
    add('V_G4_vs_sp', 'V_variants', '金10%＋世界90% を S&P500 円建てと比べる（世界の窓 1985-11〜）', g4_vs_sp(D))

    # ── X 探索 ──
    T1, T2 = targets_x1(D), targets_x2(D)
    X = {}
    X['X1'] = add('X1', 'X_exploratory', '探索: 実質の金価格が1971年からの累積平均より安い月だけ金の目標20%（gap・売らない）',
                  run_dca(D, 'sp_jpy', tgt_map=T1, detail=True))
    X['X2'] = add('X2', 'X_exploratory', '探索: 円建ての金が10か月線より上の月だけ金の目標20%（gap・売らない）',
                  run_dca(D, 'sp_jpy', tgt_map=T2, detail=True))

    # ── R 再現（JST 年次） ──
    J = jst()
    gy = gold_year_end(D['gold_me'])
    usa = us_annual(D['sp_usd'])
    R = {'R1': {}, 'R2': {}}
    for iso in JST_ALL:
        if iso != 'JPN':
            rg, re, cpi, bad = jst_series(J, iso, gy, usa, 'us', 1971, 2020)
            for w in (0.10, 0.20):
                s, _ = run_annual(rg, re, cpi, w)
                R['R1'].setdefault(iso, {})[f'w{int(w * 100)}'] = s
                tested.append({'name': f'R1_{iso}_w{int(w * 100)}', 'family': 'R_replication', 'description': f'{iso}: 金{int(w * 100)}%・gap＋S&P500 現地通貨 vs S&P500 現地通貨（年次 1971〜2020）',
                               'graded_global': False, 'angle_pass_H20': s.get('pass'), 'angle': {'H20': {'all': s}}})
        rg, re, cpi, bad = jst_series(J, iso, gy, usa, 'home', 1971, 2020)
        if len([y for y in re if 1971 <= y <= 2020]) >= 45:
            for w in (0.10, 0.20):
                s, _ = run_annual(rg, re, cpi, w)
                R['R2'].setdefault(iso, {})[f'w{int(w * 100)}'] = s
                tested.append({'name': f'R2_{iso}_w{int(w * 100)}', 'family': 'R_replication', 'description': f'{iso}: 金{int(w * 100)}%・gap＋自国株 vs 自国株（年次 1971〜2020）',
                               'graded_global': False, 'angle_pass_H20': s.get('pass'), 'angle': {'H20': {'all': s}}})
    for fam in ('R1', 'R2'):
        for w in ('w10', 'w20'):
            isos = [i for i in R[fam] if w in R[fam][i]]
            nonus = [i for i in isos if i != 'USA']
            R[fam][f'count_{w}'] = {'n': len(isos), 'pass': sum(1 for i in isos if R[fam][i][w].get('pass')),
                                    'n_nonUS': len(nonus), 'pass_nonUS': sum(1 for i in nonus if R[fam][i][w].get('pass')),
                                    'tail_gt1_nonUS': sum(1 for i in nonus if (R[fam][i][w].get('tail_ratio_q10') or 0) > 1),
                                    'median_ge097_nonUS': sum(1 for i in nonus if (R[fam][i][w].get('paired_median') or 0) >= 0.97)}
            log(fam, w, R[fam][f'count_{w}'])
    out['replication'] = R

    # ── E 参考 ──
    E = {}
    E['E1'] = e1(D)
    E2 = {}
    for iso in JST_ALL:
        rg, re, cpi, bad = jst_series(J, iso, gy, usa, 'home', 1871, 1970)
        if len(re) < 25:
            continue
        for w in (0.10, 0.20):
            s, _ = run_annual(rg, re, cpi, w, bad=bad)
            E2.setdefault(iso, {})[f'w{int(w * 100)}'] = s
            tested.append({'name': f'E2_{iso}_w{int(w * 100)}', 'family': 'E_extensions', 'description': f'{iso}: 金本位の時代 金{int(w * 100)}%・gap＋自国株 vs 自国株（年次 1871〜1970・参考）',
                           'graded_global': False, 'angle_pass_H20': s.get('pass'), 'angle': {'H20': {'all': s}}})
    for w in ('w10', 'w20'):
        isos = [i for i in E2 if w in E2[i] and E2[i][w].get('n', 0) >= 3]
        E2[f'count_{w}'] = {'n': len(isos), 'pass': sum(1 for i in isos if E2[i][w].get('pass'))}
        log('E2', w, E2[f'count_{w}'])
    E['E2'] = E2
    out['extensions'] = E
    for nm, res in E['E1'].items():
        tested.append({'name': f'E1_{nm}', 'family': 'E_extensions', 'description': f'1960年からの延長（月次・名目同額・参考） {nm}',
                       'graded_global': False, 'angle_pass_H20': res.get('H20', {}).get('pass'), 'angle': res})

    # ── 世界の線（C1〜C8） ──
    re_sp, rg = D['sp_jpy'], D['gold_jpy']
    ks_sp = sorted(set(re_sp) & set(rg))
    gl = {}
    for nm, w in (('G1', 0.10), ('G2', 0.20)):
        gs, ns, tpy = constmix(re_sp, rg, {k: w for k in ks_sp}, fee=FEE)
        gl[nm] = (gs, ns, re_sp, tpy)
    gs = annual_rebal_series(re_sp, rg, 0.20, start=ks_sp[0])
    ns = annual_rebal_series(re_sp, rg, 0.20, fee=FEE, taxed=True, cost=True, start=ks_sp[0])
    gl['G3'] = (gs, ns, re_sp, None)
    re_w = D['world_jpy']
    ks_w = sorted(set(re_w) & set(rg))
    gs, ns, tpy = constmix(re_w, rg, {k: 0.10 for k in ks_w}, fee=FEE)
    gl['G4'] = (gs, ns, re_w, tpy)
    # C5: R2 の米国外15か国（定率・年1回リバランスの金＋自国株 vs 自国株の年次の超過の平均 1971〜2020 が正か）
    def repl_for(w):
        reg = pos = 0; det = {}
        for iso in JST_ALL:
            if iso == 'USA':
                continue
            rgA, reA, _, _ = jst_series(J, iso, gy, usa, 'home', 1971, 2020)
            mu, n = annual_constmix_excess(rgA, reA, w, 1971, 2020)
            if mu is None or n < 45:
                continue
            reg += 1; pos += mu > 0; det[iso] = round(mu * 100, 2)
        return {'regions': reg, 'positive': pos, 'detail_ex_ann_pct': det}
    repl = {'G1': repl_for(0.10), 'G2': repl_for(0.20), 'G3': repl_for(0.20), 'G4': repl_for(0.10)}
    GE = {}
    for nm in ('G1', 'G2', 'G3', 'G4'):
        gs, ns, b, tpy = gl[nm]
        GE[nm] = global_eval(gs, ns, b, repl=repl[nm])
        GE[nm]['turnover_per_year'] = tpy
    hp = M.holm({nm: (GE[nm]['hold'] or {}).get('p') for nm in GE})
    for nm in GE:
        x = GE[nm]
        x['holm_p'] = hp.get(nm)
        g, c = M.grade(x['full'], x['train'], x['hold'], x['roll20'], cost_hold=x['cost_hold'], repl=x['repl'], family_holm_p=x['holm_p'])
        x['grade'], x['criteria'] = g, c
        e = P[nm]
        e['graded_global'] = True
        e['global_series'] = {'G1': '金10%/株90% 毎月リバランス（円建て）', 'G2': '金20%/株80% 毎月リバランス（円建て）',
                              'G3': '金20%/株80% 毎年1月リバランス・課税口座・1971-02 に一括（時間加重）', 'G4': '金10%/世界90% 毎月リバランス vs 世界（円建て）'}[nm]
        e['global'] = x; e['grade'] = g; e['criteria'] = c
        log(f"[格付け] {nm} {g} 訓練 {x['train']['ex_ann'] if x['train'] else None}%/年 t={x['train']['t'] if x['train'] else None} "
            f"保有 {x['hold']['ex_ann']}%/年 t={x['hold']['t']} 費用後保有 {x['cost_hold']['ex_ann']} CAGR差 {x['cost_hold']['cagr_diff']} "
            f"転がる20年 {x['roll20']['win_rate'] if x['roll20'] else None} 再現 {x['repl']['positive']}/{x['repl']['regions']} {c}")
    # X（タイミング型）
    XG = {}
    for nm, T in (('X1', T1), ('X2', T2)):
        wmap = {k: T[k] for k in ks_sp if k in T}
        gs, ns, tpy = constmix(re_sp, rg, wmap, fee=FEE)
        XG[nm] = global_eval(gs, ns, re_sp, repl=None, rf=D['rf_jpy'], timing=True)
        XG[nm]['turnover_per_year'] = tpy
        XG[nm]['avg_gold_weight'] = round(S.mean(wmap.values()), 3)
    hx = M.holm({nm: (XG[nm]['hold'] or {}).get('p') for nm in XG})
    for nm in XG:
        x = XG[nm]
        x['holm_p'] = hx.get(nm)
        sp = x['sharpe']
        g, c = M.grade(x['full'], x['train'], x['hold'], x['roll20'], cost_hold=x['cost_hold'], repl=None, family_holm_p=x['holm_p'],
                       sharpe_pair=sp, leveraged_or_timing=True)
        x['grade'], x['criteria'] = g, c
        e = X[nm]
        e['graded_global'] = True; e['exploratory'] = True
        e['global_series'] = '持ち分の切り替え版（信号の翌月は金20%/株80%・それ以外は株100%・毎月リバランス）'
        e['global'] = x; e['grade'] = g; e['criteria'] = c
        log(f"[格付け・探索] {nm} {g} 訓練 {x['train']['ex_ann']} t={x['train']['t']} 保有 {x['hold']['ex_ann']} t={x['hold']['t']} シャープ {sp} {c}")

    # ── 探索2（事前登録2） ──
    if os.path.exists(os.path.join(M.BASE, 'out', PREREG2)):
        out['part2'] = part2(D, tested, J, gy)

    if os.path.exists(os.path.join(M.BASE, 'out', PREREG3)):
        out['part3'] = part3(D, tested)

    out['n_tested'] = len(tested)
    out['n_graded_global'] = sum(1 for t in tested if t.get('graded_global'))
    out['grades'] = {t['name']: t.get('grade') for t in tested if t.get('graded_global')}
    out['tested'] = tested
    out['summary'] = make_summary(out)
    out['deviations'] = [
        '規則・線は結果を見て一つも動かしていない（第1回 6991597・第2回 f134035・第3回 110ce8a の事前登録どおり）',
        '事前登録2 の G4L_gdp の重みは、登録を書いている途中（コミット前）に JST の gdp の単位が国ごとに違う（百万・十億・兆）と分かったので World Bank の GDP（ドル）へ替えた。コミットした登録にはその経緯を書いてある',
        '事前登録3 の「発火から新高値の翌月までは金の目標0%」は文言どおり『新高値が分かった月の翌月（を含む）まで0%』と読んだ（1か月の差）',
        'mw_common.py に不具合は見つからなかった（French の CAGR 10.38%・2007〜 11.13% を再現・excess_stats/rolling/dca をそのまま使用）',
        '世界の線（C1〜C8）の格付けは、積立（gap）に一つの系列が無いので定率の近似（毎月リバランス）で付けた＝事前登録どおり。角度の問いへの答えは積立の模擬のほう'
    ]
    out['posthoc'] = [
        '事後（判定に使わない・記述だけ）: 金10%が守った窓は起点 1989〜1993年（終わり 2008〜2013年）にほぼ限られる。起点 1972〜1988年は毎年負け（−1.5〜−20%）、1994〜2006年も小さく負け（−0.1〜−6%）',
        '事後: gap は売らないので、金が上がった窓では金の比率が目標を大きく超える（相手の最悪の窓 1989-03 起点では最後に31%）。『10%を保つ』の実態は『10%以上』',
        '事後: 金も S&P500 もドル建てなので、中央の負けは円でもドルでも同じ（名目同額 G1: 円 0.947・ドル 0.950）。裾の守りは円のほうが小さい（1.025 vs 1.086）',
        '事後: 金は相手の最悪の窓を変えるが、戦略自身の最悪の窓は別の時代（1975年起点・金が高値の1980年に買った窓）に移る'
    ]
    out['log'] = LOG
    p = M.save(OUT, out)
    log('書いた', p, os.path.getsize(p))


def g4_vs_sp(D):
    """金10%＋世界90%（gap）を S&P500 円建て 100% と比べる。相手の側の株と戦略の側の株が違うので専用の模擬"""
    K = sorted(set(D['world_jpy']) & set(D['sp_jpy']) & set(D['gold_jpy']) & {k for k in D['cpi'] if ym_add(k, -1) in D['cpi']})
    RW = [D['world_jpy'][k] for k in K]; RS = [D['sp_jpy'][k] for k in K]; RG = [D['gold_jpy'][k] for k in K]
    cpi = D['cpi']
    res = {}
    for H in (20, 25, 30):
        n = H * 12
        rows = []
        for i0 in windows(K, n):
            a, z = K[i0], K[i0 + n - 1]
            base = cpi[ym_add(a, -1)]
            Cl = [0.0] * len(K)
            for j in range(i0, i0 + n):
                Cl[j] = cpi[ym_add(K[j], -1)] / base
            ws, _ = sim(K, RW, RG, Cl, i0, n, w=0.10)
            wb, _ = sim(K, RS, RG, Cl, i0, n, w=0.0)
            defl = cpi[z] / base
            rows.append((a, z, ws / defl / n, wb / defl / n))
        sub = {'all': summarize(rows), 'le1990': summarize([r for r in rows if r[0] <= 199012]), 'ge1991': summarize([r for r in rows if r[0] >= 199101])}
        sub['pass'] = sub['all'].get('pass', False)
        sub['robust'] = bool(sub['le1990'].get('pass') and sub['ge1991'].get('pass'))
        res[f'H{H}'] = sub
    return res


def e1(D):
    """1960-01 からの延長（参考）: 金 1960-01〜1968-03 は World Bank の月平均、1968-04〜 LBMA の月末。為替は 1970-12 まで 360円。名目同額"""
    wbg = wb_gold()
    px = {k: v for k, v in wbg.items() if 196001 <= k <= 196803}
    px.update({k: v for k, v in D['gold_me'].items() if k >= 196804})
    fx = {k: 360.0 for k in range(196001, 197101) if 1 <= k % 100 <= 12}
    fx.update(D['fx'])
    gj = {k: px[k] * fx[k] for k in px if k in fx}
    Dx = {'gold_e1': price_to_ret(gj), 'sp_e1': to_jpy(D['sp_usd'], fx)}
    res = {}
    for w in (0.10, 0.20):
        res[f'w{int(w * 100)}'] = run_dca(Dx, 'sp_e1', gold_key='gold_e1', w=w, contrib='nominal', cpi_key=None, start_min=196002, horizons=(20,), detail=True)
    return res


if __name__ == '__main__':
    if '--check' in sys.argv:
        check()
    else:
        main()
