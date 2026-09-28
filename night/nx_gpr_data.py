#!/usr/bin/env python3
"""night/nx_gpr_data.py — nx 角度 gpr（地政学リスク指数が跳ねた後に株を厚く持つ）の**データの取得・整形・規則の定義だけ**
（成績は計算しない・規則のリターン・市場との差・t・シャープ・勝率は一つも出さない・読むだけ・門の判定には不使用）

2026-09-28 ユーザー指示「市場に勝てる歴史検証が出るまでいろんな角度から調べて勝てる結果を出して…別のSessionで検証していない新たな分析を」。
事前登録: out/nx_gpr_prereg.json（測る前に固定）／全体の線: out/nx_prereg.json（C1〜C8・格付け）。
測る道具 night/nx_gpr.py（これから書く）はここを import する（写さない）。

データ（キャッシュ out/_nx_cache/・gitignore）
  Caldara & Iacoviello (2022, AER 112(4)『Measuring Geopolitical Risk』) の GPR 指数
    最新の版  https://www.matteoiacoviello.com/gpr_files/data_gpr_export.xls（2026-09-01 更新・1900-01〜2026-08・月次）
      GPRH/GPRHT/GPRHA = 歴史の指数（米国の3紙: NYT・Chicago Tribune・Washington Post・1900〜・1900:2019=100）と その脅威/行為の部分
      GPR/GPRT/GPRA    = 最近の指数（10紙・1985〜・1985:2019=100）
      GPRC_xxx         = 国別（10紙の中で、地政学の語とその国の名を含む記事の割合・1985〜・44か国）
      GPRHC_xxx        = 国別の歴史版（3紙・1900〜・44か国）
    古い方法の版  https://www.matteoiacoviello.com/gpr_files/gpr_web_latest.xlsx（2019 年の作業論文の語の一覧・2021-10-15 が最後の更新・
      GPRH は 1899-01〜2021-09）＝語の選び方への感度を見る報告用
    過去の版（実時間）  https://raw.githubusercontent.com/iacoviel/iacoviel.github.io/master/gpr_archive_files/data_gpr_export_YYYYMM.xls
      （YYYYMM＝更新の月。2021-10〜2026-09 のうち 2022-02 だけ無い＝59版。2021-09 以前の版は公開されていない）
  Ken French 3因子（Mkt = Mkt-RF + RF、RF・1926-07〜）
  Shiller ie_data.xls（1871〜・月平均の株価＋年率の配当 → 名目の総リターン・nx_stack_data.shiller_nominal_tr をそのまま使う）
  FRED M13002US35620M156NNBR（NBER のニューヨークの商業手形 4〜6か月・年率%・1857〜1971）＝1926年より前の現金と借入の土台
  JKP の国の市場（mkt・vw_cap・米ドル・米国の短期金利を引いた超過）＝C5（米国外の国）
  Yahoo QQQ・^NDX（報告のみ・生き残りの偏りあり）

信号（定義。事前登録 out/nx_gpr_prereg.json の signal と同じ）
  判断は月末 t。使えるのは G_{t−lag} まで（lag=1 が主: 月次の版は翌月の1日に出る＝月末 t に分かっているのは t−1 まで）。
  level  : X_m = G_m。X_m の『直近 window か月（window=None なら 1900-01 からの拡大窓）の中での中位順位の分位』
  jump   : X_m = ln((G_m + c_m) ÷ (G_{m−12..m−1} の平均 + c_m))。c_m = cfrac × (G の 1900-01〜m−1 の平均)（0 の月がある国の指数だけ cfrac=0.1・
           全体の指数は cfrac=0＝0 の月が1つも無いことを確かめて使う）。分位は拡大窓（window=None）
  分位 = (X_s < X_m の数 + 0.5 × X_s = X_m の数) ÷ 窓の中の数（s は m 以前・窓の中）。窓の中の数が min_n=120 未満なら未定義
  spike(t) = 分位(X_{t−lag}) ≥ thr（thr=0.90 が主）
  持つ月 m の状態 on(m) = spike(t) が t ∈ [m−off−H+1, m−off] のどれかで立った（H=6 が主。off=1: French の時代、off=2: Shiller の時代
           ＝月平均の株価なので1か月余分に遅らせる）。未定義の t しか無い月は規則を作らない（評価から外す）
  resid  : 探索 E10。jump を、同じ月の市場のリターン r_s と直近12か月の対数リターンの和 R12_s に拡大窓で回帰した残差の分位
           （回帰に使う組は s ≤ t−lag だけ・最低 120 組）

使い方:
  python3 night/nx_gpr_data.py            # 取得＋信号＋形の報告 → out/_nx_cache/nx_gpr_signals.json と sha256
  python3 night/nx_gpr_data.py --selftest # 定義の点検（合成データだけ。先読みの検査・持つ月の数え方・積立の帳簿）
"""
import sys, os, io, json, math, random, hashlib, datetime, bisect, csv

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import nx_common as N  # noqa: E402

OUT = os.path.join(N.CACHE, 'nx_gpr_signals.json')
VINT_OUT = os.path.join(N.CACHE, 'nx_gpr_vintages.json')
GPR_URL = 'https://www.matteoiacoviello.com/gpr_files/data_gpr_export.xls'
GPR_NAME = 'gpr_data_gpr_export.xls'
GPR_OLD_URL = 'https://www.matteoiacoviello.com/gpr_files/gpr_web_latest.xlsx'
GPR_OLD_NAME = 'gpr_old_gpr_web_latest.xlsx'
VINT_URL = 'https://raw.githubusercontent.com/iacoviel/iacoviel.github.io/master/gpr_archive_files/data_gpr_export_{}.xls'
VINT_FIRST, VINT_LAST = 202110, 202609
NBER_CP_URL = 'https://fred.stlouisfed.org/graph/fredgraph.csv?id=M13002US35620M156NNBR'
NBER_CP_NAME = 'fred_M13002US35620M156NNBR.csv'
GLOBAL_COLS = ['GPRH', 'GPRHT', 'GPRHA', 'GPR', 'GPRT', 'GPRA']
COUNTRIES = ['ARG', 'AUS', 'BEL', 'BRA', 'CAN', 'CHE', 'CHL', 'CHN', 'COL', 'DEU', 'DNK', 'EGY', 'ESP', 'FIN', 'FRA', 'GBR', 'HKG',
             'HUN', 'IDN', 'IND', 'ISR', 'ITA', 'JPN', 'KOR', 'MEX', 'MYS', 'NLD', 'NOR', 'PER', 'PHL', 'POL', 'PRT', 'RUS', 'SAU',
             'SWE', 'THA', 'TUN', 'TUR', 'TWN', 'UKR', 'USA', 'VEN', 'VNM', 'ZAF']   # GPR の44か国（USA を含む）
NON_US = [c for c in COUNTRIES if c != 'USA']
YAHOO = ['QQQ', '^NDX']

# 期間（事前登録と同じ）
PRE1926 = (191001, 192606)     # 報告（Shiller・独立の時代）
FULL = (192607, 202608)        # French の時代
TRAIN = (192607, 200612)
HOLD = (200701, 202608)
RECENT = 201307
POSTPUB_WP = 201803            # Caldara-Iacoviello の作業論文（FRB IFDP 1222・2018-02）の翌月
POSTPUB_AER = 202205           # AER 2022-04 号の翌月
POST_LEXICON = 202111          # 今の語の一覧が決まった（2021-10 改訂）後の最初の持つ月
C5_END = 202512                # JKP の終わり
C5_MIN_MONTHS = 180            # 国を数えるのに要る評価の月数（15年）


# ───────────────────────── 月の計算 ─────────────────────────
def add(ym, k):
    y, m = divmod(ym, 100)
    t = y * 12 + m - 1 + k
    return (t // 12) * 100 + t % 12 + 1


def months(a, z):
    out, m = [], a
    while m <= z:
        out.append(m)
        m = add(m, 1)
    return out


def span(d):
    ks = sorted(d)
    if not ks:
        return None
    exp = len(months(ks[0], ks[-1]))
    return {'from': ks[0], 'to': ks[-1], 'n': len(ks), 'gaps': exp - len(ks)}


# ───────────────────────── 読み手 ─────────────────────────
def parse_gpr_xls(b, keep=None):
    """GPR の xls（Sheet1: 1列目が Excel の日付・最後の2列が変数の説明）→ ({列: {yyyymm: 値}}, {列: 説明})"""
    import xlrd
    wb = xlrd.open_workbook(file_contents=b)
    s = wb.sheet_by_index(0)
    hdr = [str(h).strip() for h in s.row_values(0)]
    iv = hdr.index('var_name') if 'var_name' in hdr else None
    il = hdr.index('var_label') if 'var_label' in hdr else None
    data_cols = [h for h in hdr[1:] if h not in ('var_name', 'var_label') and h]
    if keep is not None:
        data_cols = [h for h in data_cols if h in keep]
    idx = {h: hdr.index(h) for h in data_cols}
    out = {h: {} for h in data_cols}
    labels = {}
    for r in range(1, s.nrows):
        v = s.row_values(r)
        if iv is not None and v[iv]:
            labels[str(v[iv])] = str(v[il])
        if v[0] in ('', None):
            continue
        d = xlrd.xldate_as_datetime(v[0], wb.datemode)
        ym = d.year * 100 + d.month
        for h, j in idx.items():
            x = v[j]
            if x in ('', None):
                continue
            try:
                out[h][ym] = float(x)
            except (TypeError, ValueError):
                continue
    return out, labels


def load_gpr():
    b = N.get(GPR_URL, name=GPR_NAME, max_age_days=3650)   # 固定した版（2026-09-01 更新）を使い続ける。更新は事前登録の後の別の版
    cols, labels = parse_gpr_xls(b)
    return cols, labels, hashlib.sha256(b).hexdigest()


def load_gpr_old():
    """2019 年の方法の版（gpr_web_latest.xlsx）→ {'GPRH_old','GPRHT_old','GPRHA_old','GPR_old'}"""
    import openpyxl
    b = N.get(GPR_OLD_URL, name=GPR_OLD_NAME, max_age_days=3650)
    wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
    out = {'GPRH_old': {}, 'GPRHT_old': {}, 'GPRHA_old': {}, 'GPR_old': {}}
    mon = {m: i + 1 for i, m in enumerate(['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August',
                                           'September', 'October', 'November', 'December'])}
    hdr = None
    for row in wb['GPR_HISTORICAL'].iter_rows(values_only=True):
        if hdr is None:
            if row and row[0] == 'Year':
                hdr = list(row)
            continue
        if row[0] is None or row[1] not in mon:
            continue
        ym = int(math.floor(float(row[0]) + 1e-6)) * 100 + mon[row[1]]
        for k, c in (('GPRH_old', 'GPRH'), ('GPRHT_old', 'GPRHT'), ('GPRHA_old', 'GPRHA')):
            x = row[hdr.index(c)]
            if isinstance(x, (int, float)):
                out[k][ym] = float(x)
    hdr = None
    for row in wb['GPR'].iter_rows(values_only=True):
        if hdr is None:
            if row and row[0] == 'Date':
                hdr = list(row)
            continue
        if not isinstance(row[0], datetime.datetime):
            continue
        x = row[hdr.index('GPR')]
        if isinstance(x, (int, float)):
            out['GPR_old'][row[0].year * 100 + row[0].month] = float(x)
    return out, hashlib.sha256(b).hexdigest()


def load_vintages(refresh=False):
    """実時間の版（2021-10〜2026-09・2022-02 は無い）→ {更新の月: {列: {yyyymm: 値}}}。全体の指数と GPRC_* だけ残す"""
    if os.path.exists(VINT_OUT) and not refresh:
        j = json.load(open(VINT_OUT))
        return {int(k): {c: {int(m): x for m, x in s.items()} for c, s in v.items()} for k, v in j['vintages'].items()}, j['missing'], j['sha256']
    keep = set(GLOBAL_COLS) | {'GPRC_' + c for c in COUNTRIES}
    vint, missing, shas = {}, [], {}
    for v in months(VINT_FIRST, VINT_LAST):
        try:
            b = N.get(VINT_URL.format(v), name=f'gpr_vintage_{v}.xls', max_age_days=36500, tries=2)
        except RuntimeError:
            missing.append(v)
            continue
        cols, _ = parse_gpr_xls(b, keep=keep)
        vint[v] = cols
        shas[v] = hashlib.sha256(b).hexdigest()
    json.dump({'vintages': {str(k): {c: {str(m): x for m, x in s.items()} for c, s in v.items()} for k, v in vint.items()},
               'missing': missing, 'sha256': {str(k): x for k, x in shas.items()}}, open(VINT_OUT, 'w'))
    return vint, missing, {str(k): x for k, x in shas.items()}


def shiller_tr():
    import nx_stack_data as SD
    return SD.shiller_nominal_tr()


def nber_cp():
    """NBER の商業手形の利回り（年率%・月平均）→ {yyyymm: 月の率（小数）}"""
    b = N.get(NBER_CP_URL, name=NBER_CP_NAME, max_age_days=3650)
    out = {}
    for row in csv.DictReader(io.StringIO(b.decode())):
        v = row.get('M13002US35620M156NNBR')
        if v in (None, '', '.'):
            continue
        d = row['observation_date']
        out[int(d[:4]) * 100 + int(d[5:7])] = float(v) / 100 / 12
    return out, hashlib.sha256(b).hexdigest()


def jkp_country_mkt(c):
    """JKP の国の市場（vw_cap・米ドル・米国の短期金利を引いた超過）。無ければ None"""
    import zipfile
    r = c.lower()
    url = N.JKP.format(sub='', r=r, k='mkt', f='monthly', w='vw_cap')
    try:
        b = N.get(url, name=f'jkp_factor_{r}_mkt_vw_cap_monthly.zip', tries=2)   # nx_common.jkp_rows と同じキャッシュの名前
    except RuntimeError:
        return None
    z = zipfile.ZipFile(io.BytesIO(b))
    rows = list(csv.DictReader(io.StringIO(z.read(z.namelist()[0]).decode())))
    return {N._ym(x['date']): float(x['ret']) for x in rows if x['ret'] not in ('', 'NA', 'na')}


# ───────────────────────── 信号（事前登録の定義そのもの）─────────────────────────
def jump_series(G, k=12, cfrac=0.0):
    """X_m = ln((G_m + c_m) / (mean(G_{m−k..m−1}) + c_m))、c_m = cfrac × (G の最初〜m−1 の平均)。
    cfrac=0 のとき G に 0 以下があれば止まる（欠測を0と読まない・0 の対数は作らない）"""
    ks = sorted(G)
    out, run_sum, run_n = {}, 0.0, 0
    for m in ks:
        prev = [G.get(add(m, -j)) for j in range(1, k + 1)]
        c = cfrac * (run_sum / run_n) if run_n else 0.0
        if None not in prev:
            den = sum(prev) / k + c
            num = G[m] + c
            if cfrac == 0.0 and (G[m] <= 0 or den <= 0):
                raise ValueError(f'jump: 0 以下の値（cfrac=0）{m}')
            if num > 0 and den > 0:
                out[m] = math.log(num / den)
        run_sum += G[m]
        run_n += 1
    return out


def pct_series(X, min_n=120, window=None):
    """各月 m の X_m の中位順位の分位（m 以前・直近 window か月〔None なら拡大窓〕の中）。数が min_n 未満なら入れない"""
    ks = sorted(X)
    out = {}
    if window is None:
        srt = []
        for m in ks:
            bisect.insort(srt, X[m])
            n = len(srt)
            if n < min_n:
                continue
            lo, hi = bisect.bisect_left(srt, X[m]), bisect.bisect_right(srt, X[m])
            out[m] = (lo + 0.5 * (hi - lo)) / n
        return out
    for i, m in enumerate(ks):
        # 窓は『暦の直近 window か月』（途中の欠けがあっても暦で切る）
        a = add(m, -(window - 1))
        hist = [X[k] for k in ks[max(0, i - window - 1):i + 1] if k >= a]
        if len(hist) < min_n:
            continue
        x = X[m]
        below = sum(1 for h in hist if h < x)
        eq = sum(1 for h in hist if h == x)
        out[m] = (below + 0.5 * eq) / len(hist)
    return out


def spike_decisions(G, kind, lag=1, thr=0.90, min_n=120, window=None, cfrac=0.0):
    """判断の月 t → spike か（True/False）。t で使うのは G_{t−lag} まで。定義できない t は入れない"""
    if kind == 'level':
        X = dict(G)
    elif kind == 'jump':
        X = jump_series(G, 12, cfrac)
    else:
        raise ValueError(kind)
    P = pct_series(X, min_n=min_n, window=window)
    return {add(m, lag): (p >= thr) for m, p in P.items()}


def resid_spike_decisions(G, mkt, lag=1, thr=0.90, min_n=120):
    """探索 E10: jump を [1, r_s, R12_s] に回帰した残差の分位で spike。t で使う組は s ≤ t−lag（J と市場のリターンの両方がそろう月）"""
    import numpy as np
    J = jump_series(G, 12, 0.0)
    R12 = {}
    for m in sorted(mkt):
        w = [mkt.get(add(m, -j)) for j in range(12)]
        if None not in w:
            R12[m] = sum(math.log1p(x) for x in w)
    ss = sorted(s for s in J if s in mkt and s in R12)
    out = {}
    for i, s_last in enumerate(ss):
        n = i + 1
        if n < min_n:
            continue
        S_ = ss[:n]
        Xm = np.array([[1.0, mkt[s], R12[s]] for s in S_])
        y = np.array([J[s] for s in S_])
        beta, *_ = np.linalg.lstsq(Xm, y, rcond=None)
        e = y - Xm @ beta
        x = e[-1]
        below = float((e < x).sum()); eq = float((e == x).sum())
        out[add(s_last, lag)] = ((below + 0.5 * eq) / n) >= thr
    return out


def on_state(decisions, H=6, off=1):
    """持つ月 m → on か。m の状態は判断の月 t ∈ [m−off−H+1, m−off] の spike の有無。m−off が未定義の月は入れない"""
    out = {}
    for t in sorted(decisions):
        m = add(t, off)
        out[m] = any(decisions.get(add(t, -j), False) for j in range(H))
    return out


def weights(on, w_on, w_off):
    return {m: (w_on if v else w_off) for m, v in on.items()}


# ───────────────────────── ポートフォリオと積立（定義。ここでは合成データでだけ動かす）─────────────────────────
def timing_returns(w, r, rf, spread=0.015, cost=0.001):
    """株の割合 w（月ごと）で持つ。R = rf + w(r − rf) − max(w−1,0)×spread/12（1を超える分は rf+spread で借りる）。
    売買量 = |w_m − 前月末に値動きで変わった割合|（前月末の割合 = w_{m−1}(1+r_{m−1}) ÷ (1+R_{m−1})）、最初の月は |w − 1|
    （相手＝市場を1倍で持っている状態から始める）。net = gross − 売買量 × cost。
    r・rf が無い月は作らない（0 と置かない）。戻り: (gross, net, turnover)"""
    ms = sorted(m for m in w if m in r and m in rf)
    gross, net, turn = {}, {}, {}
    prev = None
    for m in ms:
        R = rf[m] + w[m] * (r[m] - rf[m]) - max(w[m] - 1.0, 0.0) * spread / 12
        if prev is None or add(prev, 1) != m:
            tv = abs(w[m] - 1.0)
        else:
            drift = w[prev] * (1 + r[prev]) / (1 + gross[prev])
            tv = abs(w[m] - drift)
        gross[m] = R
        turn[m] = tv
        net[m] = R - tv * cost
        prev = m
    return gross, net, turn


def dca_reserve(on, r, rf, years=20, step=12, share_off=0.8, first=None, last=None):
    """積立の予備金の規則（D 族）。毎月1を出す。
      相手: 毎月1を株へ（E_b = (E_b + 1)(1 + r)）＝nx_common.dca と同じ帳簿
      規則: on の月は その月の1 と 予備金の全部 を株へ（予備金は0に）。off の月は share_off を株へ・1−share_off を予備金へ（予備金は rf で増える）
    窓は on・r・rf がそろう連続の月から years×12 か月ずつ、step か月ずつずらす（first/last で窓の始まりの範囲を絞る）。
    戻り: [(窓の始まり, 最終額の比〔(E_規則+予備金) ÷ E_相手〕, 予備金の最大〔月の拠出の何か月分〕, 予備金を投じた回数)]"""
    ks = sorted(m for m in on if m in r and m in rf)
    n = years * 12
    out = []
    i = 0
    while i + n <= len(ks):
        w = ks[i:i + n]
        if add(w[0], n - 1) != w[-1]:       # 途中に欠けがある窓は使わない
            i += step
            continue
        if (first is not None and w[0] < first) or (last is not None and w[0] > last):
            i += step
            continue
        Eb = Er = Rv = 0.0
        mx, deploys = 0.0, 0
        for m in w:
            Eb = (Eb + 1.0) * (1 + r[m])
            if on[m]:
                if Rv > 0:
                    deploys += 1
                Er = (Er + 1.0 + Rv) * (1 + r[m])
                Rv = 0.0
            else:
                Er = (Er + share_off) * (1 + r[m])
                Rv = (Rv + (1.0 - share_off)) * (1 + rf[m])
            mx = max(mx, Rv)
        out.append((w[0], round((Er + Rv) / Eb, 4), round(mx, 2), deploys))
        i += step
    return out


def dca_summary(rows):
    if not rows:
        return None
    v = sorted(x[1] for x in rows)
    return {'windows': len(rows), 'win_rate': round(sum(1 for x in v if x > 1) / len(v), 3), 'median_ratio': v[len(v) // 2],
            'worst': min(rows, key=lambda x: x[1])[:2], 'best': max(rows, key=lambda x: x[1])[:2],
            'max_reserve_months': max(x[2] for x in rows)}


# ───────────────────────── 規則の台帳（事前登録の families と同じ）─────────────────────────
SIG = {
    'jump':        {'series': 'GPRH', 'kind': 'jump', 'lag': 1, 'thr': 0.90, 'min_n': 120, 'window': None, 'cfrac': 0.0},
    'level240':    {'series': 'GPRH', 'kind': 'level', 'lag': 1, 'thr': 0.90, 'min_n': 120, 'window': 240, 'cfrac': 0.0},
    'level_exp':   {'series': 'GPRH', 'kind': 'level', 'lag': 1, 'thr': 0.90, 'min_n': 120, 'window': None, 'cfrac': 0.0},
    'jump_t80':    {'series': 'GPRH', 'kind': 'jump', 'lag': 1, 'thr': 0.80, 'min_n': 120, 'window': None, 'cfrac': 0.0},
    'jump_t95':    {'series': 'GPRH', 'kind': 'jump', 'lag': 1, 'thr': 0.95, 'min_n': 120, 'window': None, 'cfrac': 0.0},
    'jump_acts':   {'series': 'GPRHA', 'kind': 'jump', 'lag': 1, 'thr': 0.90, 'min_n': 120, 'window': None, 'cfrac': 0.0},
    'jump_threats': {'series': 'GPRHT', 'kind': 'jump', 'lag': 1, 'thr': 0.90, 'min_n': 120, 'window': None, 'cfrac': 0.0},
    'jump_lag0':   {'series': 'GPRH', 'kind': 'jump', 'lag': 0, 'thr': 0.90, 'min_n': 120, 'window': None, 'cfrac': 0.0},
    'resid':       {'series': 'GPRH', 'kind': 'resid', 'lag': 1, 'thr': 0.90, 'min_n': 120, 'window': None, 'cfrac': 0.0},
}
RULES = {
    # 主の族 P（格付け・Holm はこの4本で・C5 を測る・C8 必須）
    'P1_lev_jump':     {'family': 'P', 'sig': 'jump', 'H': 6, 'w_on': 1.5, 'w_off': 1.0},
    'P2_lev_level':    {'family': 'P', 'sig': 'level240', 'H': 6, 'w_on': 1.5, 'w_off': 1.0},
    'P3_unlev_jump':   {'family': 'P', 'sig': 'jump', 'H': 6, 'w_on': 1.0, 'w_off': 0.8},
    'P4_unlev_level':  {'family': 'P', 'sig': 'level240', 'H': 6, 'w_on': 1.0, 'w_off': 0.8},
    # 探索の族 E（格付けはするが『探索』と明記・Holm は E の中で・C5 は測らない＝N/A・C8 必須）
    'E1_H1':           {'family': 'E', 'sig': 'jump', 'H': 1, 'w_on': 1.5, 'w_off': 1.0},
    'E2_H3':           {'family': 'E', 'sig': 'jump', 'H': 3, 'w_on': 1.5, 'w_off': 1.0},
    'E3_H12':          {'family': 'E', 'sig': 'jump', 'H': 12, 'w_on': 1.5, 'w_off': 1.0},
    'E4_thr80':        {'family': 'E', 'sig': 'jump_t80', 'H': 6, 'w_on': 1.5, 'w_off': 1.0},
    'E5_thr95':        {'family': 'E', 'sig': 'jump_t95', 'H': 6, 'w_on': 1.5, 'w_off': 1.0},
    'E6_acts':         {'family': 'E', 'sig': 'jump_acts', 'H': 6, 'w_on': 1.5, 'w_off': 1.0},
    'E7_threats':      {'family': 'E', 'sig': 'jump_threats', 'H': 6, 'w_on': 1.5, 'w_off': 1.0},
    'E8_lag0':         {'family': 'E', 'sig': 'jump_lag0', 'H': 6, 'w_on': 1.5, 'w_off': 1.0},
    'E9_derisk':       {'family': 'E', 'sig': 'jump', 'H': 6, 'w_on': 0.5, 'w_off': 1.0},
    'E10_resid':       {'family': 'E', 'sig': 'resid', 'H': 6, 'w_on': 1.5, 'w_off': 1.0},
    'E11_level_exp':   {'family': 'E', 'sig': 'level_exp', 'H': 6, 'w_on': 1.5, 'w_off': 1.0},
    'E12_both':        {'family': 'E', 'sig': ('jump', 'level240'), 'H': 6, 'w_on': 1.5, 'w_off': 1.0},
    # 積立の族 D（S/A/B/C の外・事前登録の『積立の線』で裁く）
    'D1_dca_jump':     {'family': 'D', 'sig': 'jump', 'H': 6, 'share_off': 0.8},
    'D2_dca_level':    {'family': 'D', 'sig': 'level240', 'H': 6, 'share_off': 0.8},
}
COUNTRY_CFRAC = 0.1   # C5 の国の指数（GPRC_xxx）の jump だけに使う c の割合（全体の指数は 0 の月が無いので 0）
COSTS = {'spread': 0.015, 'trade': 0.001, 'spread_sens': [0.005, 0.03], 'trade_sens': [0.0005, 0.003]}


def rules_sha():
    canon = {'SIG': SIG, 'RULES': {k: {kk: (list(vv) if isinstance(vv, tuple) else vv) for kk, vv in v.items()} for k, v in RULES.items()},
             'COSTS': COSTS, 'COUNTRY_CFRAC': COUNTRY_CFRAC, 'C5_MIN_MONTHS': C5_MIN_MONTHS, 'C5_END': C5_END}
    return hashlib.sha256(json.dumps(canon, sort_keys=True, separators=(',', ':')).encode()).hexdigest()


def decisions_for(sig_key, cols, mkt=None, series_override=None):
    """SIG の1本を GPR の列から作る（resid だけ市場のリターンが要る＝測る道具の側で mkt を渡す）"""
    s = SIG[sig_key]
    G = cols[series_override or s['series']]
    if s['kind'] == 'resid':
        if mkt is None:
            return None
        return resid_spike_decisions(G, mkt, s['lag'], s['thr'], s['min_n'])
    return spike_decisions(G, s['kind'], s['lag'], s['thr'], s['min_n'], s['window'], s['cfrac'])


def rule_on(rule_key, cols, off=1, mkt=None, series_override=None):
    r = RULES[rule_key]
    sigs = r['sig'] if isinstance(r['sig'], tuple) else (r['sig'],)
    decs = [decisions_for(s, cols, mkt, series_override) for s in sigs]
    if any(d is None for d in decs):
        return None
    common = set(decs[0])
    for d in decs[1:]:
        common &= set(d)
    both = {t: all(d[t] for d in decs) for t in common}
    return on_state(both, r['H'], off)


def signals_sha(sig):
    return hashlib.sha256(json.dumps(sig, sort_keys=True, separators=(',', ':')).encode()).hexdigest()


# ───────────────────────── 点検（合成データだけ）─────────────────────────
def selftest():
    rnd = random.Random(20260928)
    G = {}
    for m in months(190001, 202608):
        G[m] = 50 + 30 * rnd.random() + (400 if rnd.random() < 0.02 else 0)
    ok = True
    # (1) 先読みの検査: t より後の G をでたらめに替えても、t までの判断は1つも変わらない（level の窓・jump・lag の各型）
    for key in ('jump', 'level240', 'level_exp', 'jump_lag0'):
        s = SIG[key]
        base = spike_decisions(G, s['kind'], s['lag'], s['thr'], s['min_n'], s['window'], s['cfrac'])
        for cut in (192606, 195012, 200612):
            G2 = {m: (v if m <= add(cut, -s['lag']) else 1 + 500 * rnd.random()) for m, v in G.items()}
            alt = spike_decisions(G2, s['kind'], s['lag'], s['thr'], s['min_n'], s['window'], s['cfrac'])
            diff = [t for t in base if t <= cut and alt.get(t) != base[t]]
            if diff:
                ok = False
                print('先読みの疑い', key, cut, diff[:5])
        # わざと1か月先を使わせた版（lag を1つ減らす）は変わること
        if s['lag'] >= 1:
            # 1か月先を使う版（lag−1）では、cut より後の値を極端に大きくすると cut の判断が spike に変わる＝検査に力がある
            base0 = spike_decisions(G, s['kind'], s['lag'] - 1, s['thr'], s['min_n'], s['window'], s['cfrac'])
            cuts = [t for t in sorted(base0) if 195001 <= t <= 200012 and not base0[t]][:20]
            hit = 0
            for cut in cuts:
                G2 = {m: (v if m <= add(cut, -s['lag']) else 1e6) for m, v in G.items()}
                alt = spike_decisions(G2, s['kind'], s['lag'] - 1, s['thr'], s['min_n'], s['window'], s['cfrac'])
                hit += int(alt.get(cut) != base0[cut])
                alt_ok = spike_decisions(G2, s['kind'], s['lag'], s['thr'], s['min_n'], s['window'], s['cfrac'])
                if alt_ok.get(cut) != base[cut]:
                    ok = False
                    print('先読みの疑い（極端な値）', key, cut)
            if hit < len(cuts):
                ok = False
                print('1か月先を使う版が変わらなかった（検査の力が弱い）', key, hit, len(cuts))
    # (2) resid も同じ
    mkt = {m: 0.01 * rnd.gauss(0.7, 4) for m in months(192607, 202608)}
    base = resid_spike_decisions(G, mkt, 1, 0.9, 120)
    for cut in (196012, 200612):
        G2 = {m: (v if m <= add(cut, -1) else 1 + 500 * rnd.random()) for m, v in G.items()}
        mk2 = {m: (v if m <= add(cut, -1) else 0.3 * rnd.random() - 0.15) for m, v in mkt.items()}
        alt = resid_spike_decisions(G2, mk2, 1, 0.9, 120)
        diff = [t for t in base if t <= cut and alt.get(t) != base[t]]
        if diff:
            ok = False
            print('resid 先読みの疑い', cut, diff[:5])
    # (3) 持つ月の数え方: spike が t=195001 の1回だけなら、H=6・off=1 で 195002〜195007 が on
    dec = {m: (m == 195001) for m in months(194801, 195212)}
    on = on_state(dec, 6, 1)
    got = sorted(m for m, v in on.items() if v)
    if got != months(195002, 195007):
        ok = False
        print('on_state がおかしい', got)
    on2 = on_state(dec, 6, 2)
    if sorted(m for m, v in on2.items() if v) != months(195003, 195008):
        ok = False
        print('on_state off=2 がおかしい')
    # (4) 帳簿: w=1 なら規則＝市場（gross）、売買量0
    r = {m: 0.01 * rnd.gauss(0.7, 4) for m in months(195001, 196012)}
    rf = {m: 0.003 for m in r}
    g, n_, tv = timing_returns({m: 1.0 for m in r}, r, rf)
    if max(abs(g[m] - r[m]) for m in r) > 1e-12 or max(tv.values()) > 1e-12:
        ok = False
        print('timing_returns w=1 が市場と違う')
    # (5) 積立: share_off=1 なら規則＝相手（比1）。いつも on でも比1
    for onv in (False, True):
        rows = dca_reserve({m: onv for m in r}, r, rf, years=5, step=12, share_off=1.0 if not onv else 0.8)
        if any(abs(x[1] - 1) > 1e-9 for x in rows):
            ok = False
            print('dca_reserve の帳簿がおかしい', onv, rows[:2])
    # (6) 中位順位: 同じ値が並ぶと 0.5 付近（0 が9割の国で 0 の月が上位にならない）
    X = {m: (0.0 if i % 10 else 1.0) for i, m in enumerate(months(190001, 192012))}
    P = pct_series(X, min_n=120)
    zeros = [p for m, p in P.items() if X[m] == 0.0]
    if max(zeros) >= 0.9:
        ok = False
        print('中位順位がおかしい', max(zeros))
    print('selftest', 'OK' if ok else 'NG')
    return ok


# ───────────────────────── 形の報告と抽出 ─────────────────────────
def on_shape(on, eras):
    out = {}
    for name, (a, z) in eras.items():
        d = [m for m in on if a <= m <= z]
        o = [m for m in d if on[m]]
        ep = sum(1 for m in o if not on.get(add(m, -1), False))
        out[name] = {'months': len(d), 'on': len(o), 'on_share': round(len(o) / len(d), 3) if d else None, 'episodes': ep,
                     'first': min(d) if d else None}
    return out


def main():
    if '--selftest' in sys.argv:
        sys.exit(0 if selftest() else 1)
    cols, labels, sha = load_gpr()
    for c in ('GPRH', 'GPRHT', 'GPRHA', 'GPR'):
        assert min(cols[c].values()) > 0, c    # cfrac=0 の前提（0 の月が無い）
    old, sha_old = load_gpr_old()
    vint, vmiss, vsha = load_vintages()
    eras = {'pre1926': PRE1926, 'train': TRAIN, 'hold': HOLD, 'recent': (RECENT, HOLD[1]),
            'post_wp': (POSTPUB_WP, HOLD[1]), 'post_aer': (POSTPUB_AER, HOLD[1]), 'post_lexicon': (POST_LEXICON, HOLD[1])}

    signals, shapes = {}, {}
    # 米国（全体の指数）: French の時代は off=1、Shiller の時代は off=2
    for rk, r in RULES.items():
        if 'resid' in (r['sig'] if isinstance(r['sig'], tuple) else (r['sig'],)):
            shapes[rk] = '市場のリターンを使うので測る道具の側で作る（定義は resid_spike_decisions）'
            continue
        on1 = rule_on(rk, cols, off=1)
        on2 = rule_on(rk, cols, off=2)
        signals[rk] = {'off1': {str(m): int(v) for m, v in sorted(on1.items())},
                       'off2_pre1926': {str(m): int(v) for m, v in sorted(on2.items()) if m <= PRE1926[1]}}
        e1 = on_shape(on1, {k: v for k, v in eras.items() if k != 'pre1926'})
        e1['pre1926_off2'] = on_shape(on2, {'pre1926': PRE1926})['pre1926']
        shapes[rk] = e1

    # C5: 国の指数（GPRC・1985〜）で同じ規則（P 族だけ）。0 の月がある国は jump に cfrac=0.1
    country = {}
    zero_months = {}
    for c in NON_US + ['USA']:
        g = cols['GPRC_' + c]
        zero_months[c] = sum(1 for v in g.values() if v == 0)
        cf = COUNTRY_CFRAC   # 国の指数はすべて同じ c（0 の月があるかどうかを全期間で見てから決めると後知恵になる）
        for rk in ('P1_lev_jump', 'P2_lev_level'):
            sk = RULES[rk]['sig']
            s = SIG[sk]
            dec = spike_decisions(g, s['kind'], s['lag'], s['thr'], s['min_n'], s['window'], cf if s['kind'] == 'jump' else 0.0)
            on = on_state(dec, RULES[rk]['H'], 1)
            country.setdefault(c, {})[sk] = {str(m): int(v) for m, v in sorted(on.items())}
    signals['_country_GPRC'] = country
    signals['_country_cfrac'] = COUNTRY_CFRAC

    # 古い方法（2019 の語の一覧）の信号（報告 R3）
    old_sig = {}
    for rk in ('P1_lev_jump', 'P2_lev_level'):
        on = rule_on(rk, old, off=1, series_override='GPRH_old')
        old_sig[rk] = {str(m): int(v) for m, v in sorted(on.items())}
    signals['_old_method_GPRH'] = old_sig

    # 実時間の版（報告 R2）: 判断の月 t（2021-10〜2026-08）に、その月までに出た最新の版だけで判断を作る
    # ★版の『名前』ではなく『中身の最後の月』で使えるかを決める: 最後の月が L の版は L+1 月の1日に出る＝判断の月 t で使えるのは L ≤ t−1 の版。
    #   （202201 という名の版は中身が 2022-01 まで＝実は 2022-02 の更新。名前どおりに使うと 2022-01 末の判断に2月の版を使う先読みになる）
    rt = {}
    vks = sorted(vint)
    vlast = {v: max(vint[v]['GPRH']) for v in vks if vint[v].get('GPRH')}
    for sk in ('jump', 'level240'):
        s = SIG[sk]
        d_rt = {}
        for t in months(VINT_FIRST, HOLD[1]):
            avail = [v for v in vks if v in vlast and vlast[v] <= add(t, -1)]
            if not avail:
                continue
            v = max(avail, key=lambda x: (vlast[x], x))
            G = vint[v].get('GPRH')
            if not G:
                continue
            dd = spike_decisions(G, s['kind'], s['lag'], s['thr'], s['min_n'], s['window'], s['cfrac'])
            # 版 v は v−1 月までのデータ。判断の月 t で使えるのは t−1 まで → 版に t−1 が無ければ最新の月で判断（遅れが1か月増える）
            last = max(G)
            tt = t if add(t, -s['lag']) <= last else add(last, s['lag'])
            if tt in dd:
                d_rt[t] = {'spike': dd[tt], 'vintage': v, 'data_to': last}
        rt[sk] = {str(t): x for t, x in d_rt.items()}
    signals['_realtime_vintage_decisions'] = rt

    sig_sha = signals_sha(signals)
    # 市場の側の形（期間・欠けだけ）
    ff = N.ff_factors()
    sh, sh_sha = shiller_tr()
    cp, cp_sha = nber_cp()
    jkp, jkp_mkt_cache = {}, {}
    for c in NON_US + ['USA']:
        m = jkp_country_mkt(c)
        jkp_mkt_cache[c] = m
        jkp[c] = span(m) if m else None
    c5_months = {}
    for c in NON_US:
        m = jkp_mkt_cache.get(c)
        c5_months[c] = {sk: (sum(1 for k, v in country[c][sk].items() if int(k) <= C5_END and m and int(k) in m) if m else 0)
                        for sk in country[c]}
        c5_months[c]['counted'] = {sk: c5_months[c][sk] >= C5_MIN_MONTHS for sk in country[c]}
    yh = {}
    for t in YAHOO:
        try:
            yh[t] = span(N.yahoo(t))
        except Exception as e:  # noqa
            yh[t] = f'取得失敗 {e}'
    # 版どうしの違い（全体の指数の値の改訂の大きさ・判断が変わった月の数）＝成績ではない
    rev = {}
    latest = cols['GPRH']
    for v in vks:
        G = vint[v].get('GPRH', {})
        common = [m for m in G if m in latest]
        if not common:
            continue
        lastm = max(G)
        d = [abs(G[m] / latest[m] - 1) for m in common]
        rev[v] = {'data_to': lastm, 'median_abs_rel_diff': round(sorted(d)[len(d) // 2], 5), 'max_abs_rel_diff': round(max(d), 4),
                  'last_month_rel_diff': round(G[lastm] / latest[lastm] - 1, 4) if lastm in latest else None}
    flips = {}
    for sk in ('jump', 'level240'):
        s = SIG[sk]
        dl = spike_decisions(latest, s['kind'], s['lag'], s['thr'], s['min_n'], s['window'], s['cfrac'])
        r_ = signals['_realtime_vintage_decisions'][sk]
        n = [t for t in r_ if int(t) in dl]
        flips[sk] = {'decisions': len(n), 'differ_from_latest': sum(1 for t in n if r_[t]['spike'] != dl[int(t)]),
                     'spikes_realtime': sum(1 for t in n if r_[t]['spike']), 'spikes_latest': sum(1 for t in n if dl[int(t)])}

    rep = {'gpr_file': {'url': GPR_URL, 'sha256': sha, 'update': '2026-09-01（ページの記載『Last update: September 1, 2026』）',
                        'span': {c: span(cols[c]) for c in GLOBAL_COLS}, 'country_GPRC': span(cols['GPRC_JPN']),
                        'country_GPRHC': span(cols['GPRHC_JPN']), 'n_columns': len(cols)},
           'gpr_old': {'url': GPR_OLD_URL, 'sha256': sha_old, 'span': {k: span(v) for k, v in old.items()}},
           'vintages': {'n': len(vint), 'missing': vmiss, 'first': vks[0] if vks else None, 'last': vks[-1] if vks else None,
                        'revision_GPRH': rev, 'decision_flips_vs_latest': flips},
           'country_zero_months_GPRC': zero_months,
           'c5_eval_months': c5_months,
           'market': {'french_mkt': span(ff['mkt']), 'french_rf': span(ff['rf']), 'shiller_tr': span(sh), 'shiller_sha1': sh_sha,
                      'nber_cp': span(cp), 'nber_cp_sha256': cp_sha, 'jkp_country_mkt': jkp, 'yahoo': yh},
           'on_shapes_US': shapes,
           'country_on_shapes': {c: {sk: on_shape({int(m): bool(v) for m, v in d.items()},
                                                  {'c5': (199501, C5_END)}) ['c5'] for sk, d in x.items()} for c, x in country.items()},
           'rules_sha256': rules_sha(), 'signals_sha256': sig_sha}
    json.dump({'signals': signals, 'shape': rep, 'signals_sha256': sig_sha, 'rules_sha256': rules_sha(),
               'generated': datetime.date.today().isoformat()}, open(OUT, 'w'), ensure_ascii=False)
    print(json.dumps(rep, ensure_ascii=False, indent=1))


if __name__ == '__main__':
    main()
