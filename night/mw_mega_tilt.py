#!/usr/bin/env python3
"""night/mw_mega_tilt.py — 『市場に勝てる歴史検証』の角度 mega_tilt（読むだけ・門の判定には不使用）

問い: 大型株の投資家が作れる形『市場（上限なしの時価加重）＋ λ ×（大型株の中の買い−売り）』で、
      S&P500 型の純粋な市場（French Mkt）に勝てたか。
事前登録: out/mw_mega_tilt_prereg.json（線は out/mw_prereg.json・判定は mw_common.grade）
出力: out/mw_mega_tilt.json

族
- P（主・Holm 34本）: P-A 米国 mega の順位加重 LS（JKP 規模別・向きを掛ける）×λ0.2 ／ P-B 米国 JKP 'vw' 三分位 LS ×λ0.2
- X1（探索）: 同じものを λ=0.33、P-A を λ_feas（訓練期間だけで決めた買いだけの上限）
- X2（探索）: French の BIG（2x3）・ME5（5x5）の時価加重 LS ×λ0.2
"""
import csv, io, json, math, os, random, statistics as S, subprocess, sys, zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import mw_common as M  # noqa: E402

_FT, _orig_ft = {}, M.french_tables


def _ft_cached(name):  # French の zip を一度だけ読む（mw_common は変えない）
    if name not in _FT:
        _FT[name] = _orig_ft(name)
    return _FT[name]


M.french_tables = _ft_cached

BASE = M.BASE
PREREG = 'mw_mega_tilt_prereg.json'
PR = json.load(open(os.path.join(BASE, 'out', PREREG)))
CH = {k: v for k, v in PR['characteristics'].items() if not k.startswith('_')}
QUAL = ['gp_at', 'ope_be', 'qmj', 'qmj_prof', 'cop_at']
REGIONS = ['developed', 'world_ex_us', 'jpn', 'emerging', 'gbr', 'deu', 'fra', 'che', 'can', 'aus']
COUNTRIES = ['jpn', 'emerging', 'gbr', 'deu', 'fra', 'che', 'can', 'aus']
LAM = 0.2
COST = 0.001
JKP_END = 202512


def sha_of(path):
    try:
        return subprocess.run(['git', 'log', '--diff-filter=A', '--format=%H', '--', path], cwd=BASE,
                              capture_output=True, text=True).stdout.strip().split('\n')[-1] or None
    except Exception:  # noqa
        return None


# ───────────────────────── データ ─────────────────────────
def directions_from_xlsx():
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(M.CACHE, 'jkp_factor_details.xlsx'), read_only=True)
    rows = list(wb['details'].iter_rows(values_only=True))
    h = rows[0]
    ia, idr = h.index('abr_jkp'), h.index('direction')
    return {r[ia]: int(r[idr]) for r in rows[1:] if r[ia] and r[idr] is not None}


def mega_raw(k, sz='mega', nmin=50):
    """JKP 規模別（characteristic-managed・順位加重・向き無し）→ ({ym: ret 向きなし}, {ym: n})"""
    u = f'https://jkpfactors-data.s3.amazonaws.com/public/factor/%5Busa%5D_%5B{k}%5D_%5B{sz}%5D.zip'
    b = M.get(u, name=f'jkp_size_usa_{k}_{sz}.zip')
    z = zipfile.ZipFile(io.BytesIO(b))
    r, n = {}, {}
    for x in csv.DictReader(io.StringIO(z.read(z.namelist()[0]).decode())):
        if x['ret'] in ('', 'NA', 'na'):
            continue
        ym = int(x['date'][:4]) * 100 + int(x['date'][5:7])
        n[ym] = int(x['n'])
        if int(x['n']) >= nmin:
            r[ym] = float(x['ret'])
    return r, n


def vw_all(region):
    """JKP all_factors（vw・符号つき）→ {特徴: {ym: ret}}。国は n_stocks_min≥10、地域は n_countries≥3"""
    out = {}
    for x in M.jkp_rows(region, 'all_factors', 'factor', 'vw'):
        if x['name'] not in CH or x['ret'] in ('', 'NA', 'na'):
            continue
        if 'n_stocks_min' in x and x['n_stocks_min'] not in ('', 'NA') and int(float(x['n_stocks_min'])) < 10:
            continue
        if 'n_countries' in x and x['n_countries'] not in ('', 'NA') and int(float(x['n_countries'])) < 3:
            continue
        out.setdefault(x['name'], {})[int(x['date'][:4]) * 100 + int(x['date'][5:7])] = float(x['ret'])
    return out


def ftab(name, want):
    for t, v in M.french_tables(name).items():
        if want.lower() in t.lower() and v['freq'] == 'monthly':
            return v
    raise KeyError(f'{name}: {want}')


def fr_raw(name, want):
    """French の社数・平均時価総額の表 → {列: {ym: 値}}（百分率ではないので 100 で割らない）"""
    v = ftab(name, want)
    out = {c: {} for c in v['cols']}
    for d, row in v['data'].items():
        for c, x in zip(v['cols'], row):
            if x is not None:
                out[c][d] = x
    return out


def fr_capshare(name, cols):
    """French の社数×平均時価総額から cols の時価総額 → {ym: $百万}"""
    n = fr_raw(name, 'Number of Firms')
    try:
        a = fr_raw(name, 'Average Market Cap')
    except KeyError:
        a = fr_raw(name, 'Average Firm Size')
    out = {}
    for d in n[cols[0]]:
        v = 0.0
        ok = True
        for c in cols:
            if d in n[c] and d in a[c]:
                v += n[c][d] * a[c][d]
            else:
                ok = False
        if ok:
            out[d] = v
    return out


def total_cap():
    return fr_capshare('Portfolios_Formed_on_ME', ['Lo 30', 'Med 40', 'Hi 30'])


def me_p80():
    b = M.get('https://mba.tuck.dartmouth.edu/pages/faculty/ken.french/ftp/ME_Breakpoints_CSV.zip', name='fr_ME_Breakpoints.zip')
    z = zipfile.ZipFile(io.BytesIO(b))
    out = {}
    for line in z.read(z.namelist()[0]).decode('latin-1').splitlines():
        c = [x.strip() for x in line.split(',')]
        if c and c[0].isdigit() and len(c[0]) == 6 and len(c) >= 22:
            out[int(c[0])] = float(c[17])  # 5,10,…,80 → 16番目の分位（列 2+15）
    return out


def prev_month(ym):
    y, m = divmod(ym, 100)
    return (y - 1) * 100 + 12 if m == 1 else ym - 1


def fr_region_mkt(name):
    """French 地域の3因子（または5因子）→ 総リターン（USD）"""
    for t, v in M.french_tables(name).items():
        if v['freq'] == 'monthly' and 'Mkt-RF' in v['cols']:
            i, j = v['cols'].index('Mkt-RF'), v['cols'].index('RF')
            return {d: (row[i] + row[j]) / 100 for d, row in v['data'].items() if row[i] is not None and row[j] is not None}
    raise KeyError(name)


# ───────────────────────── 組み立て ─────────────────────────
def blend(series_list, weights=None):
    ks = set(series_list[0])
    for s in series_list[1:]:
        ks &= set(s)
    w = weights or [1 / len(series_list)] * len(series_list)
    return {k: sum(wi * s[k] for wi, s in zip(w, series_list)) for k in sorted(ks)}


def blend_def(name):
    if name == 'Q5':
        return QUAL, None
    if name == 'QM':
        return QUAL + ['ret_12_1'], [0.5 / 5] * 5 + [0.5]
    if name == 'QMV':
        return QUAL + ['ret_12_1', 'be_me'], [1 / 15] * 5 + [1 / 3, 1 / 3]
    raise KeyError(name)


def turnover_of(comps, weights=None):
    w = weights or [1 / len(comps)] * len(comps)
    return sum(wi * CH[c]['turnover'] for wi, c in zip(w, comps))


def build(ls_by_char, name):
    """特徴 or 混合の LS を作る。混合は構成要素がすべてそろう月だけ"""
    if name in CH:
        return ls_by_char.get(name)
    comps, w = blend_def(name)
    ss = [ls_by_char.get(c) for c in comps]
    if any(s is None or not s for s in ss):
        return None
    return blend(ss, w)


# ───────────────────────── 評価 ─────────────────────────
def tilt(mkt, ls, lam, to):
    ks = sorted(k for k in set(mkt) & set(ls) if k <= JKP_END)
    s = {k: mkt[k] + lam * ls[k] for k in ks}
    b = {k: mkt[k] for k in ks}
    net = M.apply_cost(s, to * lam * 2, COST)
    net3 = M.apply_cost(s, to * lam * 2, COST * 3)
    return s, b, net, net3


def repl_test(ls_regional, lam, to, mkts=None):
    """C5: 地域ごとに 費用後の λ×LS の算術平均 > 0 か"""
    per = {}
    for reg, ls in ls_regional.items():
        if not ls or len(ls) < 60:
            continue
        c = to * lam * 2 * COST / 12
        ks = sorted(k for k in ls if k <= JKP_END)
        ex = [lam * ls[k] - c for k in ks]
        rec = {'from': ks[0], 'to': ks[-1], 'years': round(len(ks) / 12, 1), 'net_ex_ann': round(S.mean(ex) * 1200, 2),
               't': (lambda t: round(t, 2) if t is not None else None)(M.nw_t(ex)), 'positive': S.mean(ex) > 0}
        hk = [k for k in ks if k >= M.HOLD_START]
        if len(hk) >= 24:
            hx = [lam * ls[k] - c for k in hk]
            rec['hold_net_ex_ann'] = round(S.mean(hx) * 1200, 2)
            rec['hold_t'] = (lambda t: round(t, 2) if t is not None else None)(M.nw_t(hx))
        if mkts and reg in mkts:
            m = mkts[reg]
            s = {k: m[k] + lam * ls[k] - c for k in ks if k in m}
            st = M.excess_stats(s, m)
            if st:
                rec['cagr_diff'] = st['cagr_diff']
        per[reg] = rec
    n = len(per)
    k = sum(1 for v in per.values() if v['positive'])
    cn = [r for r in per if r in COUNTRIES]
    return {'regions': n, 'positive': k, 'countries_only': {'regions': len(cn), 'positive': sum(1 for r in cn if per[r]['positive'])},
            'per_region': per}


def evaluate(name, family, desc, mkt, ls, lam, to, pub=None, repl=None, primary=False, extra=None):
    if not ls:
        return {'name': name, 'family': family, 'description': desc, 'error': 'データ無し'}
    s, b, net, net3 = tilt(mkt, ls, lam, to)
    full = M.excess_stats(s, b)
    train = M.excess_stats(s, b, z=M.TRAIN_END)
    hold = M.excess_stats(s, b, a=M.HOLD_START)
    recent = M.excess_stats(s, b, a=M.RECENT_START)
    cost_hold = M.excess_stats(net, b, a=M.HOLD_START)
    cost3_hold = M.excess_stats(net3, b, a=M.HOLD_START)
    post = M.excess_stats(s, b, a=(pub + 1) * 100 + 1) if pub else None
    roll = M.rolling(net, b, 20)
    dc = M.dca(net, b, 20)
    rec = {'name': name, 'family': family, 'primary': primary, 'description': desc, 'lambda': lam,
           'turnover_leg': round(to, 3), 'annual_cost_pct': round(to * lam * 2 * COST * 100, 4),
           'full': full, 'train': train, 'hold': hold, 'recent': recent, 'net_cost_hold': cost_hold,
           'net_cost3_hold': cost3_hold, 'post_publication': post, 'pub_year': pub,
           'roll20_net': roll, 'dca20_net': dc,
           'maxdd': {'tilt': round(M.maxdd(s) * 100, 1), 'mkt': round(M.maxdd(b) * 100, 1)},
           'repl': repl}
    if extra:
        rec.update(extra)
    return rec


def finish_family(recs):
    """族の中で Holm（保有期間・費用前の p）→ 判定"""
    ps = {r['name']: r['hold']['p'] for r in recs if r.get('hold') and r['hold'].get('p') is not None}
    hp = M.holm(ps)
    for r in recs:
        if r.get('error'):
            r['grade'] = 'C'
            continue
        r['family_holm_p'] = hp.get(r['name'])
        rp = r.get('repl')
        g, c = M.grade(r['full'], r['train'], r['hold'], r['roll20_net'], cost_hold=r['net_cost_hold'],
                       repl={'regions': rp['regions'], 'positive': rp['positive']} if rp else None,
                       family_holm_p=r['family_holm_p'])
        r['grade'], r['criteria'] = g, c
    return recs


# ───────────────────────── 実行可能性（λ の上限） ─────────────────────────
def pareto_negmass(n5, a5, xmin, total, n_rank, lam, draws=30, seed=0):
    """mega の中で順位と規模が無相関と仮定し、λ×順位加重 LS を足したときの負の重みの合計（推定）"""
    if a5 <= xmin or n5 < 10:
        return None
    alpha = a5 / (a5 - xmin)
    rnd = random.Random(seed)
    n = n_rank
    p = [(r + 1) / (n + 1) for r in range(n)]
    dev = [x - 0.5 for x in p]
    sc = sum(abs(d) for d in dev) / 2
    w = [d / sc for d in dev]
    out = []
    for _ in range(draws):
        caps = [min(xmin * (1 - rnd.random()) ** (-1 / alpha), total * 0.08) for _ in range(n)]
        rnd.shuffle(caps)
        neg = 0.0
        for c, wi in zip(caps, w):
            v = c / total + lam * wi
            if v < 0:
                neg += -v
        out.append(neg)
    return S.median(out)


def feasibility(mega_n):
    tot = total_cap()
    p80 = me_p80()
    worst, typ = {}, {}
    for ym, n in mega_n.items():
        pm = prev_month(ym)
        if ym in tot and pm in p80 and n >= 50:
            base = p80[pm] / tot[ym]
            worst[ym] = base * n / 4
            typ[ym] = base * n / 2
    tr = sorted(v for k, v in worst.items() if k <= M.TRAIN_END)
    lam_feas = max(0.01, math.floor(S.median(tr) * 100) / 100) if tr else None

    def summ(d, a=None, z=None):
        v = sorted(x for k, x in d.items() if (a is None or k >= a) and (z is None or k <= z))
        if not v:
            return None
        return {'min': round(v[0], 4), 'p05': round(v[int(len(v) * 0.05)], 4), 'median': round(v[len(v) // 2], 4), 'max': round(v[-1], 4), 'months': len(v)}
    # Pareto 推定（毎年12月）
    q = fr_raw('Portfolios_Formed_on_ME', 'Number of Firms')
    try:
        qa = fr_raw('Portfolios_Formed_on_ME', 'Average Firm Size')
    except KeyError:
        qa = fr_raw('Portfolios_Formed_on_ME', 'Average Market Cap')
    neg = {}
    for ym in sorted(mega_n):
        if ym % 100 != 12 or ym not in tot or prev_month(ym) not in p80 or mega_n[ym] < 50:
            continue
        n5, a5 = q['Hi 20'].get(ym), qa['Hi 20'].get(ym)
        if not n5 or not a5:
            continue
        v = {}
        for lam in (0.05, 0.1, 0.2, 0.33):
            nm = pareto_negmass(int(n5), a5, p80[prev_month(ym)], tot[ym], mega_n[ym], lam, seed=ym)
            v[str(lam)] = round(nm * 100, 2) if nm is not None else None
        neg[ym // 100] = v
    negsum = {}
    for lam in ('0.05', '0.1', '0.2', '0.33'):
        vals = [v[lam] for v in neg.values() if v.get(lam) is not None]
        hv = [v[lam] for y, v in neg.items() if y >= 2007 and v.get(lam) is not None]
        negsum[lam] = {'median_pct': round(S.median(vals), 2) if vals else None, 'max_pct': max(vals) if vals else None,
                       'median_2007_pct': round(S.median(hv), 2) if hv else None}
    return {'mega_worst_case_bound': {'full': summ(worst), 'train': summ(worst, z=M.TRAIN_END), 'hold': summ(worst, a=M.HOLD_START)},
            'mega_typical_bound': {'full': summ(typ), 'train': summ(typ, z=M.TRAIN_END), 'hold': summ(typ, a=M.HOLD_START)},
            'lambda_feas': lam_feas,
            'lambda_feas_rule': '〜2006 の各月の最悪ケース上限（NYSE 80%点 ÷ 市場全体 × n/4・n=ret_12_1 の mega 銘柄数）の中央値を 0.01 単位で切り捨て',
            'pareto_negative_mass_pct_by_lambda': negsum,
            'pareto_negative_mass_pct_by_year': {str(y): v for y, v in sorted(neg.items()) if y % 5 == 0 or y >= 2020},
            'note': '負の重みの合計＝λ×順位加重 LS を足したとき売りが必要になる割合（ポートフォリオ全体の%）。順位と規模が mega の中で無相関という仮定の推定'}, tot


def french_lam_max(name, bad_cols, tot):
    cap = fr_capshare(name, bad_cols)
    r = {k: cap[k] / tot[k] for k in cap if k in tot and tot[k] > 0}
    v = sorted(r.values())
    if not v:
        return None
    share_ok = sum(1 for x in v if x >= LAM) / len(v)
    return {'min': round(v[0], 3), 'p05': round(v[int(len(v) * 0.05)], 3), 'median': round(v[len(v) // 2], 3),
            'share_months_ge_0.2': round(share_ok, 3), 'long_only_at_0.2': share_ok >= 0.95}


# ───────────────────────── French の答え合わせ（X2）の定義 ─────────────────────────
FR2 = {  # 2x3 BIG: (ファイル, 良い側, 悪い側, 対応する JKP 特徴, 米国外のファイル接尾辞)
    'F2_be_me': ('6_Portfolios_2x3', 'BIG HiBM', 'BIG LoBM', 'be_me', 'BE-ME'),
    'F2_ni_me': ('6_Portfolios_ME_EP_2x3', 'BIG HiEP', 'BIG LoEP', 'ni_me', None),
    'F2_ope_be': ('6_Portfolios_ME_OP_2x3', 'BIG HiOP', 'BIG LoOP', 'ope_be', 'OP'),
    'F2_at_gr1': ('6_Portfolios_ME_INV_2x3', 'BIG LoINV', 'BIG HiINV', 'at_gr1', 'INV'),
    'F2_ret_12_1': ('6_Portfolios_ME_Prior_12_2', 'BIG HiPRIOR', 'BIG LoPRIOR', 'ret_12_1', 'Prior_12_2'),
}
FR5 = {  # 5x5 ME5
    'F5_be_me': ('25_Portfolios_5x5', 'BIG HiBM', 'BIG LoBM', 'be_me', 'BE-ME'),
    'F5_ope_be': ('25_Portfolios_ME_OP_5x5', 'BIG HiOP', 'BIG LoOP', 'ope_be', 'OP'),
    'F5_at_gr1': ('25_Portfolios_ME_INV_5x5', 'BIG LoINV', 'BIG HiINV', 'at_gr1', 'INV'),
    'F5_ret_12_1': ('25_Portfolios_ME_Prior_12_2', 'BIG HiPRIOR', 'BIG LoPRIOR', 'ret_12_1', 'Prior_12_2'),
    'F5_oaccruals_at': ('25_Portfolios_ME_AC_5x5', 'BIG LoAC', 'BIG HiAC', 'oaccruals_at', None),
    'F5_chcsho_12m': ('25_Portfolios_ME_NI_5x5', 'BIG NegNI', 'BIG HiNI', 'chcsho_12m', None),
    'F5_betabab_1260d': ('25_Portfolios_ME_BETA_5x5', 'BIG LoBETA', 'BIG HiBETA', 'betabab_1260d', None),
    'F5_ivol_capm_252d': ('25_Portfolios_ME_RESVAR_5x5', 'BIG LoVAR', 'BIG HiVAR', 'ivol_capm_252d', None),
}
FR_INTL = {'Developed_ex_US': 'Developed_ex_US_3_Factors', 'Europe': 'Europe_3_Factors', 'Japan': 'Japan_3_Factors',
           'Asia_Pacific_ex_Japan': 'Asia_Pacific_ex_Japan_3_Factors', 'Emerging_Markets': 'Emerging_5_Factors'}


def fr_ls(name, good, bad):
    v = M.french_series(name, 'Value Weight')
    return {k: v[good][k] - v[bad][k] for k in v[good] if k in v[bad]}


def fr_intl_ls(suffix, good_short, bad_short):
    out = {}
    for reg in FR_INTL:
        try:
            v = M.french_series(f'{reg}_6_Portfolios_ME_{suffix}', 'Value Weight')
        except Exception as e:  # noqa
            print('  French 米国外なし', reg, suffix, e)
            continue
        g = [c for c in v if c.replace(' ', '') == good_short.replace(' ', '')]
        bcol = [c for c in v if c.replace(' ', '') == bad_short.replace(' ', '')]
        if g and bcol:
            out[reg] = {k: v[g[0]][k] - v[bcol[0]][k] for k in v[g[0]] if k in v[bcol[0]]}
    return out


# ───────────────────────── 本体 ─────────────────────────
def part1():
    out = {'angle': 'mega_tilt', 'prereg': PREREG, 'prereg_commit': sha_of(f'out/{PREREG}'),
           'global_prereg': 'mw_prereg.json', 'global_prereg_commit': sha_of('out/mw_prereg.json'),
           'benchmark': 'French Mkt（Mkt-RF+RF）', 'jkp_end': JKP_END, 'sanity': {}, 'deviations': [], 'tested': []}
    # ── 向き
    dx = directions_from_xlsx()
    bad = {k: (CH[k]['direction'], dx.get(k)) for k in CH if dx.get(k) != CH[k]['direction']}
    out['sanity']['direction_matches_xlsx'] = not bad
    if bad:
        raise SystemExit(f'向きが事前登録と xlsx で違う: {bad}')
    # ── 市場
    ff = M.ff_factors()
    mkt, mktrf = ff['mkt'], ff['mktrf']
    out['sanity']['french_mkt_cagr_full'] = round(M.cagr(mkt) * 100, 2)
    out['sanity']['french_mkt_cagr_2007'] = round(M.cagr(M.window(mkt, M.HOLD_START)) * 100, 2)
    jm = M.jkp_mkt('usa', 'vw')
    ks = sorted(set(jm) & set(mktrf))
    out['sanity']['jkp_usa_vw_mkt_minus_french_mktrf'] = {
        'full_pct_yr': round(S.mean(jm[k] - mktrf[k] for k in ks) * 1200, 2),
        'hold_pct_yr': round(S.mean(jm[k] - mktrf[k] for k in ks if k >= M.HOLD_START) * 1200, 2),
        'corr': round(M.corr([jm[k] for k in ks], [mktrf[k] for k in ks]), 4)}
    # ── 米国の LS（P-A mega・P-B vw）
    mega, mega_n_all = {}, {}
    for k in CH:
        r, n = mega_raw(k)
        mega[k] = {ym: v * CH[k]['direction'] for ym, v in r.items()}
        mega_n_all[k] = n
    vwu = vw_all('usa')
    # 検算: vw の符号つき因子 = (3−1)×direction
    chk = {}
    for k in CH:
        p = M.jkp_portfolios('usa', k, 'vw')
        if '3.0' in p and '1.0' in p and k in vwu:
            kk = sorted(set(p['3.0']) & set(p['1.0']) & set(vwu[k]))
            chk[k] = round(M.corr([(p['3.0'][m] - p['1.0'][m]) * CH[k]['direction'] for m in kk], [vwu[k][m] for m in kk]), 4)
    out['sanity']['vw_factor_vs_tercile_spread_corr'] = chk
    # 検算: mega×direction と全規模 vw の相関（正であるべき）
    out['sanity']['mega_signed_vs_vw_signed_corr'] = {
        k: round(M.corr([mega[k][m] for m in sorted(set(mega[k]) & set(vwu.get(k, {})))],
                        [vwu[k][m] for m in sorted(set(mega[k]) & set(vwu.get(k, {})))]), 3) for k in CH if k in vwu}
    neg_corr = [k for k, v in out['sanity']['mega_signed_vs_vw_signed_corr'].items() if v <= 0]
    out['sanity']['mega_sign_ok'] = not neg_corr
    # 検算: French 2x3 BIG OP と JKP mega ope_be の相関
    fop = fr_ls('6_Portfolios_ME_OP_2x3', 'BIG HiOP', 'BIG LoOP')
    kk = sorted(set(fop) & set(mega['ope_be']))
    out['sanity']['french_bigOP_vs_mega_ope_be_corr'] = round(M.corr([fop[m] for m in kk], [mega['ope_be'][m] for m in kk]), 3)
    fbm = fr_ls('6_Portfolios_2x3', 'BIG HiBM', 'BIG LoBM')
    kk = sorted(set(fbm) & set(mega['be_me']))
    out['sanity']['french_bigBM_vs_mega_be_me_corr'] = round(M.corr([fbm[m] for m in kk], [mega['be_me'][m] for m in kk]), 3)
    fpr = fr_ls('6_Portfolios_ME_Prior_12_2', 'BIG HiPRIOR', 'BIG LoPRIOR')
    kk = sorted(set(fpr) & set(mega['ret_12_1']))
    out['sanity']['french_bigPrior_vs_mega_ret_12_1_corr'] = round(M.corr([fpr[m] for m in kk], [mega['ret_12_1'][m] for m in kk]), 3)
    print('検算', json.dumps(out['sanity'], ensure_ascii=False))
    # ── 米国外（C5）
    reg_vw = {r: vw_all(r) for r in REGIONS}
    reg_mkt = {r: M.jkp_mkt(r, 'vw') for r in REGIONS}
    # ── 実行可能性
    feas, tot = feasibility(mega_n_all['ret_12_1'])
    out['feasibility'] = feas
    lam_feas = feas['lambda_feas']
    print('実行可能性', json.dumps({k: v for k, v in feas.items() if k != 'pareto_negative_mass_pct_by_year'}, ensure_ascii=False))
    # JKP vw 三分位の悪い側の時価総額（French 2x3 の BIG+SMALL の悪い側で近似）
    approx = {}
    for key, (fname, g, bcol, jk, _) in FR2.items():
        sb = bcol.replace('BIG', 'SMALL')
        approx[jk] = french_lam_max(fname, [bcol, sb], tot)
    out['feasibility']['vw_tercile_bad_side_capshare_approx_by_french_2x3'] = approx
    out['feasibility']['french_bad_side_capshare'] = {}
    for key, (fname, g, bcol, jk, _) in list(FR2.items()) + list(FR5.items()):
        out['feasibility']['french_bad_side_capshare'][key] = french_lam_max(fname, [bcol], tot)
    names = list(CH) + ['Q5', 'QM', 'QMV']

    def to_of(nm):
        if nm in CH:
            return CH[nm]['turnover']
        c, w = blend_def(nm)
        return turnover_of(c, w)

    def pub_of(nm):
        if nm in CH:
            return CH[nm]['pub']
        c, _ = blend_def(nm)
        return max(CH[x]['pub'] for x in c)

    def repl_for(nm, lam):
        regls = {r: build(reg_vw[r], nm) for r in REGIONS}
        return repl_test({r: v for r, v in regls.items() if v}, lam, to_of(nm), reg_mkt)

    fams = {}
    # P（主）
    P = []
    for nm in names:
        desc = CH[nm]['意味'] if nm in CH else PR['blends'][nm]
        P.append(evaluate(f'PA_mega_{nm}', 'P', f'米国 mega 順位加重 LS: {desc}', mkt, build(mega, nm), LAM, to_of(nm), pub_of(nm),
                          repl_for(nm, LAM), primary=True))
    for nm in names:
        desc = CH[nm]['意味'] if nm in CH else PR['blends'][nm]
        P.append(evaluate(f'PB_vw_{nm}', 'P', f'米国 vw 三分位 LS（全規模・時価加重）: {desc}', mkt, build(vwu, nm), LAM, to_of(nm), pub_of(nm),
                          repl_for(nm, LAM), primary=True))
    fams['P'] = finish_family(P)
    # X1 λ=0.33
    X1a = []
    for nm in names:
        X1a.append(evaluate(f'X1_mega_{nm}_lam033', 'X1_lam033', f'P-A と同じ・λ=0.33', mkt, build(mega, nm), 0.33, to_of(nm), pub_of(nm), repl_for(nm, 0.33)))
    for nm in names:
        X1a.append(evaluate(f'X1_vw_{nm}_lam033', 'X1_lam033', f'P-B と同じ・λ=0.33', mkt, build(vwu, nm), 0.33, to_of(nm), pub_of(nm), repl_for(nm, 0.33)))
    fams['X1_lam033'] = finish_family(X1a)
    X1b = []
    for nm in names:
        X1b.append(evaluate(f'X1_mega_{nm}_lamfeas', 'X1_lamfeas', f'P-A と同じ・λ_feas={lam_feas}（訓練期間だけで決めた買いだけの上限）', mkt,
                            build(mega, nm), lam_feas, to_of(nm), pub_of(nm), repl_for(nm, lam_feas)))
    fams['X1_lamfeas'] = finish_family(X1b)
    # X2 French
    X2 = []
    frls = {}
    for key, (fname, g, bcol, jk, suf) in list(FR2.items()) + list(FR5.items()):
        ls = fr_ls(fname, g, bcol)
        frls[key] = ls
        if suf:
            gi = g.replace('BIG ', 'BIG ')
            intl = fr_intl_ls(suf, g, bcol)
            fm = {r: fr_region_mkt(FR_INTL[r]) for r in intl}
            rp = repl_test(intl, LAM, CH[jk]['turnover'], fm)
            rp['source'] = 'French 米国外 6_Portfolios の BIG の行'
        else:
            rp = repl_for(jk, LAM)
            rp['source'] = 'JKP vw 三分位（French に米国外が無い）'
        X2.append(evaluate(key, 'X2', f'French {fname}: {g} − {bcol}（時価加重・{CH[jk]["意味"]}）', mkt, ls, LAM, CH[jk]['turnover'], CH[jk]['pub'], rp,
                           extra={'long_only_capshare': out['feasibility']['french_bad_side_capshare'].get(key)}))
    for key, parts in (('F2_QM', ['F2_ope_be', 'F2_ret_12_1']), ('F2_QMV', ['F2_ope_be', 'F2_ret_12_1', 'F2_be_me'])):
        ls = blend([frls[p] for p in parts])
        jks = [FR2[p][3] for p in parts]
        intl_parts = [fr_intl_ls(FR2[p][4], FR2[p][1], FR2[p][2]) for p in parts]
        regs = set(intl_parts[0])
        for ip in intl_parts[1:]:
            regs &= set(ip)
        intl = {r: blend([ip[r] for ip in intl_parts]) for r in regs}
        fm = {r: fr_region_mkt(FR_INTL[r]) for r in intl}
        to = turnover_of(jks)
        rp = repl_test(intl, LAM, to, fm)
        rp['source'] = 'French 米国外 6_Portfolios の BIG の行'
        X2.append(evaluate(key, 'X2', f'French 2x3 BIG の混合: {" + ".join(parts)}（等分）', mkt, ls, LAM, to, max(CH[j]['pub'] for j in jks), rp))
    fams['X2'] = finish_family(X2)
    for f, recs in fams.items():
        out['tested'].extend(recs)
    out['n_tested'] = len(out['tested'])
    out['families'] = {f: {'n': len(r), 'grades': {g: sum(1 for x in r if x.get('grade') == g) for g in 'SABC'}} for f, r in fams.items()}
    out['deviations'].append('JKP の規模別（mega）は三分位の時価加重ではなく順位加重の characteristic-managed portfolio だった（作成コード portfolios.R の cmp 節で確認）。そのため λ=0.2 は厳密な買いだけでは作れない可能性が高い（feasibility 参照）')
    out['deviations'].append('JKP の規模別は米国しか公開されていないので、P-A の C5（再現）は同じ特徴の JKP vw 三分位 LS（全規模・時価加重）の米国外10地域で代用')
    ctx = {'mkt': mkt, 'mktrf': mktrf, 'mega': mega, 'vwu': vwu, 'reg_vw': reg_vw, 'reg_mkt': reg_mkt, 'tot': tot,
           'names': names, 'to_of': to_of, 'pub_of': pub_of, 'frls': frls, 'fams': fams}
    return out, ctx


# ═════════════════════════ prereg2（探索の追加） ═════════════════════════
PREREG2 = 'mw_mega_tilt_prereg2.json'
THEME_TO = {'short_term_reversal': 10.0, 'seasonality': 3.0, 'momentum': 1.5, 'profit_growth': 1.5, 'low_risk': 0.8, 'accruals': 0.8,
            'debt_issuance': 0.8, 'investment': 0.8, 'quality': 0.4, 'profitability': 0.4, 'value': 0.4, 'low_leverage': 0.4, 'size': 0.4}


def vw_all_every(region):
    """JKP all_factors（vw・符号つき）を全特徴で → {特徴: {ym: ret}}（フィルタは vw_all と同じ）"""
    out = {}
    for x in M.jkp_rows(region, 'all_factors', 'factor', 'vw'):
        if x['ret'] in ('', 'NA', 'na'):
            continue
        if 'n_stocks_min' in x and x['n_stocks_min'] not in ('', 'NA') and int(float(x['n_stocks_min'])) < 10:
            continue
        if 'n_countries' in x and x['n_countries'] not in ('', 'NA') and int(float(x['n_countries'])) < 3:
            continue
        out.setdefault(x['name'], {})[int(x['date'][:4]) * 100 + int(x['date'][5:7])] = float(x['ret'])
    return out


def themes_vw(region):
    out = {}
    for x in M.jkp_rows(region, 'all_themes', 'factor', 'vw'):
        if x['ret'] in ('', 'NA', 'na'):
            continue
        if 'n_countries' in x and x['n_countries'] not in ('', 'NA') and int(float(x['n_countries'])) < 3:
            continue
        out.setdefault(x['name'], {})[int(x['date'][:4]) * 100 + int(x['date'][5:7])] = float(x['ret'])
    return out


def cluster_labels():
    b = M.get('https://raw.githubusercontent.com/bkelly-lab/ReplicationCrisis/master/GlobalFactors/Cluster%20Labels.csv', name='jkp_cluster_labels.csv')
    return {r['characteristic']: r['cluster'].lower().replace('-', '_').replace(' ', '_') for r in csv.DictReader(io.StringIO(b.decode()))}


def composite(series_by_char, chars, min_share=0.8):
    """選んだ特徴の等分平均。月ごとに min_share 以上そろう月だけ・そろった分で平均（0で埋めない）"""
    months = set()
    for c in chars:
        months |= set(series_by_char.get(c, {}))
    out = {}
    need = math.ceil(min_share * len(chars))
    for m in sorted(months):
        v = [series_by_char[c][m] for c in chars if m in series_by_char.get(c, {})]
        if len(v) >= need and v:
            out[m] = sum(v) / len(v)
    return out


def fr_exclusion(fname, good, bad, tot):
    ls = fr_ls(fname, good, bad)
    cap = fr_capshare(fname, [bad])
    c = {k: cap[k] / tot[k] for k in cap if k in tot and tot[k] > 0}
    out, cs = {}, []
    for k in sorted(ls):
        pk = prev_month(k)
        if pk in c:
            out[k] = c[pk] * ls[k]
            cs.append(c[pk])
    return out, (S.mean(cs) if cs else None), (sorted(cs)[len(cs) // 2] if cs else None)


def diagnostics(mkt, mktrf, ls, lam):
    ks = sorted(k for k in set(mkt) & set(ls) & set(mktrf) if M.HOLD_START <= k <= JKP_END)
    s = {k: mkt[k] + lam * ls[k] for k in ks}
    b = {k: mkt[k] for k in ks}
    d = {'D1_2007_2015': M.excess_stats(s, b, a=200701, z=201512), 'D1_2016_2025': M.excess_stats(s, b, a=201601)}
    k2 = [k for k in ks if not (200801 <= k <= 200912)]
    ex2 = [lam * ls[k] for k in k2]
    t2 = M.nw_t(ex2)
    d['D2_ex_2008_2009'] = {'ex_ann': round(S.mean(ex2) * 1200, 2), 't': round(t2, 2) if t2 is not None else None, 'months': len(k2)}
    ex = [lam * ls[k] for k in ks]
    xm = [mktrf[k] for k in ks]
    mx, me = S.mean(xm), S.mean(ex)
    beta = sum((a - mx) * (e - me) for a, e in zip(xm, ex)) / sum((a - mx) ** 2 for a in xm)
    al = [e - beta * a for e, a in zip(ex, xm)]
    ta = M.nw_t(al)
    d['D3_capm'] = {'beta_of_excess': round(beta, 3), 'alpha_ann': round(S.mean(al) * 1200, 2), 't': round(ta, 2) if ta is not None else None}
    yrs = {}
    for k in ks:
        y = k // 100
        a = yrs.setdefault(y, [1.0, 1.0])
        a[0] *= 1 + s[k]; a[1] *= 1 + b[k]
    diffs = {y: round((v[0] - v[1]) * 100, 2) for y, v in yrs.items() if sum(1 for k in ks if k // 100 == y) == 12}
    best = max(diffs, key=diffs.get) if diffs else None
    d['D4_calendar'] = {'years': len(diffs), 'years_won': sum(1 for v in diffs.values() if v > 0),
                        'mean_diff_pct': round(S.mean(diffs.values()), 2) if diffs else None,
                        'best_year': [best, diffs.get(best)] if best else None,
                        'mean_diff_ex_best_pct': round(S.mean(v for y, v in diffs.items() if y != best), 2) if len(diffs) > 1 else None,
                        'worst_year': list(min(diffs.items(), key=lambda x: x[1])) if diffs else None,
                        'by_year': diffs}
    return d


def part2(out, ctx):
    mkt, mktrf, tot = ctx['mkt'], ctx['mktrf'], ctx['tot']
    out['prereg2'] = PREREG2
    out['prereg2_commit'] = sha_of(f'out/{PREREG2}')
    fams = ctx['fams']
    # ── X3a French 除外ルール
    X3 = []
    x2rep = {r['name']: r['repl'] for r in fams['X2']}
    for key, (fname, g, bcol, jk, suf) in list(FR2.items()) + list(FR5.items()):
        ls, cmean, cmed = fr_exclusion(fname, g, bcol, tot)
        rp = x2rep.get(key)
        rec = evaluate(f'X3a_{key}_excl', 'X3', f'French {fname}: 悪い側（{bcol}）を全部売り良い側（{g}）へ移す（前月の時価総額の比・構造的に買いだけ）',
                       mkt, ls, 1.0, CH[jk]['turnover'] * (cmean or 0), CH[jk]['pub'], rp,
                       extra={'lambda_is_time_varying': True, 'lambda_mean': round(cmean, 4) if cmean else None, 'lambda_median': round(cmed, 4) if cmed else None})
        X3.append(rec)
        ctx.setdefault('ls_of', {})[rec['name']] = (ls, 1.0)
    # ── X3b JKP vw を λ_vw で
    cap = fr_capshare('6_Portfolios_ME_OP_2x3', ['BIG LoOP', 'SMALL LoOP'])
    tr = sorted(cap[k] / tot[k] for k in cap if k in tot and k <= M.TRAIN_END)
    lam_vw = max(0.01, math.floor(tr[int(len(tr) * 0.05)] * 100) / 100)
    out['feasibility']['lambda_vw'] = lam_vw
    out['feasibility']['lambda_vw_rule'] = '訓練期間（1963-07〜2006-12）の French 2x3（BIG LoOP + SMALL LoOP）の時価総額 ÷ 市場全体 の5%点を 0.01 単位で切り捨て'
    to_of, pub_of, vwu, reg_vw, reg_mkt = ctx['to_of'], ctx['pub_of'], ctx['vwu'], ctx['reg_vw'], ctx['reg_mkt']
    for nm in ctx['names']:
        regls = {r: build(reg_vw[r], nm) for r in REGIONS}
        rp = repl_test({r: v for r, v in regls.items() if v}, lam_vw, to_of(nm), reg_mkt)
        ls = build(vwu, nm)
        rec = evaluate(f'X3b_vw_{nm}_lamvw', 'X3', f'P-B と同じ・λ_vw={lam_vw}（収益性の悪い側の時価総額の近似で買いだけに収まる見込みの大きさ）', mkt, ls,
                       lam_vw, to_of(nm), pub_of(nm), rp)
        X3.append(rec)
        ctx['ls_of'][rec['name']] = (ls, lam_vw)
    fams['X3'] = finish_family(X3)
    # ── X4 合成とテーマ
    dx = directions_from_xlsx()
    avail = json.load(open(os.path.join(M.CACHE, 'jkp_availability.json')))['factor_sizes']['usa']
    mega_all = {}
    for k in avail:
        if dx.get(k) not in (1, -1):
            continue
        try:
            r, _ = mega_raw(k)
        except Exception as e:  # noqa
            print('  mega 取得失敗', k, e)
            continue
        mega_all[k] = {ym: v * dx[k] for ym, v in r.items()}
    vw_every = vw_all_every('usa')
    reg_every = {r: vw_all_every(r) for r in REGIONS}
    out['x4_inputs'] = {'mega_chars': len(mega_all), 'vw_chars': len(vw_every), 'directions_known': sum(1 for v in dx.values() if v in (1, -1))}

    def select(sd, thr):
        sel = []
        for k, s in sd.items():
            x = [v for m, v in sorted(s.items()) if m <= M.TRAIN_END]
            if len(x) < 180:
                continue
            t = M.nw_t(x)
            if S.mean(x) > 0 and t is not None and t >= thr:
                sel.append(k)
        return sorted(sel)
    X4 = []
    out['x4_selected'] = {}
    for src_name, sd in (('mega', mega_all), ('vw', vw_every)):
        for thr in (3.0, 2.0):
            sel = select(sd, thr)
            out['x4_selected'][f'{src_name}_t{thr}'] = sel
            ls = composite(sd, sel)
            regls = {r: composite(reg_every[r], [c for c in sel]) for r in REGIONS}
            rp = repl_test({r: v for r, v in regls.items() if v}, LAM, 0.8, reg_mkt)
            rec = evaluate(f'X4a_{src_name}_trainsel_t{thr:.0f}', 'X4', f'{src_name} の153特徴から訓練期間（〜2006）で平均>0 かつ t≥{thr} の {len(sel)} 本を等分に合成', mkt, ls, LAM, 0.8, None, rp,
                           extra={'n_selected': len(sel)})
            X4.append(rec)
            ctx['ls_of'][rec['name']] = (ls, LAM)
    cl = cluster_labels()
    th_vw = themes_vw('usa')
    reg_th = {r: themes_vw(r) for r in REGIONS}
    themes = sorted(set(cl.values()))
    out['x4_theme_members_mega'] = {}
    for th in themes:
        mem = [c for c, t in cl.items() if t == th and c in mega_all]
        out['x4_theme_members_mega'][th] = mem
        rp = repl_test({r: reg_th[r].get(th) for r in REGIONS if reg_th[r].get(th)}, LAM, THEME_TO[th], reg_mkt)
        ls = composite(mega_all, mem)
        rec = evaluate(f'X4b_mega_theme_{th}', 'X4', f'mega のテーマ「{th}」（{len(mem)} 特徴の順位加重 LS を等分）', mkt, ls, LAM, THEME_TO[th], None, rp)
        X4.append(rec)
        ctx['ls_of'][rec['name']] = (ls, LAM)
        ls2 = th_vw.get(th)
        rec = evaluate(f'X4b_vw_theme_{th}', 'X4', f'JKP 公開の vw テーマ「{th}」（米国・全規模・時価加重）', mkt, ls2, LAM, THEME_TO[th], None, rp)
        X4.append(rec)
        ctx['ls_of'][rec['name']] = (ls2, LAM)
    fams['X4'] = finish_family(X4)
    # ── 診断（判定に使わない）: P・X3・X4 の S/A
    for r in ctx['fams']['P']:
        nm = r['name']
        base = nm.split('_', 2)[2]
        src = ctx['mega'] if nm.startswith('PA_') else vwu
        ctx['ls_of'][nm] = (build(src, base), LAM)
    for f in ('P', 'X3', 'X4'):
        for r in fams[f]:
            if r.get('grade') in ('S', 'A') and r['name'] in ctx['ls_of']:
                ls, lam = ctx['ls_of'][r['name']]
                r['diagnostics'] = diagnostics(mkt, mktrf, ls, lam)
    out['tested'] = []
    for f, recs in fams.items():
        out['tested'].extend(recs)
    out['n_tested'] = len(out['tested'])
    out['families'] = {f: {'n': len(r), 'grades': {g: sum(1 for x in r if x.get('grade') == g) for g in 'SABC'}} for f, r in fams.items()}


def show(out):
    print('\n族  名前                         格  全期間(t)      訓練(t)       保有(t)       保有費用後  最近   20年勝率  C5')
    for r in out['tested']:
        if r.get('error'):
            print(r['family'], r['name'], 'ERR', r['error'])
            continue
        f = lambda s: f"{s['ex_ann']:+5.2f}({s['t']:+4.1f})" if s else '   —      '
        rp = r.get('repl') or {}
        print(f"{r['family']:10} {r['name'][:34]:34} {r['grade']}  {f(r['full'])}  {f(r['train'])}  {f(r['hold'])}  "
              f"{r['net_cost_hold']['ex_ann'] if r['net_cost_hold'] else '—':>6}  {r['recent']['ex_ann'] if r['recent'] else '—':>5}  "
              f"{r['roll20_net']['win_rate'] if r['roll20_net'] else '—':>5}  {rp.get('positive')}/{rp.get('regions')}")


def main():
    if '--part3-only' in sys.argv or '--part4-only' in sys.argv:
        out = json.load(open(os.path.join(BASE, 'out', 'mw_mega_tilt.json')))
        if '--part3-only' in sys.argv:
            part3(out)
        part4(out)
    else:
        out, ctx = part1()
        if '--part1-only' not in sys.argv:
            part2(out, ctx)
            part3(out)
            part4(out)
    summarize(out)
    p = M.save('mw_mega_tilt.json', out)
    show(out)
    print('書いた', p, '試した数', out['n_tested'])


# ═════════════════════════ prereg3（頑丈さ） ═════════════════════════
PREREG3 = 'mw_mega_tilt_prereg3.json'


def ols_noint(y, X):
    """定数項なしの最小二乗（列は少数）→ 係数のリスト"""
    k = len(X[0])
    A = [[sum(r[i] * r[j] for r in X) for j in range(k)] for i in range(k)]
    b = [sum(r[i] * yy for r, yy in zip(X, y)) for i in range(k)]
    # ガウスの消去
    M_ = [A[i] + [b[i]] for i in range(k)]
    for c in range(k):
        piv = max(range(c, k), key=lambda r: abs(M_[r][c]))
        M_[c], M_[piv] = M_[piv], M_[c]
        for r in range(k):
            if r != c:
                f = M_[r][c] / M_[c][c]
                M_[r] = [a - f * bb for a, bb in zip(M_[r], M_[c])]
    return [M_[i][k] / M_[i][i] for i in range(k)]


def part3(out):
    out['prereg3'] = PREREG3
    out['prereg3_commit'] = sha_of(f'out/{PREREG3}')
    out['tested'] = [r for r in out['tested'] if r.get('family') != 'X5']
    ff = M.ff_factors()
    mkt, mktrf, rf = ff['mkt'], ff['mktrf'], ff['rf']
    names = list(CH) + ['Q5', 'QM', 'QMV']

    def to_of(nm):
        if nm in CH:
            return CH[nm]['turnover']
        c, w = blend_def(nm)
        return turnover_of(c, w)

    def pub_of(nm):
        if nm in CH:
            return CH[nm]['pub']
        c, _ = blend_def(nm)
        return max(CH[x]['pub'] for x in c)
    reg_vw = {r: vw_all(r) for r in REGIONS}
    reg_every = {r: vw_all_every(r) for r in REGIONS}
    reg_mkt = {r: M.jkp_mkt(r, 'vw') for r in REGIONS}
    vwu = vw_all('usa')
    dx = directions_from_xlsx()
    X5, ls_of = [], {}
    # ── X5a large
    large = {}
    for k in CH:
        r, _ = mega_raw(k, 'large')
        large[k] = {ym: v * CH[k]['direction'] for ym, v in r.items()}
    for nm in names:
        regls = {r: build(reg_vw[r], nm) for r in REGIONS}
        rp = repl_test({r: v for r, v in regls.items() if v}, LAM, to_of(nm), reg_mkt)
        desc = CH[nm]['意味'] if nm in CH else PR['blends'][nm]
        ls = build(large, nm)
        rec = evaluate(f'X5a_large_{nm}', 'X5', f'米国 large（NYSE 50〜80%点）の順位加重 LS: {desc}', mkt, ls, LAM, to_of(nm), pub_of(nm), rp)
        X5.append(rec)
        ls_of[rec['name']] = (ls, LAM)
    large_all = {}
    for key in ('mega_t3.0', 'vw_t3.0'):
        sel = out['x4_selected'][key]
        for k in sel:
            if k not in large_all:
                r, _ = mega_raw(k, 'large')
                large_all[k] = {ym: v * dx[k] for ym, v in r.items()}
        ls = composite(large_all, sel)
        regls = {r: composite(reg_every[r], sel) for r in REGIONS}
        rp = repl_test({r: v for r, v in regls.items() if v}, LAM, 0.8, reg_mkt)
        rec = evaluate(f'X5a_large_trainsel_{key}', 'X5', f'prereg2 で {key} から選んだ {len(sel)} 特徴を選び直さずに large に当てた合成', mkt, ls, LAM, 0.8, None, rp,
                       extra={'n_selected': len(sel)})
        X5.append(rec)
        ls_of[rec['name']] = (ls, LAM)
    # ── X5b 方法の検算（French 単変量 OP 三分位）
    fv = M.french_series('Portfolios_Formed_on_OP', 'Value Weight')
    ks = sorted(k for k in fv['Lo 30'] if k in fv['Med 40'] and k in fv['Hi 30'] and k in mktrf and k in rf and k <= M.TRAIN_END)
    co = ols_noint([mktrf[k] for k in ks], [[fv[c][k] - rf[k] for c in ('Lo 30', 'Med 40', 'Hi 30')] for k in ks])
    capd = fr_capshare('Portfolios_Formed_on_OP', ['Lo 30'])
    capall = fr_capshare('Portfolios_Formed_on_OP', ['Lo 30', 'Med 40', 'Hi 30'])
    tot = total_cap()
    act_in = S.mean(capd[k] / capall[k] for k in ks if k in capd and k in capall)
    act_tot = S.mean(capd[k] / tot[k] for k in ks if k in capd and k in tot)
    out['feasibility']['x5b_method_check_french_OP'] = {
        'months': len(ks), 'coef_lo30_med40_hi30': [round(c, 3) for c in co], 'coef_sum': round(sum(co), 3),
        'actual_capshare_lo30_within_op_universe': round(act_in, 3), 'actual_capshare_lo30_of_total_market': round(act_tot, 3),
        'abs_diff_vs_total': round(abs(co[0] - act_tot), 3), 'method_ok': abs(co[0] - act_tot) <= 0.05}
    print('X5b 方法の検算', out['feasibility']['x5b_method_check_french_OP'])
    jm = M.jkp_mkt('usa', 'vw')
    lam_char, lam_info = {}, {}
    for k in CH:
        p = M.jkp_portfolios('usa', k, 'vw')
        ks = sorted(m for m in set(p.get('1.0', {})) & set(p.get('2.0', {})) & set(p.get('3.0', {})) & set(jm) if m <= M.TRAIN_END)
        co = ols_noint([jm[m] for m in ks], [[p['1.0'][m], p['2.0'][m], p['3.0'][m]] for m in ks])
        badw = co[0] if CH[k]['direction'] == 1 else co[2]
        ok = 0.9 <= sum(co) <= 1.1 and badw > 0
        lam_info[k] = {'coef_p1_p2_p3': [round(c, 3) for c in co], 'sum': round(sum(co), 3), 'bad_side_weight': round(badw, 3), 'months': len(ks), 'ok': ok}
        if ok:
            lam_char[k] = max(0.01, math.floor(0.8 * badw * 100) / 100)
            lam_info[k]['lambda_char'] = lam_char[k]
    for nm in ('Q5', 'QM', 'QMV'):
        comps, _ = blend_def(nm)
        if all(c in lam_char for c in comps):
            lam_char[nm] = min(lam_char[c] for c in comps)
    out['feasibility']['x5b_lambda_char'] = lam_info
    out['feasibility']['x5b_lambda_char_blends'] = {nm: lam_char.get(nm) for nm in ('Q5', 'QM', 'QMV')}
    print('X5b λ_char', {k: v.get('lambda_char') for k, v in lam_info.items()}, out['feasibility']['x5b_lambda_char_blends'])
    for nm in names:
        if nm not in lam_char:
            X5.append({'name': f'X5b_vw_{nm}_lamchar', 'family': 'X5', 'description': 'λ_char 推定不能（係数の和が0.9〜1.1の外）', 'error': 'λ_char 推定不能', 'grade': 'C'})
            continue
        lam = lam_char[nm]
        regls = {r: build(reg_vw[r], nm) for r in REGIONS}
        rp = repl_test({r: v for r, v in regls.items() if v}, lam, to_of(nm), reg_mkt)
        ls = build(vwu, nm)
        rec = evaluate(f'X5b_vw_{nm}_lamchar', 'X5', f'P-B と同じ・λ_char={lam}（悪い側の時価総額の推定×0.8＝買いだけに収まる見込みの大きさ）', mkt, ls, lam, to_of(nm), pub_of(nm), rp)
        X5.append(rec)
        ls_of[rec['name']] = (ls, lam)
    X5 = finish_family(X5)
    if not out['feasibility']['x5b_method_check_french_OP']['method_ok']:
        for r in X5:
            if r['name'].startswith('X5b_'):
                r['warning'] = '方法の検算に落ちた（French OP で回帰の係数が実際の時価総額の割合と 0.05 超ずれた）ので λ_char は信用できない（prereg3 の規則）'
    for r in X5:
        if r.get('grade') in ('S', 'A') and r['name'] in ls_of:
            ls, lam = ls_of[r['name']]
            r['diagnostics'] = diagnostics(mkt, mktrf, ls, lam)
    out['tested'].extend(X5)
    out['n_tested'] = len(out['tested'])
    fams = {}
    for r in out['tested']:
        fams.setdefault(r['family'], []).append(r)
    out['families'] = {f: {'n': len(r), 'grades': {g: sum(1 for x in r if x.get('grade') == g) for g in 'SABC'}} for f, r in fams.items()}


# ═════════════════════════ prereg4（2006年に知り得た特徴だけ） ═════════════════════════
PREREG4 = 'mw_mega_tilt_prereg4.json'


def pub_years():
    import openpyxl
    import re
    wb = openpyxl.load_workbook(os.path.join(M.CACHE, 'jkp_factor_details.xlsx'), read_only=True)
    rows = list(wb['details'].iter_rows(values_only=True))
    h = rows[0]
    ia, ic = h.index('abr_jkp'), h.index('cite')
    out = {}
    for r in rows[1:]:
        if r[ia]:
            m = re.findall(r'(19\d\d|20\d\d)', str(r[ic] or ''))
            if m:
                out[r[ia]] = int(m[-1])
    return out


def part4(out):
    out['prereg4'] = PREREG4
    out['prereg4_commit'] = sha_of(f'out/{PREREG4}')
    out['tested'] = [r for r in out['tested'] if r.get('family') != 'X6']
    ff = M.ff_factors()
    mkt, mktrf = ff['mkt'], ff['mktrf']
    dx = directions_from_xlsx()
    py = pub_years()
    avail = json.load(open(os.path.join(M.CACHE, 'jkp_availability.json')))['factor_sizes']['usa']
    cl = cluster_labels()
    pre = sorted(k for k in avail if py.get(k) is not None and py[k] <= 2006 and dx.get(k) in (1, -1))
    mega_all = {}
    for k in pre:
        r, _ = mega_raw(k)
        mega_all[k] = {ym: v * dx[k] for ym, v in r.items()}
    vw_every = {k: v for k, v in vw_all_every('usa').items() if k in pre}
    reg_every = {r: vw_all_every(r) for r in REGIONS}
    reg_mkt = {r: M.jkp_mkt(r, 'vw') for r in REGIONS}
    out['x6_inputs'] = {'pub_le_2006': len(pre), 'mega': len(mega_all), 'vw': len(vw_every)}

    def select(sd, thr):
        sel = []
        for k, s in sd.items():
            x = [v for m, v in sorted(s.items()) if m <= M.TRAIN_END]
            if len(x) < 180:
                continue
            t = M.nw_t(x)
            if S.mean(x) > 0 and t is not None and t >= thr:
                sel.append(k)
        return sorted(sel)
    X6, ls_of = [], {}
    out['x6_selected'] = {}
    qual06 = sorted(k for k in pre if cl.get(k) in ('quality', 'profitability'))
    out['x6_selected']['quality_profitability_pub_le_2006'] = qual06
    specs = []
    for src, sd in (('mega', mega_all), ('vw', vw_every)):
        for thr in (3.0, 2.0):
            sel = select(sd, thr)
            out['x6_selected'][f'{src}_t{thr}'] = sel
            specs.append((f'X6a_{src}_pre2006_trainsel_t{thr:.0f}', sd, sel, 0.8, f'{src}: 2006年以前に公表された特徴から訓練期間で t≥{thr} の {len(sel)} 本を等分'))
        specs.append((f'X6b_{src}_quality_pre2006', sd, qual06, 0.6, f'{src}: 2006年以前に公表された質・収益性の {len(qual06)} 本を選ばずに等分'))
        specs.append((f'X6c_{src}_all_pre2006', sd, sorted(sd), 1.0, f'{src}: 2006年以前に公表された {len(sd)} 本すべてを等分'))
    for name, sd, sel, to, desc in specs:
        ls = composite(sd, sel)
        regls = {r: composite(reg_every[r], sel) for r in REGIONS}
        rp = repl_test({r: v for r, v in regls.items() if v}, LAM, to, reg_mkt)
        rec = evaluate(name, 'X6', desc, mkt, ls, LAM, to, None, rp, extra={'n_selected': len(sel)})
        X6.append(rec)
        ls_of[name] = (ls, LAM)
    X6 = finish_family(X6)
    for r in X6:
        if r.get('grade') in ('S', 'A') and r['name'] in ls_of:
            ls, lam = ls_of[r['name']]
            r['diagnostics'] = diagnostics(mkt, mktrf, ls, lam)
    # D5 期間をそろえた訓練（判定に使わない）
    mg, _ = mega_raw('ope_be')
    vwu = vw_all('usa')
    d5 = {}
    for nm, ls in (('PA_mega_ope_be', {k: v * CH['ope_be']['direction'] for k, v in mg.items()}), ('PB_vw_ope_be', vwu['ope_be']),
                   ('F2_ope_be', fr_ls('6_Portfolios_ME_OP_2x3', 'BIG HiOP', 'BIG LoOP'))):
        s = {k: mkt[k] + LAM * ls[k] for k in ls if k in mkt and k <= JKP_END}
        b = {k: mkt[k] for k in s}
        d5[nm] = {'1963-07_2006-12': M.excess_stats(s, b, a=196307, z=M.TRAIN_END), 'start_1962-12': M.excess_stats(s, b, z=196212)}
    out['diagnostic_D5_train_aligned'] = d5
    print('D5', json.dumps({k: {kk: (vv['ex_ann'], vv['t']) if vv else None for kk, vv in v.items()} for k, v in d5.items()}, ensure_ascii=False))
    out['tested'].extend(X6)
    out['n_tested'] = len(out['tested'])
    fams = {}
    for r in out['tested']:
        fams.setdefault(r['family'], []).append(r)
    out['families'] = {f: {'n': len(r), 'grades': {g: sum(1 for x in r if x.get('grade') == g) for g in 'SABC'}} for f, r in fams.items()}


def summarize(out):
    """結果の要約（数えるだけ）と、結果を見た後に書いた解釈（判定には使わない）"""
    T = [r for r in out['tested'] if not r.get('error')]
    key = lambda r: (r['hold'] or {}).get('t') or -9
    out['summary'] = {
        'n_tested': len(out['tested']),
        'grades_by_family': out['families'],
        'S_primary': [r['name'] for r in T if r['family'] == 'P' and r['grade'] == 'S'],
        'top_primary_by_hold_t': [[r['name'], r['grade'], r['hold']['ex_ann'], r['hold']['t'], r['net_cost_hold']['ex_ann'], r.get('family_holm_p')]
                                  for r in sorted([x for x in T if x['family'] == 'P'], key=key, reverse=True)[:8]],
        'top_all_by_hold_t': [[r['name'], r['grade'], r['hold']['ex_ann'], r['hold']['t']] for r in sorted(T, key=key, reverse=True)[:12]],
    }
    out['interpretation_ja_written_after_results'] = [
        '主の族（米国・λ=0.2）で S が10本。どれも『質・収益性』（cop_at・ope_be・qmj・qmj_prof・Q5・QM）か発生主義（oaccruals_at）。割安（be_me・ni_me）・低ベータ・固有のぶれ・投資は2007年以降に負け（C）',
        '最も強い1本は vw の cop_at（現金ベースの営業利益÷総資産）: 保有期間 +1.50%/年 t3.3（族の Holm p=0.037）・費用後 +1.49・前後半とも正（+1.54/+1.47）・2008-09を除いても t3.2・公表後（2017〜）+1.74 t2.8。mega（順位加重 +0.95 t2.4）と large（別の銘柄群 +0.86 t3.2）でも S＝同じ特徴が3つの作り方で勝った',
        '★大きさの注意: λ=0.2 は厳密な買いだけでは作れない。mega の順位加重は約6%の売りが要り（Pareto 推定）、買いだけの上限は λ≈0.04（→ 保有 +0.19%/年）。vw の収益性は悪い側の時価総額が約16%（French OP）で、買いだけに収まる見込みの λ≈0.11 では cop_at +0.83%/年 t3.3・Q5 +0.68 t2.8',
        '★後知恵の漏れ: 勝った質・収益性の特徴はどれも2006年より後に公表された（gp_at 2013・ope_be 2015・cop_at 2016・qmj 2018）。2006年に知り得た特徴だけで訓練期間から選んだ合成（X6a）は保有 +0.25%/年 t2.0（S・漏れの無い形で最も強い）、2006年に知られていた質の等分（X6b）は +0.27 t1.4（B）',
        '独立の作り方（French の BIG/ME5・時価加重）では OP の訓練期間 t が 1.4/0.8 で C1 に届かない（保有 +0.93 t2.2 は同じ向き）。JKP の ope_be も訓練を 1963-2006 に揃えると mega t1.66・vw t2.05＝1950年代のデータに一部頼っている',
        '日本では再現しない: vw cop_at −0.10%/年・Q5 +0.08 t0.25（米国外10地域のうち日本だけが弱い）',
        '実在の質ETF（QUAL・SPHQ 等）は設定来 S&P500 に負けている（CLAUDE.md・mw_factor_us）。ここで勝ったのは『市場の時価の重みを保ったまま少しだけ傾ける』形で、上位だけを持つ ETF とは別物',
    ]


if __name__ == '__main__':
    main()
