#!/usr/bin/env python3
"""night/nx_leadlag_data.py — nx 角度 leadlag（業種どうしの先行・遅行）の**データの取得・整形だけ**（成績は計算しない・読むだけ・門の判定には不使用）

2026-09-28 ユーザー指示「市場に勝てる歴史検証が出るまでいろんな角度から調べて勝てる結果を出して…別のSessionで検証していない新たな分析を」。
事前登録: out/nx_leadlag_prereg.json（測る前に固定）／全体の線: out/nx_prereg.json（C1〜C8・格付け）。

何をするか
  BEA（米商務省 経済分析局）の**ベンチマーク産業連関表（Use 表・生産者価格）**を公表された年ごとに読み、
  Ken French の49業種（SIC 定義: Siccodes49）へ写して、49×49 の「業種→業種の中間財の流れ」行列 X[売り手][買い手] を作る。
  そこから Menzly & Ozbas (2010) の
     仕入れ先の重み s[G][F] = X[F][G] / Σ_{F'≠G} X[F'][G]   （業種 G が買う中間財のうち F から来る割合・自分は除く）
     顧客の重み   c[F][G] = X[F][G] / Σ_{G'≠F} X[F][G']   （業種 F が売る中間財のうち G へ行く割合・自分は除く）
  を表ごとに出し、いつから使えるか（公表の翌月末から）と一緒に out/_nx_cache/nx_leadlag_io.json に置く（gitignore）。
  C5（米国外での再現）と実在のセクターETFの答え合わせ用に、同じ X を GICS 11セクターへ畳んだ行列も置く（FR2GICS・測る前に固定）。
  **業種のリターンは一切読まない**（この道具は成績を計算しない）。

写し方（測る前に固定）
  1. できるだけ細かい表を使う（部門の大きさが流れの金額で効くように）:
       SIC 時代 = 1963（367部門）・1967（484・1977年改訂版）・1972（496）・1977（537）・1982（6桁）・1987（6桁）・1992（6桁）
       NAICS 時代 = 1997・2002（改訂版）・2007（改訂版）・2017（詳細 約400部門）
     1947・1958 は85部門しか無い（探索の「早い延長」でだけ使う・下の 3.）。
  2. 部門 → SIC: SIC 時代は各年の Sectoring Plan / SIC-IO 対照（その年の SIC 版の数字）。1982年は対照が手元に無いので
     1977年の対照（同じ1977年 SIC 基準）→ 無ければ1987年の対照（同じ6桁の番号）。NAICS 時代は表に付いた「関連 NAICS」→
     国勢調査局の対照表で 1997年 NAICS まで遡る → 1997年 NAICS → 1987年 SIC。BEA 部門の中の NAICS 6桁どうしは
     1997年経済センサスの全国の売上で割る（秘匿・0 は『分からない』として同じ部門の中央値。NAICS→SIC の1対多は等分）。
     SIC の字句は 2〜4桁の接頭辞として 1987年 SIC の有効な4桁コードへ展開する。有効コードに無い4桁（古い版の番号）は
     その数字のまま French の範囲に当てる。それでも写らない部門は、2桁の親部門の辞書（Menzly & Ozbas 2010 Internet
     Appendix Table IA.VIII の BEA 部門→SIC 辞書を土台に枝番と分かれた番号を足した IO_SIC）へ倒す。
  3. 85部門しか無い 1947・1958 年は IO_SIC 辞書で写す（1947・1958 は 74＝研究開発・飲食は 69 の中）。
  4. SIC 4桁 → French の49業種。1つの部門が複数の French 業種にまたがるときは、有効な SIC 4桁コードの数で等分
     （各段の 1対多 も等分）。どの French 業種にも入らない SIC（政府・郵便など）は落として残りで割り直す。
  5. 同じ部門どうしの取引（表の対角）は写す前に除く（部門が複数の French 業種にまたがると、対角の流れが見かけの
     業種間の結びつきになるため）。French の段の自分との取引も重みでは除く（Menzly & Ozbas と同じ）。
  6. 持ち家（owner-occupied dwellings）・政府・政府企業（78,79）・輸入・スクラップ・家計・最終需要・付加価値は入れない。

使い方: python3 night/nx_leadlag_data.py            → out/_nx_cache/nx_leadlag_io.json（sha256 を表示）
        python3 night/nx_leadlag_data.py --show     → 表ごとの被覆と、いくつかの業種の主な仕入れ先・顧客（産業連関の事実だけ）
"""
import sys, os, io, re, json, zipfile, hashlib, datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import nx_common as N  # noqa: E402

OUT = os.path.join(N.CACHE, 'nx_leadlag_io.json')
BEA = 'https://apps.bea.gov/industry/'
FILES = {
    # 85部門（探索の早い延長だけ）と 2桁表（検算用）
    '1947': ('zip/47IOexcel.zip', 'bea_47IOexcel.zip'),
    '1958': ('zip/58IOexcel.zip', 'bea_58IOexcel.zip'),
    '1992s': ('zip/ndn0180.zip', 'bea_92_2digit.zip'),
    # SIC 時代の詳細表と部門→SIC の対照（BEA「Historical Benchmark Input-Output Tables」・原本の公表値で改訂されない）
    '1963': ('zip/63IO367-leveltext.zip', 'bea_63IO367text.zip'),
    '1967': ('zip/67IO484-levelexcel.zip', 'bea_67IO484excel.zip'),
    '1972': ('zip/72IO496-levelexcel.zip', 'bea_72IO496excel.zip'),
    '1977': ('zip/77IO537-leveltext.zip', 'bea_77IO537text.zip'),
    '1982': ('zip/ndn0025.zip', 'bea_82_6digit.zip'),
    '1987': ('zip/ndn0016.zip', 'bea_87_6digit.zip'),
    '1992': ('zip/ndn0178.zip', 'bea_92_6digit.zip'),
    # NAICS 時代
    '1997': ('zip/ndn0306.zip', 'bea_97_detail.zip'),
    '2002': ('zip/2002detail.zip', 'bea_02_detail.zip'),
    '2007': ('xls/io-annual/IOUse_Before_Redefinitions_PRO_2007_Detail.xlsx', 'bea_IOUse_Before_Redefinitions_PRO_2007_Detail.xlsx'),
    '2017': ('xls/io-annual/IOUse_Before_Redefinitions_PRO_2017_Detail.xlsx', 'bea_IOUse_Before_Redefinitions_PRO_2017_Detail.xlsx'),
}
CENSUS = 'https://www.census.gov/naics/concordances/'
CFILES = ['1997_NAICS_to_1987_SIC.xls', '1987_SIC_to_1997_NAICS.xls', '1997_NAICS_to_2002_NAICS.xls',
          '2002_to_2007_NAICS.xls', '2007_to_2012_NAICS.xls', '2012_to_2017_NAICS.xlsx']
SICCODES = 'https://mba.tuck.dartmouth.edu/pages/faculty/ken.french/ftp/Siccodes49.zip'
# 1997年経済センサス（全国・NAICS 6桁の売上・出荷・収入）。NAICS 時代の「BEA 部門 → NAICS 6桁」を等分でなく規模で割るため（2000年ごろ公表＝1997年表〔2002年公表〕より前）
EC97 = 'https://www2.census.gov/econ1997/EC/sector00/E9700A1.zip'

# ── いつから使えるか（測る前に固定）。signal = その月末の業種リターンで信号を作ってよい最初の月（公表月の翌月末）、hold = 最初に持つ月。
#    公表が「年」しか分からない詳細表は、その年の12月に公表とみなす（遅い側に倒す）
SCHEDULE = [
    {'table': '1963', 'published': '1969（OBE の詳細表3巻・月は不明→12月とみなす。要約版は SCB 1969年11月号）', 'signal': 197001, 'hold': 197002},
    {'table': '1967', 'published': '1977-06（手元の484部門表は1974年2月号の当初版を NIPA に合わせて直した改訂版・BEA Staff Paper No.29）', 'signal': 197707, 'hold': 197708},
    {'table': '1972', 'published': '1979（SCB 1979年2月号に要約・詳細表は1979年の刊行物・月は不明→12月とみなす）', 'signal': 198001, 'hold': 198002},
    {'table': '1977', 'published': '1984（SCB 1984年5月号に要約・詳細表「The Detailed Input-Output Structure of the U.S. Economy: 1977」は1984年・月は不明→12月）', 'signal': 198501, 'hold': 198502},
    {'table': '1982', 'published': '1991-10-02（6桁表の書式ファイルの日付。要約は SCB 1991年7月号）', 'signal': 199111, 'hold': 199112},
    {'table': '1987', 'published': '1994-04-30（6桁表の書式ファイルの日付。SCB 1994年4月号）', 'signal': 199405, 'hold': 199406},
    {'table': '1992', 'published': '1997-11-01（6桁表の README の日付。SCB 1997年11月号）', 'signal': 199712, 'hold': 199801},
    {'table': '1997', 'published': '2002-12-11（README の日付。SCB 2002年12月号）', 'signal': 200301, 'hold': 200302},
    {'table': '2002', 'published': '2008-04-24（手元の表は改訂版 REV 4-24-08。当初版は2007年秋。改訂版の日付を採る）', 'signal': 200805, 'hold': 200806},
    {'table': '2007', 'published': '2014-11-13（手元の表の ReadMe の公表日・改訂版。当初版は2013年12月）', 'signal': 201412, 'hold': 201501},
    {'table': '2017', 'published': '2024-05-20（手元のファイルの最終更新日。2017年表の当初の公表はこれより前だが、手元の版の日付を採る）', 'signal': 202406, 'hold': 202407},
]
# 探索の「早い延長」（E_mo_early）だけが使う85部門の表
SCHEDULE_EARLY = [
    {'table': '1947', 'published': '1951（BLS の当初版）／手元の85部門の並べ直しは1970年ごろ OBE＝様式の後知恵', 'signal': 195201, 'hold': 195202},
    {'table': '1958', 'published': '1965-09（SCB 1965年9月号に係数表。取引表の初出は1964年11月号とされるが手元の書類が名指しする9月号を採る）', 'signal': 196510, 'hold': 196511},
]
NOT_SCHEDULED = {'2012': '当初版（2018年ごろ）が手元に無い。2017年ファイル内の 2012 は2024年の改訂版＝版の後知恵なので時点の表に入れない（2007年表が2024-06まで続く）'}

FR49 = ['Agric', 'Food', 'Soda', 'Beer', 'Smoke', 'Toys', 'Fun', 'Books', 'Hshld', 'Clths', 'Hlth', 'MedEq', 'Drugs', 'Chems', 'Rubbr',
        'Txtls', 'BldMt', 'Cnstr', 'Steel', 'FabPr', 'Mach', 'ElcEq', 'Autos', 'Aero', 'Ships', 'Guns', 'Gold', 'Mines', 'Coal', 'Oil',
        'Util', 'Telcm', 'PerSv', 'BusSv', 'Hardw', 'Softw', 'Chips', 'LabEq', 'Paper', 'Boxes', 'Trans', 'Whlsl', 'Rtail', 'Meals',
        'Banks', 'Insur', 'RlEst', 'Fin', 'Other']

# ── French 49 → GICS 11 セクター（C5 の他国・実在のセクターETF用。1業種1セクター・測る前に固定。'Other' は入れない）
FR2GICS = {'Agric': '30', 'Food': '30', 'Soda': '30', 'Beer': '30', 'Smoke': '30', 'Toys': '25', 'Fun': '50', 'Books': '50', 'Hshld': '30',
           'Clths': '25', 'Hlth': '35', 'MedEq': '35', 'Drugs': '35', 'Chems': '15', 'Rubbr': '15', 'Txtls': '25', 'BldMt': '15',
           'Cnstr': '20', 'Steel': '15', 'FabPr': '20', 'Mach': '20', 'ElcEq': '20', 'Autos': '25', 'Aero': '20', 'Ships': '20',
           'Guns': '20', 'Gold': '15', 'Mines': '15', 'Coal': '10', 'Oil': '10', 'Util': '55', 'Telcm': '50', 'PerSv': '25',
           'BusSv': '20', 'Hardw': '45', 'Softw': '45', 'Chips': '45', 'LabEq': '45', 'Paper': '15', 'Boxes': '15', 'Trans': '20',
           'Whlsl': '20', 'Rtail': '25', 'Meals': '25', 'Banks': '40', 'Insur': '40', 'RlEst': '60', 'Fin': '40'}
GICS11 = ['10', '15', '20', '25', '30', '35', '40', '45', '50', '55', '60']

# ── 2桁（85部門）の BEA 部門 → 1987年 SIC の範囲（Menzly & Ozbas 2010 Internet Appendix Table IA.VIII を土台）
R = lambda *p: [tuple(x) for x in p]  # noqa: E731
IO_SIC = {
    '1+2': R((100, 299)), '1': R((200, 299)), '2': R((100, 199)),
    '3': R((800, 849), (860, 919), (930, 999)),
    '4': R((700, 739), (750, 799), (850, 859), (920, 929)),
    '5+6': R((1000, 1079), (1090, 1099)), '5': R((1010, 1019)), '6': R((1000, 1009), (1020, 1079), (1090, 1099)),
    '7': R((1200, 1239), (1250, 1299)),
    '8': R((1300, 1379), (1390, 1399)),
    '9+10': R((1400, 1479), (1490, 1499)), '9': R((1400, 1469), (1490, 1499)), '10': R((1470, 1479)),
    '11': R((1080, 1089), (1240, 1249), (1380, 1389), (1480, 1489), (1500, 1799), (6550, 6559)),
    '12': R((1080, 1089), (1240, 1249), (1380, 1389), (1480, 1489), (1500, 1799), (6550, 6559)),
    '13': R((3480, 3489), (3761, 3761), (3795, 3795)),
    '14': R((2000, 2099), (5460, 5469)),
    '15': R((2100, 2199)),
    '16': R((2200, 2249), (2260, 2269), (2280, 2289)),
    '17': R((2270, 2279), (2290, 2299)),
    '18': R((2250, 2259), (2300, 2389)),
    '19': R((2390, 2399)),
    '20': R((2400, 2439), (2450, 2499)), '21': R((2440, 2449)),
    '22': R((2500, 2519)), '23': R((2520, 2599)),
    '24': R((2600, 2649), (2660, 2699)),
    '25': R((2650, 2659)),
    '26': R((2700, 2799)),
    '27': R((2800, 2819), (2860, 2899)),
    '28': R((2820, 2829)),
    '29': R((2830, 2849)),
    '30': R((2850, 2859)),
    '31': R((2900, 2999)),
    '32': R((3000, 3099)),
    '33': R((3110, 3119)), '34': R((3100, 3109), (3120, 3199)),
    '35': R((3200, 3229)),
    '36': R((3230, 3299)),
    '37': R((3300, 3329), (3390, 3399), (3462, 3462)),
    '38': R((3330, 3389), (3463, 3463)),          # IA.VIII は 3460-3469 を 38 と 41 の両方に書いていた → 鍛造(3463)だけ 38、打抜き(3465-3469)は 41
    '39': R((3400, 3419)),
    '40': R((3430, 3449)),
    '41': R((3450, 3459), (3465, 3469)),
    '42': R((3420, 3429), (3470, 3479), (3490, 3499)),
    '43': R((3500, 3519)),
    '44': R((3520, 3529)), '45': R((3530, 3533)),
    '46': R((3534, 3539)),
    '47': R((3540, 3549)),
    '48': R((3550, 3559)),
    '49': R((3560, 3569)),
    '50': R((3590, 3599)),
    '51': R((3570, 3579)),
    '52': R((3580, 3589)),
    '53': R((3600, 3629)),
    '54': R((3630, 3639)),
    '55': R((3640, 3649)),
    '56': R((3650, 3669)),
    '57': R((3670, 3679)),
    '58': R((3680, 3699)),
    '59': R((3700, 3715), (3717, 3719)),
    '60': R((3720, 3729), (3760, 3760), (3762, 3769)),
    '61': R((3716, 3716), (3730, 3759), (3770, 3794), (3796, 3799)),
    '62': R((3800, 3849)),
    '63': R((3850, 3899)),
    '64': R((3900, 3999)),
    '65': R((4000, 4299), (4400, 4799)),
    '66': R((4800, 4829), (4840, 4899)),
    '67': R((4830, 4839)),
    '68': R((4900, 4999)),
    '69': R((5000, 5459), (5470, 5799), (5900, 5999)),
    '70': R((6000, 6499), (6700, 6731), (6733, 6799)),
    '71': R((6500, 6549), (6560, 6599)),
    '72': R((7000, 7099), (7200, 7299), (7600, 7689)),
    '73': R((7300, 7399), (7690, 7699), (8100, 8199), (8700, 8732), (8734, 8799)),
    '74': R((5800, 5899)),
    '75': R((7500, 7599)),
    '76': R((7800, 7999)),
    '77': R((740, 749), (6732, 6732), (8000, 8099), (8200, 8499), (8600, 8699), (8733, 8733), (8800, 8999)),
}
# 年による違い（書類で確かめた）: 1947・1958 は 74＝研究開発・飲食は 69 の中
ERA_OVERRIDE = {'1947': {'69': R((5000, 5999)), '74': R((8731, 8733))}, '1958': {'69': R((5000, 5999)), '74': R((8731, 8733))}}
EXCLUDE_PARENT = {'78', '79'}   # 政府企業。80 以上（輸入・スクラップ・政府・家計・最終需要・付加価値）も入れない
OWNER = re.compile(r'owner[- ]occupied', re.I)


def norm_io(c):
    c = str(c).strip().strip('"').strip()
    if re.fullmatch(r'\d+\.0', c):
        c = c[:-2]
    return c


# ───────────────────────── 取得 ─────────────────────────
def fetch_all():
    for y, (p, name) in FILES.items():
        b = N.get(BEA + p, name=name, max_age_days=3650)
        if b[:2] != b'PK':
            raise RuntimeError(f'{y}: 取れたものが zip/xlsx ではない（{b[:40]!r}）')
    for f in CFILES:
        b = N.get(CENSUS + f, name='census_' + f, max_age_days=3650)
        if b[:2] != b'PK' and b[:4] != bytes.fromhex('d0cf11e0'):
            raise RuntimeError(f'{f}: 取れたものが表計算の形ではない')
    N.get(SICCODES, name='fr_Siccodes49.zip', max_age_days=3650)
    N.get(EC97, name='census_E9700A1.zip', max_age_days=3650)


def cache(name):
    return open(os.path.join(N.CACHE, name), 'rb').read()


def zmember(zname, member, inner=None):
    z = zipfile.ZipFile(io.BytesIO(cache(zname)))
    if inner:
        z = zipfile.ZipFile(io.BytesIO(z.read(inner)))
    return z.read(member)


# ───────────────────────── SIC と French 49 ─────────────────────────
def french_ranges():
    t = zipfile.ZipFile(io.BytesIO(cache('fr_Siccodes49.zip'))).read('Siccodes49.txt').decode('latin-1')
    blocks, cur = {}, None
    for line in t.splitlines():
        m = re.match(r'\s*\d+\s+(\S+)\s+', line)
        if m and not re.match(r'\s*\d{4}-\d{4}', line):
            cur = m.group(1); blocks[cur] = []; continue
        m = re.match(r'\s*(\d{4})-(\d{4})', line)
        if m and cur:
            blocks[cur].append((int(m.group(1)), int(m.group(2))))
    assert list(blocks) == FR49, list(blocks)
    return blocks


def valid_sic():
    import xlrd
    out = set()
    s = xlrd.open_workbook(file_contents=cache('census_1987_SIC_to_1997_NAICS.xls')).sheet_by_index(0)
    for i in range(1, s.nrows):
        v = str(s.cell_value(i, 0)).strip()
        if re.fullmatch(r'\d{4}', v):
            out.add(int(v))
    s = xlrd.open_workbook(file_contents=cache('census_1997_NAICS_to_1987_SIC.xls')).sheet_by_index(0)
    for i in range(1, s.nrows):
        try:
            out.add(int(float(s.cell_value(i, 3))))
        except (TypeError, ValueError):
            pass
    return out


class Mapper:
    def __init__(self):
        self.fr = french_ranges()
        self.valid = valid_sic()
        self._f = {}

    def fr_of(self, sic):
        if sic not in self._f:
            hit = [k for k, v in self.fr.items() if any(a <= sic <= b for a, b in v)]
            self._f[sic] = hit[0] if hit else None
        return self._f[sic]

    def share_ranges(self, ranges):
        """SIC の範囲 → {French 業種: 割合}（有効な4桁コードの数で等分・どこにも入らないコードは落として割り直す）"""
        cnt = {}
        for a, b in ranges:
            for s in range(a, b + 1):
                if s in self.valid:
                    f = self.fr_of(s)
                    if f:
                        cnt[f] = cnt.get(f, 0) + 1
        tot = sum(cnt.values())
        return {k: v / tot for k, v in cnt.items()} if tot else {}

    def share_sicw(self, sicw):
        cnt = {}
        for s, w in sicw.items():
            f = self.fr_of(s)
            if f:
                cnt[f] = cnt.get(f, 0) + w
        tot = sum(cnt.values())
        return {k: v / tot for k, v in cnt.items()} if tot else {}

    def share_tokens(self, toks):
        """Sectoring Plan の SIC 字句（'201'・'2041'・'2061-3'・'pt. 15'・'*0259'）→ {French: 割合}。
        接頭辞を 1987年 SIC の有効な4桁へ展開して等分。有効コードに無い4桁は数字のまま French に当てる"""
        codes = []
        for _, a, b in toks:
            prefs = [a]
            if b:
                head, lo = a[:len(a) - len(b)], int(a[len(a) - len(b):])
                prefs = [head + str(x).zfill(len(b)) for x in range(lo, int(b) + 1)]
            for p in prefs:
                if len(p) < 4:
                    codes += [s for s in self.valid if str(s).zfill(4)[:len(p)] == p]   # 2〜3桁の接頭辞 → 有効な4桁へ展開
                else:
                    codes.append(int(p))   # 4桁はそのまま（古い版の番号で有効リストに無くても、数字のまま French の範囲へ当てる）
        cnt = {}
        for s in codes:
            f = self.fr_of(s)
            if f:
                cnt[f] = cnt.get(f, 0) + 1
        tot = sum(cnt.values())
        return {k: v / tot for k, v in cnt.items()} if tot else {}


# ───────────────────────── Sectoring Plan（部門 → SIC の字句） ─────────────────────────
def _doc_text(b):
    return re.sub(rb'[^\x20-\x7e\n\r\t]', b' ', b).decode().replace('\r', '\n')


def plan_text(y):
    from striprtf.striprtf import rtf_to_text
    if y == '1963':
        return rtf_to_text(zmember('bea_63IO367text.zip', '1963 Sectoring Plan.rtf').decode('latin-1'))
    if y == '1967':
        return _doc_text(zmember('bea_67IO484excel.zip', '1967 Sectoring Plan.doc'))
    if y == '1972':
        return _doc_text(zmember('bea_72IO496excel.zip', '1972 Sectoring Plan.doc'))
    if y == '1977':
        return rtf_to_text(zmember('bea_77IO537text.zip', '1977 Sectoring Plan.rtf').decode('latin-1'))
    if y == '1987':
        return _doc_text(zmember('bea_87_6digit.zip', 'SIC-IO.DOC', inner='disk1.zip'))
    if y == '1992':
        return zmember('bea_92_6digit.zip', 'Sic-IO.txt', inner='disk1.zip').decode('latin-1')
    raise KeyError(y)


TOK = re.compile(r'(pt\.?\s*|\*)?(?<!\d)(\d{2,4})(?:-(\d{1,3}))?(?!\d)')


def parse_plan(y):
    """部門コード → (題名, SIC 字句の欄, [(pt, 接頭辞, 範囲の末尾)])"""
    t = plan_text(y)
    if y in ('1987', '1992'):
        hits = [(m.start(), m.end(), m.group(1).zfill(2) + m.group(2)) for m in re.finditer(r'(?<![\d.])(\d{1,2})\.(\d{4})(?=[ \t]{1,6}[A-Z])', t)]
    elif y == '1963':
        hits = [(m.start(), m.end(), m.group(1).zfill(2) + m.group(2)) for m in re.finditer(r'(?<![\d.])(\d{1,2})\.(\d{2})(?=[ \t]{1,6}[A-Z])', t)]
    else:
        hits = [(m.start(), m.end(), m.group(1)) for m in re.finditer(r'(\d{6})(?=[ \t]{1,6}[A-Z])', t)]
    out = {}
    for k, (a, b, code) in enumerate(hits):
        end = hits[k + 1][0] if k + 1 < len(hits) else len(t)
        seg = t[b:end]
        # 次の見出し（'9+10  Nonmetallic minerals mining:'・'29B   Cleaning and toilet preparations:'・'14  Food and kindred products:'）
        # が同じ行に紛れ込むことがあるので、見出しの形（1〜2桁の番号［+番号］［枝の英字］＋大文字で始まる語…コロン）から後ろを切る
        seg = re.split(r'(?<=\s)\d{1,2}(?:\+\d{1,2})?[A-Z]?\s+[A-Z][a-z][^:\d\n]{0,90}:', seg)[0]
        if re.search(r'-{2,}', seg):
            parts = re.split(r'-{2,}', seg)
            title, fld = parts[0], parts[-1]
        else:
            # 1987・1992 の対照は「14.0101 Meat packing plants  2011」＝題名の後の最初の数字の字句から SIC 欄
            m = re.search(r'(?<=\s)(?:pt\.\s*|\*)?\d{2,4}(?:-\d{1,3})?(?=[,\s]|$)', seg)
            title, fld = (seg[:m.start()], seg[m.start():]) if m else (seg, '')
        lines = []
        for ln in fld.split('\n'):
            s = ln.strip()
            if not s:
                continue
            if re.match(r'^\d{1,2}(?:\+\d{1,2})?[A-Z]?\s+[A-Z][a-z]', s) or (re.search(r'[A-Za-z]{4,}', s) and not re.match(r'^(pt\.|\*|\d|\()', s)):
                break                     # 次の見出し（'14 Food & Kindred products' や 'MANUFACTURING'）で止める
            lines.append(s)
        f = ' '.join(lines)
        f = re.sub(r'\((excl\.?|except)[^)]*\)', ' ', f, flags=re.I)   # 「除く」の中の数字は入れない
        f = re.split(r'[A-Za-z]{4,}', f)[0]   # 字句の後ろに続く説明文（'1. Although the SIC…'）は入れない
        out[code] = (' '.join(title.split())[:80], f.strip()[:160], TOK.findall(f))
    return out


# ───────────────────────── 表の読み手（{(売り手部門, 買い手部門): 値}） ─────────────────────────
def _acc(out, r, c, v):
    if v:
        out[(r, c)] = out.get((r, c), 0.0) + v


def read_table(y):
    import xlrd
    out = {}
    if y in ('1947', '1958'):
        member = {'1947': '1947 Transactions 85-level Data.xls', '1958': '1958 Transactions 85-level Data.xls'}[y]
        s = xlrd.open_workbook(file_contents=zmember(FILES[y][1], member)).sheet_by_index(0)
        for i in range(1, s.nrows):
            try:
                _acc(out, str(int(float(s.cell_value(i, 0)))), str(int(float(s.cell_value(i, 1)))), float(s.cell_value(i, 2)))
            except (TypeError, ValueError):
                pass
    elif y == '1963':
        for line in zmember('bea_63IO367text.zip', '1963 Transactions 367-level Data.txt').decode('latin-1').splitlines():
            p = line.split()
            if len(p) >= 3:
                _acc(out, p[0], p[1], float(p[2]))
    elif y == '1967':
        s = xlrd.open_workbook(file_contents=zmember('bea_67IO484excel.zip', '1967 Transactions 484-level Data.xls')).sheet_by_index(0)
        for i in range(1, s.nrows):
            try:
                _acc(out, norm_io(s.cell_value(i, 0)).zfill(6), norm_io(s.cell_value(i, 1)).zfill(6), float(s.cell_value(i, 2)))
            except (TypeError, ValueError):
                pass
    elif y == '1972':
        s = xlrd.open_workbook(file_contents=zmember('bea_72IO496excel.zip', '1972 Transactions 496-level Data.xls')).sheet_by_index(0)
        for i in range(1, s.nrows):
            try:
                _acc(out, norm_io(s.cell_value(i, 0)).zfill(6), norm_io(s.cell_value(i, 1)).zfill(6), float(s.cell_value(i, 2)))
            except (TypeError, ValueError):
                pass
    elif y == '1977':
        for line in zmember('bea_77IO537text.zip', '1977 Transactions 537-level Data.txt').decode('latin-1').splitlines():
            p = line.split()
            if len(p) >= 3:
                _acc(out, p[0], p[1], float(p[2]))
    elif y == '1982':
        for line in zmember('bea_82_6digit.zip', '82-6DT.DAT').decode('latin-1').splitlines():
            if len(line) >= 22 and line[12:22].strip():
                _acc(out, line[0:6], line[6:12], float(line[12:22]))
    elif y == '1987':
        for line in zmember('bea_87_6digit.zip', 'TBL2-87.DAT', inner='disk2.zip').decode('latin-1').splitlines():
            p = line.split()
            if len(p) >= 4:
                _acc(out, p[0], p[1], float(p[3]))
    elif y == '1992':
        for line in zmember('bea_92_6digit.zip', 'IOUSE.TXT', inner='disk2.zip').decode('latin-1').splitlines():
            p = line.split('\t')
            if len(p) >= 4 and p[3].strip():
                _acc(out, p[0].strip(), p[1].strip(), float(p[3]))
    elif y == '1997':
        for line in zmember('bea_97_detail.zip', 'NAICSUseDetail.txt').decode('latin-1').splitlines():
            p = line.split(',')
            if len(p) >= 5:
                try:
                    _acc(out, norm_io(p[0]), norm_io(p[1]), float(p[4]))
                except ValueError:
                    pass
    elif y == '2002':
        for line in zmember('bea_02_detail.zip', 'REV_NAICSUseDetail 4-24-08.txt').decode('latin-1').splitlines()[1:]:
            if len(line) >= 210:
                try:
                    _acc(out, norm_io(line[0:10]), norm_io(line[100:110]), float(line[200:210]))
                except ValueError:
                    pass
    elif y in ('2007', '2017'):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(cache(FILES[y][1])), read_only=True, data_only=True)
        rows = list(wb[y].iter_rows(values_only=True))
        hi = next(i for i, r in enumerate(rows) if r and r[0] == 'Code')
        cols = [norm_io(x) if x is not None else None for x in rows[hi]]
        for r in rows[hi + 1:]:
            if not r or r[0] is None:
                continue
            rc = norm_io(r[0])
            for c, v in zip(cols[2:], r[2:]):
                if c is None or v in (None, '', '...'):
                    continue
                try:
                    _acc(out, rc, c, float(v))
                except (TypeError, ValueError):
                    pass
    else:
        raise KeyError(y)
    return out


# ───────────────────────── 部門 → French の割合 ─────────────────────────
def parent_of(code):
    """詳細コード → 2桁の親（'140101' → '14'・'0101' → '1'）"""
    return str(int(code[:2])) if code[:2].isdigit() else None


def sic_era_shares(y, codes, M):
    """SIC 時代の詳細表: 部門コード → {French: 割合}。戻り値 (shares, 経路の記録)"""
    plans = {}
    if y in ('1963', '1967', '1972', '1977', '1987', '1992'):
        plans = parse_plan(y)
    elif y == '1982':
        p77, p87 = parse_plan('1977'), parse_plan('1987')
        plans = dict(p87); plans.update(p77)       # 1977 を優先（同じ1977年 SIC 基準）・無ければ 1987
    base = dict(IO_SIC); base.update(ERA_OVERRIDE.get(y, {}))
    shares, route = {}, {'plan': 0, 'parent_fallback': [], 'excluded': [], 'unmapped': []}
    for c in codes:
        par = parent_of(c)
        if par is None or par in EXCLUDE_PARENT or int(par) >= 80:
            route['excluded'].append(c); continue
        title = plans.get(c, ('', '', []))[0]
        if OWNER.search(title):
            route['excluded'].append(c); continue
        sh = M.share_tokens(plans[c][2]) if c in plans and plans[c][2] else {}
        if sh:
            route['plan'] += 1
        elif par in ('11', '12'):
            # 建設の詳細部門は対照に SIC が無いことが多い（'pt. 15' か空欄）→ 建設（SIC 1500-1799）。油井・ガス井の掘削と補修だけは石油（1381-1389）
            sh = M.share_ranges(R((1381, 1389)) if re.search(r'petroleum|natural gas well|oil', title, re.I) else R((1500, 1799)))
            route['parent_fallback'].append(c)
        else:
            sh = M.share_ranges(base.get(par, []))
            if sh:
                route['parent_fallback'].append(c)
        if sh:
            shares[c] = sh
        else:
            route['unmapped'].append(c)
    return shares, route


def era85_shares(y, codes, M):
    base = dict(IO_SIC); base.update(ERA_OVERRIDE.get(y, {}))
    shares, route = {}, {'plan': 0, 'parent_fallback': [], 'excluded': [], 'unmapped': []}
    for c in codes:
        if not c.isdigit() or int(c) >= 78:
            route['excluded'].append(c); continue
        sh = M.share_ranges(base.get(c, []))
        if sh:
            shares[c] = sh; route['parent_fallback'].append(c)
        else:
            route['unmapped'].append(c)
    return shares, route


# ── NAICS 時代
def expand_naics(tok):
    tok = tok.strip().replace('*', '').replace('.0', '').strip()
    if not tok:
        return []
    if '-' in tok:
        a, b = tok.split('-', 1)
        a, b = a.strip(), b.strip()
        if not (a.isdigit() and b.isdigit()):
            return []
        head, lo = a[:len(a) - len(b)], int(a[len(a) - len(b):])
        return [head + str(x).zfill(len(b)) for x in range(lo, int(b) + 1)]
    return [tok] if tok.isdigit() else []


class NaicsChain:
    """各版の NAICS 6桁 → 1987 SIC の重み（各段 1対多 は等分）"""

    def __init__(self):
        import xlrd, openpyxl
        pad = lambda c: (c + '0') if len(c) == 5 else c  # noqa: E731  5桁の産業は6桁の末尾0と同じ

        def code(v):
            s = str(v).strip().replace('.0', '')
            return pad(s) if re.fullmatch(r'\d{5,6}', s) else None
        s = xlrd.open_workbook(file_contents=cache('census_1997_NAICS_to_1987_SIC.xls')).sheet_by_index(0)
        self.n97_sic = {}
        for i in range(1, s.nrows):
            n = code(s.cell_value(i, 0))
            try:
                sic = int(float(s.cell_value(i, 3)))
            except (TypeError, ValueError):
                sic = None
            if n:
                self.n97_sic.setdefault(n, set())
                if sic:
                    self.n97_sic[n].add(sic)

        def rev(book_bytes, sheet, col_old, col_new, xlsx=False, start=1):
            m = {}
            if xlsx:
                wb = openpyxl.load_workbook(io.BytesIO(book_bytes), read_only=True, data_only=True)
                rows = list(wb[sheet].iter_rows(values_only=True))
            else:
                sh = xlrd.open_workbook(file_contents=book_bytes).sheet_by_name(sheet)
                rows = [sh.row_values(i) for i in range(sh.nrows)]
            for r in rows[start:]:
                o = code(r[col_old]) if r[col_old] is not None else None
                nw = code(r[col_new]) if r[col_new] is not None else None
                if o and nw:
                    m.setdefault(nw, set()).add(o)
            return m
        self.r02 = rev(cache('census_1997_NAICS_to_2002_NAICS.xls'), 'Concordance 23 US NoD', 0, 2)
        self.r07 = rev(cache('census_2002_to_2007_NAICS.xls'), '02 to 07 NAICS U.S.', 0, 2, start=3)
        self.r12 = rev(cache('census_2007_to_2012_NAICS.xls'), '2007 to 2012 NAICS U.S.', 0, 2, start=3)
        self.r17 = rev(cache('census_2012_to_2017_NAICS.xlsx'), '2012 to 2017 NAICS U.S.', 0, 2, xlsx=True, start=3)
        self.codes = {'1997': set(self.n97_sic), '2002': set(self.r02), '2007': set(self.r07), '2012': set(self.r12), '2017': set(self.r17)}
        self.prev = {'2002': ('1997', self.r02), '2007': ('2002', self.r07), '2012': ('2007', self.r12), '2017': ('2012', self.r17)}
        self._memo = {}
        self.ec97 = ec97_sales()
        self._size = {}
        # 前の版のコード → 後の版の子の数（1つの前のコードが複数に分かれたら規模を等分する）
        self.nkids = {}
        for ver, (pv, m) in self.prev.items():
            cnt = {}
            for nw, olds in m.items():
                for o in olds:
                    cnt[o] = cnt.get(o, 0) + 1
            self.nkids[pv] = cnt

    def size(self, ver, c6):
        """NAICS 6桁の1997年の売上（後の版は祖先の売上を子の数で割って足す）。分からなければ None"""
        key = (ver, c6)
        if key in self._size:
            return self._size[key]
        if ver == '1997':
            v = self.ec97.get(c6)
        else:
            pv, m = self.prev[ver]
            parts = [(self.size(pv, o), self.nkids[pv].get(o, 1)) for o in m.get(c6, set())]
            parts = [a / b for a, b in parts if a is not None]
            v = sum(parts) if parts else None
        self._size[key] = v
        return v

    def sic(self, ver, c6):
        key = (ver, c6)
        if key in self._memo:
            return self._memo[key]
        if ver == '1997':
            ss = self.n97_sic.get(c6, set())
            out = {s: 1 / len(ss) for s in ss} if ss else {}
        else:
            pv, m = self.prev[ver]
            olds = m.get(c6, set())
            out = {}
            for o in olds:
                for s, w in self.sic(pv, o).items():
                    out[s] = out.get(s, 0) + w / len(olds)
            tot = sum(out.values())
            out = {k: v / tot for k, v in out.items()} if tot else {}
        self._memo[key] = out
        return out

    def sic_of_tokens(self, ver, tokens, sized=True):
        """関連 NAICS → その版の6桁コード → SIC の重み。6桁コードどうしは1997年経済センサスの売上で割る（sized）。
        売上の分からないコードは、同じ BEA 部門の中で分かるコードの中央値を当てる（全部分からなければ等分）"""
        c6s = set()
        for t in tokens:
            for p in expand_naics(t):
                c6s |= {c for c in self.codes[ver] if c.startswith(p)}
        cs = [c for c in sorted(c6s) if self.sic(ver, c)]
        if not cs:
            return {}, len(c6s), 0
        if sized:
            sz = {c: self.size(ver, c) for c in cs}
            known = sorted(v for v in sz.values() if v)
            med = known[len(known) // 2] if known else 1.0
            wts = {c: (sz[c] if sz[c] else med) for c in cs}
        else:
            wts = {c: 1.0 for c in cs}
        tw = sum(wts.values())
        out = {}
        for c in cs:
            for s, x in self.sic(ver, c).items():
                out[s] = out.get(s, 0) + x * wts[c] / tw
        return out, len(c6s), len(cs)


def ec97_sales():
    """1997年経済センサスの全国の値（$1,000）→ {NAICS 6桁: 値}。合計の行 → OPTYPE '00' の行 → 課税・非課税（T/N）の和 の順に採る。
    秘匿（D）・0 は分からないとする（0 と読まない）"""
    z = zipfile.ZipFile(io.BytesIO(cache('census_E9700A1.zip')))
    tot, op00, tn = {}, {}, {}
    with z.open('E9700A1.dat') as h:
        head = h.readline().decode('latin-1').strip().split('|')
        for line in io.TextIOWrapper(h, encoding='latin-1'):
            if not line.startswith('01000US'):
                continue
            d = dict(zip(head, line.rstrip('\n').split('|')))
            c = d['NAICS']
            if not re.fullmatch(r'\d{6}', c):
                continue
            try:
                v = float(d['ECVALUE'])
            except ValueError:
                continue
            if v <= 0 or d['ECVALUEF'] == 'D':
                continue
            if d['TAXIND'] == '' and d['OPTYPE'] == '' and d['AUXIL'] == '':
                tot[c] = v
            elif d['OPTYPE'] == '00' and d['TAXIND'] == '' and d['AUXIL'] == '':
                op00[c] = v
            elif d['TAXIND'] in ('T', 'N') and d['OPTYPE'] == '' and d['AUXIL'] == '':
                tn[c] = tn.get(c, 0.0) + v
    out = dict(tn); out.update(op00); out.update(tot)
    return out


HOUSING_CODES = {'5310HS', '531HST', '531HSO', 'S00800'}


def naics_concordance(y):
    import xlrd, openpyxl
    out = {}
    if y == '1997':
        s = xlrd.open_workbook(file_contents=zmember('bea_97_detail.zip', 'NAICS-IO.xls')).sheet_by_index(0)
        for i in range(1, s.nrows):
            r = s.row_values(i)
            c = norm_io(r[0])
            toks = [str(x).strip().replace('.0', '') for x in r[2:] if str(x).strip()]
            if c and toks:
                out[c] = (str(r[1]).strip(), toks)
    elif y == '2002':
        s = xlrd.open_workbook(file_contents=zmember('bea_02_detail.zip', 'Appendix A_rev 4-24-08.xls')).sheet_by_index(0)
        for i in range(1, s.nrows):
            r = s.row_values(i)
            c = norm_io(r[1])
            toks = [t for t in re.split(r'[,;]', str(r[3]).replace('.0', '')) if t.strip()]
            if c and toks:
                out[c] = (str(r[2]).strip(), toks)
    elif y in ('2007', '2017'):
        wb = openpyxl.load_workbook(io.BytesIO(cache(FILES[y][1])), read_only=True, data_only=True)
        sh = wb['NAICS codes'] if y == '2007' else wb['NAICS Codes']
        for r in sh.iter_rows(values_only=True):
            r = list(r) + [None] * 8
            c, title, nc = (r[2], r[3], r[5]) if y == '2007' else (r[3], r[4], r[6])
            if c is None or nc is None:
                continue
            c = norm_io(c)
            toks = [t for t in re.split(r'[,;]', str(nc)) if t.strip()]
            if c and toks:
                out[c] = (str(title).strip(), toks)
    return out


def naics_era_shares(y, codes, M, NC):
    conc = naics_concordance(y)
    shares, route = {}, {'plan': 0, 'parent_fallback': [], 'excluded': [], 'unmapped': []}
    for c in codes:
        if c in HOUSING_CODES or OWNER.search(conc.get(c, ('', []))[0]):
            route['excluded'].append(c); continue
        if c not in conc:
            route['excluded'].append(c); continue          # 政府・輸入・スクラップ・最終需要・付加価値・合計（関連 NAICS が無い）
        w, _, _ = NC.sic_of_tokens(y, conc[c][1])
        sh = M.share_sicw(w)
        if sh:
            shares[c] = sh; route['plan'] += 1
        else:
            route['unmapped'].append(c)
    return shares, route


# ───────────────────────── 49×49 へ畳む ─────────────────────────
def fold(pairs, shares):
    """pairs: {(売り手部門, 買い手部門): 値}・shares: {部門: {French: 割合}} → X[F][G]。部門の対角は除く"""
    X = {f: {g: 0.0 for g in FR49} for f in FR49}
    tot_private = mapped = diag = 0.0
    for (r, c), v in pairs.items():
        if r in shares and c in shares:
            tot_private += v
            if r == c:
                diag += v; continue
            for f, a in shares[r].items():
                for g, b in shares[c].items():
                    X[f][g] += v * a * b
            mapped += v
    return X, {'flows_between_mapped_codes': round(tot_private, 1), 'diagonal_removed_share': round(diag / tot_private, 4) if tot_private else None}


def weights(X, names):
    sup, cus = {}, {}
    for g in names:
        den = sum(X[f][g] for f in names if f != g)
        sup[g] = {f: X[f][g] / den for f in names if f != g and X[f][g] > 0} if den > 0 else {}
    for f in names:
        den = sum(X[f][g] for g in names if g != f)
        cus[f] = {g: X[f][g] / den for g in names if g != f and X[f][g] > 0} if den > 0 else {}
    return sup, cus


def to_gics(X):
    G = {a: {b: 0.0 for b in GICS11} for a in GICS11}
    for f in FR49:
        for g in FR49:
            if f in FR2GICS and g in FR2GICS:
                G[FR2GICS[f]][FR2GICS[g]] += X[f][g]
    return G


def rnd(d, k=8):
    if isinstance(d, dict):
        return {a: rnd(b, k) for a, b in d.items()}
    return round(d, k) if isinstance(d, float) else d


def build():
    fetch_all()
    M = Mapper()
    NC = NaicsChain()
    tables, tables_modal, diag = {}, {}, {}
    for y in ['1947', '1958', '1963', '1967', '1972', '1977', '1982', '1987', '1992', '1997', '2002', '2007', '2017']:
        pairs = read_table(y)
        codes = sorted({r for r, _ in pairs} | {c for _, c in pairs})
        if y in ('1947', '1958'):
            sh, route = era85_shares(y, codes, M)
            era = 'SIC85'
        elif int(y) <= 1992:
            sh, route = sic_era_shares(y, codes, M)
            era = 'SIC'
        else:
            sh, route = naics_era_shares(y, codes, M, NC)
            era = 'NAICS'
        X, cov = fold(pairs, sh)
        # 探索の変形 E_mo_modal 用: 部門を最大の割合の French 業種へ丸ごと寄せる（割り振りの雑音を除いた頑健性の確認）
        sh_modal = {c: {max(v.items(), key=lambda kv: (kv[1], -FR49.index(kv[0])))[0]: 1.0} for c, v in sh.items()}
        Xm, _ = fold(pairs, sh_modal)
        tables_modal[y] = Xm
        single = sum(1 for v in sh.values() if max(v.values()) >= 0.999)
        diag[y] = {'era': era, 'codes_in_table': len(codes), 'codes_mapped': len(sh),
                   'codes_mapped_to_one_french_industry': single,
                   'via_plan': route['plan'], 'via_parent_fallback_n': len(route['parent_fallback']),
                   'via_parent_fallback': route['parent_fallback'][:60], 'excluded_n': len(route['excluded']),
                   'unmapped': route['unmapped'][:60], **cov}
        tables[y] = X
    out = {'generated': datetime.date.today().isoformat(),
           'what': 'BEA ベンチマーク産業連関表 → French 49業種の中間財の流れ X[売り手][買い手]（部門の対角は除いた後。French の自分との取引は X に残し、重みの段で除く）と、Menzly-Ozbas の仕入れ先・顧客の重み。業種のリターンは読んでいない',
           'industries': FR49, 'fr2gics': FR2GICS, 'gics11': GICS11,
           'schedule': SCHEDULE, 'schedule_early_exploratory': SCHEDULE_EARLY, 'not_scheduled': NOT_SCHEDULED,
           'tables': {}, 'diagnostics': diag}
    for y, X in tables.items():
        sup, cus = weights(X, FR49)
        G = to_gics(X)
        gsup, gcus = weights(G, GICS11)
        msup, mcus = weights(tables_modal[y], FR49)
        out['tables'][y] = {'X': rnd(X, 4), 'supplier_w': rnd(sup), 'customer_w': rnd(cus),
                            'supplier_w_modal': rnd(msup), 'customer_w_modal': rnd(mcus),
                            'gics_X': rnd(G, 4), 'gics_supplier_w': rnd(gsup), 'gics_customer_w': rnd(gcus),
                            'no_supplier_link': [g for g in FR49 if not sup[g]], 'no_customer_link': [f for f in FR49 if not cus[f]]}
    out['sha256_tables'] = hashlib.sha256(json.dumps(out['tables'], sort_keys=True, separators=(',', ':')).encode()).hexdigest()
    tmp = OUT + '.tmp'
    open(tmp, 'w').write(json.dumps(out, ensure_ascii=False, sort_keys=True, indent=0))
    os.replace(tmp, OUT)
    return out


def show(out):
    for y, d in out['diagnostics'].items():
        t = out['tables'][y]
        print(f"{y} {d['era']:5s} 部門 {d['codes_mapped']}/{d['codes_in_table']}（1業種に収まる {d['codes_mapped_to_one_french_industry']}・"
              f"対照 {d['via_plan']}・親へ倒す {d['via_parent_fallback_n']}・除外 {d['excluded_n']}） 部門対角 {d['diagonal_removed_share']}"
              f" 写らない {d['unmapped'][:12]} 仕入れ先なし {t['no_supplier_link']} 顧客なし {t['no_customer_link']}")
    for y in ['1963', '1972', '1987', '1997', '2017']:
        t = out['tables'][y]
        for ind in ['Autos', 'Chips', 'Softw', 'Oil', 'Soda', 'Banks']:
            s = sorted(t['supplier_w'][ind].items(), key=lambda x: -x[1])[:3]
            c = sorted(t['customer_w'][ind].items(), key=lambda x: -x[1])[:3]
            print(f'  {y} {ind:6s} 主な仕入れ先', [(a, round(b, 2)) for a, b in s], '主な顧客', [(a, round(b, 2)) for a, b in c])


if __name__ == '__main__':
    o = build()
    print('書いた:', OUT, 'sha256(tables)=', o['sha256_tables'])
    if '--show' in sys.argv:
        show(o)
